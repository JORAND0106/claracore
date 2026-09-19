"""Rutas HTTP — Planillas de Tubería (ALCANTARILLA / FILTRO)."""
from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from typing import Any, Optional
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field

from main import _es_desarrollador, _require_contract_access, get_current_user, supabase
from topografia_permissions import require_permiso_topografia
from topografia_planilla_tuberia import (
    RELACIONES_ATRAQUE,
    TIPOS_PLANILLA,
    calcular_planilla_completa,
    construir_fila_consolidado,
    validar_cartera_campo,
)
from topo_crs import gk_bogota_to_wgs84

logger = logging.getLogger("claracore.topo.planilla_tuberia")
router = APIRouter(tags=["topografia-planillas-tuberia"])


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _uid(user) -> Optional[str]:
    return (user or {}).get("sub") or (user or {}).get("id")


def _row(table: str, **eq) -> Optional[dict]:
    q = supabase.table(table).select("*")
    for k, v in eq.items():
        q = q.eq(k, v)
    data = q.limit(1).execute().data or []
    return data[0] if data else None


def _perm(user, accion: str) -> None:
    require_permiso_topografia(user, accion)  # type: ignore[arg-type]


def _assert_editable(p: dict) -> None:
    if (p.get("estado") or "").lower() in ("cerrado", "validado"):
        raise HTTPException(409, "Planilla cerrada/validada: no editable.")


def _assert_version(p: dict, version: int) -> None:
    actual = int(p.get("version") or 1)
    if int(version) != actual:
        raise HTTPException(409, f"Conflicto de versión: esperaba {actual}, recibió {version}.")


class CrearBody(BaseModel):
    tipo: str = "ALCANTARILLA"
    nombre: Optional[str] = None
    pk_id: Optional[str] = None
    costado: Optional[str] = None


class ParamsBody(BaseModel):
    version: int
    tipo: Optional[str] = None
    nombre: Optional[str] = None
    pk_id: Optional[str] = None
    costado: Optional[str] = None
    abscisa_inicial_txt: Optional[str] = None
    abscisa_final_txt: Optional[str] = None
    diametro_m: Optional[float] = None
    espesor_m: Optional[float] = None
    ancho_excavacion_m: Optional[float] = None
    relacion_atraque: Optional[str] = None
    material: Optional[str] = None
    norte_ref: Optional[float] = None
    este_ref: Optional[float] = None


class FilaBody(BaseModel):
    orden: int
    abscisa: Optional[float] = None
    terreno_natural: Optional[float] = None
    subrasante_via: Optional[float] = None
    terminado_filtro: Optional[float] = None
    cota_fondo_excavacion: Optional[float] = None
    norte: Optional[float] = None
    este: Optional[float] = None
    observacion: Optional[str] = None


class DescBody(BaseModel):
    codigo: str
    cantidad: float = 0
    nota: Optional[str] = None


class CarteraBody(BaseModel):
    version: int
    filas: list[FilaBody]
    descuentos_manuales: list[DescBody] = Field(default_factory=list)


def _filas(planilla_id: str) -> list[dict]:
    return (
        supabase.table("topo_planilla_tuberia_filas")
        .select("*").eq("planilla_id", planilla_id).order("orden").execute().data or []
    )


def _descuentos(planilla_id: str) -> list[dict]:
    return (
        supabase.table("topo_planilla_tuberia_descuentos")
        .select("*").eq("planilla_id", planilla_id).execute().data or []
    )


def _as_campo(filas_db: list[dict]) -> list[dict]:
    return [{
        "orden": f.get("orden"),
        "abscisa": f.get("abscisa"),
        "terreno_natural": f.get("terreno_natural"),
        "subrasante_via": f.get("subrasante_via"),
        "terminado_filtro": f.get("terminado_filtro"),
        "cota_fondo_excavacion": f.get("cota_fondo_excavacion"),
        "norte": f.get("norte"),
        "este": f.get("este"),
        "observacion": f.get("observacion"),
    } for f in filas_db]


def _calcular(planilla: dict, filas_db: list[dict], desc_db: list[dict]) -> dict:
    diam = float(planilla.get("diametro_m") or 0)
    ancho = float(planilla.get("ancho_excavacion_m") or 0)
    if diam <= 0 or ancho <= 0:
        raise HTTPException(422, "Configure diámetro y ancho de excavación.")
    return calcular_planilla_completa(
        tipo=planilla.get("tipo") or "ALCANTARILLA",
        diametro_m=diam,
        espesor_m=float(planilla.get("espesor_m") or 0),
        ancho_excavacion_m=ancho,
        relacion_atraque=planilla.get("relacion_atraque") or "1:3",
        filas_campo=_as_campo(filas_db),
        descuentos_manuales=[{"codigo": d["codigo"], "cantidad": d.get("cantidad")} for d in desc_db],
    )


def _audit(contrato_id: int, planilla_id: str, accion: str, user, detalle: dict) -> None:
    try:
        supabase.table("topo_planilla_tuberia_auditoria").insert({
            "planilla_id": planilla_id,
            "contrato_id": contrato_id,
            "accion": accion,
            "usuario_id": _uid(user),
            "detalle": detalle,
        }).execute()
    except Exception:
        logger.exception("audit")


def _detalle(contrato_id: int, planilla_id: str) -> dict:
    planilla = _row("topo_planillas_tuberia", id=planilla_id, contrato_id=contrato_id)
    if not planilla:
        raise HTTPException(404, "Planilla no encontrada")
    filas = _filas(planilla_id)
    descuentos = _descuentos(planilla_id)
    calculo = None
    try:
        if planilla.get("diametro_m") and planilla.get("ancho_excavacion_m"):
            calculo = _calcular(planilla, filas, descuentos)
            sec = calculo["seccion"]
            supabase.table("topo_planillas_tuberia").update({
                "altura_relleno_m": sec["altura_relleno_m"],
                "area_1_m2": sec["area_1_m2"],
                "area_2_m2": sec["area_2_m2"],
                "updated_at": _now(),
            }).eq("id", planilla_id).execute()
            planilla = _row("topo_planillas_tuberia", id=planilla_id, contrato_id=contrato_id) or planilla
    except HTTPException:
        pass
    except Exception as exc:
        logger.warning("calculo %s: %s", planilla_id, exc)

    coords = None
    if planilla.get("este_ref") is not None and planilla.get("norte_ref") is not None:
        try:
            lon, lat = gk_bogota_to_wgs84(float(planilla["este_ref"]), float(planilla["norte_ref"]))
            coords = {"lon": lon, "lat": lat}
        except Exception:
            coords = None

    return {
        "planilla": planilla,
        "filas_campo": filas,
        "descuentos_manuales": descuentos,
        "calculo": calculo,
        "coords_wgs84": coords,
        "relaciones_atraque": list(RELACIONES_ATRAQUE),
        "tipos": list(TIPOS_PLANILLA),
    }


def _replace_filas(planilla_id: str, filas: list[dict]) -> list[dict]:
    prev = _filas(planilla_id)
    prev_ids = [p["id"] for p in prev if p.get("id")]
    payloads = [{
        "id": str(uuid4()),
        "planilla_id": planilla_id,
        "orden": int(f["orden"]),
        "abscisa": f.get("abscisa"),
        "terreno_natural": f.get("terreno_natural"),
        "subrasante_via": f.get("subrasante_via"),
        "terminado_filtro": f.get("terminado_filtro"),
        "cota_fondo_excavacion": f.get("cota_fondo_excavacion"),
        "norte": f.get("norte"),
        "este": f.get("este"),
        "observacion": f.get("observacion"),
    } for f in filas]
    new_ids = [p["id"] for p in payloads]
    try:
        if payloads:
            supabase.table("topo_planilla_tuberia_filas").insert(payloads).execute()
        all_rows = (
            supabase.table("topo_planilla_tuberia_filas")
            .select("id").eq("planilla_id", planilla_id).execute().data or []
        )
        inserted = [r for r in all_rows if r["id"] in set(new_ids)]
        if len(inserted) != len(payloads):
            raise RuntimeError(f"Inserción incompleta {len(inserted)}/{len(payloads)}")
        if prev_ids:
            for i in range(0, len(prev_ids), 100):
                supabase.table("topo_planilla_tuberia_filas").delete().in_("id", prev_ids[i:i + 100]).execute()
        rows = _filas(planilla_id)
        if len(rows) != len(payloads):
            raise RuntimeError(f"Cartera inconsistente {len(rows)}/{len(payloads)}")
        return rows
    except Exception:
        try:
            if new_ids:
                for i in range(0, len(new_ids), 100):
                    supabase.table("topo_planilla_tuberia_filas").delete().in_("id", new_ids[i:i + 100]).execute()
        except Exception:
            pass
        raise


@router.get("/{contrato_id}/planillas-tuberia")
def listar(contrato_id: int, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "ver")
    return (
        supabase.table("topo_planillas_tuberia")
        .select("id,tipo,nombre,pk_id,costado,estado,version,diametro_m,relacion_atraque,created_at,updated_at,cerrado_at")
        .eq("contrato_id", contrato_id).order("created_at", desc=True).execute().data or []
    )


@router.post("/{contrato_id}/planillas-tuberia")
def crear(contrato_id: int, body: CrearBody, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "editar")
    tipo = (body.tipo or "ALCANTARILLA").upper()
    if tipo not in TIPOS_PLANILLA:
        raise HTTPException(422, f"Tipo inválido: {tipo}")
    row = supabase.table("topo_planillas_tuberia").insert({
        "contrato_id": contrato_id,
        "tipo": tipo,
        "nombre": body.nombre or f"Planilla {tipo}",
        "pk_id": body.pk_id,
        "costado": body.costado,
        "estado": "borrador",
        "version": 1,
        "relacion_atraque": "1:3",
        "espesor_m": 0,
        "creado_por": _uid(current_user),
        "created_at": _now(),
        "updated_at": _now(),
    }).execute().data
    if not row:
        raise HTTPException(500, "No se pudo crear la planilla")
    return _detalle(contrato_id, row[0]["id"])


@router.get("/{contrato_id}/planillas-tuberia/{planilla_id}")
def obtener(contrato_id: int, planilla_id: str, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "ver")
    return _detalle(contrato_id, planilla_id)


@router.put("/{contrato_id}/planillas-tuberia/{planilla_id}/params")
def actualizar_params(contrato_id: int, planilla_id: str, body: ParamsBody, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "editar")
    p = _row("topo_planillas_tuberia", id=planilla_id, contrato_id=contrato_id)
    if not p:
        raise HTTPException(404, "Planilla no encontrada")
    _assert_editable(p)
    _assert_version(p, body.version)
    patch: dict[str, Any] = {"updated_at": _now(), "version": int(p.get("version") or 1) + 1}
    if body.tipo is not None:
        tipo_up = str(body.tipo).upper()
        if tipo_up not in TIPOS_PLANILLA:
            raise HTTPException(422, f"Tipo inválido: {tipo_up}")
        patch["tipo"] = tipo_up
    for field in (
        "nombre", "pk_id", "costado", "abscisa_inicial_txt", "abscisa_final_txt",
        "diametro_m", "espesor_m", "ancho_excavacion_m", "material", "norte_ref", "este_ref",
    ):
        val = getattr(body, field)
        if val is not None:
            patch[field] = val
    if body.relacion_atraque is not None:
        if body.relacion_atraque not in RELACIONES_ATRAQUE:
            raise HTTPException(422, "Relación de atraque inválida")
        patch["relacion_atraque"] = body.relacion_atraque
    diam = patch.get("diametro_m", p.get("diametro_m"))
    esp = patch.get("espesor_m", p.get("espesor_m") or 0)
    ancho = patch.get("ancho_excavacion_m", p.get("ancho_excavacion_m"))
    rel = patch.get("relacion_atraque", p.get("relacion_atraque") or "1:3")
    tipo_calc = patch.get("tipo", p.get("tipo") or "ALCANTARILLA")
    if diam and ancho:
        calc = calcular_planilla_completa(
            tipo=tipo_calc, diametro_m=float(diam), espesor_m=float(esp or 0),
            ancho_excavacion_m=float(ancho), relacion_atraque=rel, filas_campo=[],
        )
        sec = calc["seccion"]
        patch.update({
            "altura_relleno_m": sec["altura_relleno_m"],
            "area_1_m2": sec["area_1_m2"],
            "area_2_m2": sec["area_2_m2"],
        })
    supabase.table("topo_planillas_tuberia").update(patch).eq("id", planilla_id).execute()
    return _detalle(contrato_id, planilla_id)


@router.put("/{contrato_id}/planillas-tuberia/{planilla_id}/cartera")
def guardar_cartera(contrato_id: int, planilla_id: str, body: CarteraBody, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "editar")
    p = _row("topo_planillas_tuberia", id=planilla_id, contrato_id=contrato_id)
    if not p:
        raise HTTPException(404, "Planilla no encontrada")
    _assert_editable(p)
    _assert_version(p, body.version)

    filas_util = []
    for f in body.filas:
        d = f.model_dump()
        if any(d.get(k) is not None for k in (
            "abscisa", "terreno_natural", "subrasante_via", "terminado_filtro", "cota_fondo_excavacion"
        )):
            filas_util.append(d)

    valid = validar_cartera_campo(filas_util, p.get("tipo") or "ALCANTARILLA")
    if not valid.get("ok"):
        raise HTTPException(422, {
            "mensaje": "Validación restrictiva: corrija errores antes de guardar.",
            "errores": valid.get("errores") or [],
            "infos": valid.get("infos") or [],
        })

    try:
        rows = _replace_filas(planilla_id, filas_util)
    except Exception as exc:
        raise HTTPException(500, f"Error al guardar cartera: {exc}") from exc

    supabase.table("topo_planilla_tuberia_descuentos").delete().eq("planilla_id", planilla_id).execute()
    desc_rows = [
        {"planilla_id": planilla_id, "codigo": d.codigo, "cantidad": float(d.cantidad), "nota": d.nota}
        for d in (body.descuentos_manuales or []) if d.codigo
    ]
    if desc_rows:
        supabase.table("topo_planilla_tuberia_descuentos").insert(desc_rows).execute()

    nueva_v = int(p.get("version") or 1) + 1
    supabase.table("topo_planillas_tuberia").update({"version": nueva_v, "updated_at": _now()}).eq("id", planilla_id).execute()
    detalle = _detalle(contrato_id, planilla_id)
    if len(detalle.get("filas_campo") or []) != len(rows):
        raise HTTPException(500, "Verificación post-guardado falló.")
    return {**detalle, "verified": True, "count": len(rows), "version": nueva_v,
            "validacion": {"infos": valid.get("infos") or []}}


@router.post("/{contrato_id}/planillas-tuberia/{planilla_id}/calcular")
def calcular(contrato_id: int, planilla_id: str, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "ver")
    return _detalle(contrato_id, planilla_id)


@router.post("/{contrato_id}/planillas-tuberia/{planilla_id}/cerrar")
def cerrar(contrato_id: int, planilla_id: str, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "editar")
    p = _row("topo_planillas_tuberia", id=planilla_id, contrato_id=contrato_id)
    if not p:
        raise HTTPException(404, "Planilla no encontrada")
    if (p.get("estado") or "").lower() != "borrador":
        raise HTTPException(422, "Solo se cierran planillas en borrador.")
    filas = _filas(planilla_id)
    if not filas:
        raise HTTPException(422, "No hay cartera para cerrar.")
    calc = _calcular(p, filas, _descuentos(planilla_id))
    now = _now()
    consol = construir_fila_consolidado(
        {**p, "estado": "cerrado", "cerrado_at": now, "contrato_id": contrato_id}, calc
    )
    supabase.table("topo_planillas_tuberia").update({
        "estado": "cerrado", "cerrado_at": now, "cerrado_por": _uid(current_user),
        "calculo_snapshot": calc, "version": int(p.get("version") or 1) + 1, "updated_at": now,
    }).eq("id", planilla_id).execute()
    supabase.table("topo_planilla_tuberia_consolidado").delete().eq("planilla_id", planilla_id).execute()
    supabase.table("topo_planilla_tuberia_consolidado").insert({
        "contrato_id": contrato_id,
        "planilla_id": planilla_id,
        "c01_planilla_id": consol.get("c01_planilla_id"),
        "c02_tipo": consol.get("c02_tipo"),
        "c03_pk_id": consol.get("c03_pk_id"),
        "c04_nombre": consol.get("c04_nombre"),
        "c05_costado": consol.get("c05_costado"),
        "c06_abscisa_inicial": consol.get("c06_abscisa_inicial"),
        "c07_abscisa_final": consol.get("c07_abscisa_final"),
        "c08_longitud_m": consol.get("c08_longitud_m"),
        "c09_diametro_m": consol.get("c09_diametro_m"),
        "c10_espesor_m": consol.get("c10_espesor_m"),
        "c11_relacion_atraque": consol.get("c11_relacion_atraque"),
        "c12_ancho_excavacion_m": consol.get("c12_ancho_excavacion_m"),
        "c13_vol_excavacion_m3": consol.get("c13_vol_excavacion_m3"),
        "c14_vol_triturado_m3": consol.get("c14_vol_triturado_m3"),
        "c15_vol_relleno_m3": consol.get("c15_vol_relleno_m3"),
        "c16_area_geotextil_m2": consol.get("c16_area_geotextil_m2"),
        "c17_long_tuberia_m": consol.get("c17_long_tuberia_m"),
        "c18_norte_ref": consol.get("c18_norte_ref"),
        "c19_este_ref": consol.get("c19_este_ref"),
        "c20_estado": "cerrado",
        "c21_cerrado_at": now,
        "c22_contrato_id": contrato_id,
    }).execute()
    _audit(contrato_id, planilla_id, "CERRAR", current_user, {"consolidado": consol})
    return _detalle(contrato_id, planilla_id)


@router.post("/{contrato_id}/planillas-tuberia/{planilla_id}/reabrir")
def reabrir(contrato_id: int, planilla_id: str, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    if not _es_desarrollador(current_user):
        raise HTTPException(403, "Solo Desarrollador puede reabrir planillas.")
    p = _row("topo_planillas_tuberia", id=planilla_id, contrato_id=contrato_id)
    if not p:
        raise HTTPException(404, "Planilla no encontrada")
    if (p.get("estado") or "").lower() not in ("cerrado", "validado"):
        raise HTTPException(422, "La planilla no está cerrada.")
    now = _now()
    supabase.table("topo_planillas_tuberia").update({
        "estado": "borrador", "nivel_validacion": 0, "validado_at": None, "validado_por": None,
        "reabierto_at": now, "reabierto_por": _uid(current_user),
        "version": int(p.get("version") or 1) + 1, "updated_at": now,
    }).eq("id", planilla_id).execute()
    supabase.table("topo_planilla_tuberia_consolidado").delete().eq("planilla_id", planilla_id).execute()
    _audit(contrato_id, planilla_id, "REABRIR", current_user, {"estado_previo": p.get("estado")})
    return _detalle(contrato_id, planilla_id)


@router.post("/{contrato_id}/planillas-tuberia/{planilla_id}/revocar-validacion")
def revocar(contrato_id: int, planilla_id: str, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    if not _es_desarrollador(current_user):
        raise HTTPException(403, "Solo Desarrollador puede revocar validación.")
    p = _row("topo_planillas_tuberia", id=planilla_id, contrato_id=contrato_id)
    if not p:
        raise HTTPException(404, "Planilla no encontrada")
    now = _now()
    supabase.table("topo_planillas_tuberia").update({
        "nivel_validacion": 0,
        "estado": "cerrado" if p.get("cerrado_at") else "borrador",
        "validado_at": None, "validado_por": None,
        "validacion_revocada_at": now, "validacion_revocada_por": _uid(current_user),
        "version": int(p.get("version") or 1) + 1, "updated_at": now,
    }).eq("id", planilla_id).execute()
    _audit(contrato_id, planilla_id, "REVOCAR_VALIDACION", current_user, {})
    return _detalle(contrato_id, planilla_id)


@router.get("/{contrato_id}/planillas-tuberia-consolidado")
def consolidado(contrato_id: int, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "ver")
    return (
        supabase.table("topo_planilla_tuberia_consolidado").select("*")
        .eq("contrato_id", contrato_id).order("created_at", desc=True).execute().data or []
    )


@router.get("/{contrato_id}/planillas-tuberia/{planilla_id}/excel")
def excel(contrato_id: int, planilla_id: str, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "exportar")
    det = _detalle(contrato_id, planilla_id)
    calc = det.get("calculo")
    if not calc:
        raise HTTPException(422, "Sin cálculo para exportar.")
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise HTTPException(500, "openpyxl no disponible") from exc

    p = det["planilla"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilla"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showZeros = False
    ws["A1"] = "PLANILLA DE TUBERÍA"
    ws["A2"] = "Tipo:"
    ws["B2"] = p.get("tipo")
    ws["C2"] = "PK:"
    ws["D2"] = p.get("pk_id") or ""
    ws["A3"] = (
        f"Ø={p.get('diametro_m')} esp={p.get('espesor_m')} B={p.get('ancho_excavacion_m')} "
        f"Rel={p.get('relacion_atraque')} H.Relleno={p.get('altura_relleno_m')}"
    )
    ws["A4"] = "Relación atraque:"
    ws["B4"] = p.get("relacion_atraque") or "1:3"

    dv_tipo = DataValidation(type="list", formula1='"ALCANTARILLA,FILTRO"', allow_blank=False)
    dv_rel = DataValidation(
        type="list",
        formula1='"' + ",".join(RELACIONES_ATRAQUE) + '"',
        allow_blank=False,
    )
    ws.add_data_validation(dv_tipo)
    ws.add_data_validation(dv_rel)
    dv_tipo.add(ws["B2"])
    dv_rel.add(ws["B4"])

    for c, h in enumerate(["#", "Abscisa", "TN", "Nivel", "CFE", "H.Exc", "H.Trit", "H.Rell", "Ancho Geo"], 1):
        ws.cell(5, c, h)
    d_ext = calc["seccion"]["diametro_externo_m"]
    b = calc["seccion"]["ancho_excavacion_m"]
    h_atr = p.get("altura_relleno_m")
    r = 6
    for f in calc["cartera"]["filas"]:
        if f.get("vacio"):
            continue
        ws.cell(r, 1, f.get("orden"))
        ws.cell(r, 2, f.get("abscisa"))
        ws.cell(r, 3, f.get("terreno_natural"))
        ws.cell(r, 4, f.get("nivel_referencia"))
        ws.cell(r, 5, f.get("cota_fondo_excavacion"))
        ws.cell(r, 6, f'=IF(OR(C{r}="",E{r}=""),"",C{r}-E{r})')
        ws.cell(r, 7, h_atr)
        ws.cell(r, 8, f'=IF(OR(D{r}="",E{r}=""),"",D{r}-(E{r}+{d_ext}))')
        if p.get("tipo") == "FILTRO":
            ws.cell(r, 9, f'=IF(F{r}="","",{b}+2*(G{r}+MAX(0,H{r})))')
        else:
            ws.cell(r, 9, f'=IF(F{r}="","",{b}+2*F{r})')
        r += 1

    ws2 = wb.create_sheet("Cantidades")
    ws2.sheet_view.showGridLines = False
    ws2.sheet_view.showZeros = False
    ws2.append(["Código", "Nombre", "Unidad", "Bruto", "Descuentos", "Neto"])
    for i, n in enumerate(calc.get("netos") or [], 2):
        ws2.cell(i, 1, n["codigo"]); ws2.cell(i, 2, n["nombre"]); ws2.cell(i, 3, n["unidad"])
        ws2.cell(i, 4, n["bruto"]); ws2.cell(i, 5, n["descuentos"]); ws2.cell(i, 6, f"=D{i}-E{i}")

    ws3 = wb.create_sheet("_Perfil"); ws3.sheet_state = "hidden"
    ws3.append(["Abscisa", "TN", "Nivel", "CFE"])
    for a, tn, nv, cfe in zip(
        calc["perfil"]["abscisas"], calc["perfil"]["terreno_natural"],
        calc["perfil"]["nivel_referencia"], calc["perfil"]["cota_fondo_excavacion"],
    ):
        ws3.append([a, tn, nv, cfe])

    ws4 = wb.create_sheet("_Consolidado"); ws4.sheet_state = "hidden"
    fila_c = construir_fila_consolidado(p, calc)
    ws4.append(list(fila_c.keys())); ws4.append(list(fila_c.values()))

    buf = io.BytesIO(); wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="planilla_tuberia_{planilla_id[:8]}.xlsx"'},
    )


@router.get("/{contrato_id}/planillas-tuberia/{planilla_id}/pdf")
def pdf(contrato_id: int, planilla_id: str, current_user=Depends(get_current_user)):
    _require_contract_access(current_user, contrato_id)
    _perm(current_user, "exportar")
    det = _detalle(contrato_id, planilla_id)
    calc = det.get("calculo")
    if not calc:
        raise HTTPException(422, "Sin cálculo para exportar.")
    p = det["planilla"]

    def fmt(v, d=3):
        if v is None:
            return "—"
        try:
            return f"{float(v):.{d}f}"
        except Exception:
            return str(v)

    rows = "".join(
        f"<tr><td>{f.get('orden')}</td><td>{fmt(f.get('abscisa'),2)}</td>"
        f"<td>{fmt(f.get('terreno_natural'))}</td><td>{fmt(f.get('nivel_referencia'))}</td>"
        f"<td>{fmt(f.get('cota_fondo_excavacion'))}</td><td>{fmt(f.get('altura_excavacion'))}</td>"
        f"<td>{fmt(f.get('altura_triturado'))}</td><td>{fmt(f.get('altura_relleno'))}</td>"
        f"<td>{fmt(f.get('ancho_geotextil'))}</td></tr>"
        for f in calc["cartera"]["filas"] if not f.get("vacio")
    )
    cants = "".join(
        f"<tr><td>{n['codigo']}</td><td>{n['nombre']}</td><td>{n['unidad']}</td>"
        f"<td>{fmt(n['bruto'])}</td><td>{fmt(n['descuentos'])}</td><td>{fmt(n['neto'])}</td></tr>"
        for n in calc.get("netos") or []
    )
    descs = "".join(
        f"<tr><td>{d['codigo']}</td><td>{d['nombre']}</td><td>{d.get('item_cant_codigo','')}</td>"
        f"<td>{fmt(d['cantidad'])}</td></tr>"
        for d in calc.get("descuentos") or []
    )
    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"/>
    <style>body{{font-family:Arial,sans-serif;font-size:11px}}
    table{{border-collapse:collapse;width:100%;margin-bottom:12px}} th,td{{border:1px solid #cbd5e1;padding:3px 5px}}
    th{{background:#e2e8f0}} .meta{{color:#475569}}</style></head><body>
    <h1>Planilla de Tubería — {p.get('tipo')}</h1>
    <p class="meta">{p.get('nombre') or ''} · PK {p.get('pk_id') or ''} · Costado {p.get('costado') or '—'}
    · Ø={fmt(p.get('diametro_m'))} · Rel={p.get('relacion_atraque') or ''}
    · H.Relleno={fmt(p.get('altura_relleno_m'),4)} · A1={fmt(p.get('area_1_m2'),4)} · A2={fmt(p.get('area_2_m2'),4)}</p>
    <table><thead><tr><th>#</th><th>Abs</th><th>TN</th><th>Nivel</th><th>CFE</th>
    <th>H.Exc</th><th>H.Trit</th><th>H.Rell</th><th>Geo</th></tr></thead><tbody>{rows}</tbody></table>
    <h2>Resumen de cantidades</h2>
    <table><thead><tr><th>Cód</th><th>Ítem</th><th>Ud</th><th>Bruto</th><th>Desc</th><th>Neto</th></tr></thead>
    <tbody>{cants}</tbody></table>
    <h2>Descuentos específicos (por código de ítem)</h2>
    <table><thead><tr><th>Cód</th><th>Nombre</th><th>Ítem cant.</th><th>Cantidad</th></tr></thead>
    <tbody>{descs}</tbody></table>
    <p class="meta">Estado: {p.get('estado') or ''} · Firmas resueltas por rol en ClaraCore al validar.</p>
    </body></html>"""
    try:
        from topografia_utils import to_pdf_bytes
        content, media = to_pdf_bytes(html), "application/pdf"
    except Exception:
        content, media = html.encode("utf-8"), "text/html; charset=utf-8"
    return Response(
        content=content, media_type=media,
        headers={"Content-Disposition": f'attachment; filename="planilla_tuberia_{planilla_id[:8]}.pdf"'},
    )
