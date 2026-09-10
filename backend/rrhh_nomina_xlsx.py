"""Exportación Excel — relación consolidada de nómina con fórmulas reales."""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from rrhh_nomina_params import (
    ARL_PCT,
    FSP_PCT_BASE,
    FSP_UMBRAL_SMMLV,
    PCT_CAJA,
    PCT_CESANTIAS,
    PCT_ICBF,
    PCT_INTERES_CESANTIAS_ANUAL,
    PCT_PENSION_EMPLEADO,
    PCT_PENSION_PATRONAL,
    PCT_PRIMA,
    PCT_SALUD_EMPLEADO,
    PCT_SALUD_PATRONAL,
    PCT_SENA,
    PCT_VACACIONES,
    UMBRAL_EXONERACION_SMMLV,
    params_for_year,
)

_FILL_HDR = PatternFill("solid", fgColor="0077B6")
_FONT_HDR = Font(bold=True, color="FFFFFF", size=10)
_FONT_TITLE = Font(bold=True, size=14, color="0077B6")
_SIDE = Side(style="thin", color="CCCCCC")
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
_AL_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_AL_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_AL_RIGHT = Alignment(horizontal="right", vertical="center")
_BOGOTA = ZoneInfo("America/Bogota")

# Columnas base (valores) y calculadas (fórmulas).
# A..K base; L..AE calculadas + email; AF tasa ARL (base auxiliar).
HEADERS = [
    "Documento",                 # A
    "Colaborador",               # B
    "Cargo",                     # C
    "Salario base",              # D
    "Salario periodo",           # E
    "Horas extras",              # F
    "Recargos",                  # G
    "Bonificaciones",            # H
    "Descuento novedades",       # I
    "Detalle novedades",         # J
    "Subsidio transporte",       # K
    "IBC",                       # L  fórmula
    "Total devengado",           # M  fórmula
    "Ded. salud",                # N  fórmula
    "Ded. pensión",              # O  fórmula
    "Ded. FSP",                  # P  fórmula
    "Total deducciones",         # Q  fórmula
    "Neto a pagar",              # R  fórmula
    "Aporte salud patronal",     # S  fórmula
    "Aporte pensión patronal",   # T  fórmula
    "ARL",                       # U  fórmula
    "Caja compensación",         # V  fórmula
    "SENA",                      # W  fórmula
    "ICBF",                      # X  fórmula
    "Total aportes patronales",  # Y  fórmula
    "Prov. cesantías",           # Z  fórmula
    "Prov. interés cesantías",   # AA fórmula
    "Prov. prima",               # AB fórmula
    "Prov. vacaciones",          # AC fórmula
    "Total provisiones",         # AD fórmula
    "Email",                     # AE
    "Tasa ARL",                  # AF (base auxiliar para fórmula ARL)
]

# Índices 1-based omitidos: las fórmulas se generan en `_formulas_fila`.


def _nov_detalle(item: dict) -> str:
    det = item.get("detalle_json") or {}
    if isinstance(det, str):
        return ""
    parts = []
    for n in det.get("novedades") or []:
        parts.append(
            f"{n.get('tipo')} {n.get('dias')}d ({n.get('fecha_inicio')}–{n.get('fecha_fin')})"
        )
    return "; ".join(parts)


def _nombre(item: dict, trabajadores_by_id: Dict[int, dict]) -> str:
    det = item.get("detalle_json") or {}
    if isinstance(det, dict) and det.get("nombre"):
        return det["nombre"]
    t = trabajadores_by_id.get(int(item.get("trabajador_id") or 0)) or {}
    return f"{t.get('nombres') or ''} {t.get('apellidos') or ''}".strip()


def _doc(item: dict, trabajadores_by_id: Dict[int, dict]) -> str:
    det = item.get("detalle_json") or {}
    if isinstance(det, dict) and det.get("documento"):
        return det["documento"]
    t = trabajadores_by_id.get(int(item.get("trabajador_id") or 0)) or {}
    return f"{t.get('tipo_documento') or 'CC'} {t.get('numero_documento') or ''}".strip()


def _arl_tasa(item: dict, trabajadores_by_id: Dict[int, dict]) -> float:
    det = item.get("detalle_json") or {}
    nivel = None
    if isinstance(det, dict):
        nivel = det.get("arl_nivel")
    if not nivel:
        t = trabajadores_by_id.get(int(item.get("trabajador_id") or 0)) or {}
        nivel = t.get("arl_nivel_riesgo")
    key = str(nivel or "I").strip().upper()
    return float(ARL_PCT.get(key, ARL_PCT["I"]))


def _formulas_fila(row: int) -> Dict[int, str]:
    """
    Fórmulas por columna (1-based) para una fila de datos.
    Referencias a Parametros!$B$2 (SMMLV) y tasas en Parametros.
    IBC = salario_periodo - descuento + extras + recargos + bonificaciones
    """
    r = row
    # Parametros sheet cells:
    # B2 SMMLV, B3 salud emp, B4 pension emp, B5 fsp base, B6 umbral fsp (×SMMLV)
    # B7 salud pat, B8 pension pat, B9 caja, B10 sena, B11 icbf
    # B12 umbral exoneración (×SMMLV)
    # B13 cesantías, B14 interés anual, B15 prima, B16 vacaciones
    ibc = f"MAX(0,E{r}-I{r}+F{r}+G{r}+H{r})"
    return {
        12: f"={ibc}",  # IBC
        13: f"=E{r}-I{r}+F{r}+G{r}+H{r}+K{r}",  # Total devengado
        14: f"=ROUND(L{r}*Parametros!$B$3,2)",  # salud emp
        15: f"=ROUND(L{r}*Parametros!$B$4,2)",  # pensión emp
        16: (
            f'=IF(L{r}<Parametros!$B$6*Parametros!$B$2,0,'
            f"ROUND(L{r}*Parametros!$B$5,2))"
        ),  # FSP (base 1% desde 4 SMMLV)
        17: f"=N{r}+O{r}+P{r}",  # total deducciones
        18: f"=M{r}-Q{r}",  # neto
        19: (
            f"=IF(L{r}<Parametros!$B$12*Parametros!$B$2,0,"
            f"ROUND(L{r}*Parametros!$B$7,2))"
        ),  # salud patronal
        20: f"=ROUND(L{r}*Parametros!$B$8,2)",  # pensión patronal
        21: f"=ROUND(L{r}*AF{r},2)",  # ARL
        22: f"=ROUND(L{r}*Parametros!$B$9,2)",  # caja
        23: (
            f"=IF(L{r}<Parametros!$B$12*Parametros!$B$2,0,"
            f"ROUND(L{r}*Parametros!$B$10,2))"
        ),  # SENA
        24: (
            f"=IF(L{r}<Parametros!$B$12*Parametros!$B$2,0,"
            f"ROUND(L{r}*Parametros!$B$11,2))"
        ),  # ICBF
        25: f"=S{r}+T{r}+U{r}+V{r}+W{r}+X{r}",  # total aportes
        26: f"=ROUND(L{r}*Parametros!$B$13,2)",  # cesantías (sobre IBC/base prov)
        27: f"=ROUND(Z{r}*Parametros!$B$14/12,2)",  # interés cesantías
        28: f"=ROUND(L{r}*Parametros!$B$15,2)",  # prima
        29: f"=ROUND(L{r}*Parametros!$B$16,2)",  # vacaciones
        30: f"=Z{r}+AA{r}+AB{r}+AC{r}",  # total provisiones
    }


def _write_parametros_sheet(wb: Workbook, *, anio: int, smmlv: float) -> None:
    ws = wb.create_sheet("Parametros", 0)
    ws["A1"] = "Parámetro"
    ws["B1"] = "Valor"
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)

    rows = [
        (2, "SMMLV", smmlv),
        (3, "% Salud empleado", PCT_SALUD_EMPLEADO),
        (4, "% Pensión empleado", PCT_PENSION_EMPLEADO),
        (5, "% FSP base", FSP_PCT_BASE),
        (6, "Umbral FSP (× SMMLV)", FSP_UMBRAL_SMMLV),
        (7, "% Salud patronal", PCT_SALUD_PATRONAL),
        (8, "% Pensión patronal", PCT_PENSION_PATRONAL),
        (9, "% Caja compensación", PCT_CAJA),
        (10, "% SENA", PCT_SENA),
        (11, "% ICBF", PCT_ICBF),
        (12, "Umbral exoneración (× SMMLV)", UMBRAL_EXONERACION_SMMLV),
        (13, "% Cesantías", PCT_CESANTIAS),
        (14, "% Interés cesantías anual", PCT_INTERES_CESANTIAS_ANUAL),
        (15, "% Prima", PCT_PRIMA),
        (16, "% Vacaciones", PCT_VACACIONES),
        (17, "Año parámetros", anio),
    ]
    for r, label, val in rows:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=val)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws["A19"] = (
        "Las columnas calculadas de la hoja Nómina referencian estas celdas. "
        "La tasa ARL por colaborador está en la columna «Tasa ARL» de Nómina."
    )
    ws.merge_cells("A19:B21")


def build_nomina_xlsx(
    nomina: dict,
    items: Sequence[dict],
    *,
    trabajadores: Optional[Sequence[dict]] = None,
    contrato_label: str = "",
) -> bytes:
    by_id = {int(t["id"]): t for t in (trabajadores or []) if t.get("id") is not None}
    wb = Workbook()
    # Quitar hoja default; Parametros primero, luego Nómina
    default = wb.active
    wb.remove(default)

    anio = int(nomina.get("anio") or datetime.now(_BOGOTA).year)
    smmlv = float(nomina.get("smmlv_usado") or params_for_year(anio).smmlv)
    _write_parametros_sheet(wb, anio=anio, smmlv=smmlv)

    ws = wb.create_sheet("Nómina")

    per = nomina.get("periodicidad") or ""
    q = nomina.get("quincena")
    titulo = (
        f"Relación de nómina — {contrato_label or 'Contrato'} — "
        f"{per} {nomina.get('mes')}/{nomina.get('anio')}"
        + (f" Q{q}" if q else "")
    )
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = _FONT_TITLE
    generado = datetime.now(_BOGOTA).strftime("%d/%m/%Y %H:%M")
    ws.cell(
        row=2,
        column=1,
        value=(
            f"Periodo: {nomina.get('fecha_inicio')} — {nomina.get('fecha_fin')} · "
            f"Estado: {nomina.get('estado')} · Generado: {generado} · "
            "Columnas calculadas = fórmulas Excel (ver hoja Parametros)"
        ),
    )

    hr = 4
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=hr, column=col, value=h)
        cell.fill = _FILL_HDR
        cell.font = _FONT_HDR
        cell.alignment = _AL_CENTER
        cell.border = _BORDER

    money_cols = set(range(4, 31)) | {32}
    items_list = list(items)

    for i, item in enumerate(items_list, hr + 1):
        tid = int(item.get("trabajador_id") or 0)
        t = by_id.get(tid) or {}
        det = item.get("detalle_json") or {}
        if not isinstance(det, dict):
            det = {}

        # Valores base (no calculados)
        base_vals: Dict[int, Any] = {
            1: _doc(item, by_id),
            2: _nombre(item, by_id),
            3: det.get("cargo") or t.get("cargo_aspira") or "",
            4: float(item.get("salario_base") or 0),
            5: float(item.get("salario_periodo") or 0),
            6: float(item.get("valor_extras") or 0),
            7: float(item.get("valor_recargos") or 0),
            8: float(item.get("valor_bonificaciones") or 0),
            9: float(item.get("descuento_novedades") or 0),
            10: _nov_detalle(item),
            11: float(item.get("subsidio_transporte") or 0),
            31: det.get("email") or t.get("email") or "",
            32: _arl_tasa(item, by_id),
        }
        for col, v in base_vals.items():
            cell = ws.cell(row=i, column=col, value=v)
            cell.border = _BORDER
            cell.alignment = (
                _AL_RIGHT if col in money_cols and col != 10 else _AL_LEFT
            )

        for col, formula in _formulas_fila(i).items():
            cell = ws.cell(row=i, column=col, value=formula)
            cell.border = _BORDER
            cell.alignment = _AL_RIGHT

    for idx in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 16 if idx > 3 else 22
    ws.column_dimensions["J"].width = 36
    ws.column_dimensions["AF"].width = 12

    # Totales con SUM de fórmulas
    if items_list:
        first = hr + 1
        last = hr + len(items_list)
        tot_row = last + 1
        ws.cell(row=tot_row, column=1, value="TOTALES").font = Font(bold=True)
        for col_idx in (13, 17, 18, 25, 30):  # devengado, ded, neto, aportes, prov
            letter = get_column_letter(col_idx)
            cell = ws.cell(
                row=tot_row,
                column=col_idx,
                value=f"=SUM({letter}{first}:{letter}{last})",
            )
            cell.font = Font(bold=True)
            cell.border = _BORDER
            cell.alignment = _AL_RIGHT

    # Ocultar columna auxiliar Tasa ARL (sigue referenciable por fórmulas)
    ws.column_dimensions["AF"].hidden = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
