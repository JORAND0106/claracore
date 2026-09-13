"""
Exportación Excel de poligonal con fórmulas vivas (cierres alineados al backend).

Ejecutar:
  cd backend && python3 -m unittest tests.test_poligonal_excel_formulas -v
"""
from __future__ import annotations

import io
import math
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

from tests.test_poligonal_cierre_lineal_bowditch import _cuadrado  # noqa: E402
from topografia_poligonal_xlsx import build_poligonal_xlsx_bytes  # noqa: E402
from topografia_utils import (  # noqa: E402
    ajustar_poligonal_armadas,
    calcular_cierre_poligonal,
    radiar_armadas,
)


def _mod360(x: float) -> float:
    return float(x) % 360.0


def _simular_excel_cierre(
    *,
    azimuts: list[float],
    distancias: list[float],
    base_az: float,
    antihorario: bool = True,
    tiene_orientacion: bool = False,
    ang_orientacion: float = 0.0,
):
    """Replica fórmulas Resumen/Cartera para cierre angular + lineal."""
    n = len(azimuts)
    ang_deriv = []
    for i, az in enumerate(azimuts):
        if i == 0:
            ang_deriv.append(_mod360(az - base_az))
        else:
            ang_deriv.append(_mod360(az - azimuts[i - 1] - 180.0))
    en_suma = []
    for i in range(n):
        if i == 0 and tiene_orientacion:
            en_suma.append(0)
        else:
            en_suma.append(1)
    suma_obs = sum(a for a, f in zip(ang_deriv, en_suma) if f) + (
        ang_orientacion if tiene_orientacion else 0.0
    )
    teorico = ((n - 2) if antihorario else (n + 2)) * 180.0
    diff = suma_obs - teorico
    n_ang = n  # con orientación: excluye 1º y suma orientación → mismo conteo
    peri = sum(distancias)
    # Proyecciones con azimut de campo (cierre lineal de campo / preliminar)
    dn = [d * math.cos(math.radians(a)) for d, a in zip(distancias, azimuts)]
    de = [d * math.sin(math.radians(a)) for d, a in zip(distancias, azimuts)]
    err_n, err_e = sum(dn), sum(de)
    return {
        "ang_deriv": ang_deriv,
        "suma_obs": suma_obs,
        "teorico": teorico,
        "diff": diff,
        "diff_seg": diff * 3600.0,
        "n_ang": n_ang,
        "peri": peri,
        "error_lineal": math.hypot(err_n, err_e),
        "delta_norte": -err_n,
        "delta_este": -err_e,
    }


def _simular_cartera_excel(
    *,
    azimuts: list[float],
    distancias: list[float],
    base_az: float,
    norte0: float,
    este0: float,
    cota0: float | None,
    antihorario: bool = True,
    tiene_orientacion: bool = False,
    ang_orientacion: float = 0.0,
    hi: float = 1.5,
    ht: float = 0.0,
    ang_v: float = 90.0,
):
    """Cascada completa (cierre angular con exclusión + Bowditch)."""
    n = len(azimuts)
    cierre = _simular_excel_cierre(
        azimuts=azimuts,
        distancias=distancias,
        base_az=base_az,
        antihorario=antihorario,
        tiene_orientacion=tiene_orientacion,
        ang_orientacion=ang_orientacion,
    )
    ang_deriv = cierre["ang_deriv"]
    diff = cierre["diff"]
    n_ang = max(cierre["n_ang"], 1)
    ang_corr = [a - diff / n_ang for a in ang_deriv]
    az_corr = []
    for i, az in enumerate(azimuts):
        if i == 0:
            az_corr.append(_mod360(az))
        else:
            az_corr.append(_mod360(az_corr[i - 1] + 180.0 + ang_corr[i]))
    dn = [d * math.cos(math.radians(a)) for d, a in zip(distancias, az_corr)]
    de = [d * math.sin(math.radians(a)) for d, a in zip(distancias, az_corr)]
    sin_v = math.sin(math.radians(ang_v))
    dz = [
        0.0 if abs(sin_v) < 1e-9 else hi + d * math.cos(math.radians(ang_v)) / sin_v - ht
        for d in distancias
    ]
    err_n, err_e, err_z = sum(dn), sum(de), sum(dz)
    peri = sum(distancias)
    cn = [0.0 if peri <= 0 else -err_n * d / peri for d in distancias]
    ce = [0.0 if peri <= 0 else -err_e * d / peri for d in distancias]
    cz = [0.0 if peri <= 0 else -err_z * d / peri for d in distancias]
    nortes, estes, cotas = [], [], []
    nn, ee, zz = norte0, este0, cota0
    for i in range(n):
        nn = nn + dn[i] + cn[i]
        ee = ee + de[i] + ce[i]
        if zz is not None:
            zz = zz + dz[i] + cz[i]
        nortes.append(nn)
        estes.append(ee)
        cotas.append(zz)
    return {
        **cierre,
        "ang_corr": ang_corr,
        "az_corr": az_corr,
        "err_n": err_n,
        "err_e": err_e,
        "nortes": nortes,
        "estes": estes,
        "cotas": cotas,
        "cn": cn,
        "ce": ce,
    }


class TestPoligonalExcelFormulas(unittest.TestCase):
    def _build_square(self, dists=None, con_orientacion=False):
        dists = dists or [100.0, 100.0, 100.0, 100.5]
        armadas, estaciones, amarres, pi = _cuadrado(dists, con_orientacion=con_orientacion)
        arms, _, flat = radiar_armadas(armadas, estaciones, amarres)
        cierre = calcular_cierre_poligonal(
            arms, pi, tipo_pol="cerrada", tol_relativa=500, precision_angular_seg=10.0
        )
        raw = build_poligonal_xlsx_bytes(
            contrato={"numero": "C-1", "objeto": "Demo", "contratista": "A", "interventoria": "B"},
            pol={
                "nombre": "Cuadrado demo",
                "tipo": "cerrada",
                "sentido": "antihorario",
                "tolerancia_relativa": 500,
                "precision_angular_seg": 10,
                "equipo_marca": "Leica",
                "operador": "Topo",
                "fecha_campo": "2026-09-13",
            },
            estaciones=flat,
            punto_inicial=pi,
            cierre=cierre,
            armadas=arms,
        )
        return raw, arms, flat, pi, cierre, armadas, estaciones, amarres

    def test_xlsx_pestanas_separadas(self):
        raw, *_ = self._build_square()
        self.assertEqual(raw[:2], b"PK")
        wb = load_workbook(io.BytesIO(raw))
        self.assertEqual(wb.sheetnames[:3], ["Resumen", "Cartera", "Esquema"])
        wr = wb["Resumen"]
        self.assertIn("CIERRE ANGULAR", str(wr["A23"].value))
        self.assertIn("CIERRE LINEAL", str(wr["A39"].value))
        self.assertTrue(str(wr["B26"].value).startswith("="))  # Σ Observada
        self.assertTrue(str(wr["B28"].value).startswith("="))  # Diff = obs − teor
        self.assertIn("SUMIF", str(wr["B26"].value))
        ws = wb["Cartera"]
        self.assertEqual(ws["D4"].value, "Az campo °")
        self.assertEqual(ws["U4"].value, "En Σ")
        self.assertTrue(str(ws["I6"].value).startswith("="))
        self.assertIn("Resumen!", str(ws["J6"].value))
        self.assertIn("Resumen!", str(ws["O6"].value))
        # 1º tramo: En Σ depende de orientación
        self.assertIn("Resumen!$B$18", str(ws["U6"].value))
        ws2 = wb["Esquema"]
        self.assertIn("Cartera!", str(ws2["C5"].value))
        self.assertTrue(len(ws2._charts) >= 1)

    def test_base_azimut_y_params_en_resumen(self):
        raw, arms, *_ = self._build_square()
        wb = load_workbook(io.BytesIO(raw))
        wr = wb["Resumen"]
        self.assertAlmostEqual(float(wr["B15"].value), float(arms[0]["base_azimut"]), places=4)
        self.assertEqual(int(wr["B17"].value), 1)  # azimut directo
        self.assertEqual(int(wr["B18"].value), 0)  # sin orientación

    def test_cascada_coincide_con_plataforma(self):
        raw, arms, flat, pi, cierre, armadas, estaciones, amarres = self._build_square()
        wb = load_workbook(io.BytesIO(raw))
        ws = wb["Cartera"]
        wr = wb["Resumen"]
        base = float(wr["B15"].value)
        azs, dists = [], []
        for r in range(6, 10):
            azs.append(float(ws[f"D{r}"].value))
            dists.append(float(ws[f"E{r}"].value))

        sim = _simular_cartera_excel(
            azimuts=azs,
            distancias=dists,
            base_az=base,
            norte0=float(pi["norte"]),
            este0=float(pi["este"]),
            cota0=float(pi["cota"]),
        )
        self.assertAlmostEqual(sim["diff"], float(cierre["error_angular"] or 0), places=4)
        self.assertAlmostEqual(sim["suma_obs"], float(cierre["suma_observada"]), places=4)
        self.assertAlmostEqual(sim["teorico"], float(cierre["suma_teorica"]), places=4)
        self.assertAlmostEqual(sim["error_lineal"], float(cierre["error_lineal"]), places=3)
        self.assertAlmostEqual(sim["delta_este"], float(cierre["delta_este"]), places=3)

        pol = {
            "sentido": "antihorario",
            "tolerancia_relativa": 500,
            "tolerancia_cota_mm_km": 12,
            "precision_angular_seg": 10,
            "tipo": "cerrada",
        }
        adj = ajustar_poligonal_armadas(pol, armadas, estaciones, amarres, pi)
        updates = {u["id"]: u for u in adj["updates"] if u.get("norte_ajustado") is not None}
        legs = [e for e in flat if float(e.get("distancia") or 0) > 0]
        for i, e in enumerate(legs):
            u = updates[e["id"]]
            self.assertAlmostEqual(sim["nortes"][i], float(u["norte_ajustado"]), places=3)
            self.assertAlmostEqual(sim["estes"][i], float(u["este_ajustado"]), places=3)
            self.assertAlmostEqual(sim["az_corr"][i], float(u["azimut"]), places=4)

    def test_orientacion_excluye_amarre_como_plataforma(self):
        """Excluir amarre + sumar orientación evita inflar Σ (bug 5579 vs 5400)."""
        d1 = {"norte": 5000.0, "este": 6000.0, "cota": 100.0}

        def gms(d, m, s):
            return d + m / 60.0 + s / 3600.0

        az_arranque = gms(164, 19, 53)
        az_cierre = gms(164, 18, 58)
        az_backsight = (az_arranque + 90) % 360
        gps2 = {
            "norte": d1["norte"] + 80.0 * math.cos(math.radians(az_backsight)),
            "este": d1["este"] + 80.0 * math.sin(math.radians(az_backsight)),
            "cota": 100.0,
        }
        azs_lados = [
            az_arranque,
            (az_arranque + 90) % 360,
            (az_arranque + 180) % 360,
            (az_arranque + 270) % 360,
        ]
        armadas = []
        estaciones = []
        for i, (est, fwd, vis) in enumerate(
            zip(["D1", "D2", "D3", "D4"], ["D2", "D3", "D4", "D1"], ["GPS2", "D1", "D2", "D3"])
        ):
            armadas.append({
                "id": f"a{i+1}", "orden": i + 1,
                "estacion_nombre": est, "visado_nombre": vis, "altura_instrumento": 1.5,
            })
            estaciones.append({
                "id": f"p{i+1}", "armada_id": f"a{i+1}", "orden": 1,
                "nombre_punto": fwd, "tipo_punto": "estacion",
                "angulo_medido": azs_lados[i], "angulo_vertical": 90.0,
                "distancia": 100.0, "altura_objetivo": 0,
            })
        armadas.append({
            "id": "a5", "orden": 5,
            "estacion_nombre": "D1", "visado_nombre": "GPS2", "altura_instrumento": 1.5,
        })
        estaciones.append({
            "id": "p5", "armada_id": "a5", "orden": 1,
            "nombre_punto": "CIERRE", "tipo_punto": "estacion",
            "angulo_medido": az_cierre, "angulo_vertical": 90.0,
            "distancia": 0.0, "altura_objetivo": 0,
        })
        arms, _, flat = radiar_armadas(armadas, estaciones, {"D1": d1, "GPS2": gps2})
        cierre = calcular_cierre_poligonal(
            arms, {"nombre": "D1", **d1}, sentido="horario", tipo_pol="cerrada",
            precision_angular_seg=10.0,
        )
        self.assertTrue(cierre["tiene_orientacion"])
        amarre = float(cierre["angulo_amarre_inicial_excluido"])
        orient_det = next(d for d in cierre["angulos_cierre_detalle"] if d.get("orientacion"))

        raw = build_poligonal_xlsx_bytes(
            contrato={"numero": "C-1"},
            pol={"nombre": "Orient", "tipo": "cerrada", "tolerancia_relativa": 500},
            estaciones=flat,
            punto_inicial={"nombre": "D1", **d1},
            cierre=cierre,
            armadas=arms,
        )
        wb = load_workbook(io.BytesIO(raw))
        wr = wb["Resumen"]
        ws = wb["Cartera"]
        self.assertEqual(int(wr["B18"].value), 1)
        self.assertEqual(int(wr["B12"].value), 0)  # horario → anti=0
        base = float(wr["B15"].value)
        azs = [float(ws[f"D{r}"].value) for r in range(6, 10)]
        dists = [float(ws[f"E{r}"].value) for r in range(6, 10)]
        ang_or = float(wr["B19"].value)
        self.assertAlmostEqual(ang_or, float(orient_det["angulo_cierre"]), places=4)

        ok = _simular_excel_cierre(
            azimuts=azs, distancias=dists, base_az=base,
            antihorario=False, tiene_orientacion=True, ang_orientacion=ang_or,
        )
        malo = _simular_excel_cierre(
            azimuts=azs, distancias=dists, base_az=base,
            antihorario=False, tiene_orientacion=False,
        )
        self.assertAlmostEqual(ok["suma_obs"], float(cierre["suma_observada"]), places=4)
        self.assertAlmostEqual(ok["teorico"], float(cierre["suma_teorica"]), places=4)
        self.assertAlmostEqual(ok["diff_seg"], float(cierre["error_angular_seg"]), delta=1.0)
        # Incluir el amarre sin sustituir por orientación infla Σ (~+ángulo_amarre − orient)
        self.assertGreater(abs(malo["suma_obs"] - ok["suma_obs"]), 0.01)
        self.assertAlmostEqual(malo["suma_obs"] - ok["suma_obs"], amarre - ang_or, places=3)
        # Diff del bug (~cientos de miles de ″) no debe aparecer cuando se excluye bien
        self.assertLess(abs(ok["diff_seg"]), 120.0)

    def test_diff_obs_menos_teorica_no_invertido(self):
        raw, *_rest = self._build_square()
        wb = load_workbook(io.BytesIO(raw))
        # B28 = B26 - B27  (obs − teor), no teor − obs
        self.assertEqual(str(wb["Resumen"]["B28"].value), "=$B$26-$B$27")

    def test_editar_distancia_recalcula_en_cascada(self):
        _, arms, flat, pi, *_rest = self._build_square([100, 100, 100, 100])
        base = float(arms[0]["base_azimut"])
        azs = [float(e["azimut"]) for e in flat if float(e.get("distancia") or 0) > 0]
        ok = _simular_cartera_excel(
            azimuts=azs,
            distancias=[100.0, 100.0, 100.0, 100.0],
            base_az=base,
            norte0=pi["norte"],
            este0=pi["este"],
            cota0=pi["cota"],
        )
        self.assertAlmostEqual(ok["error_lineal"], 0.0, places=4)
        edited = _simular_cartera_excel(
            azimuts=azs,
            distancias=[100.0, 100.0, 100.0, 100.8],
            base_az=base,
            norte0=pi["norte"],
            este0=pi["este"],
            cota0=pi["cota"],
        )
        self.assertGreater(edited["error_lineal"], 0.5)
        self.assertAlmostEqual(edited["nortes"][-1], pi["norte"], places=3)
        self.assertAlmostEqual(edited["estes"][-1], pi["este"], places=3)


if __name__ == "__main__":
    unittest.main()
