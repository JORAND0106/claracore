"""Resumen de Cantidades / Descuentos: presentación a 2 decimales (FE/PDF/Excel)."""
from __future__ import annotations

import io
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


class TestResumen2DecSource(unittest.TestCase):
    def test_pdf_fmt_2_en_resumen_y_descuentos(self):
        src = (ROOT / "topografia_planilla_tuberia_routes.py").read_text(encoding="utf-8")
        self.assertIn("fmt(n.get('long'), 2)", src)
        self.assertIn("fmt(n.get('ancho'), 2)", src)
        self.assertIn("fmt(n.get('espesor'), 2)", src)
        self.assertIn("fmt(n.get('descuentos'), 2)", src)
        self.assertIn("fmt(d.get('cantidad'), 2)", src)
        # Cartera sigue con default (3) — no forzar 2 en abscisas de cartera
        self.assertIn("fmt(f.get('terreno_natural'))", src)

    def test_excel_num_fmt_y_round(self):
        src = (ROOT / "topografia_planilla_tuberia_excel.py").read_text(encoding="utf-8")
        self.assertIn('NUM_FMT_RESUMEN_CANTIDADES = "0.00"', src)
        self.assertIn("_num_resumen_2", src)
        self.assertIn("number_format = NUM_FMT_RESUMEN_CANTIDADES", src)
        self.assertIn("=ROUND(PRODUCT(K46:M46),2)", src)
        # SICOE dims 3 no deben cambiarse aquí
        self.assertNotIn("_dim_sicoe_3", src)

    def test_sicoe_dims_siguen_en_3(self):
        src = (ROOT / "topografia_planilla_tuberia.py").read_text(encoding="utf-8")
        self.assertIn("def _dim_sicoe_3", src)
        self.assertIn("Longitud / Ancho / Espesor → 3 decimales", src)


class TestExcelNumFmtRuntime(unittest.TestCase):
    def test_celdas_resumen_con_formato_2dec(self):
        try:
            from openpyxl import load_workbook
            from topografia_planilla_tuberia_excel import (
                NUM_FMT_RESUMEN_CANTIDADES,
                build_planilla_tuberia_xlsx,
            )
        except ImportError:
            self.skipTest("openpyxl no disponible")

        raw = build_planilla_tuberia_xlsx(
            planilla={
                "tipo": "ALCANTARILLA",
                "diametro_m": 0.9,
                "espesor_m": 0.05,
                "ancho_excavacion_m": 1.5,
                "relacion_atraque": "1:3",
                "meta_cabecera": {"cama_triturado_m": 0.1},
            },
            calculo={
                "cartera": {"filas": []},
                "netos": [
                    {
                        "codigo": "EXC_ROC",
                        "nombre": "Excavación Roca",
                        "unidad": "m³",
                        "long": 10.1234,
                        "ancho": 1.5678,
                        "espesor": 0.12345,
                        "neto": 2.34567,
                    }
                ],
                "descuentos": [
                    {
                        "codigo": "DESC_OTROS_1",
                        "nombre": "Otros: Pozo",
                        "long": 3.14159,
                        "ancho": 2.71828,
                        "espesor": 0.99999,
                        "cantidad": 8.5397,
                    }
                ],
            },
            vacia=False,
        )
        wb = load_workbook(io.BytesIO(raw))
        ws = wb["planilla"]
        # Compactado: única línea de netos y descuentos en fila 45
        self.assertEqual(ws["D45"].number_format, NUM_FMT_RESUMEN_CANTIDADES)
        self.assertEqual(ws["H45"].number_format, NUM_FMT_RESUMEN_CANTIDADES)
        self.assertEqual(ws["K45"].number_format, NUM_FMT_RESUMEN_CANTIDADES)
        self.assertEqual(ws["N45"].number_format, NUM_FMT_RESUMEN_CANTIDADES)
        # EXC_ROC: literales de dims/cantidad redondeados a 2
        self.assertEqual(ws["D45"].value, 10.12)
        self.assertEqual(ws["E45"].value, 1.57)
        self.assertEqual(ws["F45"].value, 0.12)
        self.assertEqual(ws["H45"].value, 2.35)
        # DESC_OTROS_*: dims redondeadas; cantidad vía ROUND(PRODUCT…,2)
        self.assertEqual(ws["K45"].value, 3.14)
        self.assertEqual(ws["L45"].value, 2.72)
        self.assertEqual(ws["M45"].value, 1.0)
        self.assertEqual(ws["N45"].value, "=ROUND(PRODUCT(K45:M45),2)")


if __name__ == "__main__":
    unittest.main()
