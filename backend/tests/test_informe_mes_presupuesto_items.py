"""Informe mensual: solo ítems del presupuesto, sin cruzar capítulos."""
from __future__ import annotations

import sys
from io import BytesIO
from types import ModuleType, SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from acta_mes_conciliacion import build_resumen_contrato_4cols
from corte_sub_conciliacion import valor_por_cantidad_vu
from excel_formula_eval import FormulaBook
from informe_mes_items import (
    clave_item_exacta,
    construir_cuadro_mensual,
    elegir_version_obra_vigente,
    fuente_de_version,
    indice_listado_exacto,
    resumir_registros_fuera,
)

CAP_PAV = "2. ESTRUCTURA DE PAVIMENTO Y ADOQUINES"
CAP_SENAL = "2. SEÑALIZACIÓN Y CANALIZACIÓN"
CAP_HORIZ = "4. SEÑALIZACIÓN HORIZONTAL Y VERTICAL"
CAP_PMT = "4. OPERACIÓN PMT"


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]
    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake
    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)
    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    return inf


def test_version_vigente_es_la_actualizada_mas_reciente_o_v0():
    assert elegir_version_obra_vigente([]) is None
    v0 = {"id": "a", "numero_version": 0, "etiqueta": "V0", "es_vigente": False}
    assert elegir_version_obra_vigente([v0])["id"] == "a"
    v1 = {"id": "b", "numero_version": 1, "etiqueta": "Act 1", "es_vigente": False}
    v2 = {"id": "c", "numero_version": 2, "etiqueta": "Act 2", "es_vigente": True}
    elegida = elegir_version_obra_vigente([v2, v0, v1])
    assert elegida["id"] == "c"
    assert fuente_de_version(elegida) == "presupuesto_vivo"
    assert fuente_de_version(v1) == "presupuesto_version_items"
    assert fuente_de_version(None) == "presupuesto_vivo"


def test_no_cruza_capitulos_ni_sufijos_y_lista_todo_el_presupuesto():
    vu_subbase = 297349.0
    cant_subbase = 1363.79
    costo_subbase = valor_por_cantidad_vu(cant_subbase, vu_subbase)
    cant_41 = 11623.24
    vu_41 = 1500.0
    costo_41 = valor_por_cantidad_vu(cant_41, vu_41)

    filas = [
        _fila(CAP_PAV, "2.1", "Subbase granular", "M3", vu_subbase, cant_subbase, costo_subbase),
        _fila(CAP_PAV, "2.1._", "Código con sufijo", "UND", 10, 5, 50),
        _fila(CAP_PAV, "2.2", "Base granular", "M3", 1000, 10, 10000),
        _fila(CAP_PAV, "2.3", "Adoquín", "M2", 2000, 4, 8000),
        _fila(CAP_HORIZ, "4.1", "Línea de demarcación", "ML", vu_41, cant_41, costo_41),
        _fila(CAP_HORIZ, "4.2", "Tachas", "UND", 800, 3, 2400),
        _fila(CAP_HORIZ, "4.3", "Señal vertical", "UND", 900, 2, 1800),
        _fila(CAP_HORIZ, "4.4", "Poste", "UND", 700, 1, 700),
        # Dos tramos del mismo ítem: el costo almacenado no coincide con ROUND(Σcant × VU).
        _fila(CAP_PAV, "2.9", "Riego", "M2", 100, 1.006, 101),
        _fila(CAP_PAV, "2.9", "Riego", "M2", 100, 1.006, 101),
    ]
    # El mismo número en otro capítulo del presupuesto es otro ítem.
    filas.append(_fila(CAP_HORIZ, "2.1", "Flecha", "UND", 50, 7, 350))

    listado = indice_listado_exacto(
        [
            {
                "id": 1,
                "capitulo": CAP_PAV,
                "item_numero": "2.1",
                "descripcion": "SUBBASE GRANULAR",
                "unidad": "M3",
                "precio_unitario": vu_subbase,
                "especificacion_tecnica": "INVIAS 320",
            },
            {
                "id": 2,
                "capitulo": CAP_SENAL,
                "item_numero": "2.1._",
                "descripcion": "Señal SR-30",
                "unidad": "UND",
                "precio_unitario": 50000,
                "especificacion_tecnica": "PMT-30",
            },
            {
                "id": 3,
                "capitulo": CAP_HORIZ,
                "item_numero": "4.1",
                "descripcion": "LÍNEA DE DEMARCACIÓN",
                "unidad": "ML",
                "precio_unitario": vu_41,
                "especificacion_tecnica": "INVIAS 700",
            },
        ]
    )
    registros = [
        {
            "capitulo": CAP_SENAL,
            "item_numero": "2.1._",
            "item_descripcion": "Señal SR-30",
            "cantidad_total": 12,
            "vlr_unitario": 50000,
            "costo_directo": 600000,
        },
        {
            "capitulo": CAP_PMT,
            "item_numero": "4.1._",
            "item_descripcion": "Instalación inicial PMT",
            "cantidad_total": 3,
            "vlr_unitario": 1000,
            "costo_directo": 11623.24,  # cantidad ajena; no debe pegarse al 4.1 del presupuesto
        },
        {
            "capitulo": CAP_PAV,
            "item_numero": "2.2",
            "item_descripcion": "copia vieja",
            "cantidad_total": 1.5,
            "vlr_unitario": 1000,
            "costo_directo": 1500,
        },
    ]
    from informe_mes_items import sumar_cantidad_por_clave

    cuadro = construir_cuadro_mensual(
        filas,
        cant_presente=sumar_cantidad_por_clave(registros),
        cant_anterior={},
        listado_por_clave=listado,
        claves_con_registros_acta=sumar_cantidad_por_clave(registros).keys(),
    )
    items = cuadro["items"]
    codigos = [(it["capitulo"], it["item_numero"]) for it in items]

    assert (CAP_PAV, "2.1") in codigos
    assert (CAP_PAV, "2.1._") in codigos
    assert (CAP_PAV, "2.2") in codigos
    assert (CAP_PAV, "2.3") in codigos
    assert (CAP_HORIZ, "4.1") in codigos
    assert (CAP_HORIZ, "4.2") in codigos
    assert (CAP_HORIZ, "4.3") in codigos
    assert (CAP_HORIZ, "4.4") in codigos
    assert (CAP_HORIZ, "2.1") in codigos
    assert (CAP_SENAL, "2.1._") not in codigos
    assert (CAP_PMT, "4.1._") not in codigos
    assert all(it["capitulo"] not in (CAP_SENAL, CAP_PMT) for it in items)

    subbase = next(it for it in items if it["capitulo"] == CAP_PAV and it["item_numero"] == "2.1")
    assert subbase["item_descripcion"] == "SUBBASE GRANULAR"
    assert subbase["unidad"] == "M3"
    assert subbase["norma_tecnica"] == "INVIAS 320"
    assert subbase["vlr_unitario"] == vu_subbase
    assert subbase["cant_actualizadas"] == 1363.79
    assert subbase["cant_presente"] == 0
    assert subbase["cant_acumulado"] == 0
    assert subbase["cant_saldo"] == 1363.79
    assert subbase["valor_actualizadas"] == costo_subbase
    assert subbase["valor_actualizadas_literal"] is False
    sufijo = next(it for it in items if it["capitulo"] == CAP_PAV and it["item_numero"] == "2.1._")
    assert sufijo["cant_actualizadas"] == 5
    assert sufijo["vlr_unitario"] == 10
    assert sufijo["item_descripcion"] == "Código con sufijo"

    horiz = next(it for it in items if it["capitulo"] == CAP_HORIZ and it["item_numero"] == "4.1")
    assert horiz["cant_actualizadas"] == 11623.24
    assert horiz["item_descripcion"] == "LÍNEA DE DEMARCACIÓN"
    assert horiz["cant_presente"] == 0
    assert horiz["vlr_unitario"] == vu_41

    otro_21 = next(it for it in items if it["capitulo"] == CAP_HORIZ and it["item_numero"] == "2.1")
    assert otro_21["cant_actualizadas"] == 7
    assert otro_21["vlr_unitario"] == 50
    assert otro_21["item_descripcion"] == "Flecha"

    base = next(it for it in items if it["item_numero"] == "2.2")
    assert base["cant_presente"] == 1.5
    assert base["item_descripcion"] == "Base granular"

    riego = next(it for it in items if it["item_numero"] == "2.9")
    assert riego["cant_actualizadas"] == 2.01
    assert riego["valor_actualizadas"] == 202
    assert riego["valor_actualizadas_literal"] is True

    assert sum(it["valor_actualizadas"] for it in items) == cuadro["costo_directo_presupuesto"]

    fuera = resumir_registros_fuera(registros, cuadro["claves"])
    assert fuera["n_registros"] == 2
    por_codigo = {clave_item_exacta(it["capitulo"], it["item_numero"]): it for it in fuera["items"]}
    assert por_codigo[(CAP_SENAL, "2.1._")]["valor"] == 600000
    assert por_codigo[(CAP_SENAL, "2.1._")]["item_descripcion"] == "Señal SR-30"
    assert por_codigo[(CAP_PMT, "4.1._")]["item_descripcion"] == "Instalación inicial PMT"
    assert por_codigo[(CAP_PMT, "4.1._")]["valor"] == 11623
    assert fuera["valor_total"] == 600000 + 11623


def _fila(cap, item, desc, und, vu, cant, costo):
    return {
        "capitulo": cap,
        "item": item,
        "descripcion": desc,
        "und": und,
        "vlr_unitario": vu,
        "cant_total": cant,
        "costo_directo": costo,
    }


def test_vista_previa_pdf_y_excel_muestran_los_mismos_items():
    inf = _import_informes_with_stubs()
    vu = 297349.0
    costo = valor_por_cantidad_vu(1363.79, vu)
    filas = [
        _fila(CAP_PAV, "2.1", "Subbase granular", "M3", vu, 1363.79, costo),
        _fila(CAP_PAV, "2.2", "Base granular", "M3", 1000, 10, 10000),
        _fila(CAP_PAV, "2.3", "Adoquín", "M2", 2000, 4, 8000),
        _fila(CAP_HORIZ, "4.1", "Línea", "ML", 1500, 11623.24, valor_por_cantidad_vu(11623.24, 1500)),
        _fila(CAP_HORIZ, "4.2", "Tachas", "UND", 800, 3, 2400),
        _fila(CAP_HORIZ, "4.3", "Señal vertical", "UND", 900, 2, 1800),
        _fila(CAP_HORIZ, "4.4", "Poste", "UND", 700, 1, 700),
        _fila(CAP_PAV, "2.9", "Riego", "M2", 100, 1.006, 101),
        _fila(CAP_PAV, "2.9", "Riego", "M2", 100, 1.006, 101),
    ]
    listado = indice_listado_exacto(
        [
            {
                "id": 9,
                "capitulo": CAP_PAV,
                "item_numero": "2.1",
                "descripcion": "SUBBASE GRANULAR",
                "unidad": "M3",
                "precio_unitario": vu,
                "especificacion_tecnica": "INVIAS 320",
            }
        ]
    )
    registros_acta = [
        {
            "capitulo": CAP_PAV,
            "item_numero": "2.2",
            "item_descripcion": "Base granular",
            "cantidad_total": 1.25,
            "vlr_unitario": 1000,
            "costo_directo": 1250,
            "numero_registro": 1,
            "unidad": "M3",
            "longitud": 1.25,
            "ancho": "",
            "espesor": "",
            "cantidad": "",
            "observacion": "",
            "abs_inicio": "",
            "abs_final": "",
            "infraestructura": "",
        }
    ]
    from informe_mes_items import sumar_cantidad_por_clave

    cuadro = construir_cuadro_mensual(
        filas,
        cant_presente=sumar_cantidad_por_clave(registros_acta),
        listado_por_clave=listado,
        claves_con_registros_acta=[(CAP_PAV, "2.2")],
    )
    items = cuadro["items"]
    assert [it["item_numero"] for it in items] == ["2.1", "2.2", "2.3", "2.9", "4.1", "4.2", "4.3", "4.4"]
    r4 = build_resumen_contrato_4cols(items=items)
    contrato = {"numero": "ICCU-CTO-1614-2025", "contratista": "Ctor", "nit": "900", "interventoria": "Int"}
    acta = {"numero_rpo": "1", "consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"}
    html = inf._html_cc_mes_001_v1(
        contrato,
        acta,
        items,
        sum(it["valor_presente"] for it in items),
        "Ana",
        "Residente",
        resumen_4cols=r4,
        c3_value="1",
        c4_value="1",
    )
    assert "SUBBASE GRANULAR" in html
    assert CAP_PAV in html
    assert CAP_HORIZ in html
    assert CAP_SENAL not in html
    assert CAP_PMT not in html
    assert "2.1._" not in html
    assert "4.1._" not in html
    for it in items:
        assert it["item_numero"] in html
        assert inf._fm_informe(it["valor_actualizadas"]) in html
        assert inf._fn_cant_informe(it["cant_actualizadas"]) in html

    ctx = {
        "contrato": contrato,
        "acta": acta,
        "items": items,
        "total_costo": sum(it["valor_presente"] for it in items),
        "resumen_4cols": r4,
        "otros_conceptos": [],
        "usuario_nombre": "Ana",
        "usuario_cargo": "Residente",
        "costo_directo_presupuesto": cuadro["costo_directo_presupuesto"],
    }
    user = {"nombre": "Ana", "apellidos": "Pérez", "cargo_nombre": "Residente"}

    def _fetch(_sb, _cid, inum, **_kw):
        return [r for r in registros_acta if str(r.get("item_numero")) == str(inum) and str(r.get("capitulo")) == CAP_PAV]

    with (
        patch.object(inf, "_acta_pertenece_contrato", return_value=True),
        patch.object(inf, "_contexto_acta_mes_conciliacion", return_value=ctx),
        patch.object(inf, "_row", return_value=contrato),
        patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
        patch.object(inf, "_sub_corte_dummy_memoria", return_value=({}, {})),
        patch.object(inf, "fetch_registros_memoria_cc_mes_alineado_acta", side_effect=_fetch),
    ):
        raw_001 = inf._cc_mes_001_excel_bytes(3, 620, user, nivel_aprobacion=4)
        raw_int = inf._cc_mes_integral_excel_bytes(3, 620, user, nivel_aprobacion=4)

    wb = load_workbook(BytesIO(raw_001))
    ws = wb["CC-MES-001"]
    book = FormulaBook(wb)
    filas_excel = _filas_item(ws)
    assert len(filas_excel) == len(items)
    for row, it in zip(filas_excel, items):
        assert ws.cell(row, 1).value == it["item_numero"]
        assert ws.cell(row, 2).value == it["item_descripcion"]
        assert float(ws.cell(row, 5).value) == float(it["cant_actualizadas"])
        if it.get("valor_actualizadas_literal"):
            assert ws.cell(row, 6).value == it["valor_actualizadas"]
        else:
            assert "ROUND(ROUND(" in str(ws.cell(row, 6).value)
        assert book.eval_cell("CC-MES-001", row, 6) == it["valor_actualizadas"]
        assert book.eval_cell("CC-MES-001", row, 7) == it["cant_presente"]
        assert book.eval_cell("CC-MES-001", row, 8) == it["valor_presente"]
    texto = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert CAP_SENAL not in texto
    assert CAP_PMT not in texto

    cd_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Costo Directo")
    assert book.eval_cell("CC-MES-001", cd_row, 6) == cuadro["costo_directo_presupuesto"]

    wb_i = load_workbook(BytesIO(raw_int))
    assert wb_i.sheetnames[0] == "CC-MES-001"
    # Solo el ítem con registros del acta abre pestaña de memoria.
    assert len(wb_i.sheetnames) == 2
    assert "2.1._" not in " ".join(wb_i.sheetnames)
    assert CAP_PMT not in " ".join(wb_i.sheetnames)
    book_i = FormulaBook(wb_i)
    ws_i = wb_i["CC-MES-001"]
    for row, it in zip(_filas_item(ws_i), items):
        assert ws_i.cell(row, 1).value == it["item_numero"]
        assert (ws_i.cell(row, 4).value or "") == (it.get("norma_tecnica") or "")
        assert book_i.eval_cell("CC-MES-001", row, 7) == it["valor_actualizadas"]


def _filas_item(ws):
    limite = next(
        r for r in range(9, ws.max_row + 1) if ws.cell(r, 1).value == "RESUMEN DE CONCILIACIÓN"
    )
    rows = []
    for r in range(9, limite):
        val = ws.cell(r, 1).value
        if isinstance(val, str) and val.startswith("Subtotal"):
            continue
        if val in (None, ""):
            continue
        rows.append(r)
    return rows
