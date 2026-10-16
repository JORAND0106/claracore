"""Informe mensual: opción de agrupar ítems NP- en el capítulo NO PREVISTOS."""
from __future__ import annotations

import inspect
import sys
from io import BytesIO
from pathlib import Path
from types import ModuleType, SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from acta_mes_conciliacion import build_resumen_contrato_4cols
from corte_sub_conciliacion import valor_por_cantidad_vu
from excel_formula_eval import FormulaBook
from informe_mes_items import (
    CAPITULO_NO_PREVISTOS,
    agrupar_items_np_en_capitulo,
    clave_orden_np,
    es_item_np,
    informe_tiene_items_np,
)

CAP1 = "1. PRELIMINARES"
CAP2 = "2. PAVIMENTO"
CAP3 = "3. OTROS"


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]
    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake
    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)
    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    return inf


def _item(cap, num, desc, und, vu, cant_act, cant_pres, norma=""):
    cant_acum = cant_pres
    cant_saldo = round(cant_act - cant_acum, 2)
    return {
        "capitulo": cap,
        "item_numero": num,
        "item_descripcion": desc,
        "unidad": und,
        "norma_tecnica": norma,
        "especificacion_tecnica": norma,
        "vlr_unitario": vu,
        "vlr_unitario_sub": vu,
        "cantidad": cant_pres,
        "costo_directo": valor_por_cantidad_vu(cant_pres, vu),
        "cant_actualizadas": cant_act,
        "valor_actualizadas": valor_por_cantidad_vu(cant_act, vu),
        "valor_actualizadas_literal": False,
        "cant_presente": cant_pres,
        "valor_presente": valor_por_cantidad_vu(cant_pres, vu),
        "cant_acum_anterior": 0,
        "cant_acumulado": cant_acum,
        "valor_acumulado": valor_por_cantidad_vu(cant_acum, vu),
        "cant_saldo": cant_saldo,
        "valor_saldo": valor_por_cantidad_vu(cant_saldo, vu),
        "sin_precio": vu <= 0,
    }


def _cuadro():
    """NP-05 y NP-08 en el 1, NP-10/NP-01/NP-09/NP-02 en el 2, NP-07 en el 3."""
    return [
        _item(CAP1, "1.1", "Localización", "M2", 1000, 10, 1, "NTC 1"),
        _item(CAP1, "NP-08", "No previsto ocho", "UND", 80, 3, 1, "NTC NP8"),
        _item(CAP1, "NP-05", "No previsto cinco", "ML", 50, 4, 2, "NTC NP5"),
        _item(CAP1, "1.2", "Cerramiento", "ML", 200, 5, 0, "NTC 2"),
        _item(CAP2, "2.1", "Subbase", "M3", 300, 8, 1, "INVIAS 320"),
        _item(CAP2, "NP-10", "No previsto diez", "UND", 10, 9, 1, "NTC NP10"),
        _item(CAP2, "NP-01", "No previsto uno", "UND", 11, 2, 0, "NTC NP1"),
        _item(CAP2, "NP-09", "No previsto nueve", "M2", 90, 1.5, 0.5, "NTC NP9"),
        _item(CAP2, "NP-02", "No previsto dos", "UND", 12, 6, 1, "NTC NP2"),
        _item(CAP3, "3.1", "Señal", "UND", 40, 7, 1, "NTC 3"),
        _item(CAP3, "NP-07", "No previsto siete", "GL", 70, 2, 1, "NTC NP7"),
    ]


def _filas_item(ws):
    limite = next(
        r for r in range(9, ws.max_row + 1) if ws.cell(r, 1).value == "RESUMEN DE CONCILIACIÓN"
    )
    rows = []
    for r in range(9, limite):
        val = ws.cell(r, 1).value
        if isinstance(val, str) and val.startswith("Subtotal"):
            continue
        if val in (None, ""):
            continue
        rows.append(r)
    return rows


def _subtotales(ws):
    limite = next(
        r for r in range(9, ws.max_row + 1) if ws.cell(r, 1).value == "RESUMEN DE CONCILIACIÓN"
    )
    out = []
    for r in range(9, limite):
        val = ws.cell(r, 1).value
        if isinstance(val, str) and val.startswith("Subtotal"):
            out.append((r, val))
    return out


def _linea(ws, nombre):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == nombre)


def test_orden_numerico_y_capitulo_aparte():
    assert es_item_np("np-09")
    assert es_item_np(" NP-10 ")
    assert not es_item_np("1.1")
    assert not es_item_np("NPA-01")
    assert clave_orden_np("NP-09") < clave_orden_np("NP-10")
    assert clave_orden_np("NP-2") < clave_orden_np("NP-09")
    assert es_item_np("NP-06.")
    assert clave_orden_np("NP-06.") < clave_orden_np("NP-09")
    assert clave_orden_np("NP-06.") < clave_orden_np("NP-10")

    origen = _cuadro()
    copia_np = next(it for it in origen if it["item_numero"] == "NP-05")
    agrupados = agrupar_items_np_en_capitulo(origen)
    assert [it["item_numero"] for it in agrupados] == [
        "1.1",
        "1.2",
        "2.1",
        "3.1",
        "NP-01",
        "NP-02",
        "NP-05",
        "NP-07",
        "NP-08",
        "NP-09",
        "NP-10",
    ]
    assert all(it["capitulo"] == CAPITULO_NO_PREVISTOS for it in agrupados if es_item_np(it["item_numero"]))
    np5 = next(it for it in agrupados if it["item_numero"] == "NP-05")
    assert np5 is not copia_np
    assert copia_np["capitulo"] == CAP1
    assert np5["capitulo_origen"] == CAP1
    assert np5["item_descripcion"] == "No previsto cinco"
    assert np5["unidad"] == "ML"
    assert np5["norma_tecnica"] == "NTC NP5"
    assert np5["vlr_unitario"] == 50
    assert np5["cant_actualizadas"] == 4
    assert np5["valor_presente"] == copia_np["valor_presente"]
    assert next(it["capitulo_origen"] for it in agrupados if it["item_numero"] == "NP-10") == CAP2
    assert informe_tiene_items_np(origen)
    assert not informe_tiene_items_np([_item(CAP1, "1.1", "A", "UND", 1, 1, 0)])
    sin_np = [_item(CAP1, "1.1", "A", "UND", 1, 1, 0)]
    assert agrupar_items_np_en_capitulo(sin_np)[0] is sin_np[0]


def test_pdf_html_subtotales_y_resumen_identico():
    inf = _import_informes_with_stubs()
    items = _cuadro()
    agrupados = agrupar_items_np_en_capitulo(items)
    r4 = build_resumen_contrato_4cols(items=items)
    contrato = {"numero": "C-NP", "contratista": "Ctor", "nit": "900", "interventoria": "Int"}
    acta = {"numero_rpo": "1", "consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"}
    comun = dict(
        total_costo=sum(it["valor_presente"] for it in items),
        usuario_nombre="Ana",
        usuario_cargo="Residente",
        resumen_4cols=r4,
        c3_value="1",
        c4_value="1",
    )
    html_no = inf._html_cc_mes_001_v1(contrato, acta, items, **comun)
    html_si = inf._html_cc_mes_001_v1(contrato, acta, agrupados, **comun)
    marca = "RESUMEN DE CONCILIACIÓN"
    assert html_no.split(marca, 1)[1] == html_si.split(marca, 1)[1]
    assert f"Subtotal {CAPITULO_NO_PREVISTOS}" in html_si
    assert f"Subtotal {CAPITULO_NO_PREVISTOS}" not in html_no
    plan_no = inf._cc_sub_001_plan_filas_capitulo(items)
    plan_si = inf._cc_sub_001_plan_filas_capitulo(agrupados)
    subs_no = {cap: sums for kind, cap, sums in plan_no if kind == "subcap"}
    subs_si = {cap: sums for kind, cap, sums in plan_si if kind == "subcap"}
    assert list(subs_si)[-1] == CAPITULO_NO_PREVISTOS
    for cap in (CAP1, CAP2, CAP3):
        for clave in ("valor_actualizadas", "valor_presente", "valor_acumulado", "valor_saldo"):
            assert subs_si[cap][clave] < subs_no[cap][clave]
    np_items = [it for it in items if es_item_np(it["item_numero"])]
    for clave in ("valor_actualizadas", "valor_presente", "valor_acumulado", "valor_saldo"):
        assert subs_si[CAPITULO_NO_PREVISTOS][clave] == sum(it[clave] for it in np_items)
    suma_no = sum(s["valor_presente"] for s in subs_no.values())
    suma_si = sum(s["valor_presente"] for s in subs_si.values())
    assert suma_no == suma_si == sum(it["valor_presente"] for it in items)
    for it in np_items:
        assert it["item_descripcion"] in html_si
        assert inf._fm_informe(it["valor_actualizadas"]) in html_si


def _excel_par(inf, *, integral=False, norma=False):
    items = _cuadro()
    r4 = build_resumen_contrato_4cols(items=items)
    contrato = {"numero": "C-NP", "contratista": "Ctor", "nit": "900", "interventoria": "Int"}
    acta = {"numero_rpo": "1", "consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"}
    ctx = {
        "contrato": contrato,
        "acta": acta,
        "items": items,
        "total_costo": sum(it["valor_presente"] for it in items),
        "resumen_4cols": r4,
        "otros_conceptos": [],
        "usuario_nombre": "Ana",
        "usuario_cargo": "Residente",
    }
    user = {"nombre": "Ana", "apellidos": "Pérez", "cargo_nombre": "Residente"}
    registro_np = {
        "capitulo": CAP1,
        "item_numero": "NP-05",
        "item_descripcion": "No previsto cinco",
        "cantidad_total": 2,
        "vlr_unitario": 50,
        "costo_directo": 100,
        "numero_registro": 5,
        "unidad": "ML",
        "longitud": 2,
        "ancho": "",
        "espesor": "",
        "cantidad": "",
        "observacion": "",
        "abs_inicio": "",
        "abs_final": "",
        "infraestructura": "",
    }

    def _fetch(_sb, _cid, inum, **_kw):
        if str(inum) == "NP-05":
            return [registro_np]
        return []

    def _generar(agrupar):
        with (
            patch.object(inf, "_acta_pertenece_contrato", return_value=True),
            patch.object(inf, "_contexto_acta_mes_conciliacion", return_value=ctx),
            patch.object(inf, "_row", return_value=contrato),
            patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
            patch.object(inf, "_sub_corte_dummy_memoria", return_value=({}, {})),
            patch.object(inf, "fetch_registros_memoria_cc_mes_alineado_acta", side_effect=_fetch),
        ):
            if integral:
                return inf._cc_mes_integral_excel_bytes(
                    3, 620, user, nivel_aprobacion=4, agrupar_np=agrupar
                )
            return inf._cc_mes_001_excel_bytes(
                3, 620, user, nivel_aprobacion=4, agrupar_np=agrupar
            )

    return _generar(False), _generar(True), items, norma


def _assert_excel_si_vs_no(raw_no, raw_si, items, *, norma):
    wb_no = load_workbook(BytesIO(raw_no))
    wb_si = load_workbook(BytesIO(raw_si))
    ws_no = wb_no["CC-MES-001"]
    ws_si = wb_si["CC-MES-001"]
    book_no = FormulaBook(wb_no)
    book_si = FormulaBook(wb_si)
    shift = 1 if norma else 0
    col_desc, col_und, col_norma, col_vu = 2, 3, 4, 4 + shift
    col_cant_act, col_val_act = 5 + shift, 6 + shift
    col_cant_pres, col_val_pres = 7 + shift, 8 + shift
    col_val_acum, col_val_saldo = 10 + shift, 12 + shift

    filas_no = _filas_item(ws_no)
    filas_si = _filas_item(ws_si)
    assert [ws_no.cell(r, 1).value for r in filas_no] == [it["item_numero"] for it in items]
    assert [ws_si.cell(r, 1).value for r in filas_si] == [
        it["item_numero"] for it in agrupar_items_np_en_capitulo(items)
    ]
    por_num = {it["item_numero"]: it for it in items}
    for row in filas_si:
        num = ws_si.cell(row, 1).value
        it = por_num[num]
        assert ws_si.cell(row, col_desc).value == it["item_descripcion"]
        assert ws_si.cell(row, col_und).value == it["unidad"]
        if norma:
            assert (ws_si.cell(row, col_norma).value or "") == it["norma_tecnica"]
        assert float(ws_si.cell(row, col_vu).value) == float(it["vlr_unitario"])
        assert book_si.eval_cell("CC-MES-001", row, col_cant_act) == it["cant_actualizadas"]
        assert book_si.eval_cell("CC-MES-001", row, col_val_act) == it["valor_actualizadas"]
        assert book_si.eval_cell("CC-MES-001", row, col_val_pres) == it["valor_presente"]
        assert "ROUND(ROUND(" in str(ws_si.cell(row, col_val_act).value)
        assert ws_si.cell(row, col_val_act).number_format == ws_no.cell(filas_no[0], col_val_act).number_format

    subs_si = _subtotales(ws_si)
    assert subs_si[-1][1] == f"Subtotal {CAPITULO_NO_PREVISTOS}"
    assert f"Subtotal {CAPITULO_NO_PREVISTOS}" not in [n for _, n in _subtotales(ws_no)]
    cap1_no = next(r for r, n in _subtotales(ws_no) if n == f"Subtotal {CAP1}")
    cap1_si = next(r for r, n in subs_si if n == f"Subtotal {CAP1}")
    np_row = subs_si[-1][0]
    assert book_si.eval_cell("CC-MES-001", cap1_si, col_val_act) < book_no.eval_cell(
        "CC-MES-001", cap1_no, col_val_act
    )
    for col in (col_val_act, col_val_pres, col_val_acum, col_val_saldo):
        assert isinstance(ws_si.cell(np_row, col).value, str) and str(ws_si.cell(np_row, col).value).startswith("=")
        assert ws_si.cell(np_row, col).number_format == ws_no.cell(cap1_no, col).number_format

    for nombre in (
        "Costo Directo",
        "Costo Directo + AIU",
    ):
        rn = _linea(ws_no, nombre)
        rs = _linea(ws_si, nombre)
        for col in (col_val_act, col_val_pres, col_val_acum, col_val_saldo):
            assert book_si.eval_cell("CC-MES-001", rs, col) == book_no.eval_cell("CC-MES-001", rn, col)
    cd_si = _linea(ws_si, "Costo Directo")
    # El Costo Directo suma cada subtotal de capítulo, incluido NO PREVISTOS.
    formula_cd = str(ws_si.cell(cd_si, col_val_act).value)
    letra = get_column_letter(col_val_act)
    for row_sub, _nombre in subs_si:
        assert f"{letra}{row_sub}" in formula_cd


def test_excel_sin_flag_igual_y_con_flag_agrupa():
    inf = _import_informes_with_stubs()
    raw_no, raw_si, items, _norma = _excel_par(inf, integral=False)
    _assert_excel_si_vs_no(raw_no, raw_si, items, norma=False)
    with patch("informe_mes_items.agrupar_items_np_en_capitulo") as agrup:
        ctx = {
            "contrato": {"numero": "C"},
            "acta": {"numero_rpo": "1", "consecutivo": 1},
            "items": items,
            "total_costo": 0,
            "resumen_4cols": build_resumen_contrato_4cols(items=items),
            "otros_conceptos": [],
            "usuario_nombre": "Ana",
            "usuario_cargo": "Residente",
        }
        with (
            patch.object(inf, "_acta_pertenece_contrato", return_value=True),
            patch.object(inf, "_contexto_acta_mes_conciliacion", return_value=ctx),
            patch.object(inf, "_row", return_value=ctx["contrato"]),
            patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
        ):
            inf._cc_mes_001_excel_bytes(1, 1, {}, agrupar_np=False)
        agrup.assert_not_called()


def test_excel_integral_conserva_memoria_del_capitulo_origen():
    inf = _import_informes_with_stubs()
    raw_no, raw_si, items, _norma = _excel_par(inf, integral=True, norma=True)
    _assert_excel_si_vs_no(raw_no, raw_si, items, norma=True)
    wb_no = load_workbook(BytesIO(raw_no))
    wb_si = load_workbook(BytesIO(raw_si))
    assert wb_no.sheetnames == wb_si.sheetnames
    assert all(CAPITULO_NO_PREVISTOS not in nombre for nombre in wb_si.sheetnames)
    assert any("NP-05" in nombre and CAP1.split(".")[0] in nombre for nombre in wb_si.sheetnames)
    ws = wb_si["CC-MES-001"]
    fila_np = next(r for r in _filas_item(ws) if ws.cell(r, 1).value == "NP-05")
    # Con norma técnica, la cantidad de Presente acta es la columna 8.
    formula = str(ws.cell(fila_np, 8).value)
    assert "NP-05" in formula
    assert CAPITULO_NO_PREVISTOS not in formula
    assert "ROUND" in formula.upper()


def test_mismo_np_en_dos_capitulos_conserva_cada_memoria():
    """NP-06. existe en dos capítulos (acta real 625). Cada uno sigue su pestaña."""
    inf = _import_informes_with_stubs()
    from openpyxl import Workbook

    cap1 = _item(CAP1, "NP-06.", "Excavación extra", "M3", 100, 2, 1, "N1")
    cap3 = _item(CAP3, "NP-06.", "Alcantarilla extra", "ML", 200, 3, 1, "N3")
    base = _item(CAP1, "1.1", "Localización", "M2", 10, 1, 0)
    agrupados = agrupar_items_np_en_capitulo([base, cap1, cap3])
    nps = [it for it in agrupados if es_item_np(it["item_numero"])]
    assert [it["capitulo_origen"] for it in nps] == [CAP1, CAP3]
    assert nps[0]["item_descripcion"] == "Excavación extra"
    assert nps[1]["item_descripcion"] == "Alcantarilla extra"
    links = {
        f"{CAP1}||NP-06.": {"sheet": "1. PRELIMINARES|NP-06.", "tot_row": 20},
        f"{CAP3}||NP-06.": {"sheet": "3. OTROS|NP-06.", "tot_row": 30},
        "NP-06.": {"sheet": "3. OTROS|NP-06.", "tot_row": 30},
    }
    r4 = build_resumen_contrato_4cols(items=[base, cap1, cap3])
    wb = Workbook()
    ws = wb.active
    ws.title = "CC-MES-001"
    inf._fill_corte_sub_001_excel_ws(
        ws,
        {"numero": "C", "contratista": "Ctor", "nit": "", "interventoria": "Int"},
        {"razon_social": "1", "nombre_contacto": "—"},
        {"consecutivo": 1},
        agrupados,
        500,
        "Ana",
        "Residente",
        {},
        resumen_4cols=r4,
        memoria_links=links,
        incluir_norma_tecnica=True,
    )
    filas = [r for r in _filas_item(ws) if ws.cell(r, 1).value == "NP-06."]
    assert len(filas) == 2
    assert "1. PRELIMINARES|NP-06." in str(ws.cell(filas[0], 8).value)
    assert "3. OTROS|NP-06." in str(ws.cell(filas[1], 8).value)


def test_rutas_y_ui_preguntan_solo_al_descargar():
    inf = _import_informes_with_stubs()
    paths = [getattr(r, "path", "") for r in inf.router.routes]
    assert any(p.endswith("/tiene-items-np") for p in paths)
    for fn in (
        inf.pdf_cc_mes_001_acta,
        inf.pdf_cc_mes_001_acta_con_sello_firma,
        inf.excel_cc_mes_001_acta,
        inf.excel_cc_mes_integral_acta,
    ):
        default = inspect.signature(fn).parameters["agrupar_np"].default
        assert default.default is False
    src = Path(__file__).resolve().parents[2] / "frontend" / "src" / "ModuloInformes.jsx"
    text = src.read_text(encoding="utf-8")
    assert "Este informe tiene ítems No Previstos (NP). ¿Agruparlos en un capítulo aparte?" in text
    assert "tiene-items-np" in text
    assert "agrupar_np" in text
    assert "descargarPdfInformeMes" in text
    assert "elegirAgruparNp" in text
