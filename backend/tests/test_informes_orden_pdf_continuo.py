"""
Causa raíz (2.º ciclo — orden + espacios muertos PDF):

1. Orden: ``_contexto_*`` y ``enriquecer_items_bloques`` ordenaban con
   ``sort_items_por_orden_listado`` (índice Admin ``orden_listado``), no por
   capítulo+ítem numérico. Eso desordenaba 2.9/2.10 y mezclaba capítulos.

2. PDF: ``_html_cc_sub_v1_plain`` partía el plan en chunks de 18/22 filas e
   insertaba ``<pdf:nextpage />`` entre tablas, dejando páginas a medias.
   Además ``.cc001-resumen { page-break-inside:avoid }`` empujaba el resumen
   entero a la página siguiente.

Corrección: ``sort_items_capitulo_item_asc`` único (preview/PDF/Excel) y
una sola tabla continua con thead repetido.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from corte_sub_conciliacion import (
    enriquecer_items_bloques,
    filtrar_items_con_cantidades,
    sort_items_capitulo_item_asc,
    sort_items_por_orden_listado,
)
from test_corte_sub_001_excel_filtro import _import_informes_with_stubs


def _items_desordenados_listado():
    """Simula listado Admin con orden_listado que NO es capítulo+ítem asc."""
    return [
        {
            "item_numero": "2.10",
            "capitulo": "2. MOVIMIENTO",
            "orden_listado": 0,  # listado lo pone primero (mal)
            "item_descripcion": "Base granular",
            "unidad": "M3",
            "vlr_unitario_sub": 1000.0,
            "cant_actualizadas": 1,
            "cant_presente": 1,
            "cant_acumulado": 1,
            "cant_saldo": 0,
            "valor_actualizadas": 1000,
            "valor_presente": 1000,
            "valor_acumulado": 1000,
            "valor_saldo": 0,
            "costo_directo": 1000,
        },
        {
            "item_numero": "NP-02",
            "capitulo": "1. PRELIMINARES",
            "orden_listado": 1,
            "item_descripcion": "NP cap 1",
            "unidad": "GL",
            "vlr_unitario_sub": 500,
            "cant_actualizadas": 1,
            "cant_presente": 0,
            "cant_acumulado": 0,
            "cant_saldo": 1,
            "valor_actualizadas": 500,
            "valor_presente": 0,
            "valor_acumulado": 0,
            "valor_saldo": 500,
            "costo_directo": 0,
        },
        {
            "item_numero": "1.2",
            "capitulo": "1. PRELIMINARES",
            "orden_listado": 2,
            "item_descripcion": "Descapote",
            "unidad": "M3",
            "vlr_unitario_sub": 2000,
            "cant_actualizadas": 2,
            "cant_presente": 1,
            "cant_acumulado": 1,
            "cant_saldo": 1,
            "valor_actualizadas": 4000,
            "valor_presente": 2000,
            "valor_acumulado": 2000,
            "valor_saldo": 2000,
            "costo_directo": 2000,
        },
        {
            "item_numero": "2.9",
            "capitulo": "2. MOVIMIENTO",
            "orden_listado": 3,
            "item_descripcion": "Excavación",
            "unidad": "M3",
            "vlr_unitario_sub": 3000,
            "cant_actualizadas": 3,
            "cant_presente": 1,
            "cant_acumulado": 1,
            "cant_saldo": 2,
            "valor_actualizadas": 9000,
            "valor_presente": 3000,
            "valor_acumulado": 3000,
            "valor_saldo": 6000,
            "costo_directo": 3000,
        },
        {
            "item_numero": "1.1",
            "capitulo": "1. PRELIMINARES",
            "orden_listado": 4,
            "item_descripcion": "Replanteo",
            "unidad": "M2",
            "vlr_unitario_sub": 1000,
            "cant_actualizadas": 4,
            "cant_presente": 1,
            "cant_acumulado": 1,
            "cant_saldo": 3,
            "valor_actualizadas": 4000,
            "valor_presente": 1000,
            "valor_acumulado": 1000,
            "valor_saldo": 3000,
            "costo_directo": 1000,
        },
        {
            "item_numero": "10.1",
            "capitulo": "10. SEÑALIZACIÓN",
            "orden_listado": 5,
            "item_descripcion": "Señal",
            "unidad": "UND",
            "vlr_unitario_sub": 800,
            "cant_actualizadas": 1,
            "cant_presente": 1,
            "cant_acumulado": 1,
            "cant_saldo": 0,
            "valor_actualizadas": 800,
            "valor_presente": 800,
            "valor_acumulado": 800,
            "valor_saldo": 0,
            "costo_directo": 800,
        },
        {
            "item_numero": "NP-01",
            "capitulo": "2. MOVIMIENTO",
            "orden_listado": 6,
            "item_descripcion": "NP cap 2",
            "unidad": "GL",
            "vlr_unitario_sub": 700,
            "cant_actualizadas": 1,
            "cant_presente": 0,
            "cant_acumulado": 0,
            "cant_saldo": 1,
            "valor_actualizadas": 700,
            "valor_presente": 0,
            "valor_acumulado": 0,
            "valor_saldo": 700,
            "costo_directo": 0,
        },
    ]


ESPERADO = ["1.1", "1.2", "NP-02", "2.9", "2.10", "NP-01", "10.1"]


def test_causa_orden_listado_desordena_capitulo_item():
    """Evidencia: orden_listado NO produce 1.1→1.2→… ni 2.9 antes de 2.10."""
    items = _items_desordenados_listado()
    by_listado = [it["item_numero"] for it in sort_items_por_orden_listado(items)]
    assert by_listado == ["2.10", "NP-02", "1.2", "2.9", "1.1", "10.1", "NP-01"]
    assert by_listado != ESPERADO


def test_sort_capitulo_item_asc_numerico_y_np_al_final_del_cap():
    out = sort_items_capitulo_item_asc(_items_desordenados_listado())
    assert [it["item_numero"] for it in out] == ESPERADO
    # Capítulo 2 antes que 10; 2.9 antes que 2.10; NP tras numéricos del cap
    assert ESPERADO.index("2.9") < ESPERADO.index("2.10")
    assert ESPERADO.index("2.10") < ESPERADO.index("NP-01")
    assert ESPERADO.index("1.2") < ESPERADO.index("NP-02")
    assert ESPERADO.index("NP-02") < ESPERADO.index("2.9")
    assert ESPERADO.index("2.9") < ESPERADO.index("10.1")


def test_enriquecer_usa_orden_capitulo_item_no_listado():
    presente = [
        {
            "item_numero": "2.10",
            "item_descripcion": "Base",
            "unidad": "M3",
            "capitulo": "2. MOVIMIENTO",
            "cantidad": 1.0,
            "vlr_unitario_sub": 1000.0,
        }
    ]
    meta = {
        "1.1": {
            "capitulo": "1. PRELIMINARES",
            "descripcion": "A",
            "unidad": "M2",
            "vlr_unitario": 100,
            "orden_listado": 99,
        },
        "2.9": {
            "capitulo": "2. MOVIMIENTO",
            "descripcion": "B",
            "unidad": "M3",
            "vlr_unitario": 200,
            "orden_listado": 0,
        },
        "2.10": {
            "capitulo": "2. MOVIMIENTO",
            "descripcion": "Base",
            "unidad": "M3",
            "vlr_unitario": 1000,
            "orden_listado": 1,
        },
        "NP-01": {
            "capitulo": "2. MOVIMIENTO",
            "descripcion": "NP",
            "unidad": "GL",
            "vlr_unitario": 50,
            "orden_listado": 2,
        },
    }
    items = enriquecer_items_bloques(
        presente,
        cant_actualizadas={"1.1": 1, "2.9": 1, "2.10": 1, "NP-01": 1},
        cant_acum_anterior={},
        vu_por_item={k: v["vlr_unitario"] for k, v in meta.items()},
        meta_por_item=meta,
    )
    items = filtrar_items_con_cantidades(items)
    assert [it["item_numero"] for it in items] == ["1.1", "2.9", "2.10", "NP-01"]


def test_html_pdf_continuo_sin_nextpage_y_misma_secuencia():
    inf = _import_informes_with_stubs()
    items = sort_items_capitulo_item_asc(_items_desordenados_listado())
    html = inf._html_cc_sub_v1_plain(
        {"numero": "ICCU-CTO-1614-2025", "contratista": "Ctor", "nit": "900", "interventoria": "Int"},
        {"razon_social": "M.E OBRAS Y PAVIMENTOS SAS", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        items,
        7800.0,
        "Tester",
        "Desarrollador",
    )
    assert "<pdf:nextpage" not in html.lower()
    assert html.count('class="cc001-tabla-items"') == 1
    assert "display: table-header-group" in html
    assert "page-break-inside: avoid" in html
    # Resumen no fuerza bloque entero a otra página
    assert "cc001-resumen" in html
    assert "page-break-inside: auto" in html

    # Secuencia de ítems en el HTML (sin filas Subtotal)
    nums = []
    for it in items:
        assert f">{it['item_numero']}</td>" in html or f">{it['item_numero']}<" in html
        nums.append(it["item_numero"])
    assert nums == ESPERADO

    # Orden relativo en el HTML: 1.1 antes de 1.2 antes de NP-02 … 2.9 antes de 2.10
    pos = {n: html.find(n) for n in ESPERADO}
    assert all(p >= 0 for p in pos.values())
    for a, b in zip(ESPERADO, ESPERADO[1:]):
        assert pos[a] < pos[b], f"{a} debe aparecer antes que {b}"


def test_plan_chunk_es_un_solo_bloque_continuo():
    inf = _import_informes_with_stubs()
    items = sort_items_capitulo_item_asc(_items_desordenados_listado())
    plan = inf._cc_sub_001_plan_filas_capitulo(items)
    chunks = inf._cc_sub_001_chunk_plan(plan)
    assert len(chunks) == 1
    assert len(chunks[0]) == len(plan)
    # Muchas filas ya no generan múltiples chunks
    big = items * 20
    plan_big = inf._cc_sub_001_plan_filas_capitulo(big)
    assert len(inf._cc_sub_001_chunk_plan(plan_big)) == 1


def test_excel_misma_secuencia_que_sort():
    inf = _import_informes_with_stubs()
    items = sort_items_capitulo_item_asc(_items_desordenados_listado())
    from corte_sub_conciliacion import build_resumen_conciliacion_4cols

    r4 = build_resumen_conciliacion_4cols(
        items=items, tributos={"administracion": 10}, otros_presente=0
    )
    raw = inf._corte_sub_001_excel_bytes(
        {"numero": "ICCU-CTO-1614-2025", "contratista": "Ctor", "nit": "900", "interventoria": "Int"},
        {"razon_social": "M.E OBRAS Y PAVIMENTOS SAS", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        items,
        7800.0,
        "Tester",
        "Desarrollador",
        {},
        resumen_4cols=r4,
    )
    ws = load_workbook(BytesIO(raw)).active
    excel_nums = []
    for r in range(9, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None:
            continue
        s = str(v)
        if s.startswith("Subtotal"):
            continue
        if s in ("Costo Directo", "Concepto", "FIRMAS", "Elaboró:", "Revisó:", "Aprobó (subcontratista):"):
            break
        if s.startswith("RESUMEN") or s.startswith("Período"):
            break
        excel_nums.append(s)
    assert excel_nums == ESPERADO


def test_sort_inplace_informes_alineado():
    inf = _import_informes_with_stubs()
    items = list(_items_desordenados_listado())
    inf._sort_items_corte_por_item_numero_asc(items)
    assert [it["item_numero"] for it in items] == ESPERADO
