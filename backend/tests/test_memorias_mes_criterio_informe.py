"""CC-MES-002 usa el mismo universo de ítems que el informe CC-MES-001."""
from __future__ import annotations

import sys
from contextlib import ExitStack
from io import BytesIO
from types import ModuleType, SimpleNamespace
from unittest.mock import patch

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from acta_mes_conciliacion import build_resumen_contrato_4cols
from informe_mes_items import (
    construir_cuadro_mensual,
    filtrar_registros_memorias_mes,
    sumar_cantidad_por_clave,
)

CAP_PAV = "2. ESTRUCTURA DE PAVIMENTO Y ADOQUINES"
CAP_SENAL = "2. SEÑALIZACIÓN Y CANALIZACIÓN"
CAP_HORIZ = "4. SEÑALIZACIÓN HORIZONTAL Y VERTICAL"


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]
    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake
    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)
    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    return inf


def _listado(cap, item, desc, *, estado="Aprobado", id_=1, vu=1000):
    return {
        "id": id_,
        "capitulo": cap,
        "item_numero": item,
        "descripcion": desc,
        "unidad": "M3",
        "precio_unitario": vu,
        "estado_precio": estado,
    }


def _reg(cap, item, desc, numero, cant=1.0, marca=""):
    return {
        "capitulo": cap,
        "item_numero": item,
        "item_descripcion": desc,
        "unidad": "M3",
        "numero_registro": numero,
        "cantidad_total": cant,
        "vlr_unitario": 1000,
        "costo_directo": cant * 1000,
        "longitud": cant,
        "ancho": "",
        "espesor": "",
        "cantidad": "",
        "observacion": marca,
        "abs_inicio": "",
        "abs_final": "",
        "infraestructura": "",
    }


def _catalogo():
    filas = [
        _listado(CAP_PAV, "2.1", "SUBBASE GRANULAR", id_=1),
        _listado(CAP_PAV, "2.1._", "Código con sufijo", id_=2),
        _listado(CAP_PAV, "2.2", "Base sin ejecución", id_=3),
        _listado(CAP_HORIZ, "2.1", "Flecha", id_=4),
        _listado(CAP_PAV, "2.7", "Pendiente de precio", estado="Pendiente", id_=5),
        _listado(CAP_SENAL, "2.1", "Señal del mismo código", id_=6),
        _listado(CAP_SENAL, "2.1._", "Señal SR-30", id_=7),
        _listado("4.  OPERACIÓN   PMT", "4.9._", "Desmonte", id_=8),
        _listado(CAP_SENAL, "9.9._", "Señal SR-30", id_=9),
    ]
    registros = [
        _reg(CAP_SENAL, "2.1._", "Señal SR-30", 1, 4, "marca-senal-sufijo"),
        _reg(CAP_PAV, "2.7", "Pendiente de precio", 2, 3, "marca-pendiente"),
        _reg(CAP_PAV, "2.1", "SUBBASE GRANULAR", 3, 1.5, "marca-subbase"),
        _reg("4.  OPERACIÓN   PMT", "4.9._", "Desmonte", 4, 2, "marca-desmonte"),
        _reg(CAP_HORIZ, "2.1", "Flecha", 5, 0.25, "marca-flecha"),
        _reg(CAP_PAV, "2.1._", "Código con sufijo", 6, 0.5, "marca-sufijo-pav"),
        _reg(CAP_SENAL, "2.1", "Señal del mismo código", 7, 9, "marca-senal-mismo"),
    ]
    return filas, registros


def test_filtro_conserva_solo_claves_del_informe_y_el_orden():
    filas, registros = _catalogo()
    quedaron = filtrar_registros_memorias_mes(registros, filas)
    assert [r["numero_registro"] for r in quedaron] == [3, 5, 6]
    assert [(r["capitulo"], r["item_numero"]) for r in quedaron] == [
        (CAP_PAV, "2.1"),
        (CAP_HORIZ, "2.1"),
        (CAP_PAV, "2.1._"),
    ]
    cuadro = construir_cuadro_mensual(
        filas,
        cant_presente=sumar_cantidad_por_clave(registros),
        claves_con_registros_acta=sumar_cantidad_por_clave(registros).keys(),
    )
    con_ejecucion = {
        (it["capitulo"], it["item_numero"])
        for it in cuadro["items"]
        if it.get("tiene_registros_acta")
    }
    assert {(r["capitulo"], r["item_numero"]) for r in quedaron} == con_ejecucion
    assert (CAP_PAV, "2.2") in {(it["capitulo"], it["item_numero"]) for it in cuadro["items"]}
    assert (CAP_PAV, "2.2") not in {(r["capitulo"], r["item_numero"]) for r in quedaron}


def test_cambiar_el_estado_en_el_listado_saca_la_memoria_y_el_informe():
    filas, registros = _catalogo()
    for fila in filas:
        if fila["item_numero"] == "2.1._" and fila["capitulo"] == CAP_PAV:
            fila["estado_precio"] = "Rechazado"
    quedaron = filtrar_registros_memorias_mes(registros, filas)
    assert (CAP_PAV, "2.1._") not in {(r["capitulo"], r["item_numero"]) for r in quedaron}
    cuadro = construir_cuadro_mensual(filas, cant_presente=sumar_cantidad_por_clave(registros))
    assert (CAP_PAV, "2.1._") not in {(it["capitulo"], it["item_numero"]) for it in cuadro["items"]}
    assert [r["numero_registro"] for r in quedaron] == [3, 5]


def _parches_salida(inf, filas, registros):
    contrato = {"numero": "C-1", "contratista": "Ctor", "nit": "1", "interventoria": "Int"}
    acta = {"numero_rpo": "1", "consecutivo": 1, "contrato_id": 3}

    def _row(table, *_a, **_k):
        if table == "actas":
            return acta
        if table == "contratos":
            return contrato
        return None

    def _fetch_item(_sb, _cid, inum, **_kw):
        return [r for r in registros if str(r.get("item_numero")) == str(inum)]

    return (
        patch.object(inf, "_perm_informes_ccd"),
        patch.object(inf, "_perm_informes_ccd_lectura"),
        patch.object(inf, "_acta_pertenece_contrato", return_value=True),
        patch.object(inf, "_nivel_aprobacion_mes_query", return_value=None),
        patch.object(inf, "_row", side_effect=_row),
        patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
        patch.object(inf, "_firma_data_uri_para_slot_contexto", return_value=None),
        patch.object(inf, "_sub_corte_dummy_memoria", return_value=({}, {})),
        patch.object(inf, "fetch_registros_informe_cc_mes_por_acta", return_value=list(registros)),
        patch.object(inf, "fetch_registros_memoria_cc_mes_acta_todos", return_value=list(registros)),
        patch.object(inf, "fetch_registros_memoria_cc_mes_alineado_acta", side_effect=_fetch_item),
        patch("informe_mes_items.fetch_listado_contrato", return_value=filas),
    )


def test_vista_previa_pdf_y_excel_omiten_pmt_y_no_aprobados():
    inf = _import_informes_with_stubs()
    filas, registros = _catalogo()
    user = {"nombre": "Ana", "apellidos": "Pérez", "cargo_nombre": "Residente"}
    htmls = []

    def _to_pdf(html):
        htmls.append(html)
        return b"%PDF-1.4"

    with ExitStack() as stack:
        for parche in _parches_salida(inf, filas, registros):
            stack.enter_context(parche)
        stack.enter_context(patch.object(inf, "_to_pdf", side_effect=_to_pdf))
        lista = inf.ccd_items_conciliacion_acta(3, 620, None, user)
        resp = inf.pdf_cc_mes_002_acta(3, 620, "2.1", None, user)
        resp_sufijo = inf.pdf_cc_mes_002_acta(3, 620, "2.1._", None, user)
        with pytest.raises(HTTPException) as pmt:
            inf.pdf_cc_mes_002_acta(3, 620, "4.9._", None, user)
        with pytest.raises(HTTPException) as pend:
            inf.pdf_cc_mes_002_acta_con_sello_firma(3, 620, "2.7", None, user)
        ctx = inf._cc_mes_002_acta_completo_ctx(3, 620, user)
        raw_item = inf._cc_mes_002_item_excel_bytes(3, 620, "2.1", user)
        raw_todos = inf._cc_mes_002_acta_completo_excel_bytes(3, 620, user)
        with pytest.raises(HTTPException):
            inf._cc_mes_002_item_excel_bytes(3, 620, "4.9._", user)

    assert {it["item_numero"] for it in lista["items"]} == {"2.1", "2.1._"}
    assert lista["registros_raw"] == 3
    assert resp.body == b"%PDF-1.4"
    assert resp_sufijo.body == b"%PDF-1.4"
    html_mismo, html_sufijo = htmls
    assert "SUBBASE GRANULAR" in html_mismo
    assert "marca-subbase" in html_mismo and "marca-flecha" in html_mismo
    assert "marca-senal-mismo" not in html_mismo
    assert "Código con sufijo" in html_sufijo and "marca-sufijo-pav" in html_sufijo
    assert "marca-senal-sufijo" not in html_sufijo
    assert "marca-pendiente" not in html_mismo + html_sufijo
    assert "marca-desmonte" not in html_mismo + html_sufijo
    assert pmt.value.status_code == 404
    assert pend.value.status_code == 404
    assert ctx["n_items"] == 2
    assert ctx["n_regs"] == 3
    assert set(ctx["numeros"]) == {"2.1", "2.1._"}
    joined = ctx["html_joined"]
    assert "marca-subbase" in joined and "marca-flecha" in joined and "marca-sufijo-pav" in joined
    assert "marca-senal-sufijo" not in joined
    assert "marca-senal-mismo" not in joined
    assert "marca-pendiente" not in joined
    assert "marca-desmonte" not in joined

    wb_item = load_workbook(BytesIO(raw_item))
    numeros = [
        wb_item.active.cell(r, 1).value
        for r in range(1, wb_item.active.max_row + 1)
        if isinstance(wb_item.active.cell(r, 1).value, int)
    ]
    assert numeros == [3, 5]
    wb = load_workbook(BytesIO(raw_todos))
    assert len(wb.sheetnames) == 2
    texto = " ".join(
        str(c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    )
    assert "Señal SR-30" not in texto
    assert "Pendiente de precio" not in texto
    assert "Desmonte" not in texto
    assert "SUBBASE GRANULAR" in texto and "Código con sufijo" in texto


def test_libro_integral_pestañas_y_formulas_siguen_al_informe():
    inf = _import_informes_with_stubs()
    filas, registros = _catalogo()
    cuadro = construir_cuadro_mensual(
        filas,
        cant_presente=sumar_cantidad_por_clave(registros),
        claves_con_registros_acta=sumar_cantidad_por_clave(registros).keys(),
    )
    items = list(cuadro["items"])
    # Deriva: el contexto trae también un PMT y un pendiente. El filtro compartido los deja fuera.
    items.append(
        {
            "capitulo": CAP_SENAL,
            "item_numero": "9.9._",
            "item_descripcion": "Señal SR-30",
            "unidad": "UND",
            "vlr_unitario": 50000,
            "cant_presente": 4,
            "valor_presente": 200000,
            "cant_actualizadas": 4,
            "valor_actualizadas": 200000,
            "cant_acum_anterior": 0,
            "cant_acumulado": 4,
            "cant_saldo": 0,
            "valor_acumulado": 200000,
            "valor_saldo": 0,
            "norma_tecnica": "",
        }
    )
    items.append(
        {
            "capitulo": CAP_PAV,
            "item_numero": "2.7",
            "item_descripcion": "Pendiente de precio",
            "unidad": "UND",
            "vlr_unitario": 80,
            "cant_presente": 3,
            "valor_presente": 240,
            "cant_actualizadas": 3,
            "valor_actualizadas": 240,
            "cant_acum_anterior": 0,
            "cant_acumulado": 3,
            "cant_saldo": 0,
            "valor_acumulado": 240,
            "valor_saldo": 0,
            "norma_tecnica": "",
        }
    )
    r4 = build_resumen_contrato_4cols(items=cuadro["items"])
    contrato = {"numero": "C-1", "contratista": "Ctor", "nit": "1", "interventoria": "Int"}
    ctx = {
        "contrato": contrato,
        "acta": {"numero_rpo": "1", "consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        "items": items,
        "total_costo": sum(it["valor_presente"] for it in cuadro["items"]),
        "resumen_4cols": r4,
        "otros_conceptos": [],
        "usuario_nombre": "Ana",
        "usuario_cargo": "Residente",
    }
    user = {"nombre": "Ana", "apellidos": "Pérez", "cargo_nombre": "Residente"}

    def _fetch(_sb, _cid, inum, **_kw):
        if str(inum) == "9.9._":
            return [_reg(CAP_SENAL, "9.9._", "Señal SR-30", 1, 4, "marca-senal-sufijo")]
        return [r for r in registros if str(r.get("item_numero")) == str(inum)]

    with (
        patch.object(inf, "_acta_pertenece_contrato", return_value=True),
        patch.object(inf, "_contexto_acta_mes_conciliacion", return_value=ctx),
        patch.object(inf, "_row", return_value=contrato),
        patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
        patch.object(inf, "_sub_corte_dummy_memoria", return_value=({}, {})),
        patch.object(inf, "fetch_registros_memoria_cc_mes_alineado_acta", side_effect=_fetch),
        patch("informe_mes_items.fetch_listado_contrato", return_value=filas),
    ):
        raw = inf._cc_mes_integral_excel_bytes(3, 620, user)

    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames[0] == "CC-MES-001"
    memorias = wb.sheetnames[1:]
    assert len(memorias) == 3
    texto_hojas = " ".join(memorias)
    assert "2.1._" in texto_hojas
    assert CAP_SENAL not in texto_hojas
    assert "2.7" not in texto_hojas
    assert "9.9" not in texto_hojas
    ws = wb["CC-MES-001"]
    formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("=ROUND(") and "!H" in cell.value:
                formulas.append(cell.value)
    assert len(formulas) == 3
    for formula in formulas:
        nombre = formula.split("'", 2)[1]
        assert nombre in memorias
    assert all("2.7" not in f and CAP_SENAL not in f for f in formulas)
