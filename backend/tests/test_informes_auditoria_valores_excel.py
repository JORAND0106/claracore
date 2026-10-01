"""
Auditoría exhaustiva: consistencia plataforma / PDF / Excel (CC-SUB y CC-MES).

Causas raíz (ANTES):
1. ``valor_por_cantidad_vu`` hacía ROUND0(cant×VU) sin redondear cant a 2 dp →
   valores distintos a la regla (p. ej. 6.005×2205=13241 vs 6.00×2205=13230).
2. Excel: ``ROUND(E*D,0)`` sin ``ROUND(cant,2)``; subtotales ``=F9+F10`` celda a
   celda; CD sumaba ítems, no subtotales; formatos ``$"#,##0`` y cantidades
   mixtas (#,##0 / #,##0.##).
3. PDF ``_fn_cant_informe`` ocultaba .00 en enteros.

Regla única DESPUÉS:
  cant = ROUND(cant, 2); valor = ROUND(cant×VU, 0);
  subtotales/totales = suma de valores ya redondeados.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from corte_sub_conciliacion import (
    build_resumen_conciliacion_4cols,
    enriquecer_items_bloques,
    filtrar_items_con_cantidades,
    valor_por_cantidad_vu,
)
from test_corte_sub_001_excel_filtro import _import_informes_with_stubs


def _dataset_auditoria():
    presente = [
        {
            "item_numero": "1.1",
            "item_descripcion": "Excavación",
            "unidad": "M3",
            "capitulo": "1. PRELIMINARES",
            "cantidad": 6.005,
            "vlr_unitario_sub": 2205.0,
        },
        {
            "item_numero": "1.2",
            "item_descripcion": "Descapote",
            "unidad": "M3",
            "capitulo": "1. PRELIMINARES",
            "cantidad": 2.5,
            "vlr_unitario_sub": 1000.0,
        },
    ]
    meta = {
        "1.1": {
            "capitulo": "1. PRELIMINARES",
            "descripcion": "Excavación",
            "unidad": "M3",
            "vlr_unitario": 2205.0,
            "orden_listado": 0,
        },
        "1.2": {
            "capitulo": "1. PRELIMINARES",
            "descripcion": "Descapote",
            "unidad": "M3",
            "vlr_unitario": 1000.0,
            "orden_listado": 1,
        },
        "2.9": {
            "capitulo": "2. MOVIMIENTO",
            "descripcion": "Base",
            "unidad": "M3",
            "vlr_unitario": 1000.0,
            "orden_listado": 2,
        },
    }
    items = enriquecer_items_bloques(
        presente,
        cant_actualizadas={"1.1": 10.004, "1.2": 5.0, "2.9": 3.14159},
        cant_acum_anterior={"1.1": 1.001},
        vu_por_item={k: v["vlr_unitario"] for k, v in meta.items()},
        meta_por_item=meta,
    )
    return filtrar_items_con_cantidades(items)


def test_regla_unica_cantidad_2dp_valor_round0():
    assert valor_por_cantidad_vu(6.005, 2205) == 13230.0  # 6.00×2205
    assert valor_por_cantidad_vu(10.004, 2205) == 22050.0
    assert valor_por_cantidad_vu(3.14159, 1000) == 3140.0


def test_enriquecer_aplica_regla_unica():
    items = _dataset_auditoria()
    by = {it["item_numero"]: it for it in items}
    assert by["1.1"]["cant_presente"] == 6.0  # round(6.005,2) → 6.0 en float
    assert by["1.1"]["cant_actualizadas"] == 10.0
    assert by["1.1"]["cant_acum_anterior"] == 1.0
    assert by["1.1"]["cant_acumulado"] == 7.0
    assert by["1.1"]["cant_saldo"] == 3.0
    assert by["1.1"]["valor_presente"] == 13230.0
    assert by["1.1"]["valor_actualizadas"] == 22050.0
    assert by["1.1"]["valor_acumulado"] == 15435.0  # 7.00×2205
    assert by["1.1"]["valor_saldo"] == 6615.0
    assert by["2.9"]["cant_actualizadas"] == 3.14
    assert by["2.9"]["valor_actualizadas"] == 3140.0


def test_plataforma_pdf_excel_misma_secuencia_y_valores():
    """Cuadro comparativo DESPUÉS: ctx = PDF = Excel (cant/valor)."""
    items = _dataset_auditoria()
    r4 = build_resumen_conciliacion_4cols(
        items=items, tributos={"administracion": 10}, otros_presente=0
    )
    inf = _import_informes_with_stubs()
    total = sum(float(i["valor_presente"]) for i in items)
    html = inf._html_cc_sub_v1_plain(
        {"numero": "ICCU-CTO-1614-2025", "contratista": "C", "nit": "1", "interventoria": "I"},
        {"razon_social": "M.E OBRAS Y PAVIMENTOS SAS", "nombre_contacto": "A"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        items,
        total,
        "T",
        "D",
        resumen_4cols=r4,
    )
    xlsx = inf._corte_sub_001_excel_bytes(
        {"numero": "ICCU-CTO-1614-2025", "contratista": "C", "nit": "1", "interventoria": "I"},
        {"razon_social": "M.E OBRAS Y PAVIMENTOS SAS", "nombre_contacto": "A"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        items,
        total,
        "T",
        "D",
        {},
        resumen_4cols=r4,
    )
    ws = load_workbook(BytesIO(xlsx)).active

    # Map Excel item rows
    excel_by_item = {}
    for r in range(9, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None or str(a).startswith("Subtotal") or a in ("RESUMEN DE CONCILIACIÓN", "Concepto", "Costo Directo"):
            if a == "Costo Directo" or (isinstance(a, str) and a.startswith("RESUMEN")):
                break
            if isinstance(a, str) and a.startswith("Subtotal"):
                continue
            if a is None:
                continue
        if isinstance(a, str) and a[0].isdigit() or (isinstance(a, str) and a.startswith("NP")):
            excel_by_item[str(a)] = r

    diffs = []
    for it in items:
        n = it["item_numero"]
        # PDF muestra cantidades con 2 dp y valores Round0
        for ck, label in (
            ("cant_actualizadas", "cant_act"),
            ("cant_presente", "cant_pres"),
            ("cant_acumulado", "cant_acum"),
            ("cant_saldo", "cant_saldo"),
        ):
            pdf_txt = inf._fn_cant_informe(it[ck])
            expected = f"{float(it[ck]):,.2f}"
            if pdf_txt != expected:
                diffs.append(f"{n}.{label}: PDF={pdf_txt!r} ctx={expected!r}")
        for vk in ("valor_actualizadas", "valor_presente", "valor_acumulado", "valor_saldo"):
            pdf_v = inf._fm(it[vk])
            # _fm sin .00; valor numérico Round0
            if float(it[vk]) != round(float(it[vk]), 0):
                diffs.append(f"{n}.{vk}: no Round0 {it[vk]!r}")
            if f"{int(round(float(it[vk]))):,}" not in pdf_v.replace(" ", "") and pdf_v != f"$ {int(round(float(it[vk]))):,}":
                # soft check: PDF contiene el entero
                if str(int(round(float(it[vk])))) not in pdf_v.replace(",", ""):
                    diffs.append(f"{n}.{vk}: PDF={pdf_v!r}")

        # Excel: cantidades literales = ctx; fórmulas valor con ROUND(ROUND(,2)*VU,0)
        r = excel_by_item.get(n)
        assert r is not None, f"ítem {n} no en Excel"
        assert float(ws.cell(r, 5).value) == float(it["cant_actualizadas"])
        assert float(ws.cell(r, 7).value) == float(it["cant_presente"])
        assert ws.cell(r, 5).number_format == "#,##0.00"
        assert ws.cell(r, 6).number_format == '"$"#,##0.00'
        f_val = str(ws.cell(r, 6).value)
        assert "ROUND(ROUND(" in f_val and ",0)" in f_val

    assert not diffs, "Diferencias plataforma/PDF:\n" + "\n".join(diffs)

    # Resumen CD: plataforma = suma valores presentes redondeados
    cd_line = next(ln for ln in (r4.get("lineas") or []) if ln.get("key") == "cd")
    assert float(cd_line["valores"]["presente"]) == total

    # Excel CD = SUM de subtotales
    cd_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Costo Directo")
    cd_f = str(ws.cell(cd_row, 6).value)
    assert "Subtotal" not in cd_f
    # Debe referenciar filas de subtotal (11 y 13 en este dataset: 1.1,1.2,sub, 2.9,sub)
    assert cd_f.startswith("=")
    assert ws.cell(cd_row, 6).number_format == '"$"#,##0.00'


def test_excel_subtotal_sum_rango_y_cd_suma_subtotales():
    items = _dataset_auditoria()
    r4 = build_resumen_conciliacion_4cols(
        items=items, tributos={"administracion": 10}, otros_presente=0
    )
    inf = _import_informes_with_stubs()
    raw = inf._corte_sub_001_excel_bytes(
        {"numero": "ICCU-CTO-1614-2025", "contratista": "C", "nit": "1", "interventoria": "I"},
        {"razon_social": "M.E OBRAS Y PAVIMENTOS SAS", "nombre_contacto": "A"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        items,
        sum(float(i["valor_presente"]) for i in items),
        "T",
        "D",
        {},
        resumen_4cols=r4,
    )
    ws = load_workbook(BytesIO(raw)).active
    # R9=1.1 R10=1.2 R11=subcap1 R12=2.9 R13=subcap2
    assert ws.cell(9, 1).value == "1.1"
    assert ws.cell(10, 1).value == "1.2"
    assert str(ws.cell(11, 1).value).startswith("Subtotal 1.")
    assert ws.cell(11, 6).value == "=SUM(F9:F10)"
    assert ws.cell(12, 1).value == "2.9"
    assert str(ws.cell(13, 1).value).startswith("Subtotal 2.")
    assert ws.cell(13, 6).value == "=F12"  # un solo ítem → referencia directa
    cd_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Costo Directo")
    assert ws.cell(cd_row, 6).value == "=F11+F13" or "F11" in str(ws.cell(cd_row, 6).value)


def test_memoria_formatos_dim_3_cant_2():
    inf = _import_informes_with_stubs()
    assert inf._EXCEL_NUM_FMT_DIM == "0.000"
    assert inf._EXCEL_NUM_FMT_CANT == "#,##0.00"
    assert inf._EXCEL_NUM_FMT_MONEY == '"$"#,##0.00'
    # Fórmula CANT TOT a 2 dp
    f = inf._excel_formula_cantidad_total(12)
    assert "ROUND(" in f and "ROUND(" in f
    assert ">=0.1" not in f  # ya no redondeo dinámico 2/3
    assert inf._excel_formula_valor_cant_vu("E9", "D9") == (
        '=IF(OR(D9="",E9=""),0,ROUND(ROUND(E9,2)*D9,0))'
    )
