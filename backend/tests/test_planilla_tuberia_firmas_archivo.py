"""Pie de firmas (quién validó y cuándo) y nombre de archivo con el reporte vigente."""
from __future__ import annotations

import io
import unittest
from pathlib import Path

from openpyxl import load_workbook

from topografia_planilla_tuberia import (
    aplicar_firmas_validacion_export,
    conservar_links_sicoe_vigentes,
    html_pie_firmas_planilla,
    nombre_archivo_planilla_tuberia,
)
from topografia_planilla_tuberia_excel import build_planilla_tuberia_xlsx

_NOMBRES = {10: "Ana Gómez", 20: "Luis Pérez"}
_N1 = "2026-10-06T19:35:42.123456+00:00"  # 14:35 en Colombia
_N2 = "2026-10-06T21:10:07.999999Z"  # 16:10 en Colombia
_ELABORO = "Ana Gómez · 06/10/2026 14:35"
_APROBO = "Luis Pérez · 06/10/2026 16:10"


def _planilla(**extra):
    base = {
        "id": "6fa2ebee-1111-2222-3333-444455556666",
        "tipo": "ALCANTARILLA",
        "firmas": {
            "elaboro_nombre": "Nombre configurado",
            "aprobo_nombre": "Otro configurado",
        },
        "meta_cabecera": {
            "sicoe_reportes": [
                {"reporte_id": 9, "numero_reporte": 1},
                {"reporte_id": 4, "numero_reporte": 99},
            ]
        },
    }
    base.update(extra)
    return base


def _excel(planilla):
    raw = build_planilla_tuberia_xlsx(
        planilla=planilla,
        calculo={"cartera": {"filas": []}},
        vacia=True,
    )
    return load_workbook(io.BytesIO(raw))["planilla"]


class TestPieFirmasYNombreArchivo(unittest.TestCase):
    def test_ambos_niveles_aprobados_con_reporte_vigente(self):
        p = _planilla(
            nivel1_estado="Aprobado",
            nivel1_usuario_id=10,
            nivel1_fecha=_N1,
            nivel2_estado="Aprobado",
            nivel2_usuario_id=20,
            nivel2_fecha=_N2,
        )
        exportada = aplicar_firmas_validacion_export(p, _NOMBRES)
        self.assertEqual(exportada["firmas"]["elaboro_nombre"], _ELABORO)
        self.assertEqual(exportada["firmas"]["aprobo_nombre"], _APROBO)
        self.assertNotIn("42", exportada["firmas"]["elaboro_nombre"])
        self.assertNotIn("123456", exportada["firmas"]["elaboro_nombre"])
        self.assertEqual(p["firmas"]["elaboro_nombre"], "Nombre configurado")

        links = conservar_links_sicoe_vigentes(
            p["meta_cabecera"]["sicoe_reportes"],
            [{"id": 9, "numero_reporte": 128}],
        )
        self.assertEqual([lk["numero_reporte"] for lk in links], [128])
        self.assertEqual(
            nombre_archivo_planilla_tuberia(links, extension="xlsx"),
            "planilla_tuberia_128.xlsx",
        )
        self.assertEqual(
            nombre_archivo_planilla_tuberia(links, extension="pdf"),
            "planilla_tuberia_128.pdf",
        )
        self.assertNotIn("6fa2ebee", nombre_archivo_planilla_tuberia(links, extension="xlsx"))

        html = html_pie_firmas_planilla(_ELABORO, _APROBO)
        self.assertIn("Elaboró", html)
        self.assertIn(_ELABORO, html)
        self.assertIn("Aprobó:", html)
        self.assertIn(_APROBO, html)
        self.assertIn("Topografo de Obra (Contratista)", html)
        self.assertIn("Topografo Interventoria", html)

        ws = _excel(exportada)
        self.assertEqual(ws["A65"].value, _ELABORO)
        self.assertEqual(ws["H65"].value, _APROBO)
        self.assertEqual(ws["A64"].value, "Elaboró")
        self.assertEqual(ws["H64"].value, "Aprobó:")
        self.assertEqual(ws["A66"].value, "Topografo de Obra (Contratista)")
        self.assertEqual(ws["H66"].value, "Topografo Interventoria")

    def test_solo_contratista_aprobado(self):
        p = _planilla(
            nivel1_estado="Aprobado",
            nivel1_usuario_id=10,
            nivel1_fecha="2026-10-06 19:35:42+00:00",
            nivel2_estado="Pendiente",
            nivel2_usuario_id=20,
            nivel2_fecha=_N2,
        )
        exportada = aplicar_firmas_validacion_export(p, _NOMBRES)
        self.assertEqual(exportada["firmas"]["elaboro_nombre"], _ELABORO)
        self.assertEqual(exportada["firmas"]["aprobo_nombre"], "")
        self.assertEqual(exportada["firmas"]["aprobo"], "")

        links = conservar_links_sicoe_vigentes(
            [{"reporte_id": 3, "numero_reporte": 7}],
            [{"id": 3, "numero_reporte": 7}],
        )
        self.assertEqual(
            nombre_archivo_planilla_tuberia(links, extension="pdf"),
            "planilla_tuberia_7.pdf",
        )
        self.assertEqual(
            nombre_archivo_planilla_tuberia(links, extension="xlsx"),
            "planilla_tuberia_7.xlsx",
        )

        html = html_pie_firmas_planilla(_ELABORO, "")
        self.assertIn(_ELABORO, html)
        self.assertNotIn("Luis Pérez", html)
        self.assertNotIn("Otro configurado", html)
        self.assertIn("<b>Aprobó:</b><br/><br/>", html)

        ws = _excel(exportada)
        self.assertEqual(ws["A65"].value, _ELABORO)
        self.assertFalse(ws["H65"].value)
        self.assertEqual(ws["H64"].value, "Aprobó:")
        self.assertEqual(ws["H66"].value, "Topografo Interventoria")

    def test_sin_reporte_vigente(self):
        p = _planilla(
            nivel1_estado="Aprobado",
            nivel1_usuario_id=10,
            nivel1_fecha=_N1,
            nivel2_estado="Aprobado",
            nivel2_usuario_id=20,
            nivel2_fecha=_N2,
            meta_cabecera={
                "sicoe_reportes": [{"reporte_id": 4, "numero_reporte": 99}],
            },
        )
        links = conservar_links_sicoe_vigentes(
            p["meta_cabecera"]["sicoe_reportes"],
            [],
        )
        self.assertEqual(links, [])
        for ext in ("pdf", "xlsx"):
            nombre = nombre_archivo_planilla_tuberia(links, extension=ext)
            self.assertEqual(nombre, f"planilla_tuberia.{ext}")
            self.assertNotIn("6fa2ebee", nombre)
            self.assertNotIn("99", nombre)

        exportada = aplicar_firmas_validacion_export(p, _NOMBRES)
        html = html_pie_firmas_planilla(
            exportada["firmas"]["elaboro_nombre"],
            exportada["firmas"]["aprobo_nombre"],
        )
        self.assertIn(_ELABORO, html)
        self.assertIn(_APROBO, html)
        ws = _excel(exportada)
        self.assertEqual(ws["A65"].value, _ELABORO)
        self.assertEqual(ws["H65"].value, _APROBO)

    def test_escapa_html_y_varios_reportes(self):
        html = html_pie_firmas_planilla('Ana & "Gómez" <obra>', "")
        self.assertIn("Ana &amp; &quot;Gómez&quot; &lt;obra&gt;", html)
        self.assertNotIn("<obra>", html)
        links = [
            {"reporte_id": 1, "numero_reporte": 12},
            {"reporte_id": 2, "numero_reporte": 12},
            {"reporte_id": 3, "numero_reporte": 15},
        ]
        self.assertEqual(
            nombre_archivo_planilla_tuberia(links, extension="xlsx", plantilla=True),
            "planilla_tuberia_plantilla_12_15.xlsx",
        )

    def test_la_ruta_usa_el_reporte_vigente_y_no_el_id(self):
        src = Path(__file__).resolve().parents[1].joinpath(
            "topografia_planilla_tuberia_routes.py"
        ).read_text(encoding="utf-8")
        self.assertNotIn("planilla_id[:8]", src)
        self.assertIn("aplicar_firmas_validacion_export", src)
        self.assertIn("html_pie_firmas_planilla", src)
        self.assertIn("nombre_archivo_planilla_tuberia", src)
        self.assertIn("_filtrar_sicoe_reportes_vigentes", src)
        self.assertIn("conservar_links_sicoe_vigentes", src)


if __name__ == "__main__":
    unittest.main()
