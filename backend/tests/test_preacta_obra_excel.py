"""Tabla preacta_obra del Excel integral mensual: nombre fijo, una fila por ítem, fórmulas."""
from __future__ import annotations

import sys
import zipfile
from io import BytesIO
from types import ModuleType, SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from corte_sub_conciliacion import build_resumen_conciliacion_4cols, valor_por_cantidad_vu
from excel_formula_eval import FormulaBook


HEADERS = (
    "Clave",
    "Capitulo",
    "Item",
    "Descripcion",
    "Unidad",
    "VrUnitario",
    "ActualizadaCant",
    "ActualizadaValor",
    "AnteriorCant",
    "AnteriorValor",
    "PresenteCant",
    "PresenteValor",
    "AcumuladoCant",
    "AcumuladoValor",
    "SaldoCant",
    "SaldoValor",
)


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]

    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake

    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)

    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    return inf


def _item(cap, num, desc, und, vu, act, pres, ant):
    acum = round(float(ant) + float(pres), 2)
    saldo = round(float(act) - acum, 2)
    return {
        "capitulo": cap,
        "item_numero": num,
        "item_descripcion": desc,
        "unidad": und,
        "vlr_unitario_sub": vu,
        "vlr_unitario": vu,
        "cant_actualizadas": act,
        "cant_presente": pres,
        "cant_acum_anterior": ant,
        "cantidad": pres,
        "costo_directo": valor_por_cantidad_vu(pres, vu),
        "cant_acumulado": acum,
        "cant_saldo": saldo,
        "valor_actualizadas": valor_por_cantidad_vu(act, vu),
        "valor_presente": valor_por_cantidad_vu(pres, vu),
        "valor_acumulado": valor_por_cantidad_vu(acum, vu),
        "valor_saldo": valor_por_cantidad_vu(saldo, vu),
    }


CAP3 = "3. OBRAS DE ARTE (ALCANTARILLA)"
CAP4 = "4. PAVIMENTOS"


def _dataset():
    """Mismo número de ítem en dos capítulos, más un ítem solo con anterior."""
    return [
        _item(CAP3, "3.4.", "Alcantarilla Doble", "M", 1000, 10, 2.5, 1.5),
        _item(CAP3, "3.5.", "Cuneta", "M", 2205, 8, 0, 3),
        _item(CAP4, "3.4.", "Base Granular", "M3", 500, 4, 1.25, 0.25),
    ]


def _regs_for(items):
    regs = []
    for it in items:
        if float(it["cant_presente"]) <= 0:
            continue
        regs.append(
            {
                "numero_registro": len(regs) + 1,
                "capitulo": it["capitulo"],
                "item_numero": it["item_numero"],
                "item_descripcion": it["item_descripcion"],
                "unidad": it["unidad"],
                "longitud": it["cant_presente"],
                "ancho": "",
                "espesor": "",
                "cantidad": "",
                "cantidad_total": it["cant_presente"],
                "observacion": "",
                "abs_inicio": "",
                "abs_final": "",
                "infraestructura": "",
            }
        )
    return regs


def _snapshot(ws):
    return [
        (cell.coordinate, cell.value)
        for row in ws.iter_rows()
        for cell in row
    ]


def test_append_no_modifica_la_hoja_del_informe_y_crea_la_tabla():
    inf = _import_informes_with_stubs()
    items = _dataset()
    r4 = build_resumen_conciliacion_4cols(items=items, tributos={"administracion": 0})
    wb = Workbook()
    ws = wb.active
    ws.title = "CC-MES-001"
    filas = inf._fill_corte_sub_001_excel_ws(
        ws,
        {"numero": "ICCU-CTO-1614-2025", "contratista": "C", "nit": "1", "interventoria": "I"},
        {"razon_social": "RPO 1", "nombre_contacto": "A"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        items,
        sum(it["valor_presente"] for it in items),
        "Ana",
        "Residente",
        {},
        resumen_4cols=r4,
    )
    antes = _snapshot(ws)
    assert len(filas) == 3
    assert [f["row"] for f in filas] == [9, 10, 12]
    assert all(f["subtotal_row"] in (11, 13) for f in filas)
    inf._append_preacta_obra_sheet(wb, "CC-MES-001", filas)
    assert _snapshot(ws) == antes
    assert ws.max_column == 12
    assert "Clave" not in {c.value for row in ws.iter_rows(max_row=8) for c in row}

    buf = BytesIO()
    wb.save(buf)
    raw = buf.getvalue()
    with zipfile.ZipFile(BytesIO(raw)) as zf:
        tablas = [n for n in zf.namelist() if n.startswith("xl/tables/")]
        assert tablas
        xml = b"".join(zf.read(n) for n in tablas)
    assert b'displayName="preacta_obra"' in xml
    assert b'name="Clave"' in xml
    assert b'name="SaldoValor"' in xml

    loaded = load_workbook(BytesIO(raw))
    assert loaded.sheetnames[0] == "CC-MES-001"
    assert loaded.sheetnames[1] == "preacta_obra"
    hoja = loaded["preacta_obra"]
    tabla = hoja.tables["preacta_obra"]
    assert tabla.displayName == "preacta_obra"
    assert tuple(tabla.column_names) == HEADERS
    assert tabla.ref == "A1:P4"
    assert hoja.freeze_panes == "A2"
    # Sin fila de totales de la tabla ni filas vacías bajo los ítems.
    assert tabla.totalsRowCount is None
    for col in range(1, 17):
        assert hoja.cell(5, col).value is None


def test_valores_coinciden_con_informe_y_pdf_sin_subtotales():
    inf = _import_informes_with_stubs()
    items = _dataset()
    r4 = build_resumen_conciliacion_4cols(items=items)
    wb = Workbook()
    ws = wb.active
    ws.title = "CC-MES-001"
    filas = inf._fill_corte_sub_001_excel_ws(
        ws,
        {"numero": "ICCU-CTO-1614-2025", "contratista": "C", "nit": "1", "interventoria": "I"},
        {"razon_social": "RPO 1", "nombre_contacto": "A"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        items,
        sum(it["valor_presente"] for it in items),
        "Ana",
        "Residente",
        {},
        resumen_4cols=r4,
    )
    inf._append_preacta_obra_sheet(wb, "CC-MES-001", filas)
    html = inf._html_cc_mes_001_v1(
        {"numero": "ICCU-CTO-1614-2025", "contratista": "C", "nit": "1", "interventoria": "I"},
        {"numero_rpo": "1", "consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"},
        items,
        sum(it["valor_presente"] for it in items),
        "Ana",
        "Residente",
        {},
        c3_value="1",
        c4_value="1",
    )
    book = FormulaBook(wb)
    hoja = wb["preacta_obra"]
    claves = []
    for i, it in enumerate(items):
        row = 2 + i
        got = {HEADERS[c - 1]: book.eval_cell("preacta_obra", row, c) for c in range(1, 17)}
        clave = f"{it['capitulo']}|{it['item_numero']}"
        claves.append(clave)
        assert got["Clave"] == clave
        assert got["Capitulo"] == it["capitulo"]
        assert got["Item"] == it["item_numero"]
        assert got["Descripcion"] == str(it["item_descripcion"]).lower()
        assert got["Unidad"] == it["unidad"]
        assert got["VrUnitario"] == it["vlr_unitario_sub"]
        assert got["ActualizadaCant"] == it["cant_actualizadas"]
        assert got["ActualizadaValor"] == it["valor_actualizadas"]
        assert got["AnteriorCant"] == it["cant_acum_anterior"]
        assert got["AnteriorValor"] == valor_por_cantidad_vu(it["cant_acum_anterior"], it["vlr_unitario_sub"])
        assert got["PresenteCant"] == it["cant_presente"]
        assert got["PresenteValor"] == it["valor_presente"]
        assert got["AcumuladoCant"] == round(it["cant_acum_anterior"] + it["cant_presente"], 2)
        assert got["AcumuladoValor"] == it["valor_acumulado"]
        assert got["SaldoCant"] == round(it["cant_actualizadas"] - got["AcumuladoCant"], 2)
        assert got["SaldoValor"] == it["valor_saldo"]
        # La fórmula de cada columna apunta a la fila de ítem, no a la de subtotal.
        for col in range(1, 17):
            formula = str(hoja.cell(row, col).value)
            assert formula.startswith("=")
            assert f"!A{filas[i]['subtotal_row']}" not in formula or col <= 2
            assert "Subtotal " not in got["Clave"]
        assert it["item_descripcion"].lower() in html
        assert inf._fn_cant_informe(it["cant_actualizadas"]) in html
        assert inf._fn_cant_informe(it["cant_presente"]) in html
        assert inf._fn_cant_informe(it["cant_acumulado"]) in html
        assert inf._fn_cant_informe(it["cant_saldo"]) in html
        assert inf._fm_informe(it["valor_acumulado"]) in html
        assert inf._fm_informe(it["valor_saldo"]) in html
        assert inf._fm_informe(it["valor_presente"]) in html
        assert inf._fm_informe(it["valor_actualizadas"]) in html
        assert f"Subtotal {it['capitulo']}" in html

    assert len(claves) == len(set(claves)) == 3
    # Ejemplo del requerimiento: capítulo e ítem unidos por barra vertical.
    assert claves[0] == "3. OBRAS DE ARTE (ALCANTARILLA)|3.4."
    assert claves[2] == "4. PAVIMENTOS|3.4."


def test_integral_mensual_incluye_preacta_formulada_desde_memorias():
    inf = _import_informes_with_stubs()
    items = _dataset()
    regs = _regs_for(items)
    r4 = build_resumen_conciliacion_4cols(items=items)
    contrato = {
        "numero": "ICCU-CTO-1614-2025",
        "contratista": "Ctor",
        "nit": "900",
        "interventoria": "Int",
    }
    ctx = {
        "contrato": contrato,
        "acta": {
            "numero_rpo": "1",
            "consecutivo": 1,
            "fecha_inicio": "2026-01-01",
            "fecha_fin": "2026-01-31",
        },
        "items": items,
        "total_costo": sum(it["valor_presente"] for it in items),
        "resumen_4cols": r4,
        "otros_conceptos": [],
    }
    user = {"nombre": "Ana", "apellidos": "Pérez", "cargo_nombre": "Residente"}

    def _fetch(_sb, _cid, inum, **_kw):
        return [r for r in regs if str(r.get("item_numero")) == str(inum)]

    with (
        patch.object(inf, "_acta_pertenece_contrato", return_value=True),
        patch.object(inf, "_contexto_acta_mes_conciliacion", return_value=ctx),
        patch.object(inf, "_row", return_value=contrato),
        patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
        patch.object(inf, "_sub_corte_dummy_memoria", return_value=({}, {})),
        patch.object(inf, "fetch_registros_memoria_cc_mes_alineado_acta", side_effect=_fetch),
    ):
        raw = inf._cc_mes_integral_excel_bytes(3, 620, user, nivel_aprobacion=4)

    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames[0] == "CC-MES-001"
    assert wb.sheetnames[1] == "preacta_obra"
    assert wb["preacta_obra"].tables["preacta_obra"].displayName == "preacta_obra"
    assert tuple(wb["preacta_obra"].tables["preacta_obra"].column_names) == HEADERS
    informe = wb["CC-MES-001"]
    assert informe.max_column == 12
    assert informe.cell(1, 4).value.startswith("INFORME EJECUCIÓN MENSUAL")
    assert any(
        isinstance(informe.cell(r, 1).value, str) and str(informe.cell(r, 1).value).startswith("Subtotal ")
        for r in range(9, 20)
    )
    book = FormulaBook(wb)
    claves = []
    for i, it in enumerate(items):
        got_clave = book.eval_cell("preacta_obra", 2 + i, 1)
        claves.append(got_clave)
        assert got_clave == f"{it['capitulo']}|{it['item_numero']}"
        assert book.eval_cell("preacta_obra", 2 + i, 7) == it["cant_actualizadas"]
        assert book.eval_cell("preacta_obra", 2 + i, 9) == it["cant_acum_anterior"]
        # Presente sale de la memoria (fórmula en el informe) y coincide con el ítem.
        assert book.eval_cell("preacta_obra", 2 + i, 11) == it["cant_presente"]
        assert book.eval_cell("preacta_obra", 2 + i, 12) == it["valor_presente"]
        pres_formula = str(wb["preacta_obra"].cell(2 + i, 11).value)
        assert pres_formula.startswith("='CC-MES-001'!G")
    assert len(set(claves)) == 3


def test_integral_sin_items_deja_tabla_solo_con_encabezados():
    inf = _import_informes_with_stubs()
    wb = Workbook()
    ws = wb.active
    ws.title = "CC-MES-001"
    inf._append_preacta_obra_sheet(wb, "CC-MES-001", [])
    buf = BytesIO()
    wb.save(buf)
    loaded = load_workbook(buf)
    tabla = loaded["preacta_obra"].tables["preacta_obra"]
    assert tabla.displayName == "preacta_obra"
    assert tuple(tabla.column_names) == HEADERS
    assert tabla.ref == "A1:P1"
