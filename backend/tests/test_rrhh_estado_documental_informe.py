"""Tests — informe Estado documental de colaboradores (RRHH)."""
from __future__ import annotations

from io import BytesIO

from rrhh_estado_documental_informe import (
    ESTADO_DOCUMENTAL_FORMATO,
    LEYENDA_ICONOS,
    _CELDA_FALTA,
    _CELDA_OK,
    _CELDA_OPCIONAL,
    _celda_doc,
    _columnas_checklist,
    _documento_identidad,
    _es_aprobado,
    _filtrar_trabajadores,
    _filtro_label,
    _nombre_informe,
    _slug_tipo,
    build_estado_documental_pdf,
    build_estado_documental_xlsx,
)


def test_formato_calidad_ajustable():
    assert ESTADO_DOCUMENTAL_FORMATO["codigo"] == "CC-RRHH-DOC-001"
    assert ESTADO_DOCUMENTAL_FORMATO["version"] == "1.0"
    assert "documental" in ESTADO_DOCUMENTAL_FORMATO["nombre"].lower()
    assert _CELDA_OK in LEYENDA_ICONOS
    assert _CELDA_FALTA in LEYENDA_ICONOS
    assert _CELDA_OPCIONAL in LEYENDA_ICONOS


def test_slug_tipo_otro():
    assert _slug_tipo("Carta de recomendación") == "ext_carta_de_recomendacion"
    assert _slug_tipo("  ").startswith("ext_")


def test_filtrar_trabajadores_tres_opciones():
    rows = [
        {"id": 1, "doc_validacion_estado": "aprobado"},
        {"id": 2, "doc_validacion_estado": "pendiente"},
        {"id": 3, "doc_validacion_estado": "rechazado"},
        {"id": 4, "doc_validacion_estado": "Aprobado"},
    ]
    assert [r["id"] for r in _filtrar_trabajadores(rows, "todos")] == [1, 2, 3, 4]
    assert [r["id"] for r in _filtrar_trabajadores(rows, "aprobado")] == [1, 4]
    assert [r["id"] for r in _filtrar_trabajadores(rows, "sin_aprobar")] == [2, 3]
    assert _filtro_label("sin_aprobar") == "Personal sin aprobar"


def test_nombre_corto_informe():
    t = {
        "nombres": "Carlos Andrés",
        "apellidos": "Pérez Gómez",
        "tipo_documento": "CC",
        "numero_documento": "123",
        "doc_validacion_estado": "aprobado",
    }
    assert _nombre_informe(t) == "Carlos Pérez"
    assert _documento_identidad(t) == "CC 123"
    assert _es_aprobado(t) is True


def test_columnas_checklist_incluye_custom_otro_sin_slot_ui():
    custom = [
        ("soporte", "ext_carta", "Carta laboral"),
        ("ingreso", "ext_carnet", "Carnet obra"),
    ]
    cols = _columnas_checklist(custom)
    keys = [c["key"] for c in cols]
    assert keys[0] == "contrato_laboral"
    assert "soporte:cedula" in keys
    assert "soporte:ext_carta" in keys
    assert "ingreso:ext_carnet" in keys
    assert "afiliacion:cert_eps" in keys
    assert not any(c["tipo"] == "otro" for c in cols)


def test_celda_iconos_obligatorio_y_opcional():
    trab = {"eps": "", "pension": "", "arl": "", "cesantias": "", "caja_compensacion": ""}
    col_contrato = {"categoria": "contrato_laboral", "tipo": "contrato_laboral"}
    col_cedula = {"categoria": "soporte", "tipo": "cedula"}
    col_hv = {"categoria": "soporte", "tipo": "hoja_vida"}
    col_eps = {"categoria": "afiliacion", "tipo": "cert_eps"}
    col_caja = {"categoria": "afiliacion", "tipo": "cert_caja"}

    assert _celda_doc(trab, col_contrato, vigentes=set(), tiene_contrato=True, aprobado=False) == _CELDA_OK
    assert _celda_doc(trab, col_contrato, vigentes=set(), tiene_contrato=False, aprobado=False) == _CELDA_FALTA
    assert _celda_doc(
        trab, col_cedula, vigentes={("soporte", "cedula")}, tiene_contrato=False, aprobado=False,
    ) == _CELDA_OK
    assert _celda_doc(trab, col_cedula, vigentes=set(), tiene_contrato=False, aprobado=False) == _CELDA_FALTA
    # Opcional vacío → 🔘 (nunca ❌)
    assert _celda_doc(trab, col_hv, vigentes=set(), tiene_contrato=False, aprobado=False) == _CELDA_OPCIONAL
    assert _celda_doc(trab, col_caja, vigentes=set(), tiene_contrato=False, aprobado=False) == _CELDA_OPCIONAL
    # EPS sin entidad y sin archivo → ❌ (nunca N/A / 🔘)
    assert _celda_doc(trab, col_eps, vigentes=set(), tiene_contrato=False, aprobado=False) == _CELDA_FALTA
    # Aprobado → ✅
    assert _celda_doc(trab, col_cedula, vigentes=set(), tiene_contrato=False, aprobado=True) == _CELDA_OK


def _sample_matrix():
    return {
        "formato": dict(ESTADO_DOCUMENTAL_FORMATO),
        "filtro": "todos",
        "fecha_calidad": "15/10/2026",
        "leyenda": LEYENDA_ICONOS,
        "contrato": {
            "numero": "C-001",
            "contratista": "Consorcio Demo",
            "nit": "900",
            "objeto": "Obra demo",
        },
        "identity_columns": [
            {"key": "nombre", "label": "Nombre"},
            {"key": "empresa", "label": "Empresa"},
            {"key": "cargo", "label": "Cargo"},
            {"key": "documento", "label": "Documento de identidad"},
        ],
        "doc_columns": [
            {"key": "contrato_laboral", "label": "Contrato laboral", "categoria": "contrato_laboral", "tipo": "contrato_laboral"},
            {"key": "soporte:cedula", "label": "Cédula", "categoria": "soporte", "tipo": "cedula"},
            {"key": "soporte:hoja_vida", "label": "Hoja de vida", "categoria": "soporte", "tipo": "hoja_vida"},
        ],
        "rows": [
            {
                "id": 1,
                "nombre": "Ana López",
                "empresa": "CD",
                "cargo": "Auxiliar",
                "documento": "CC 1",
                "celdas": {
                    "contrato_laboral": _CELDA_OK,
                    "soporte:cedula": _CELDA_FALTA,
                    "soporte:hoja_vida": _CELDA_OPCIONAL,
                },
            },
        ],
        "total": 1,
    }


def test_xlsx_iconos_leyenda_anchos():
    data, fname = build_estado_documental_xlsx(_sample_matrix())
    assert fname.startswith("CC-RRHH-DOC-001_")
    assert fname.endswith(".xlsx")
    assert data[:2] == b"PK"

    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert "ESTADO DOCUMENTAL" in str(ws.cell(1, 2).value or "").upper()
    cal = str(ws.cell(1, 5).value or "")
    assert "CC-RRHH-DOC-001" in cal
    assert _CELDA_OK in str(ws.cell(5, 2).value or "")
    assert ws.cell(7, 1).value == "Nombre"
    assert ws.cell(7, 2).value == "Empresa"
    assert ws.cell(8, 1).value == "Ana López"
    assert ws.cell(8, 2).value == "CD"
    assert ws.cell(8, 5).value == _CELDA_OK
    assert ws.cell(8, 6).value == _CELDA_FALTA
    assert ws.cell(8, 7).value == _CELDA_OPCIONAL
    # Nombre más ancho que columnas de documento
    assert ws.column_dimensions["A"].width >= 30
    assert ws.column_dimensions[get_column_letter(5)].width <= 8


def test_pdf_genera_bytes_con_leyenda():
    data, fname = build_estado_documental_pdf(_sample_matrix())
    assert fname.endswith(".pdf")
    assert data[:4] == b"%PDF"
    assert len(data) > 500
