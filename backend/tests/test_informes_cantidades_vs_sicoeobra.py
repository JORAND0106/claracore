"""
Evidencia de la causa raíz: re-redondeo del agregado vs suma de registros SicoeObra.

Cuadro comparativo (ejemplo sintético, misma regla que ICCU/ME Obras en producción):

| Ítem | Registros aprobados (cantidad_total) | SicoeObra sumCant | Informe ANTES (re-redondeo) | Informe DESPUÉS |
|------|--------------------------------------|-------------------|-----------------------------|-----------------|
| 1.1  | 0.044 + 0.044 + 0.044                | 0.132             | 0.13                        | 0.132           |
| 2.1  | 0.055 + 0.055                        | 0.110             | 0.11                        | 0.110           |
| 3.1  | 1.255 + 1.255                        | 2.510             | 2.51                        | 2.510           |

Causa en código (antes):
- SicoeObra: ``agruparRegistrosPorItem`` suma ``cantidad_total`` sin re-redondear
  (frontend/src/modules/sicoe-obra/sicoeReporteItemsTablaHelpers.js).
- Informes: ``enriquecer_items_bloques`` + ``cantidades_por_item_*`` aplicaban
  ``redondear_cantidad_total_dinamico`` sobre la suma (2 dp si ≥0.10).

Filtro de aprobación (sin cambio de eje):
- CC-MES: cascada matriz ``_registro_aprobado_matriz_panel`` (igual KPI SicoeObra).
- CC-SUB: ``sub_estado=Aprobado`` (eje subcontratista; distinto de capas nivel*).
"""
from __future__ import annotations

from ccd_conciliacion import aggregate_items_conciliacion
from corte_sub_conciliacion import (
    enriquecer_items_bloques,
    filtrar_items_con_cantidades,
    valor_por_cantidad_vu,
)
from sicoe_cantidad_redondeo import redondear_cantidad_total_dinamico


def test_suma_registros_no_se_re_redondea_como_sicoeobra_sumcant():
    """0.044×3 → SicoeObra 0.132; re-redondeo dinámico daba 0.13."""
    regs = [
        {"item_numero": "1.1", "cantidad_total": 0.044, "vlr_unitario": 1000, "costo_directo": 44},
        {"item_numero": "1.1", "cantidad_total": 0.044, "vlr_unitario": 1000, "costo_directo": 44},
        {"item_numero": "1.1", "cantidad_total": 0.044, "vlr_unitario": 1000, "costo_directo": 44},
    ]
    items, _ = aggregate_items_conciliacion(regs)
    assert abs(items[0]["cantidad"] - 0.132) < 1e-9
    # Evidencia de la divergencia histórica
    assert redondear_cantidad_total_dinamico(0.132) == 0.13

    enriched = enriquecer_items_bloques(
        items,
        cant_actualizadas={"1.1": 10.0},
        cant_acum_anterior={},
        vu_por_item={"1.1": 1000.0},
    )
    assert abs(enriched[0]["cant_presente"] - 0.132) < 1e-9
    assert abs(enriched[0]["cant_acumulado"] - 0.132) < 1e-9
    assert enriched[0]["valor_presente"] == valor_por_cantidad_vu(0.132, 1000)


def test_aggregate_normaliza_espacios_item_como_sicoeobra():
    regs = [
        {"item_numero": " 1.1 ", "cantidad_total": 2.0, "vlr_unitario": 100, "costo_directo": 200},
        {"item_numero": "1.1", "cantidad_total": 3.0, "vlr_unitario": 100, "costo_directo": 300},
    ]
    items, _ = aggregate_items_conciliacion(regs)
    assert len(items) == 1
    assert items[0]["item_numero"] == "1.1"
    assert items[0]["cantidad"] == 5.0


def test_filtrar_incluye_solo_actualizadas():
    items = [
        {"item_numero": "a", "cant_presente": 0, "cant_acumulado": 0, "cant_actualizadas": 12},
        {"item_numero": "b", "cant_presente": 1, "cant_acumulado": 0, "cant_actualizadas": 0},
        {"item_numero": "c", "cant_presente": 0, "cant_acumulado": 0, "cant_actualizadas": 0},
    ]
    out = filtrar_items_con_cantidades(items)
    assert {i["item_numero"] for i in out} == {"a", "b"}


def test_fuente_sin_fila_cd_presente_acta_bajo_cuadro():
    from pathlib import Path

    text = (Path(__file__).resolve().parents[1] / "informes.py").read_text(encoding="utf-8")
    assert "COSTO DIRECTO (PRESENTE ACTA):" not in text
    assert "RESUMEN DE CONCILIACIÓN" in text


def test_excel_cd_resumen_suma_items_sin_fila_intermedia():
    from io import BytesIO

    from openpyxl import load_workbook
    from corte_sub_conciliacion import build_resumen_conciliacion_4cols
    from test_corte_sub_001_excel_filtro import _import_informes_with_stubs

    inf = _import_informes_with_stubs()
    items = [
        {
            "capitulo": "I",
            "item_numero": "1.1",
            "item_descripcion": "Exc",
            "unidad": "M3",
            "vlr_unitario_sub": 1000.0,
            "cant_actualizadas": 100.0,
            "cant_presente": 10.0,
            "cant_acum_anterior": 0.0,
            "cant_acumulado": 10.0,
            "cant_saldo": 90.0,
            "valor_presente": 10000.0,
            "valor_actualizadas": 100000.0,
            "valor_acumulado": 10000.0,
            "valor_saldo": 90000.0,
            "costo_directo": 10000.0,
        }
    ]
    r4 = build_resumen_conciliacion_4cols(items=items, tributos={"administracion": 10}, otros_presente=0)
    raw = inf._corte_sub_001_excel_bytes(
        {"numero": "C-1", "contratista": "Ctor", "nit": "900", "interventoria": "Int"},
        {"razon_social": "Sub SA", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-15"},
        items,
        10000.0,
        "Tester",
        "Desarrollador",
        {},
        resumen_4cols=r4,
    )
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    labels = []
    cd_resumen_formula = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=12):
        a = row[0].value
        if isinstance(a, str):
            labels.append(a)
            # Línea Costo Directo del resumen (exacta; no "Costo Directo + AIU")
            if a.strip() == "Costo Directo":
                cd_resumen_formula = row[5].value  # col F (valor Actualizadas)
    assert "COSTO DIRECTO:" not in labels
    assert any(x == "RESUMEN DE CONCILIACIÓN" for x in labels)
    assert cd_resumen_formula is not None
    assert str(cd_resumen_formula).startswith("=")
    assert "F" in str(cd_resumen_formula)
