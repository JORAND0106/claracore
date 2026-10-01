"""Informe Excel de Logs del Sistema: encabezado, usuario filtrado, grilla y BarChart."""
from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from logs_export_xlsx import (
    GRID_COLUMNS,
    aggregate_activity_by_day,
    build_filtros_resumen,
    build_logs_informe_xlsx,
    day_key_bogota,
    format_log_fecha_bogota,
)


def _row(**overrides):
    base = {
        "created_at": "2026-03-15T20:00:00Z",
        "usuario_nombre": "Ana Pérez",
        "cargo_nombre": "Ingeniera",
        "rol_nombre": "Contratista",
        "modulo": "SICOE",
        "accion": "VALIDAR",
        "severidad": "AUDIT",
        "entidad_tipo": "registro",
        "entidad_id": 10,
        "contrato_numero": "IDU-1551",
        "ip": "10.0.0.1",
        "resultado": "ok",
    }
    base.update(overrides)
    return base


class TestLogsExportXlsx(unittest.TestCase):
    def test_aggregate_activity_by_day_bogota(self):
        rows = [
            _row(created_at="2026-03-15T20:00:00Z"),  # 15:00 COT
            _row(created_at="2026-03-16T04:30:00Z"),  # 23:30 COT del 15
            _row(created_at="2026-03-16T15:00:00Z"),  # 10:00 COT del 16
        ]
        series = aggregate_activity_by_day(rows)
        self.assertEqual(series, [("2026-03-15", 2), ("2026-03-16", 1)])
        self.assertEqual(day_key_bogota("2026-03-15T20:00:00Z"), "2026-03-15")

    def test_format_fecha_bogota(self):
        txt = format_log_fecha_bogota("2026-03-15T20:00:00Z")
        self.assertIn("15/03/2026", txt)
        self.assertNotEqual(txt, "—")

    def test_filtros_resumen(self):
        s = build_filtros_resumen(
            usuario_nombre="Ana",
            accion="VALIDAR",
            excluir_rutina_auth=True,
        )
        self.assertIn("Usuario: Ana", s)
        self.assertIn("Acción: VALIDAR", s)
        # Con acción explícita no se anuncia exclusión de LOGIN
        self.assertNotIn("LOGIN", s)

        s2 = build_filtros_resumen(excluir_rutina_auth=True)
        self.assertIn("Sin LOGIN/LOGIN_FAIL", s2)

    def test_workbook_has_header_user_grid_and_barchart(self):
        rows = [
            _row(),
            _row(created_at="2026-03-15T21:00:00Z", entidad_id=11),
            _row(created_at="2026-03-17T12:00:00Z", accion="EDITAR", entidad_id=12),
        ]
        blob = build_logs_informe_xlsx(
            rows,
            contrato_meta={
                "numero": "IDU-1551",
                "contratista": "Constructora X",
                "interventoria": "Interventoría Y",
                "objeto": "Mantenimiento vial",
            },
            usuario_filtrado={
                "id": 42,
                "nombre": "Ana Pérez",
                "cargo": "Ingeniera",
                "rol": "Contratista",
                "email": "ana@example.com",
                "contrato_numero": "IDU-1551",
            },
            filtros_resumen="Usuario: Ana Pérez · Acción: VALIDAR",
            accion_filtro="VALIDAR",
            descargado_por="Admin Test",
            gen_ts="01/04/2026 10:00 AM",
        )
        self.assertGreater(len(blob), 2000)
        wb = load_workbook(io.BytesIO(blob))
        ws = wb.active
        self.assertEqual(ws.title, "Logs")
        self.assertEqual(ws["B2"].value, "IDU-1551")
        self.assertEqual(ws["B3"].value, "Constructora X")

        labels = [ws.cell(row=r, column=1).value for r in range(1, 40)]
        self.assertIn("Usuario filtrado", labels)
        self.assertTrue(any(isinstance(v, str) and v.startswith("Registros filtrados (3)") for v in labels))
        self.assertTrue(
            any(isinstance(v, str) and v.startswith("Actividad por fecha") for v in labels)
        )

        # Encabezados de grilla
        found_hdr = False
        for r in range(1, 40):
            if ws.cell(row=r, column=1).value == GRID_COLUMNS[0][1]:
                for ci, (_k, label) in enumerate(GRID_COLUMNS, start=1):
                    self.assertEqual(ws.cell(row=r, column=ci).value, label)
                found_hdr = True
                break
        self.assertTrue(found_hdr)

        self.assertEqual(len(ws._charts), 1)
        chart = ws._charts[0]
        self.assertEqual(chart.type, "col")
        self.assertIn("VALIDAR", str(chart.title))

    def test_without_user_filter_omits_user_block(self):
        blob = build_logs_informe_xlsx(
            [_row()],
            contrato_meta={"numero": "C-1", "contratista": "A", "interventoria": "", "objeto": ""},
            usuario_filtrado=None,
            filtros_resumen="Módulo: SICOE",
            accion_filtro=None,
            descargado_por="Dev",
        )
        wb = load_workbook(io.BytesIO(blob))
        ws = wb.active
        labels = [ws.cell(row=r, column=1).value for r in range(1, 30)]
        self.assertNotIn("Usuario filtrado", labels)
        self.assertTrue(any(isinstance(v, str) and "todas las acciones" in v for v in labels))
        self.assertEqual(len(ws._charts), 1)

    def test_empty_rows_still_builds_chart_placeholder(self):
        blob = build_logs_informe_xlsx(
            [],
            contrato_meta={"numero": "C-1"},
            filtros_resumen="Sin filtros",
        )
        wb = load_workbook(io.BytesIO(blob))
        ws = wb.active
        labels = [ws.cell(row=r, column=1).value for r in range(1, 40)]
        self.assertTrue(any(isinstance(v, str) and v.startswith("Registros filtrados (0)") for v in labels))
        # Sin series no se crea gráfico OOXML (solo mensaje)
        self.assertEqual(len(ws._charts), 0)


class TestLogsExportRouteWired(unittest.TestCase):
    def test_main_declares_xlsx_route(self):
        from pathlib import Path

        src = Path(__file__).resolve().parents[1] / "main.py"
        text = src.read_text(encoding="utf-8")
        self.assertIn('@app.get("/logs/export.xlsx")', text)
        self.assertIn("build_logs_informe_xlsx", text)
        self.assertIn("export_logs_xlsx", text)


if __name__ == "__main__":
    unittest.main()
