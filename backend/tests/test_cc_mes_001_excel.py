"""CC-MES-001 Excel: ruta y builder equivalentes a CC-SEM-001 / PDF mensual."""
from __future__ import annotations

import sys
from io import BytesIO
from types import ModuleType, SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]

    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake

    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)

    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    return inf


def test_cc_mes_001_excel_bytes_builds_workbook_with_items():
    """El Excel mensual reutiliza _fill_cc_conc_001_excel_ws con datos del acta."""
    inf = _import_informes_with_stubs()

    contrato = {
        "numero": "C-1",
        "contratista": "Ctor",
        "nit": "900",
        "interventoria": "Int",
    }
    items = [
        {
            "capitulo": "1",
            "item_numero": "1.01",
            "item_descripcion": "Excavación",
            "unidad": "m3",
            "vlr_unitario": 1000.0,
            "cantidad": 2.0,
            "costo_directo": 2000.0,
        }
    ]
    user = {"nombre": "Ana", "apellidos": "Pérez", "cargo_nombre": "Residente"}

    with (
        patch.object(inf, "_acta_pertenece_contrato", return_value=True),
        patch.object(
            inf,
            "fetch_registros_informe_cc_mes_por_acta",
            return_value=[{"item_numero": "1.01"}],
        ),
        patch.object(inf, "aggregate_items_conciliacion", return_value=(items, 2000.0)),
        patch.object(inf, "_sort_items_corte_por_item_numero_asc"),
        patch.object(
            inf,
            "_row",
            side_effect=lambda table, *_a, **_k: (
                {"id": 10, "numero_rpo": "RPO-7", "consecutivo": "3"}
                if table == "actas"
                else contrato
            ),
        ),
        patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
    ):
        raw = inf._cc_mes_001_excel_bytes(1, 10, user, nivel_aprobacion=3)

    assert isinstance(raw, (bytes, bytearray))
    assert len(raw) > 100
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    assert ws.title == "CC-MES-001"
    assert "MENSUAL" in str(ws.cell(1, 1).value or "").upper()
    assert "CC-MES-001" in str(ws.cell(2, 1).value or "")


def test_excel_cc_mes_001_route_registered():
    """La ruta Excel mensual debe existir (antes solo estaba el botón disabled en UI)."""
    inf = _import_informes_with_stubs()
    paths = [getattr(r, "path", "") for r in inf.router.routes]
    assert any(
        "excel/cc-mes-001/acta" in p for p in paths
    ), f"Falta ruta excel/cc-mes-001; rutas: {[p for p in paths if 'cc-mes' in p]}"


def test_formatos_ccd_sem_mes_001_acceso_interventoria():
    """SEM/MES-001 quedan habilitados para interventoría en el catálogo CCD."""
    inf = _import_informes_with_stubs()
    sem = inf.FORMATOS_CCD[inf.CODIGO_FORMATO_CCD_CC_SEM_001]
    mes = inf.FORMATOS_CCD[inf.CODIGO_FORMATO_CCD_CC_MES_001]
    assert sem.get("acceso_interventoria") is True
    assert mes.get("acceso_interventoria") is True
    assert "ejecución semanal" in (sem.get("titulo") or "").lower()
    assert "ejecución mensual" in (mes.get("titulo") or "").lower()


def test_ui_no_longer_hard_disables_mes_001_excel():
    """Regresión: el botón Excel de CC-MES-001 no debe quedar disabled permanente."""
    from pathlib import Path

    src = Path(__file__).resolve().parents[2] / "frontend" / "src" / "ModuloInformes.jsx"
    text = src.read_text(encoding="utf-8")
    assert "Exportación Excel aún no disponible para el informe mensual" not in text
    assert "descargarExcelCcMes001" in text
    assert "excel/cc-mes-001/acta" in text
    assert "Informe ejecución semanal (CC-SEM-001)" in text
