"""Excel planilla tubería: oculta líneas con cantidad 0 en Resumen/Descuentos."""
from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from topografia_planilla_tuberia_excel import (
    _descuentos_visibles_excel,
    _excel_cantidad_no_cero,
    _netos_visibles_excel,
    build_planilla_tuberia_xlsx,
)


def _ws_con_calculo(calculo: dict, tipo: str = "ALCANTARILLA"):
    raw = build_planilla_tuberia_xlsx(
        planilla={
            "tipo": tipo,
            "diametro_m": 0.9,
            "espesor_m": 0.05,
            "ancho_excavacion_m": 1.5,
            "relacion_atraque": "1:3",
            "meta_cabecera": {"cama_triturado_m": 0.1},
        },
        calculo=calculo,
        vacia=False,
    )
    return load_workbook(io.BytesIO(raw))["planilla"]


class TestExcelOcultarCeros(unittest.TestCase):
    def test_helpers_filtran_cero(self):
        self.assertFalse(_excel_cantidad_no_cero(0))
        self.assertFalse(_excel_cantidad_no_cero(0.0))
        self.assertFalse(_excel_cantidad_no_cero(None))
        self.assertTrue(_excel_cantidad_no_cero(0.01))
        calc = {
            "netos": [
                {"codigo": "EXC", "nombre": "Excavación Varias", "neto": 12.5, "unidad": "m³"},
                {"codigo": "EXC_ROC", "nombre": "Excavación Roca", "neto": 0, "unidad": "m³"},
                {"codigo": "OTROS_1", "nombre": "Otros: ____", "neto": 0, "unidad": "m³"},
                {"codigo": "TUB", "nombre": "Long Tubería", "neto": 10, "unidad": "ml"},
            ],
            "descuentos": [
                {"codigo": "DESC_A1", "nombre": "Area 1", "cantidad": 1.2},
                {"codigo": "DESC_A2", "nombre": "Area 2", "cantidad": 0},
                {"codigo": "DESC_OTROS", "nombre": "Otros", "cantidad": 0},
            ],
        }
        netos = _netos_visibles_excel(calc)
        self.assertEqual([n["codigo"] for n in netos], ["EXC", "TUB"])
        descs = _descuentos_visibles_excel(calc)
        self.assertEqual([d["codigo"] for d in descs], ["DESC_A1"])

    def test_excel_no_incluye_lineas_en_cero(self):
        calc = {
            "cartera": {"filas": []},
            "netos": [
                {"codigo": "EXC", "nombre": "Excavación Varias", "neto": 5.0, "unidad": "m³"},
                {"codigo": "TUB", "nombre": "Long Tubería", "neto": 0, "unidad": "ml"},
                {"codigo": "TRI", "nombre": "Atraque mat. filtrante", "neto": 2.0, "unidad": "m³"},
                {"codigo": "REL", "nombre": "Relleno Gran.", "neto": 0, "unidad": "m³"},
                {"codigo": "GEO", "nombre": "Geotextil", "neto": 0, "unidad": "m²"},
                {"codigo": "EXC_ROC", "nombre": "Excavación Roca", "neto": 0, "unidad": "m³"},
                {"codigo": "OTROS_1", "nombre": "Otros: ____", "neto": 0, "unidad": "m³"},
            ],
            "descuentos": [
                {"codigo": "DESC_A1", "nombre": "Area 1", "cantidad": 0.5, "unidad": "m³"},
                {"codigo": "DESC_A2", "nombre": "Area 2", "cantidad": 0, "unidad": "m³"},
            ],
        }
        ws = _ws_con_calculo(calc)
        # Compactado: EXC, TRI en resumen; Area 1 en descuentos
        self.assertEqual(ws["B45"].value, "Excavación Varias")
        self.assertEqual(ws["B46"].value, "Atraque mat. filtrante")
        self.assertIn(ws["B47"].value, (None, ""))
        self.assertNotEqual(ws["B50"].value, "Excavación Roca")
        self.assertNotEqual(ws["B51"].value, "Otros: ____")
        self.assertEqual(ws["I45"].value, "Area 1")
        self.assertIn(ws["I46"].value, (None, ""))
        # Fórmula viva de TRI apunta a N45 (Area1 compactada)
        self.assertEqual(ws["G46"].value, "=N45")
        self.assertIn("PRODUCT(D46:F46)", str(ws["H46"].value))
        self.assertIn("-G46", str(ws["H46"].value))

    def test_plantilla_vacia_conserva_filas_fijas(self):
        raw = build_planilla_tuberia_xlsx(
            planilla={"tipo": "ALCANTARILLA", "meta_cabecera": {}},
            calculo={"cartera": {"filas": []}},
            vacia=True,
        )
        ws = load_workbook(io.BytesIO(raw))["planilla"]
        self.assertEqual(ws["B50"].value, "Excavación Roca")
        self.assertEqual(ws["B47"].value, "Atraque mat. filtrante")
        self.assertEqual(ws["I46"].value, "Area 1")
        self.assertEqual(ws["G47"].value, "=N46")

        raw_fil = build_planilla_tuberia_xlsx(
            planilla={"tipo": "FILTRO", "meta_cabecera": {}},
            calculo={"cartera": {"filas": []}},
            vacia=True,
        )
        ws_fil = load_workbook(io.BytesIO(raw_fil))["planilla"]
        self.assertEqual(ws_fil["B47"].value, "Mat. Granular Filtrante")


if __name__ == "__main__":
    unittest.main()
