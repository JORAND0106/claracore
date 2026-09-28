"""Memorias: abscisas K0+000.00, orden, subtotales por tramo y unidad en totales."""
from __future__ import annotations

import sys
from io import BytesIO
from types import ModuleType, SimpleNamespace

from openpyxl import Workbook, load_workbook


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]

    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake

    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)

    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    inf._memoria_prepare_media_uri = lambda url: ""  # type: ignore
    inf._memoria_prefetch_media_uris = lambda urls: {
        str(u or "").strip(): "" for u in urls if str(u or "").strip()
    }  # type: ignore
    return inf


def test_fmt_abscisa_k_casos():
    inf = _import_informes_with_stubs()
    assert inf._fmt_abscisa_k(11090.0) == "K11+090.00"
    assert inf._fmt_abscisa_k(1314.5) == "K1+314.50"
    assert inf._fmt_abscisa_k("1+000") == "K1+000.00"
    assert inf._fmt_abscisa_k("K2+015.5") == "K2+015.50"
    assert inf._fmt_abscisa_k(None) == ""
    assert "K11+090.00" in inf._memoria_abscisas_txt(
        {"abs_inicio": 11090.0, "abs_final": 11200.0}
    )
    assert "K11+200.00" in inf._memoria_abscisas_txt(
        {"abs_inicio": 11090.0, "abs_final": 11200.0}
    )


def test_plan_filas_ordena_tramos_y_abscisas():
    inf = _import_informes_with_stubs()
    regs = [
        {
            "numero_registro": 3,
            "tramo": "7",
            "abs_inicio": 11200,
            "abs_final": 11250,
            "cantidad_total": 5,
        },
        {
            "numero_registro": 1,
            "tramo": "5",
            "abs_inicio": 12000,
            "abs_final": 12010,
            "cantidad_total": 2,
        },
        {
            "numero_registro": 2,
            "tramo": "7",
            "abs_inicio": 11090,
            "abs_final": 11100,
            "cantidad_total": 10,
        },
        {
            "numero_registro": 4,
            "tramo": "5",
            "abs_inicio": 12000,
            "abs_final": 12005,
            "cantidad_total": 3,
        },
    ]
    plan = inf._memoria_plan_filas_detalle(regs)
    kinds = [p["kind"] for p in plan]
    assert kinds.count("reg") == 4
    assert kinds.count("subtramo") == 2
    # Tramo 7 tiene abs mín. 11090 < tramo 5 (12000) → primero
    assert plan[0]["kind"] == "reg" and plan[0]["reg"]["numero_registro"] == 2
    assert plan[1]["reg"]["numero_registro"] == 3
    assert plan[2]["kind"] == "subtramo" and plan[2]["tramo"] == "7"
    # Dentro de tramo 5: misma abs_inicio → abs_final menor primero (12005 antes que 12010)
    assert plan[3]["reg"]["numero_registro"] == 4
    assert plan[4]["reg"]["numero_registro"] == 1
    assert plan[5]["kind"] == "subtramo" and plan[5]["tramo"] == "5"


def test_excel_subtotales_formulas_y_unidad_sin_duplicar():
    inf = _import_informes_with_stubs()
    wb = Workbook()
    ws = wb.active
    regs = [
        {
            "numero_registro": 1,
            "tramo": "7",
            "abs_inicio": 11200.0,
            "abs_final": 11250.0,
            "longitud": 2,
            "ancho": 1,
            "espesor": 1,
            "cantidad": None,
            "cantidad_total": 10.0,
            "observacion": "",
            "item_numero": "1.1",
            "item_descripcion": "Excavación",
            "unidad": "M3",
            "capitulo": "I",
        },
        {
            "numero_registro": 2,
            "tramo": "7",
            "abs_inicio": 11090.0,
            "abs_final": 11100.0,
            "longitud": 3,
            "ancho": 1,
            "espesor": 1,
            "cantidad": None,
            "cantidad_total": 26.16,
            "observacion": "",
            "item_numero": "1.1",
            "unidad": "M3",
            "capitulo": "I",
        },
        {
            "numero_registro": 3,
            "tramo": "5",
            "abs_inicio": 1314.5,
            "abs_final": 1316.0,
            "longitud": 1,
            "ancho": 1,
            "espesor": 1,
            "cantidad": None,
            "cantidad_total": 4.0,
            "observacion": "",
            "item_numero": "1.1",
            "unidad": "M3",
            "capitulo": "I",
        },
    ]
    item_info = inf._item_info_desde_registros(regs, "1.1")
    tot_r = inf._fill_memoria_excel_ws(
        ws,
        {"numero": "C-1"},
        {"razon_social": "Sub", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-07"},
        item_info,
        regs,
        None,
        conc_meta={
            "titulo": "MEMORIA SEMANAL",
            "codigo": "CC-SEM-002",
            "cells": [("CONTRATO", "C-1"), ("SEMANA", "1"), ("VIGENCIA", "—"), ("REFERENCIA", "—")],
        },
        aprobo_interventoria_desde_config=True,
    )
    buf = BytesIO()
    wb.save(buf)
    ws2 = load_workbook(BytesIO(buf.getvalue())).active

    # Orden por abscisado del tramo: tramo 5 (abs 1314…) antes que tramo 7 (11090…)
    assert ws2["B9"].value == "K1+314.50 – K1+316.00"
    assert str(ws2["A10"].value) == "Total tramo 5"
    assert isinstance(ws2["H10"].value, str) and "H9" in ws2["H10"].value
    assert ws2["I10"].value == "M3"
    assert ws2["B11"].value == "K11+090.00 – K11+100.00"
    assert ws2["B12"].value == "K11+200.00 – K11+250.00"
    assert str(ws2["A13"].value) == "Total tramo 7"
    assert isinstance(ws2["H13"].value, str) and "H11" in ws2["H13"].value and "H12" in ws2["H13"].value
    assert ws2["I13"].value == "M3"
    assert tot_r == 14
    assert str(ws2["A14"].value) == "CANTIDAD TOTAL DEL ÍTEM"
    tot_f = str(ws2["H14"].value)
    assert tot_f.startswith("=")
    assert "H9" in tot_f and "H11" in tot_f and "H12" in tot_f
    assert "H10" not in tot_f and "H13" not in tot_f
    assert ws2["I14"].value == "M3"


def test_html_memoria_abscisas_subtotales_y_unidad():
    inf = _import_informes_with_stubs()
    regs = [
        {
            "numero_registro": 2,
            "tramo": "7",
            "abs_inicio": 11090.0,
            "abs_final": 11100.0,
            "longitud": 1,
            "ancho": 1,
            "espesor": 1,
            "cantidad": None,
            "cantidad_total": 10.0,
            "observacion": "",
            "foto_url": None,
            "item_numero": "1.1",
            "item_descripcion": "Exc",
            "unidad": "M3",
            "capitulo": "I",
        },
        {
            "numero_registro": 1,
            "tramo": "5",
            "abs_inicio": 1314.5,
            "abs_final": 1316.0,
            "longitud": 1,
            "ancho": 1,
            "espesor": 1,
            "cantidad": None,
            "cantidad_total": 4.0,
            "observacion": "",
            "foto_url": None,
            "item_numero": "1.1",
            "item_descripcion": "Exc",
            "unidad": "M3",
            "capitulo": "I",
        },
    ]
    item_info = inf._item_info_desde_registros(regs, "1.1")
    html = inf._html_memoria_item(
        {"numero": "C-1"},
        {"razon_social": "Sub", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-07"},
        item_info,
        regs,
        "Usuario",
        "Cargo",
        conc_meta={
            "titulo": "MEMORIA MENSUAL",
            "codigo": "CC-MES-002",
            "cells": [("CONTRATO", "C-1"), ("ACTA RPO", "1"), ("CONSECUTIVO", "1"), ("FECHA ACTA", "—")],
        },
        aprobo_interventoria_desde_config=True,
        pie_fotos_contexto="Acta RPO 1",
    )
    assert "K11+090.00" in html and "K11+100.00" in html
    assert "K1+314.50" in html and "K1+316.00" in html
    assert "Total tramo 7" in html
    assert "Total tramo 5" in html
    assert "mem002-subtramo" in html
    assert "CANTIDAD TOTAL DEL ÍTEM" in html
    assert ">M3<" in html or ">M3</" in html
    # Total del ítem = 14 (10+4), no suma subtotales
    assert "14" in html or "14.00" in html or "14,00" in html


def test_integral_link_sigue_apuntando_a_tot_r_con_subtotales():
    """La celda H del total del ítem (tot_r) sigue siendo la referenciada por el informe de ejecución."""
    inf = _import_informes_with_stubs()
    wb = Workbook()
    ws = wb.active
    regs = [
        {
            "numero_registro": 1,
            "tramo": "A",
            "abs_inicio": 100,
            "abs_final": 110,
            "longitud": 1,
            "ancho": 1,
            "espesor": 1,
            "cantidad": None,
            "cantidad_total": 1.0,
            "unidad": "M3",
            "item_numero": "1.1",
            "capitulo": "I",
        },
        {
            "numero_registro": 2,
            "tramo": "B",
            "abs_inicio": 200,
            "abs_final": 210,
            "longitud": 1,
            "ancho": 1,
            "espesor": 1,
            "cantidad": None,
            "cantidad_total": 2.0,
            "unidad": "M3",
            "item_numero": "1.1",
            "capitulo": "I",
        },
    ]
    tot_r = inf._fill_memoria_excel_ws(
        ws,
        {"numero": "C"},
        {},
        {"consecutivo": 1},
        {"item_numero": "1.1", "item_descripcion": "x", "unidad": "M3", "capitulo": "I"},
        regs,
        None,
        conc_meta={"titulo": "T", "codigo": "CC-SEM-002", "cells": [("A", "1"), ("B", "2"), ("C", "3"), ("D", "4")]},
        aprobo_interventoria_desde_config=True,
    )
    # 2 regs + 2 subtotales + 1 total → tot_r = data_row(9) + 4 = 13
    assert tot_r == 13
    assert ws.cell(tot_r, 1).value == "CANTIDAD TOTAL DEL ÍTEM"
    assert isinstance(ws.cell(tot_r, 8).value, str)
    assert ws.cell(tot_r, 8).value.startswith("=")
