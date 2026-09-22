"""Excel planilla tubería: anclas L20:N40 / B52:N63 y cuadrícula del panel gráfico."""
from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor

from topografia_planilla_tuberia_excel import (
    PANEL_RESULTADO_CELDAS,
    PERFIL_CHART_FROM,
    PERFIL_CHART_TO,
    SECCION_IMG_FROM,
    SECCION_IMG_TO,
    build_planilla_tuberia_xlsx,
)


def _ws(tipo: str = "ALCANTARILLA"):
    raw = build_planilla_tuberia_xlsx(
        planilla={
            "tipo": tipo,
            "diametro_m": 0.9,
            "espesor_m": 0.05,
            "ancho_excavacion_m": 1.5,
            "relacion_atraque": "1:3",
            "meta_cabecera": {"cama_triturado_m": 0.1},
        },
        calculo={"cartera": {"filas": []}},
        vacia=True,
    )
    return load_workbook(io.BytesIO(raw))["planilla"]


class TestExcelGraficosAnclas(unittest.TestCase):
    def test_una_sola_imagen_en_L20_N40(self):
        ws = _ws()
        self.assertEqual(len(ws._images), 1)
        anc = ws._images[0].anchor
        self.assertIsInstance(anc, TwoCellAnchor)
        self.assertEqual((anc._from.col, anc._from.row), SECCION_IMG_FROM)
        self.assertEqual((anc.to.col, anc.to.row), SECCION_IMG_TO)
        # L=12ª letra → col index 11; row 20 → index 19
        self.assertEqual(SECCION_IMG_FROM, (11, 19))
        self.assertEqual(SECCION_IMG_TO, (14, 40))

    def test_un_solo_chart_en_B52_N63_sin_invadir_firmas(self):
        ws = _ws()
        self.assertEqual(len(ws._charts), 1)
        anc = ws._charts[0].anchor
        self.assertIsInstance(anc, TwoCellAnchor)
        self.assertEqual((anc._from.col, anc._from.row), PERFIL_CHART_FROM)
        self.assertEqual((anc.to.col, anc.to.row), PERFIL_CHART_TO)
        # B52 → (1,51); termina en fila 63 (< firmas en 64)
        self.assertEqual(PERFIL_CHART_FROM, (1, 51))
        self.assertEqual(PERFIL_CHART_TO, (14, 63))
        self.assertLess(anc.to.row, 64)
        # Firmas intactas
        self.assertEqual(ws["A64"].value, "Elaboró")
        self.assertEqual(ws["H64"].value, "Aprobó:")

    def test_panel_grafico_sin_cuadricula_salvo_resultados(self):
        ws = _ws()
        # Interior L20:N40 sin borde
        cell = ws["M30"]
        self.assertTrue(
            cell.border is None
            or (
                cell.border.left is None
                and cell.border.right is None
                and cell.border.top is None
                and cell.border.bottom is None
            )
            or (
                getattr(cell.border.left, "style", None) is None
                and getattr(cell.border.right, "style", None) is None
            ),
            msg="M30 no debería tener cuadrícula",
        )
        for addr in PANEL_RESULTADO_CELDAS:
            b = ws[addr].border
            self.assertIsNotNone(b)
            self.assertIsNotNone(getattr(b.left, "style", None) or getattr(b.right, "style", None)
                                 or getattr(b.top, "style", None) or getattr(b.bottom, "style", None),
                                 msg=f"{addr} debe conservar borde")

    def test_fuente_sin_doble_chart_ni_ancla_vieja(self):
        from pathlib import Path
        import topografia_planilla_tuberia_excel as mod

        src = Path(mod.__file__).read_text(encoding="utf-8")
        self.assertEqual(src.count("_add_profile_chart("), 2)  # def + 1 call
        self.assertNotIn('add_chart(chart, "A52")', src)
        self.assertNotIn('add_image(img, "K17")', src)
        self.assertIn("L20:N40", src)
        self.assertIn("B52:N63", src)

    def test_perfil_ejes_y_grilla_tenue(self):
        """Etiquetas X/Y presentes; grilla major tenue/delgada; ancla B52:N63 intacta."""
        import re
        import zipfile

        from topografia_planilla_tuberia_excel import (
            PERFIL_GRID_LINE_COLOR,
            PERFIL_GRID_LINE_WIDTH_EMU,
            PERFIL_X_AXIS_TITLE,
            PERFIL_Y_AXIS_TITLE,
        )

        ws = _ws()
        ch = ws._charts[0]
        self.assertEqual(str(ch.x_axis.title.tx.rich.p[0].r[0].t), PERFIL_X_AXIS_TITLE)
        self.assertEqual(str(ch.y_axis.title.tx.rich.p[0].r[0].t), PERFIL_Y_AXIS_TITLE)
        self.assertEqual(ch.x_axis.axPos, "b")
        self.assertEqual(ch.y_axis.axPos, "l")
        self.assertIsNotNone(ch.x_axis.majorGridlines)
        self.assertIsNotNone(ch.y_axis.majorGridlines)
        self.assertIsNone(ch.x_axis.minorGridlines)
        self.assertIsNone(ch.y_axis.minorGridlines)
        # Ancla sin cambios
        self.assertEqual((ch.anchor._from.col, ch.anchor._from.row), PERFIL_CHART_FROM)
        self.assertEqual((ch.anchor.to.col, ch.anchor.to.row), PERFIL_CHART_TO)

        raw = build_planilla_tuberia_xlsx(
            planilla={"tipo": "ALCANTARILLA", "meta_cabecera": {}},
            calculo={"cartera": {"filas": []}},
            vacia=True,
        )
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            xml = z.read("xl/charts/chart1.xml").decode("utf-8")
        self.assertIn(PERFIL_X_AXIS_TITLE, xml)
        self.assertIn(PERFIL_Y_AXIS_TITLE, xml)
        self.assertIn(PERFIL_GRID_LINE_COLOR, xml)
        self.assertRegex(xml, rf'w="{PERFIL_GRID_LINE_WIDTH_EMU}"')
        # Dos bloques majorGridlines con spPr/ln (X e Y); openpyxl omite prefijo c:
        grids = re.findall(r"<majorGridlines>.*?</majorGridlines>", xml, re.S)
        self.assertEqual(len(grids), 2)
        for g in grids:
            self.assertIn("a:ln", g)
            self.assertIn(PERFIL_GRID_LINE_COLOR, g)


if __name__ == "__main__":
    unittest.main()
