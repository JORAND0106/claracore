"""Memorias SEM/MES/SUB-002: capítulo, grilla 3×2, dedupe, enlaces soporte, fórmulas."""
from __future__ import annotations

import sys
from io import BytesIO
from types import ModuleType, SimpleNamespace

from openpyxl import load_workbook


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]

    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake

    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)

    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    # Evitar HTTP real en tests de HTML (prefetch de fotos/gráficos).
    inf._memoria_prepare_media_uri = lambda url: ""  # type: ignore
    inf._memoria_prefetch_media_uris = lambda urls: {str(u or "").strip(): "" for u in urls if str(u or "").strip()}  # type: ignore
    return inf


def _regs_base(n_fotos=3, with_grafico=False, with_enlace=False, shared_foto_url=None):
    regs = []
    for i in range(n_fotos):
        url = shared_foto_url if shared_foto_url else f"https://example.com/f{i}.jpg"
        r = {
            "numero_registro": i + 1,
            "abs_inicio": "1+000",
            "abs_final": "1+010",
            "pk_ids": {"pk_id": "PK1"},
            "calzada": "Izq",
            "infraestructura": "Ciclorruta",
            "enlace_soporte": '["https://drive.google.com/file/d/abc/view"]' if with_enlace else None,
            "longitud": 10,
            "ancho": 2,
            "espesor": 0.1,
            "cantidad": None,
            "cantidad_total": 2.0,
            "observacion": "Obs " + ("larga " * 20),
            "foto_url": url,
            "foto_numero": 132 if shared_foto_url else i + 1,
            "item_numero": "5.01",
            "item_descripcion": "Excavación",
            "unidad": "m3",
            "capitulo": "3. SANITARIO",
        }
        if with_grafico and i == 0:
            r["grafico_url"] = "https://example.com/g1.png"
            r["grafico_numero"] = 7
        regs.append(r)
    return regs


def test_memoria_cells_con_capitulo_agrega_sin_duplicar():
    inf = _import_informes_with_stubs()
    cells = inf._memoria_cells_con_capitulo(
        [("CONTRATO", "C1"), ("ACTA RPO", "10"), ("CONSECUTIVO", "1"), ("FECHA ACTA", "—")],
        "2. PLUVIAL",
    )
    assert len(cells) == 5
    assert cells[4] == ("CAPÍTULO", "2. PLUVIAL")


def test_excel_formula_cantidad_total_columnas_defg():
    """CANT TOT en H = producto D–G (Long, Ancho, Esp, Cant) tras quitar PK ID."""
    inf = _import_informes_with_stubs()
    f = inf._excel_formula_cantidad_total(12)
    assert f.startswith("=")
    assert 'IF(D12="",1,D12)' in f
    assert "E12" in f and "F12" in f and "G12" in f
    assert "ROUND(" in f


def test_memoria_media_fit_pt_contain_sin_deformar():
    """xhtml2pdf ignora CSS max-*; el contain se fija en pt conservando proporción."""
    import base64
    import io

    from PIL import Image

    inf = _import_informes_with_stubs()

    def _uri(w, h, color=(30, 30, 50)):
        im = Image.new("RGB", (w, h), color)
        buf = io.BytesIO()
        im.save(buf, format="JPEG")
        return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()

    box_w, box_h = inf._MEMORIA_MEDIA_BOX_W_PT, inf._MEMORIA_MEDIA_BOX_H_PT
    # Ancha
    uri_w = _uri(1600, 400)
    w, h = inf._memoria_fit_media_pt(uri_w, box_w, box_h)
    assert w <= box_w + 0.01 and h <= box_h + 0.01
    assert abs(w / h - 4.0) < 0.08
    # Alta
    uri_t = _uri(400, 1200, (200, 200, 200))
    w2, h2 = inf._memoria_fit_media_pt(uri_t, box_w, box_h)
    assert w2 <= box_w + 0.01 and h2 <= box_h + 0.01
    assert abs(w2 / h2 - 400 / 1200) < 0.08
    # HTML con dims explícitas (attrs + style pt)
    tag = inf._memoria_media_img_html("https://cdn/x.jpg", {"https://cdn/x.jpg": uri_w})
    assert f"width:{w}pt" in tag and f"height:{h}pt" in tag
    assert 'width="' in tag and 'height="' in tag
    assert "data:image/jpeg;base64," in tag


def test_enlaces_soporte_parse_labels_y_html():
    inf = _import_informes_with_stubs()
    r = {
        "enlace_soporte": (
            '["https://drive.google.com/file/d/aaa/view",'
            '"https://contoso.sharepoint.com/sites/obra/doc.pdf",'
            '"https://example.com/folder/plano.dwg"]'
        )
    }
    links = inf._memoria_enlaces_soporte(r)
    assert len(links) == 3
    assert links[0]["label"] == "Drive 1"
    assert links[1]["label"] == "doc.pdf"
    assert links[2]["label"] == "plano.dwg"
    html = inf._memoria_enlaces_html_cell(r)
    assert 'href="https://drive.google.com/file/d/aaa/view"' in html
    assert "Drive 1" in html
    assert "plano.dwg" in html
    assert html.count("<a href=") == 3
    assert inf._memoria_enlaces_html_cell({"enlace_soporte": None}) == ""
    assert inf._memoria_enlace_label("https://drive.google.com/open?id=1", 2) == "Drive 2"


def test_dedupe_fotos_y_graficos_compartidos():
    inf = _import_informes_with_stubs()
    regs = [
        {
            "numero_registro": 181,
            "foto_url": "https://cdn/foto132.jpg",
            "foto_numero": 132,
            "observacion": "a",
            "grafico_url": "https://cdn/g.png",
            "grafico_numero": 3,
        },
        {
            "numero_registro": 295,
            "foto_url": "https://cdn/foto132.jpg",
            "foto_numero": 132,
            "observacion": "b",
            "grafico_url": "https://cdn/g.png",
            "grafico_numero": 3,
        },
        {
            "numero_registro": 297,
            "foto_url": "https://cdn/foto132.jpg",
            "foto_numero": 132,
            "observacion": "",
        },
        {
            "numero_registro": 300,
            "foto_url": "https://cdn/otra.jpg",
            "foto_numero": 140,
        },
    ]
    fotos = inf._dedupe_fotos_memoria(regs)
    assert len(fotos) == 2
    assert fotos[0]["foto_numero"] == 132
    assert fotos[0]["numeros_registro"] == [181, 295, 297]
    assert fotos[1]["foto_numero"] == 140
    graficos = inf._dedupe_graficos_memoria(regs)
    assert len(graficos) == 1
    assert graficos[0]["numeros_registro"] == [181, 295]
    top = inf._memoria_media_caption_top("foto", 132, [181, 295, 297])
    assert top == "Foto 132 — Reg. 181, 295, 297"


def test_html_memoria_encabezado_grilla6_columnas_sin_pk_y_dedupe():
    inf = _import_informes_with_stubs()
    contrato = {"numero": "IDU-1", "logo_contratista": None}
    sub = {"razon_social": "Sub SA", "nombre_contacto": "Ana"}
    corte = {"consecutivo": 3, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-07"}
    registros = _regs_base(n_fotos=7, with_grafico=True, with_enlace=True)
    registros.extend(
        [
            {
                **registros[0],
                "numero_registro": 181,
                "foto_url": "https://example.com/shared132.jpg",
                "foto_numero": 132,
                "grafico_url": "https://example.com/g1.png",
                "grafico_numero": 7,
                "enlace_soporte": (
                    '["https://drive.google.com/file/d/aaa/view",'
                    '"https://contoso.sharepoint.com/sites/x"]'
                ),
            },
            {
                **registros[0],
                "numero_registro": 295,
                "foto_url": "https://example.com/shared132.jpg",
                "foto_numero": 132,
                "grafico_url": "https://example.com/g1.png",
                "grafico_numero": 7,
            },
            {
                **registros[0],
                "numero_registro": 297,
                "foto_url": "https://example.com/shared132.jpg",
                "foto_numero": 132,
            },
        ]
    )
    item_info = inf._item_info_desde_registros(registros, "5.01")
    html = inf._html_memoria_item_body(
        contrato,
        sub,
        corte,
        item_info,
        registros,
        "Usuario",
        "Cargo",
        conc_meta={
            "titulo": "MEMORIA SEMANAL",
            "codigo": "CC-SEM-002",
            "cells": [
                ("CONTRATO", "IDU-1"),
                ("SEMANA", "N° 1"),
                ("VIGENCIA", "2026-01-01 — 2026-01-07"),
                ("REFERENCIA", "Cantidades"),
            ],
        },
        pie_fotos_contexto="Semana N° 1",
    )
    assert "CAPÍTULO" in html
    assert "3. SANITARIO" in html
    assert html.count('class="mem002-foto-slot"') == 18
    assert html.count("REGISTRO FOTOGRÁFICO — ÍTEM") == 2
    assert html.count("MEMORIA SEMANAL") >= 4
    assert "height:6.85cm" in html
    assert "height:4.55cm" not in html
    assert 'class="mem002-media-img"' in html
    # Contención: dims explícitas en pt y/o tope max-*; nunca max-width:0
    assert ("pt;" in html and "mem002-media-img" in html) or "max-height:" in html
    assert "max-width:0" not in html
    assert "font-size:0" not in html
    assert "ABSCISAS" in html
    assert "ENLACE" in html
    assert "INFRAESTRUCTURA" in html
    assert "PK ID" not in html
    assert "COSTADO" not in html
    assert "ABS INI" not in html
    assert "K1+000.00 – K1+010.00" in html
    assert "Ciclorruta" in html
    assert 'href="https://drive.google.com/file/d/abc/view"' in html
    assert "Drive 1" in html
    assert "SharePoint 2" in html
    assert html.index("CANT TOT") < html.index(">ENLACE<")
    assert html.index(">ENLACE<") < html.index("OBSERVACIÓN")
    assert "REGISTRO GRÁFICO — ÍTEM" in html
    assert "Gráfico 7 — Reg. 1, 181, 295" in html
    assert "Foto 132 — Reg. 181, 295, 297" in html
    assert html.count("https://example.com/shared132.jpg") == 1


def test_html_memoria_sin_graficos_no_genera_seccion():
    inf = _import_informes_with_stubs()
    contrato = {"numero": "IDU-1", "logo_contratista": None}
    sub = {"razon_social": "Sub SA", "nombre_contacto": "Ana"}
    corte = {"consecutivo": 3, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-07"}
    registros = _regs_base(n_fotos=2, with_grafico=False)
    item_info = inf._item_info_desde_registros(registros, "5.01")
    html = inf._html_memoria_item_body(
        contrato,
        sub,
        corte,
        item_info,
        registros,
        "Usuario",
        "Cargo",
        pie_fotos_contexto="Semana N° 1",
    )
    assert "REGISTRO FOTOGRÁFICO" in html
    assert "REGISTRO GRÁFICO" not in html
    assert html.count('class="mem002-foto-slot"') == 6


def test_fill_memoria_excel_ws_formulas_columnas_y_grafico():
    inf = _import_informes_with_stubs()
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    contrato = {"numero": "IDU-1"}
    sub = {"razon_social": "Sub SA", "nombre_contacto": "Ana"}
    corte = {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-07"}
    registros = [
        {
            "numero_registro": 1,
            "abs_inicio": "1+000",
            "abs_final": "1+010",
            "pk_ids": {"pk_id": "A"},
            "infraestructura": "Andén",
            "enlace_soporte": (
                '["https://drive.google.com/file/d/aaa/view",'
                '"https://contoso.sharepoint.com/sites/obra/doc.pdf"]'
            ),
            "longitud": 2,
            "ancho": 3,
            "espesor": 4,
            "cantidad": None,
            "cantidad_total": 24,
            "observacion": "ok",
            "foto_url": "https://example.com/shared.jpg",
            "foto_numero": 132,
            "grafico_url": "https://example.com/g.png",
            "grafico_numero": 2,
            "item_numero": "1.01",
            "item_descripcion": "Item",
            "unidad": "m3",
            "capitulo": "1. PRELIMINARES",
        },
        {
            "numero_registro": 2,
            "abs_inicio": None,
            "abs_final": None,
            "pk_ids": {},
            "infraestructura": "",
            "longitud": None,
            "ancho": None,
            "espesor": None,
            "cantidad": 5,
            "cantidad_total": 5,
            "observacion": "",
            "foto_url": "https://example.com/shared.jpg",
            "foto_numero": 132,
            "grafico_url": "https://example.com/g.png",
            "grafico_numero": 2,
            "item_numero": "1.01",
            "item_descripcion": "Item",
            "unidad": "m3",
            "capitulo": "1. PRELIMINARES",
        },
    ]
    item_info = inf._item_info_desde_registros(registros, "1.01")
    inf._fill_memoria_excel_ws(
        ws,
        contrato,
        sub,
        corte,
        item_info,
        registros,
        None,
        conc_meta={
            "titulo": "MEMORIA MENSUAL",
            "codigo": "CC-MES-002",
            "cells": [
                ("CONTRATO", "IDU-1"),
                ("ACTA RPO", "72"),
                ("CONSECUTIVO", "1"),
                ("FECHA ACTA", "—"),
            ],
        },
        pie_fotos_contexto="Acta RPO 72",
        aprobo_interventoria_desde_config=True,
    )
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    ws2 = load_workbook(buf)[ws.title]

    assert ws2["B8"].value == "ABSCISAS"
    assert ws2["C8"].value == "INFRAESTRUCTURA"
    assert ws2["D8"].value == "LONG"
    assert ws2["H8"].value == "CANT TOT"
    assert ws2["I8"].value == "ENLACE"
    assert ws2["J8"].value == "OBSERVACIÓN"
    assert "PK ID" not in [ws2.cell(row=8, column=c).value for c in range(1, 11)]
    assert isinstance(ws2["H9"].value, str) and ws2["H9"].value.startswith("=")
    assert "D9" in ws2["H9"].value and "G9" in ws2["H9"].value
    assert ws2["D9"].value == 2
    assert ws2["G10"].value == 5
    # Subtotal por tramo (ambos sin tramo → grupo «—») + total del ítem (solo filas de registro)
    assert str(ws2["A11"].value or "").startswith("Total tramo")
    assert isinstance(ws2["H11"].value, str) and ("H9" in ws2["H11"].value and "H10" in ws2["H11"].value)
    assert ws2["I11"].value == "m3"
    assert isinstance(ws2["H12"].value, str) and ("H9" in ws2["H12"].value and "H10" in ws2["H12"].value)
    assert "H11" not in str(ws2["H12"].value)  # total no incluye subtotales
    assert ws2["I12"].value == "m3"
    assert ws2["B9"].value == "K1+000.00 – K1+010.00"
    # Enlace: etiquetas cortas (no URL completa) + hipervínculo al primero
    assert ws2["I9"].value == "Drive 1\ndoc.pdf"
    assert ws2["I9"].hyperlink is not None
    assert "drive.google.com" in (ws2["I9"].hyperlink.target or "")
    assert ws2["C9"].value == "Andén"
    assert ws2["I10"].value in ("", None)

    flat = []
    for row in ws2.iter_rows(min_row=1, max_row=40, max_col=11, values_only=True):
        flat.extend([str(v) for v in row if v is not None])
    assert any("REGISTRO GRÁFICO" in v for v in flat)
    assert any("1. PRELIMINARES" in v for v in flat)
    foto_rows = [v for v in flat if "shared.jpg" in v]
    assert len(foto_rows) == 1
    assert any("1, 2" in v or "1,2" in v for v in flat)


def test_lista_graficos_desde_historial():
    inf = _import_informes_with_stubs()
    r = {
        "numero_registro": 9,
        "observacion": "x",
        "graficos_historial": [
            {"url": "https://a/1.png", "numero": 1, "creado_en": "2026-01-02"},
            {"url": "https://a/2.png", "numero": 2, "creado_en": "2026-01-01"},
        ],
    }
    g = inf._lista_graficos_memoria_registro(r)
    assert [x["grafico_numero"] for x in g] == [2, 1]
    assert g[0]["numero_registro"] == 9
