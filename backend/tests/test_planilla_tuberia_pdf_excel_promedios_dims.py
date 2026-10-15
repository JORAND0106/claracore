"""PDF/Excel: dimensiones bajo gráfico, notas de descuento y fila de promedios."""
from __future__ import annotations

import io
import importlib.util
import sys
import types
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def _load_pdf_mod():
    tu = types.ModuleType("topografia_utils")
    tu.svg_embed_pdf = lambda s, w, h: f'<img width="{w}" height="{h}"/>'
    tu._html_logo_pdf = lambda *a, **k: "<div>LOGO</div>"
    sys.modules["topografia_utils"] = tu
    spec = importlib.util.spec_from_file_location(
        "topografia_planilla_tuberia_pdf", ROOT / "topografia_planilla_tuberia_pdf.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _st_alc():
    return {
        "tipo": "ALCANTARILLA",
        "ancho_excavacion_m": 1.50,
        "diametro_externo_m": 1.000,
        "prom_altura_excavacion": 2.1234,
        "altura_relleno_m": 0.456,
        "prom_altura_triturado": 0.556,
        "relacion_atraque": "1:3",
        "area_1_m2": 0.1234,
        "area_2_m2": 0.5678,
        "cama_triturado_m": 0.1,
    }


def _calculo_sample():
    return {
        "seccion_tipica": _st_alc(),
        "seccion": _st_alc(),
        "perfil": {},
        "cartera": {
            "filas": [
                {
                    "orden": 1,
                    "vacio": False,
                    "abscisa": 100,
                    "terreno_natural": 105,
                    "subrasante_via": 104,
                    "cota_lomo": 103.5,
                    "cota_fondo_excavacion": 103,
                    "altura_excavacion": 2.0,
                    "altura_triturado": 0.5,
                    "altura_relleno": 1.5,
                    "ancho_geotextil": 0,
                },
                {
                    "orden": 2,
                    "vacio": False,
                    "abscisa": 130,
                    "terreno_natural": 104.7,
                    "subrasante_via": 103.7,
                    "cota_lomo": 103.2,
                    "cota_fondo_excavacion": 102.7,
                    "altura_excavacion": 2.0,
                    "altura_triturado": 0.5,
                    "altura_relleno": 1.5,
                    "ancho_geotextil": 0,
                },
            ],
            "totales": {
                "longitud_m": 30.0,
                "prom_altura_excavacion": 2.1234,
                "prom_altura_triturado": 0.556,
                "prom_altura_relleno": 1.5678,
                "prom_ancho_geotextil": 0.0,
            },
        },
        "notas_descuento_volumen": [
            "Excavación Roca: Vol 1.50 m³ descontado de Excavación Varias",
        ],
        "netos": [],
        "descuentos": [],
    }


class TestCaptionsDimensiones(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.mod = _load_pdf_mod()

    def test_captions_match_screen_alc(self):
        l1, l2 = self.mod.captions_dimensiones_seccion(_st_alc(), tipo="ALCANTARILLA")
        self.assertIn("B=1.50 m", l1)
        self.assertIn("Øext=1.000 m", l1)
        self.assertIn("h_exc=2.12 m", l1)
        self.assertIn("h_atr=0.456 m", l1)
        self.assertIn("(1:3)", l1)
        self.assertIn("A1=0.1234 m²", l2)
        self.assertIn("A2=0.5678 m²", l2)

    def test_bloque_graficos_incluye_dims_con_png(self):
        html = self.mod.html_bloque_graficos_pdf(_calculo_sample(), tipo="ALCANTARILLA")
        self.assertIn("B=1.50 m", html)
        self.assertIn("Øext=1.000 m", html)
        self.assertIn("A1=0.1234", html)
        self.assertIn("seccion-dims", html)
        self.assertIn("data:image/png", html)

    def test_notas_descuento_html(self):
        html = self.mod.html_notas_descuento_volumen_pdf(_calculo_sample())
        self.assertIn("Notas de descuento de volumen", html)
        self.assertIn("Excavación Roca", html)
        self.assertEqual(self.mod.html_notas_descuento_volumen_pdf({}), "")


class TestRoutesPdfPromediosNotas(unittest.TestCase):
    def test_routes_tienen_promedios_y_notas(self):
        src = (ROOT / "topografia_planilla_tuberia_routes.py").read_text(encoding="utf-8")
        self.assertIn("html_notas_descuento_volumen_pdf", src)
        self.assertIn("class='promedios'", src)
        self.assertIn("prom_altura_excavacion", src)
        self.assertIn("prom_altura_triturado", src)
        self.assertIn("prom_altura_relleno", src)
        self.assertIn("prom_ancho_geotextil", src)
        self.assertIn("height:158px", src)


class TestExcelDimsPromediosNotas(unittest.TestCase):
    def test_excel_dims_bajo_grafico_y_promedios_del_calculo(self):
        from topografia_planilla_tuberia_excel import (
            SECCION_IMG_FROM,
            SECCION_IMG_TO,
            build_planilla_tuberia_xlsx,
        )

        calc = _calculo_sample()
        raw = build_planilla_tuberia_xlsx(
            planilla={
                "tipo": "ALCANTARILLA",
                "diametro_m": 0.9,
                "espesor_m": 0.05,
                "ancho_excavacion_m": 1.5,
                "relacion_atraque": "1:3",
                "meta_cabecera": {"cama_triturado_m": 0.1},
            },
            calculo=calc,
            vacia=False,
        )
        ws = load_workbook(io.BytesIO(raw))["planilla"]
        # Ancla del gráfico intacta
        self.assertEqual(len(ws._images), 1)
        anc = ws._images[0].anchor
        self.assertEqual((anc._from.col, anc._from.row), SECCION_IMG_FROM)
        self.assertEqual((anc.to.col, anc.to.row), SECCION_IMG_TO)
        # Dimensiones bajo el gráfico (K41)
        k41 = str(ws["K41"].value or "")
        self.assertIn("B=1.50 m", k41)
        self.assertIn("Øext=", k41)
        self.assertIn("A1=", k41)
        # Promedios = totales del cálculo (3 dec), no AVERAGE genérico
        self.assertEqual(ws["G41"].value, 2.123)
        self.assertEqual(ws["H41"].value, 0.556)
        self.assertEqual(ws["I41"].value, 1.568)
        # Notas en bloque reservado
        found = False
        for row in ws.iter_rows(min_row=53, max_row=62, min_col=1, max_col=1):
            val = str(row[0].value or "")
            if "Excavación Roca" in val:
                found = True
                break
        self.assertTrue(found, msg="notas de descuento deben ir al bloque Notas")

    def test_plantilla_vacia_conserva_formulas_average(self):
        from topografia_planilla_tuberia_excel import build_planilla_tuberia_xlsx

        raw = build_planilla_tuberia_xlsx(
            planilla={"tipo": "ALCANTARILLA", "meta_cabecera": {}},
            calculo={"cartera": {"filas": []}},
            vacia=True,
        )
        ws = load_workbook(io.BytesIO(raw))["planilla"]
        self.assertIn("AVERAGE(G17:G36)", str(ws["G41"].value or ""))


class TestFrontendFilaPromedios(unittest.TestCase):
    def test_form_tiene_fila_promedios(self):
        form = (
            ROOT.parents[0]
            / "frontend"
            / "src"
            / "components"
            / "topografia"
            / "planillaTuberia"
            / "PlanillaTuberiaForm.jsx"
        ).read_text(encoding="utf-8")
        self.assertIn("data-cartera-promedios", form)
        self.assertIn("prom_altura_excavacion", form)
        self.assertIn("prom_altura_triturado", form)
        self.assertIn("prom_altura_relleno", form)
        self.assertIn("prom_ancho_geotextil", form)
        self.assertNotIn("prom H.Exc=", form)


class TestSelladaUsaCalculoCongelado(unittest.TestCase):
    """PDF/Excel leen del cálculo (snapshot); no inventan cifras nuevas."""

    def test_export_helpers_leen_totales_y_notas_del_calculo(self):
        mod = _load_pdf_mod()
        # Snapshot congelado (valores distintos a los defaults del helper).
        frozen = _calculo_sample()
        frozen["seccion_tipica"]["ancho_excavacion_m"] = 2.25
        frozen["seccion_tipica"]["area_1_m2"] = 0.9999
        frozen["notas_descuento_volumen"] = ["Nota congelada sellada XYZ"]
        frozen["cartera"]["totales"]["prom_altura_excavacion"] = 3.3333

        l1, l2 = mod.captions_dimensiones_seccion(
            frozen["seccion_tipica"], tipo="ALCANTARILLA"
        )
        self.assertIn("B=2.25 m", l1)
        self.assertIn("A1=0.9999", l2)
        self.assertIn("Nota congelada sellada XYZ", mod.html_notas_descuento_volumen_pdf(frozen))

        from topografia_planilla_tuberia_excel import build_planilla_tuberia_xlsx

        raw = build_planilla_tuberia_xlsx(
            planilla={"tipo": "ALCANTARILLA", "meta_cabecera": {}},
            calculo=frozen,
            vacia=False,
        )
        ws = load_workbook(io.BytesIO(raw))["planilla"]
        self.assertEqual(ws["G41"].value, 3.333)
        self.assertIn("B=2.25 m", str(ws["K41"].value or ""))
        blob = " ".join(
            str(ws.cell(row=r, column=1).value or "")
            for r in range(53, 63)
        )
        self.assertIn("Nota congelada sellada XYZ", blob)


if __name__ == "__main__":
    unittest.main()
