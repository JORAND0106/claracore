"""Excel planilla tubería: bordes + layout fijo por tipo (sin selector)."""
from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from topografia_planilla_tuberia_excel import (
    TITULO_ALC,
    TITULO_FIL,
    build_planilla_tuberia_xlsx,
)


def _wb(tipo: str):
    raw = build_planilla_tuberia_xlsx(
        planilla={
            "tipo": tipo,
            "diametro_m": 0.9,
            "espesor_m": 0.05,
            "ancho_excavacion_m": 1.5,
            "relacion_atraque": "1:3",
            "meta_cabecera": {"cama_triturado_m": 0.1},
        },
        calculo={"cartera": {"filas": []}},
        vacia=True,
    )
    return load_workbook(io.BytesIO(raw))


class TestExcelTipoFijo(unittest.TestCase):
    def test_sin_dropdown_tipo_ni_p1_p2(self):
        for tipo in ("ALCANTARILLA", "FILTRO"):
            ws = _wb(tipo)["planilla"]
            self.assertIsNone(ws["P1"].value)
            self.assertIsNone(ws["P2"].value)
            formulas = " ".join(
                str(dv.formula1) for dv in ws.data_validations.dataValidation
            )
            self.assertNotIn("$P$1", formulas)
            self.assertEqual(ws["M13"].value, tipo)

    def test_layout_alcantarilla(self):
        ws = _wb("ALCANTARILLA")["planilla"]
        self.assertEqual(ws["F1"].value, TITULO_ALC)
        self.assertEqual(ws["E16"].value, "Cota Lomo")
        self.assertEqual(ws["D16"].value, "Subrasante de Vía")
        self.assertEqual(ws["F14"].value, "Cama Triturado")
        self.assertIn("$E$15+$F$15", str(ws["H17"].value))
        self.assertEqual(ws["G48"].value, "=N46")
        self.assertEqual(ws["I46"].value, "Area 1")
        self.assertTrue(not ws["I45"].value)
        self.assertEqual(len(ws._images), 1)

    def test_layout_filtro(self):
        ws = _wb("FILTRO")["planilla"]
        self.assertEqual(ws["F1"].value, TITULO_FIL)
        self.assertEqual(ws["E16"].value, "Terminado Filtro")
        self.assertTrue(not ws["D16"].value)
        self.assertTrue(not ws["F14"].value)
        self.assertIn("E17-F17", str(ws["H17"].value))
        self.assertEqual(ws["G48"].value, "=N45")
        self.assertEqual(ws["I45"].value, "Tubería Filtro")
        self.assertTrue(not ws["I46"].value)
        self.assertEqual(len(ws._images), 1)

    def test_imagenes_distintas_por_tipo(self):
        alc = _wb("ALCANTARILLA")["planilla"]
        fil = _wb("FILTRO")["planilla"]
        self.assertEqual(len(alc._images), 1)
        self.assertEqual(len(fil._images), 1)
        # openpyxl Image.path / ref differs by file
        p_alc = getattr(alc._images[0], "path", None) or str(alc._images[0].ref)
        p_fil = getattr(fil._images[0], "path", None) or str(fil._images[0].ref)
        # At least both present; binary blobs differ
        self.assertNotEqual(alc._images[0]._data(), fil._images[0]._data())


class TestExcelBordes(unittest.TestCase):
    def test_bloques_con_borde(self):
        ws = _wb("ALCANTARILLA")["planilla"]
        samples = ("B5", "B16", "B17", "J36", "B44", "H50", "I44", "N50", "L18")
        for addr in samples:
            cell = ws[addr]
            self.assertIsNotNone(cell.border)
            self.assertIsNotNone(cell.border.left)
            self.assertIsNotNone(cell.border.left.style, msg=f"{addr} sin borde")

    def test_fuente_sin_condicional_tipo(self):
        from pathlib import Path
        import topografia_planilla_tuberia_excel as mod

        src = Path(mod.__file__).read_text(encoding="utf-8")
        self.assertNotIn('formula1="$P$1:$P$2"', src)
        self.assertNotIn("IF($F$1=$P$1", src)
        self.assertNotIn("IF(F1=P1", src)


if __name__ == "__main__":
    unittest.main()
