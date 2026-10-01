"""Agrupación por capítulo + subtotales en CC-SUB/MES-001 (sin columna CAP.)."""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from corte_sub_conciliacion import build_resumen_conciliacion_4cols
from test_corte_sub_001_excel_filtro import _import_informes_with_stubs


def _sample_items():
    return [
        {
            "capitulo": "1. PRELIMINARES",
            "item_numero": "1.1",
            "item_descripcion": "Localización y replanteo",
            "unidad": "M2",
            "vlr_unitario_sub": 1000.0,
            "cant_actualizadas": 10.0,
            "cant_presente": 2.0,
            "cant_acum_anterior": 1.0,
            "cant_acumulado": 3.0,
            "cant_saldo": 7.0,
            "valor_actualizadas": 10000.0,
            "valor_presente": 2000.0,
            "valor_acumulado": 3000.0,
            "valor_saldo": 7000.0,
            "costo_directo": 2000.0,
        },
        {
            "capitulo": "1. PRELIMINARES",
            "item_numero": "1.2",
            "item_descripcion": "Descapote",
            "unidad": "M3",
            "vlr_unitario_sub": 2000.0,
            "cant_actualizadas": 5.0,
            "cant_presente": 1.0,
            "cant_acum_anterior": 0.0,
            "cant_acumulado": 1.0,
            "cant_saldo": 4.0,
            "valor_actualizadas": 10000.0,
            "valor_presente": 2000.0,
            "valor_acumulado": 2000.0,
            "valor_saldo": 8000.0,
            "costo_directo": 2000.0,
        },
        {
            "capitulo": "3. OBRAS DE ARTE (ALCANTARILLA)",
            "item_numero": "3.1",
            "item_descripcion": "Alcantarilla box",
            "unidad": "ML",
            "vlr_unitario_sub": 5000.0,
            "cant_actualizadas": 4.0,
            "cant_presente": 0.0,
            "cant_acum_anterior": 0.0,
            "cant_acumulado": 0.0,
            "cant_saldo": 4.0,
            "valor_actualizadas": 20000.0,
            "valor_presente": 0.0,
            "valor_acumulado": 0.0,
            "valor_saldo": 20000.0,
            "costo_directo": 0.0,
        },
    ]


def test_plan_agrupa_y_subtotal_por_capitulo():
    inf = _import_informes_with_stubs()
    plan = inf._cc_sub_001_plan_filas_capitulo(_sample_items())
    kinds = [p[0] for p in plan]
    assert kinds == ["item", "item", "subcap", "item", "subcap"]
    sub1 = plan[2]
    assert sub1[1] == "1. PRELIMINARES"
    assert sub1[2]["valor_presente"] == 4000.0
    assert sub1[2]["valor_actualizadas"] == 20000.0
    sub2 = plan[4]
    assert "ALCANTARILLA" in sub2[1]
    assert sub2[2]["valor_actualizadas"] == 20000.0
    assert sub2[2]["valor_presente"] == 0.0


def test_html_sin_columna_cap_con_subtotal_capitulo():
    inf = _import_informes_with_stubs()
    items = _sample_items()
    html = inf._html_cc_sub_v1_plain(
        {"numero": "C-1", "contratista": "Ctor", "nit": "900", "interventoria": "Int"},
        {"razon_social": "Sub SA", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-15"},
        items,
        4000.0,
        "Tester",
        "Desarrollador",
    )
    assert ">CAP.<" not in html
    assert "DESCRIPCIÓN" in html
    assert "width:25%" in html  # ancho liberado a descripción
    assert "Subtotal 1. PRELIMINARES" in html
    assert "Subtotal 3. OBRAS DE ARTE (ALCANTARILLA)" in html
    assert html.count('class="cc001-cap-sub"') == 2
    assert "<pdf:nextpage" not in html.lower()
    assert html.count('class="cc001-tabla-items"') == 1


def test_excel_subtotales_formulas_y_cd_solo_items():
    inf = _import_informes_with_stubs()
    items = _sample_items()
    r4 = build_resumen_conciliacion_4cols(
        items=items, tributos={"administracion": 10}, otros_presente=0
    )
    raw = inf._corte_sub_001_excel_bytes(
        {"numero": "C-1", "contratista": "Ctor", "nit": "900", "interventoria": "Int"},
        {"razon_social": "Sub SA", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-15"},
        items,
        4000.0,
        "Tester",
        "Desarrollador",
        {},
        resumen_4cols=r4,
    )
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    # Sin CAP.; bloques parten en col 5
    assert ws.cell(7, 1).value == "ÍTEM"
    assert ws.cell(7, 2).value == "DESCRIPCIÓN"
    assert ws.cell(7, 5).value == "ACTUALIZADAS"
    assert ws.cell(7, 7).value == "PRESENTE ACTA"
    # Ítems
    assert ws.cell(9, 1).value == "1.1"
    assert ws.cell(10, 1).value == "1.2"
    # Subtotal cap 1
    assert str(ws.cell(11, 1).value).startswith("Subtotal 1. PRELIMINARES")
    assert str(ws.cell(11, 6).value).startswith("=")  # fórmula suma valores
    assert "F9" in str(ws.cell(11, 6).value) and "F10" in str(ws.cell(11, 6).value)
    # Cantidad de subtotal vacía
    assert ws.cell(11, 5).value in (None, "")
    # Ítem cap 3 + subtotal
    assert ws.cell(12, 1).value == "3.1"
    assert "ALCANTARILLA" in str(ws.cell(13, 1).value)
    # CD del resumen suma solo ítems (no filas 11/13 de subtotal)
    cd_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "Costo Directo":
            cd_row = r
            break
    assert cd_row is not None
    cd_f = str(ws.cell(cd_row, 6).value)
    assert cd_f.startswith("=")
    assert "F9" in cd_f and "F10" in cd_f and "F12" in cd_f
    assert "F11" not in cd_f and "F13" not in cd_f
    # Memorias: Presente cant en col G; acum = G+ant
    assert ws.cell(9, 7).value == 2.0
    assert str(ws.cell(9, 9).value) == "=G9+1.0"
    assert str(ws.cell(9, 11).value) == "=E9-I9"
