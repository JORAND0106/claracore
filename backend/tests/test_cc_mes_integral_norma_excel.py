"""Excel integral mensual: sin preacta_obra, nombre con acta y columna Norma técnica."""
from __future__ import annotations

import sys
from io import BytesIO
from types import ModuleType, SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from corte_sub_conciliacion import build_resumen_conciliacion_4cols, valor_por_cantidad_vu
from excel_formula_eval import FormulaBook


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]

    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake

    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)

    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    return inf


def _item(cap, num, desc, und, vu, act, pres, ant, norma=""):
    acum = round(float(ant) + float(pres), 2)
    saldo = round(float(act) - acum, 2)
    return {
        "capitulo": cap,
        "item_numero": num,
        "item_descripcion": desc,
        "unidad": und,
        "norma_tecnica": norma,
        "vlr_unitario_sub": vu,
        "vlr_unitario": vu,
        "cant_actualizadas": act,
        "cant_presente": pres,
        "cant_acum_anterior": ant,
        "cantidad": pres,
        "costo_directo": valor_por_cantidad_vu(pres, vu),
        "cant_acumulado": acum,
        "cant_saldo": saldo,
        "valor_actualizadas": valor_por_cantidad_vu(act, vu),
        "valor_presente": valor_por_cantidad_vu(pres, vu),
        "valor_acumulado": valor_por_cantidad_vu(acum, vu),
        "valor_saldo": valor_por_cantidad_vu(saldo, vu),
    }


CAP3 = "3. OBRAS DE ARTE (ALCANTARILLA)"
CAP4 = "4. PAVIMENTOS"


def _dataset():
    return [
        _item(CAP3, "3.4.", "Alcantarilla Doble", "M", 1000, 10, 2.5, 1.5, "NTC 174"),
        _item(CAP3, "3.5.", "Cuneta", "M", 2205, 8, 0, 3, ""),
        _item(CAP4, "3.4.", "Base Granular", "M3", 500, 4, 1.25, 0.25, "INVIAS 300"),
    ]


def _regs_for(items):
    regs = []
    for it in items:
        if float(it["cant_presente"]) <= 0:
            continue
        regs.append(
            {
                "numero_registro": len(regs) + 1,
                "capitulo": it["capitulo"],
                "item_numero": it["item_numero"],
                "item_descripcion": it["item_descripcion"],
                "unidad": it["unidad"],
                "longitud": it["cant_presente"],
                "ancho": "",
                "espesor": "",
                "cantidad": "",
                "cantidad_total": it["cant_presente"],
                "observacion": "",
                "abs_inicio": "",
                "abs_final": "",
                "infraestructura": "",
            }
        )
    return regs


def _libro(items=None):
    inf = _import_informes_with_stubs()
    items = _dataset() if items is None else items
    regs = _regs_for(items)
    r4 = build_resumen_conciliacion_4cols(items=items)
    contrato = {
        "numero": "ICCU-CTO-1614-2025",
        "contratista": "Ctor",
        "nit": "900",
        "interventoria": "Int",
    }
    ctx = {
        "contrato": contrato,
        "acta": {
            "numero_rpo": "1",
            "consecutivo": 1,
            "fecha_inicio": "2026-01-01",
            "fecha_fin": "2026-01-31",
        },
        "items": items,
        "total_costo": sum(it["valor_presente"] for it in items),
        "resumen_4cols": r4,
        "otros_conceptos": [],
        "usuario_nombre": "Ana Pérez",
        "usuario_cargo": "Residente",
    }
    user = {"nombre": "Ana", "apellidos": "Pérez", "cargo_nombre": "Residente"}

    def _fetch(_sb, _cid, inum, **_kw):
        return [r for r in regs if str(r.get("item_numero")) == str(inum)]

    with (
        patch.object(inf, "_acta_pertenece_contrato", return_value=True),
        patch.object(inf, "_contexto_acta_mes_conciliacion", return_value=ctx),
        patch.object(inf, "_row", return_value=contrato),
        patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
        patch.object(inf, "_sub_corte_dummy_memoria", return_value=({}, {})),
        patch.object(inf, "fetch_registros_memoria_cc_mes_alineado_acta", side_effect=_fetch),
    ):
        raw = inf._cc_mes_integral_excel_bytes(3, 620, user, nivel_aprobacion=4)
    return inf, items, raw, load_workbook(BytesIO(raw))


def _fila_resumen(ws):
    for r in range(9, ws.max_row + 1):
        if ws.cell(r, 1).value == "RESUMEN DE CONCILIACIÓN":
            return r
    raise AssertionError("falta el resumen de conciliación")


def _filas_item(ws):
    limite = _fila_resumen(ws)
    rows = []
    for r in range(9, limite):
        val = ws.cell(r, 1).value
        if isinstance(val, str) and str(val).startswith("Subtotal"):
            continue
        if val in (None, ""):
            continue
        rows.append(r)
    return rows


def test_integral_no_trae_preacta_y_corre_norma_tecnica():
    _inf, items, _raw, wb = _libro()
    assert wb.sheetnames[0] == "CC-MES-001"
    assert "preacta_obra" not in wb.sheetnames
    assert all(not ws.tables for ws in wb.worksheets)
    # La memoria sigue siendo la segunda hoja.
    assert len(wb.sheetnames) >= 2

    ws = wb["CC-MES-001"]
    assert ws.cell(7, 3).value == "UND"
    assert ws.cell(7, 4).value == "NORMA TÉCNICA"
    assert ws.cell(7, 5).value == "V. UNIT."
    assert ws.cell(7, 6).value == "ACTUALIZADAS"
    assert ws.cell(7, 8).value == "PRESENTE ACTA"
    assert ws.cell(7, 10).value == "ACUMULADO"
    assert ws.cell(7, 12).value == "SALDO"
    assert ws.max_column == 13

    filas = _filas_item(ws)
    assert len(filas) == len(items)
    book = FormulaBook(wb)
    for row, it in zip(filas, items):
        esperado = str(it.get("norma_tecnica") or "").strip()
        assert (ws.cell(row, 4).value or "") == esperado
        assert ws.cell(row, 5).value == it["vlr_unitario_sub"]
        assert book.eval_cell("CC-MES-001", row, 7) == it["valor_actualizadas"]
        assert book.eval_cell("CC-MES-001", row, 8) == it["cant_presente"]
        assert book.eval_cell("CC-MES-001", row, 9) == it["valor_presente"]
        assert book.eval_cell("CC-MES-001", row, 10) == it["cant_acumulado"]
        assert book.eval_cell("CC-MES-001", row, 11) == it["valor_acumulado"]
        assert book.eval_cell("CC-MES-001", row, 12) == it["cant_saldo"]
        assert book.eval_cell("CC-MES-001", row, 13) == it["valor_saldo"]
        if float(it["cant_presente"]) > 0:
            pres = str(ws.cell(row, 8).value)
            assert pres.startswith("=ROUND('") and "!H" in pres

    # Ítem sin norma: celda vacía. Mismo ítem en otro capítulo conserva la suya.
    assert ws.cell(filas[1], 4).value in (None, "")
    assert ws.cell(filas[0], 4).value == "NTC 174"
    assert ws.cell(filas[2], 4).value == "INVIAS 300"
    assert ws.cell(filas[0], 1).value == ws.cell(filas[2], 1).value == "3.4."


def test_subtotales_y_costo_directo_siguen_la_columna_corrida():
    _inf, items, _raw, wb = _libro()
    ws = wb["CC-MES-001"]
    book = FormulaBook(wb)
    limite = _fila_resumen(ws)
    sub_rows = [
        r
        for r in range(9, limite)
        if isinstance(ws.cell(r, 1).value, str) and str(ws.cell(r, 1).value).startswith("Subtotal")
    ]
    assert len(sub_rows) >= 2
    # Valor actualizadas del subtotal está en G, no en F.
    assert "G" in str(ws.cell(sub_rows[0], 7).value)
    assert ws.cell(sub_rows[0], 6).value in (None, "")
    presente = sum(book.eval_cell("CC-MES-001", r, 9) for r in sub_rows)
    assert presente == sum(it["valor_presente"] for it in items)

    cd_row = next(
        r
        for r in range(sub_rows[-1] + 1, ws.max_row + 1)
        if ws.cell(r, 1).value == "Costo Directo"
    )
    assert book.eval_cell("CC-MES-001", cd_row, 9) == presente


def test_norma_se_cruza_por_capitulo_e_item():
    inf = _import_informes_with_stubs()
    items = [
        {"capitulo": "1. PRELIMINARES", "item_numero": "1.1"},
        {"capitulo": "2. EXCAVACIONES", "item_numero": "1.1"},
        {"capitulo": "1. PRELIMINARES", "item_numero": "1.2"},
    ]
    meta = {
        ("1.PRELIMINARES", "1.1"): {"especificacion_tecnica": "NTC 174"},
        ("2.EXCAVACIONES", "1.1"): {"especificacion_tecnica": "  INVIAS 300  "},
        ("1.PRELIMINARES", "1.2"): {"especificacion_tecnica": "   "},
    }
    inf._aplicar_norma_tecnica_desde_listado(items, meta)
    assert items[0]["norma_tecnica"] == "NTC 174"
    assert items[1]["norma_tecnica"] == "INVIAS 300"
    assert items[2]["norma_tecnica"] == ""


def test_nombre_descarga_incluye_numero_de_acta():
    inf = _import_informes_with_stubs()
    with (
        patch.object(inf, "_perm_informes_ccd"),
        patch.object(inf, "_nivel_aprobacion_mes_query", return_value=4),
        patch.object(inf, "_cc_mes_integral_excel_bytes", return_value=b"PK"),
        patch.object(inf, "_row", return_value={"numero_rpo": "1", "consecutivo": 9}),
    ):
        resp = inf.excel_cc_mes_integral_acta(3, 620, nivel_aprobacion=4, current_user={})
    cd = resp.headers["content-disposition"]
    assert 'filename="CC-MES-integral_Acta_1.xlsx"' in cd
