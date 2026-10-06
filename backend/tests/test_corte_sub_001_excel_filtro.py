"""CC-SUB-001 Excel: bloques PDF + fórmulas + solo ítems con cantidades."""
from __future__ import annotations

import sys
from io import BytesIO
from types import ModuleType, SimpleNamespace

from openpyxl import load_workbook


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]

    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake

    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)

    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    return inf


def test_filtrar_items_con_cantidades():
    from corte_sub_conciliacion import filtrar_items_con_cantidades

    items = [
        {"item_numero": "1.1", "cant_presente": 5, "cant_acumulado": 0},
        {"item_numero": "1.2", "cant_presente": 0, "cant_acumulado": 3},
        {"item_numero": "1.3", "cant_presente": 0, "cant_acumulado": 0, "cant_actualizadas": 10},
        {"item_numero": "1.4", "cantidad": 2, "cant_acumulado": 0},
        {"item_numero": "1.5", "cant_presente": 0, "cant_acumulado": 0, "cant_actualizadas": 0},
    ]
    out = filtrar_items_con_cantidades(items)
    nums = {i["item_numero"] for i in out}
    assert nums == {"1.1", "1.2", "1.3", "1.4"}
    assert "1.5" not in nums


def test_corte_sub_001_excel_bloques_y_formulas():
    inf = _import_informes_with_stubs()
    items = [
        {
            "capitulo": "I",
            "item_numero": "1.1",
            "item_descripcion": "Excavación",
            "unidad": "M3",
            "vlr_unitario_sub": 1000.0,
            "cant_actualizadas": 100.0,
            "cant_presente": 10.0,
            "cant_acum_anterior": 40.0,
            "cant_acumulado": 50.0,
            "cant_saldo": 50.0,
            "valor_presente": 10000.0,
            "valor_actualizadas": 100000.0,
            "valor_acumulado": 50000.0,
            "valor_saldo": 50000.0,
            "costo_directo": 10000.0,
        },
        {
            "capitulo": "I",
            "item_numero": "1.2",
            "item_descripcion": "Relleno",
            "unidad": "M3",
            "vlr_unitario_sub": 2000.0,
            "cant_actualizadas": 20.0,
            "cant_presente": 0.0,
            "cant_acum_anterior": 5.0,
            "cant_acumulado": 5.0,
            "cant_saldo": 15.0,
            "valor_presente": 0.0,
            "valor_actualizadas": 40000.0,
            "valor_acumulado": 10000.0,
            "valor_saldo": 30000.0,
            "costo_directo": 0.0,
        },
    ]
    from corte_sub_conciliacion import build_resumen_conciliacion_4cols

    r4 = build_resumen_conciliacion_4cols(items=items, tributos={"administracion": 10}, otros_presente=0)
    raw = inf._corte_sub_001_excel_bytes(
        {"numero": "C-1", "contratista": "Ctor", "nit": "900", "interventoria": "Int"},
        {"razon_social": "Sub SA", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-15"},
        items,
        10000.0,
        "Tester",
        "Desarrollador",
        {},
        resumen_4cols=r4,
    )
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    assert ws.title == "CC-SUB-001"
    assert ws.page_setup.orientation == "landscape"
    assert ws.cell(1, 4).value == "INFORME CORTE DE SUB CONTRATISTA"
    assert ws.cell(7, 5).value == "ACTUALIZADAS"
    assert ws.cell(7, 7).value == "PRESENTE ACTA"
    assert ws.cell(7, 9).value == "ACUMULADO"
    assert ws.cell(7, 11).value == "SALDO"
    # Fila de datos (row 9 = data0): sin CAP.; ÍTEM en col 1
    assert ws.cell(9, 1).value == "1.1"
    assert ws.cell(9, 4).value == 1000.0
    assert ws.cell(9, 5).value == 100.0
    assert str(ws.cell(9, 6).value).startswith("=")
    assert ws.cell(9, 7).value == 10.0
    assert str(ws.cell(9, 9).value) == "=ROUND(G9+40.0,2)"
    assert str(ws.cell(9, 11).value) == "=ROUND(E9-I9,2)"
    # Subtotal del capítulo I
    assert str(ws.cell(11, 1).value).startswith("Subtotal")
    assert "F9" in str(ws.cell(11, 6).value)


def test_popup_y_excel_usan_filtro_solo_aprobados_en_fuente():
    from pathlib import Path

    popup = (Path(__file__).resolve().parents[1].parent / "frontend/src/informes/CorteSubConciliacionPopup.jsx").read_text(
        encoding="utf-8"
    )
    mod = (Path(__file__).resolve().parents[1].parent / "frontend/src/ModuloInformes.jsx").read_text(
        encoding="utf-8"
    )
    assert "pathConFiltroSubAprobacion" in popup
    assert "filtroSubAprobacion" in popup
    assert "filtroSubAprobacion={filtroSubAprobacion}" in mod
    assert "puedeVistaPreviaCorteSinConciliar" in mod
    assert "Vista sin conciliar" in mod
    assert "sinConciliar" in mod
