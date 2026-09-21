"""Verificación bloque a bloque vs inventario XLSM (sin Excel)."""
from __future__ import annotations

import json
import math
import unittest
from pathlib import Path

from topografia_planilla_tuberia import (
    area_1_m2,
    area_2_m2,
    area_tuberia_m2,
    calcular_planilla_completa,
    altura_relleno_atraque_m,
)
from topografia_planilla_tuberia_excel import build_planilla_tuberia_xlsx, TITULO_ALC, TITULO_FIL, CODIGO_DOC

ROOT = Path(__file__).resolve().parents[2]
INV = ROOT / "docs" / "topografia" / "planillas_tuberia" / "inventario_planilla_tuberia_xlsm.json"
LAYOUT = ROOT / "docs" / "topografia" / "planillas_tuberia" / "layout_map_planilla_tuberia.json"
INV_RUNTIME = (
    ROOT / "backend" / "data" / "planillas_tuberia" / "inventario_planilla_tuberia_xlsm.json"
)


class TestInventarioPresente(unittest.TestCase):
    def test_artefactos(self):
        self.assertTrue(INV.exists())
        self.assertTrue(LAYOUT.exists())
        self.assertTrue(
            INV_RUNTIME.exists(),
            "Copia runtime del inventario en backend/data (sin depender del .xlsm)",
        )
        inv = json.loads(INV.read_text())
        self.assertEqual(inv["sheetnames"], ["Tbl_Auxiliares", "planilla", "Resumen_BASE"])


class TestFormulasVsInventario(unittest.TestCase):
    def test_k13_area_tuberia(self):
        # K13 = ROUND(PI()*((I13/2)+J13)^2,3)
        self.assertAlmostEqual(area_tuberia_m2(0.9, 0.05), round(math.pi * 0.5 ** 2, 3))

    def test_e15_altura_relleno(self):
        self.assertAlmostEqual(altura_relleno_atraque_m(0.9, 0.05, "1:3"), 0.333)

    def test_b15_segmento_circular(self):
        r = 0.5
        h = 0.333
        esperado = round(
            r ** 2 * math.acos((r - h) / r) - (r - h) * math.sqrt(2 * r * h - h ** 2),
            3,
        )
        self.assertAlmostEqual(area_1_m2(0.9, 0.05, 1.5, "1:3"), esperado, places=3)
        self.assertAlmostEqual(
            area_1_m2(0.9, 0.05, 1.5, "1:3") + area_2_m2(0.9, 0.05, "1:3"),
            area_tuberia_m2(0.9, 0.05),
            places=3,
        )

    def test_items_resumen_literales(self):
        r = calcular_planilla_completa(
            tipo="ALCANTARILLA",
            diametro_m=0.9,
            espesor_m=0.05,
            ancho_excavacion_m=1.5,
            relacion_atraque="1:3",
            filas_campo=[
                {"orden": 1, "abscisa": 100, "terreno_natural": 105, "subrasante_via": 104, "cota_fondo_excavacion": 103},
                {"orden": 2, "abscisa": 130, "terreno_natural": 104.7, "subrasante_via": 103.7, "cota_fondo_excavacion": 102.7},
            ],
            cama_triturado_m=0.1,
        )
        nombres = [n["nombre"] for n in r["netos"]]
        self.assertEqual(
            nombres,
            [
                "Excavación Varias",
                "Excavación Roca",
                "Long Tubería",
                "Triturado / Atraque",
                "Relleno Gran.",
                "Geotextil",
            ],
        )
        desc = [d["nombre"] for d in r["descuentos"]]
        self.assertIn("Area 1", desc)
        self.assertIn("Area 2", desc)


class TestExcelExportFidelidad(unittest.TestCase):
    def test_titulos_y_formulas_vivas(self):
        from openpyxl import load_workbook
        import io

        raw = build_planilla_tuberia_xlsx(
            planilla={
                "tipo": "ALCANTARILLA",
                "diametro_m": 0.9,
                "espesor_m": 0.05,
                "ancho_excavacion_m": 1.5,
                "relacion_atraque": "1:3",
                "pk_id": "K0+000",
                "meta_cabecera": {"cama_triturado_m": 0.3},
            },
            calculo={
                "cartera": {
                    "filas": [
                        {
                            "orden": 1,
                            "abscisa": 100,
                            "terreno_natural": 105,
                            "subrasante_via": 104,
                            "cota_fondo_excavacion": 103,
                            "vacio": False,
                        }
                    ]
                }
            },
        )
        wb = load_workbook(io.BytesIO(raw))
        ws = wb["planilla"]
        self.assertEqual(ws["F1"].value, TITULO_ALC)
        self.assertEqual(ws["M13"].value, "ALCANTARILLA")
        self.assertIsNone(ws["P1"].value)
        self.assertIsNone(ws["P2"].value)
        self.assertEqual(ws["M1"].value, CODIGO_DOC)
        self.assertEqual(ws["E16"].value, "Cota Lomo")
        self.assertEqual(len(ws._images), 1)
        self.assertTrue(str(ws["G17"].value).startswith("=IFERROR"))
        self.assertTrue(str(ws["H48"].value).startswith("=ROUND(PRODUCT"))
        self.assertGreaterEqual(len(ws._charts), 1)

    def test_xlsx_sin_macros_ni_xlsm(self):
        """Exportación vacía (Dev): .xlsx OOXML sin VBA; módulo no abre .xlsm."""
        import io
        from openpyxl import load_workbook
        from topografia_planilla_tuberia_excel import __file__ as excel_mod

        mod_src = Path(excel_mod).read_text(encoding="utf-8")
        self.assertNotIn("load_workbook", mod_src)
        self.assertNotIn("Planilla_Tuberia_original.xlsm", mod_src)

        raw = build_planilla_tuberia_xlsx(
            planilla={"tipo": "ALCANTARILLA", "meta_cabecera": {}},
            calculo={"cartera": {"filas": []}},
            vacia=True,
        )
        self.assertEqual(raw[:2], b"PK")
        self.assertNotIn(b"macroEnabled", raw)
        self.assertNotIn(b"vbaProject", raw)
        wb = load_workbook(io.BytesIO(raw))
        self.assertIn("planilla", wb.sheetnames)
        self.assertIsNone(wb.vba_archive)
        self.assertGreaterEqual(len(wb["planilla"]._charts), 1)
        self.assertEqual(wb["planilla"]["G5"].value, "INFORMACION DEL CONTRATO")


class TestPdfBloquesInventario(unittest.TestCase):
    def test_cabecera_y_graficos_vacios(self):
        from topografia_planilla_tuberia_pdf import (
            html_bloque_graficos_pdf,
            html_cabecera_planilla_tuberia,
        )

        cab = html_cabecera_planilla_tuberia(
            contrato={"numero": "CTO-1", "objeto": "Obra demo", "contratista": "ACME"},
            meta={},
            titulo=TITULO_ALC,
        )
        self.assertIn("INFORMACION DEL CONTRATO", cab)
        self.assertIn("CTO-1", cab)
        self.assertIn(TITULO_ALC, cab)

        graf = html_bloque_graficos_pdf({})
        # Sección típica = PNG tipado del XLSM; perfil = SVG (aunque sin series).
        self.assertIn("data:image/png", graf)
        self.assertIn("graficos-wrap", graf)
        self.assertTrue(
            "Filtro" in graf or "Tuber" in graf or "GRAFICO" in graf
        )

        graf_fil = html_bloque_graficos_pdf({"seccion_tipica": {"tipo": "FILTRO"}})
        graf_alc = html_bloque_graficos_pdf({"seccion_tipica": {"tipo": "ALCANTARILLA"}})
        self.assertIn("Filtro", graf_fil)
        self.assertIn("Tuber", graf_alc)
        self.assertNotEqual(graf_fil, graf_alc)

    def test_franja_tramo_tablas_estructura(self):
        from topografia_planilla_tuberia_pdf import html_franja_tramo_tuberia

        html = html_franja_tramo_tuberia(
            planilla={
                "pk_id": "PK-01",
                "costado": "Izquierdo",
                "diametro_m": 0.9,
                "espesor_m": 0.075,
                "material": "Concreto",
                "tipo": "ALCANTARILLA",
                "ancho_excavacion_m": 1.8,
                "relacion_atraque": "1:3",
            },
            calculo={
                "seccion": {
                    "area_tuberia_m2": 0.636,
                    "area_1_m2": 0.45,
                    "area_2_m2": 0.32,
                    "altura_relleno_m": 0.8,
                    "cama_triturado_m": 0.1,
                },
                "cartera": {"totales": {"abscisa_inicial": 100.0, "abscisa_final": 120.0}},
            },
            meta={
                "abscisa_inicial": 100.0,
                "abscisa_final": 120.0,
                "norte_abs_inicial": 1.0,
                "este_abs_inicial": 2.0,
                "norte_abs_final": 3.0,
                "este_abs_final": 4.0,
            },
        )
        # Tres tablas (Abs/PK, geo+tubo, áreas/sección) — no texto corrido.
        self.assertGreaterEqual(html.count("<table"), 3)
        for label in (
            "Abs Inicial",
            "Abs Final",
            "PK_ID",
            "Costado",
            "Abscisa Inicial",
            "Norte Abs Inicial",
            "Este Abs Incial",
            "Abscisa Final",
            "Norte Abs Final",
            "Este Abs Final",
            "D. TUBERÍA",
            "ESP. TUBERÍA",
            "AREA TUBERÍA",
            "MATERIAL",
            "TIPO DE RED",
            "Area 1",
            "Area 2",
            "Altura Atraque",
            "Altura Relleno",
            "Cama Triturado",
            "Anc. Excavación",
        ):
            self.assertIn(label, html)
        self.assertIn("PK-01", html)
        self.assertIn("Izquierdo", html)
        self.assertIn("ALCANTARILLA", html)
        self.assertIn("#BDD7EE", html)  # franja Abs/PK
        self.assertIn("#D9D9D9", html)  # cabeceras geo/params
        self.assertNotIn("θ=", html)
        self.assertNotIn(" · ", html)
        # Valores de georref inicio/fin deben aparecer tal cual en el HTML.
        self.assertIn("1.000", html)  # norte_abs_inicial
        self.assertIn("2.000", html)  # este_abs_inicial
        self.assertIn("3.000", html)  # norte_abs_final
        self.assertIn("4.000", html)  # este_abs_final


class TestExcelGeorrefInicioFin(unittest.TestCase):
    def test_excel_muestra_ambos_pares_coordenadas(self):
        from openpyxl import load_workbook
        import io

        raw = build_planilla_tuberia_xlsx(
            planilla={
                "tipo": "ALCANTARILLA",
                "diametro_m": 0.9,
                "espesor_m": 0.05,
                "ancho_excavacion_m": 1.5,
                "relacion_atraque": "1:3",
                "norte_ref": 111.1,  # no debe preferirse sobre meta
                "este_ref": 222.2,
                "meta_cabecera": {
                    "norte_abs_inicial": 1000001.25,
                    "este_abs_inicial": 900002.5,
                    "norte_abs_final": 1000100.75,
                    "este_abs_final": 900150.0,
                },
            },
            calculo={"cartera": {"filas": []}},
        )
        ws = load_workbook(io.BytesIO(raw))["planilla"]
        self.assertEqual(ws["C13"].value, 1000001.25)
        self.assertEqual(ws["D13"].value, 900002.5)
        self.assertEqual(ws["F13"].value, 1000100.75)
        self.assertEqual(ws["G13"].value, 900150.0)

    def test_excel_fallback_inicio_a_norte_este_ref(self):
        from openpyxl import load_workbook
        import io

        raw = build_planilla_tuberia_xlsx(
            planilla={
                "tipo": "FILTRO",
                "norte_ref": 55.5,
                "este_ref": 66.6,
                "meta_cabecera": {},
            },
            calculo={"cartera": {"filas": []}},
            vacia=True,
        )
        ws = load_workbook(io.BytesIO(raw))["planilla"]
        self.assertEqual(ws["C13"].value, 55.5)
        self.assertEqual(ws["D13"].value, 66.6)
        self.assertIn(ws["F13"].value, (None, ""))
        self.assertIn(ws["G13"].value, (None, ""))


class TestExcelDescuentosFijosPorTipo(unittest.TestCase):
    def _ws(self, tipo: str):
        from openpyxl import load_workbook
        import io
        raw = build_planilla_tuberia_xlsx(
            planilla={
                "tipo": tipo,
                "diametro_m": 0.9,
                "espesor_m": 0.05,
                "ancho_excavacion_m": 1.5,
                "relacion_atraque": "1:3",
                "meta_cabecera": {"cama_triturado_m": 0.1},
            },
            calculo={"cartera": {"filas": []}},
        )
        return load_workbook(io.BytesIO(raw))["planilla"]

    def test_descuentos_fijos_por_tipo(self):
        ws_alc = self._ws("ALCANTARILLA")
        self.assertEqual(ws_alc["G48"].value, "=N46")
        self.assertEqual(ws_alc["G49"].value, "=N47")
        self.assertEqual(ws_alc["I46"].value, "Area 1")
        self.assertEqual(ws_alc["I47"].value, "Area 2")
        self.assertIn(ws_alc["I45"].value, (None, ""))
        self.assertNotIn("$F$1=$P$1", str(ws_alc["G48"].value))

        ws_fil = self._ws("FILTRO")
        self.assertEqual(ws_fil["G48"].value, "=N45")
        self.assertEqual(ws_fil["G49"].value, 0)
        self.assertEqual(ws_fil["I45"].value, "Tubería Filtro")
        self.assertIn(ws_fil["I46"].value, (None, ""))
        self.assertIn(ws_fil["I47"].value, (None, ""))
        # Cartera tipada: H ALC = E15+F15; H FIL = E-F
        self.assertIn("$E$15+$F$15", str(ws_alc["H17"].value))
        self.assertIn("E17-F17", str(ws_fil["H17"].value))

if __name__ == "__main__":
    unittest.main()
