"""Resumen de conciliación 4 columnas + Excel CC-SUB-001 alineado al PDF."""
from __future__ import annotations

import sys
from io import BytesIO
from types import ModuleType, SimpleNamespace

from openpyxl import load_workbook

from corte_sub_conciliacion import (
    build_resumen_conciliacion_4cols,
    enriquecer_items_bloques,
    filtrar_items_con_cantidades,
    valor_por_cantidad_vu,
)


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]

    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake

    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)

    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    return inf


TRIBUTOS = {
    "administracion": 10,
    "imprevistos": 5,
    "utilidad": 7,
    "iva": {"porcentaje": 19},
}


def _sample_items():
    items = enriquecer_items_bloques(
        [
            {
                "item_numero": "1.1",
                "item_descripcion": "Excavación",
                "unidad": "M3",
                "cantidad": 10,
                "vlr_unitario_sub": 1000,
                "capitulo": "I",
            }
        ],
        cant_actualizadas={"1.1": 100},
        cant_acum_anterior={"1.1": 40},
    )
    return filtrar_items_con_cantidades(items)


def test_resumen_4cols_reglas_basicas():
    items = _sample_items()
    r = build_resumen_conciliacion_4cols(
        items=items,
        tributos=TRIBUTOS,
        anticipo=50_000,
        amortizacion_pct=20,
        amortizado_anterior=10_000,
        otros_presente=2_000,
        aiu_otros_anterior={
            "costo_directo": 40_000,
            "valor_administracion": 4_000,
            "valor_imprevistos": 2_000,
            "valor_utilidad": 2_800,
            "valor_iva_utilidad": 532,
            "costo_directo_mas_aiu": 49_332,
            "total_otros": 1_000,
            "amortizacion_presente": 10_000,
        },
    )
    by = {ln["key"]: ln["valores"] for ln in r["lineas"]}
    assert by["cd"]["presente"] == valor_por_cantidad_vu(10, 1000)
    assert by["cd"]["actualizadas"] == valor_por_cantidad_vu(100, 1000)
    assert by["cd"]["acumulado"] == valor_por_cantidad_vu(50, 1000)
    assert by["cd"]["saldo"] == by["cd"]["actualizadas"] - by["cd"]["acumulado"]
    # A presente = 10% CD presente
    assert by["a"]["presente"] == round(by["cd"]["presente"] * 0.10)
    # Acumulado A = ant + presente
    assert by["a"]["acumulado"] == 4_000 + by["a"]["presente"]
    # Amortización: Act=anticipo, Pres=amort corte, Acum=ant+pres, Saldo=anticipo-acum
    assert by["amort"]["actualizadas"] == 50_000
    assert by["amort"]["acumulado"] == by["amort"]["presente"] + 10_000
    assert by["amort"]["saldo"] == 50_000 - by["amort"]["acumulado"]
    # Otros: solo presente y acumulado
    assert by["otros"]["actualizadas"] is None
    assert by["otros"]["saldo"] is None
    assert by["otros"]["presente"] == 2_000
    assert by["otros"]["acumulado"] == 3_000
    # Gran total presente = sub_amort + otros
    assert by["gran_total"]["presente"] == by["sub_amort"]["presente"] + 2_000


def test_resumen_4cols_segundo_corte_amort_acumula():
    """Simula 2.º corte: amort acumulada crece y saldo de anticipo baja."""
    items = _sample_items()
    r1 = build_resumen_conciliacion_4cols(
        items=items,
        tributos=TRIBUTOS,
        anticipo=100_000,
        amortizacion_pct=30,
        amortizado_anterior=0,
        otros_presente=0,
    )
    amort1 = r1["amortizacion"]["amortizacion_presente"]
    assert amort1 > 0
    r2 = build_resumen_conciliacion_4cols(
        items=items,
        tributos=TRIBUTOS,
        anticipo=100_000,
        amortizacion_pct=30,
        amortizado_anterior=amort1,
        otros_presente=0,
        aiu_otros_anterior={
            "costo_directo": r1["aiu_presente"]["costo_directo"],
            "valor_administracion": r1["aiu_presente"]["valor_administracion"],
            "valor_imprevistos": r1["aiu_presente"]["valor_imprevistos"],
            "valor_utilidad": r1["aiu_presente"]["valor_utilidad"],
            "valor_iva_utilidad": r1["aiu_presente"]["valor_iva_utilidad"],
            "costo_directo_mas_aiu": r1["aiu_presente"]["costo_directo_mas_aiu"],
            "total_otros": 0,
            "amortizacion_presente": amort1,
        },
    )
    by2 = {ln["key"]: ln["valores"] for ln in r2["lineas"]}
    assert by2["amort"]["acumulado"] == amort1 + by2["amort"]["presente"]
    assert by2["amort"]["saldo"] == 100_000 - by2["amort"]["acumulado"]
    assert by2["amort"]["saldo"] < 100_000 - amort1 or by2["amort"]["presente"] == 0


def test_excel_encabezado_y_resumen_4cols():
    inf = _import_informes_with_stubs()
    items = _sample_items()
    r4 = build_resumen_conciliacion_4cols(
        items=items,
        tributos=TRIBUTOS,
        anticipo=50_000,
        amortizacion_pct=20,
        amortizado_anterior=0,
        otros_presente=1_500,
    )
    raw = inf._corte_sub_001_excel_bytes(
        {"numero": "C-1", "contratista": "Ctor", "nit": "900", "interventoria": "Int"},
        {"razon_social": "Sub SA", "nombre_contacto": "Ana"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-15"},
        items,
        float(r4["cd_bloques"]["presente"]),
        "Tester",
        "Desarrollador",
        {},
        resumen_4cols=r4,
        otros_conceptos=[
            {
                "descripcion": "Flete",
                "unidad": "UN",
                "cantidad": 1,
                "valor_unitario": 1500,
                "costo_total": 1500,
            }
        ],
    )
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    assert ws.title == "CC-SUB-001"
    assert ws.page_setup.orientation == "landscape"
    assert ws.cell(1, 4).value == "INFORME CORTE DE SUB CONTRATISTA"
    assert ws.cell(1, 10).value == "CC-SUB-001"
    assert ws.cell(3, 1).value == "CONTRATO"
    assert ws.cell(7, 6).value == "ACTUALIZADAS"
    # Resumen
    found = False
    for r in range(1, 40):
        if ws.cell(r, 1).value == "RESUMEN DE CONCILIACIÓN":
            found = True
            assert ws.cell(r + 1, 7).value == "Actualizadas"
            assert ws.cell(r + 1, 9).value == "Presente acta"
            # CD row uses formulas referencing item totals
            assert str(ws.cell(r + 2, 7).value).startswith("=")
            break
    assert found


def test_pdf_html_incluye_resumen_4cols():
    inf = _import_informes_with_stubs()
    items = _sample_items()
    r4 = build_resumen_conciliacion_4cols(
        items=items,
        tributos=TRIBUTOS,
        anticipo=0,
        amortizacion_pct=None,
        otros_presente=0,
    )
    html = inf._html_cc_sub_v1_plain(
        {"numero": "C-1", "contratista": "Ctor", "interventoria": "Int"},
        {"razon_social": "Sub"},
        {"consecutivo": 1, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-15"},
        items,
        float(r4["cd_bloques"]["presente"]),
        "T",
        "D",
        aiu_resumen=r4["aiu_presente"],
        amortizacion=r4["amortizacion"],
        resumen_4cols=r4,
    )
    assert "Actualizadas" in html
    assert "Presente acta" in html
    assert "Acumulado" in html
    assert "Saldo" in html
    assert "Costo Directo + AIU" in html
    assert "Gran total" in html
