"""CC-MES-002: descripción completa + Periodo (en lugar de Fecha acta)."""
from __future__ import annotations

import sys
from types import ModuleType, SimpleNamespace

from openpyxl import Workbook


def _import_informes_with_stubs():
    if "informes" in sys.modules:
        return sys.modules["informes"]

    stubs = {
        "main": SimpleNamespace(
            get_current_user=lambda: None,
            get_current_user_optional=lambda: None,
        ),
        "mail_smtp": SimpleNamespace(try_send_text_email=lambda *a, **k: None),
    }
    saved = {}
    for name, mod in stubs.items():
        saved[name] = sys.modules.get(name)
        fake = ModuleType(name)
        for k, v in vars(mod).items():
            if not k.startswith("_"):
                setattr(fake, k, v)
        sys.modules[name] = fake

    for extra in ("passlib", "passlib.context", "jose", "python_jose"):
        if extra not in sys.modules:
            sys.modules[extra] = ModuleType(extra)

    try:
        import informes as inf  # noqa: WPS433
    finally:
        for name, prev in saved.items():
            if prev is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = prev
    inf._memoria_prepare_media_uri = lambda url: ""  # type: ignore
    inf._memoria_prefetch_media_uris = lambda urls: {
        str(u or "").strip(): "" for u in urls if str(u or "").strip()
    }  # type: ignore
    return inf


def test_periodo_acta_admin_fmt_igual_admin_actas():
    inf = _import_informes_with_stubs()
    assert (
        inf._periodo_acta_admin_fmt(
            {"fecha_inicio": "2025-03-01", "fecha_fin": "2025-03-31T00:00:00"}
        )
        == "2025-03-01 → 2025-03-31"
    )
    assert inf._periodo_acta_admin_fmt({}) == "—"
    assert inf._periodo_acta_admin_fmt(None) == "—"
    assert inf._periodo_acta_admin_fmt({"fecha_inicio": "2025-01-15"}) == "2025-01-15 → —"


def test_conc_meta_cc_mes_002_usa_periodo_no_fecha_acta():
    inf = _import_informes_with_stubs()
    meta = inf._conc_meta_cc_mes_002(
        "ICCU-CTO-1614-2025",
        {
            "numero_rpo": 1,
            "consecutivo": 1,
            "fecha_inicio": "2025-01-01",
            "fecha_fin": "2025-01-31",
        },
    )
    assert meta["codigo"] == "CC-MES-002"
    labels = [c[0] for c in meta["cells"]]
    assert "FECHA ACTA" not in labels
    assert "PERÍODO" in labels
    assert meta["cells"][3] == ("PERÍODO", "2025-01-01 → 2025-01-31")


def test_html_memoria_muestra_descripcion_completa_y_periodo():
    inf = _import_informes_with_stubs()
    desc = (
        "SUMINISTRO E INSTALACIÓN DE GEOCELDA DE HDPE DE 150 MM DE ALTURA, "
        "INCLUYE SUMINISTRO, ANCLAJE, EXPANSIÓN, FIJACIÓN Y RELLENO CON MATERIAL "
        "SELECCIONADO SEGÚN ESPECIFICACIÓN TÉCNICA DEL LISTADO DE PRECIOS"
    )
    contrato = {"numero": "ICCU-CTO-1614-2025", "logo_contratista": None}
    item_info = {
        "item_numero": "NP-04",
        "item_descripcion": desc,
        "unidad": "M2",
        "capitulo": "1.ACTIVIDADES PRELIMINARES",
    }
    regs = [
        {
            "numero_registro": 1,
            "cantidad_total": 1,
            "capitulo": item_info["capitulo"],
            "item_numero": "NP-04",
            "item_descripcion": desc,
            "unidad": "M2",
        }
    ]
    conc_meta = inf._conc_meta_cc_mes_002(
        contrato["numero"],
        {
            "numero_rpo": 1,
            "consecutivo": 1,
            "fecha_inicio": "2025-02-01",
            "fecha_fin": "2025-02-28",
        },
    )
    sub, corte = inf._sub_corte_dummy_memoria()
    html = inf._html_memoria_item(
        contrato,
        sub,
        corte,
        item_info,
        regs,
        "QA",
        "QA",
        conc_meta=conc_meta,
        pie_fotos_contexto="Acta RPO 1",
        estilo_formato_codigo="CC-MES-002",
    )
    assert "FECHA ACTA" not in html
    assert "PERÍODO" in html
    assert "2025-02-01 → 2025-02-28" in html
    assert "mem002-item-desc" in html
    assert desc in html
    # Sin truncar a un prefijo: el texto completo del listado está en el HTML.
    assert html.count(desc) >= 1


def test_excel_memoria_descripcion_wrap_y_altura_y_periodo():
    inf = _import_informes_with_stubs()
    desc = (
        "EXCAVACION MANUAL EN MATERIAL COMUN H=0.0-2.0M (INCLUYE RETIRO DE SOBRANTES "
        "A UNA DISTANCIA MENOR DE 5 KM) SEGUN LISTADO DE PRECIOS VIGENTE DEL CONTRATO"
    )
    contrato = {"numero": "ICCU-CTO-1614-2025"}
    item_info = {
        "item_numero": "NP-09",
        "item_descripcion": desc,
        "unidad": "M3",
        "capitulo": "1.ACTIVIDADES",
    }
    regs = [
        {
            "numero_registro": 10,
            "cantidad_total": 2.5,
            "capitulo": "1.ACTIVIDADES",
            "item_numero": "NP-09",
            "longitud": 1,
            "ancho": 1,
            "espesor": 1,
            "cantidad": 2.5,
        }
    ]
    conc_meta = inf._conc_meta_cc_mes_002(
        contrato["numero"],
        {
            "numero_rpo": 1,
            "consecutivo": 1,
            "fecha_inicio": "2025-01-01",
            "fecha_fin": "2025-01-31",
        },
    )
    wb = Workbook()
    ws = wb.active
    sub, corte = inf._sub_corte_dummy_memoria()
    inf._fill_memoria_excel_ws(
        ws,
        contrato,
        sub,
        corte,
        item_info,
        regs,
        {},
        conc_meta=conc_meta,
        pie_fotos_contexto="Acta RPO 1",
        aprobo_interventoria_desde_config=True,
    )
    # Periodo en banda (fila 4, celda B4 = 4º campo)
    assert ws["A4"].value == "PERÍODO"
    assert ws["B4"].value == "2025-01-01 → 2025-01-31"
    assert ws["E5"].value == desc
    assert ws["E5"].alignment.wrap_text is True
    assert (ws.row_dimensions[5].height or 0) >= 22


def test_excel_row_height_crece_con_descripcion_larga():
    inf = _import_informes_with_stubs()
    corto = inf._excel_row_height_wrapped_text("Corta")
    largo = inf._excel_row_height_wrapped_text("X" * 220)
    assert largo > corto


def test_html_encabezado_periodo_y_desc_en_fotos_y_graficos():
    """El encabezado repetido en fotos/gráficos también lleva Periodo y descripción completa."""
    inf = _import_informes_with_stubs()
    desc = (
        "SUMINISTRO E INSTALACION DE TUBERIA EN PVC CORRUGADO DOBLE PARED DE DIAMETRO "
        "NOMINAL 24 PULGADAS TIPO S INCLUYE UNIONES SELLOS Y PRUEBAS SEGUN LISTADO"
    )
    contrato = {"numero": "ICCU-CTO-1614-2025", "logo_contratista": None}
    item_info = {
        "item_numero": "NP-08",
        "item_descripcion": desc,
        "unidad": "ML",
        "capitulo": "2. DRENAJE",
    }
    regs = [
        {
            "numero_registro": 1,
            "cantidad_total": 12.5,
            "capitulo": "2. DRENAJE",
            "item_numero": "NP-08",
            "item_descripcion": desc,
            "unidad": "ML",
            "foto_url": "https://cdn.example/f1.jpg",
            "foto_numero": 1,
            "grafico_url": "https://cdn.example/g1.png",
            "grafico_numero": 1,
        }
    ]
    conc_meta = inf._conc_meta_cc_mes_002(
        contrato["numero"],
        {
            "numero_rpo": 1,
            "consecutivo": 1,
            "fecha_inicio": "2025-01-01",
            "fecha_fin": "2025-01-31",
        },
    )
    sub, corte = inf._sub_corte_dummy_memoria()
    html = inf._html_memoria_item(
        contrato,
        sub,
        corte,
        item_info,
        regs,
        "QA",
        "QA",
        conc_meta=conc_meta,
        pie_fotos_contexto="Acta RPO 1",
        estilo_formato_codigo="CC-MES-002",
    )
    assert "REGISTRO FOTOGRÁFICO" in html
    assert "REGISTRO GRÁFICO" in html
    assert html.count("PERÍODO") >= 3
    assert html.count("2025-01-01 → 2025-01-31") >= 3
    assert html.count("mem002-item-desc") >= 3
    assert html.count(desc) >= 3
    assert "FECHA ACTA" not in html
