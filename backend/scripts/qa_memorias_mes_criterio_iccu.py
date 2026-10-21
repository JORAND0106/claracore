#!/usr/bin/env python3
"""Memorias CC-MES-002 del acta real ICCU-CTO-1614-2025, Acta RPO 1.

Comprueba que la vista previa, el PDF, el Excel y las pestañas del libro
integral salen de los mismos ítems con ejecución que el informe CC-MES-001.
"""
from __future__ import annotations

import os
import sys
import time
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))


def _prime_env() -> None:
    url = os.environ.get("SUPABASE_URL") or os.environ.get("VITE_SUPABASE_URL") or ""
    key = (
        os.environ.get("SUPABASE_KEY")
        or os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
        or os.environ.get("SUPABASE_ANON_KEY")
        or os.environ.get("VITE_SUPABASE_ANON_KEY")
        or ""
    )
    if not url or not key:
        env_path = ROOT / "frontend" / ".env.development"
        for line in env_path.read_text().splitlines():
            if line.startswith("VITE_SUPABASE_URL=") and not url:
                url = line.split("=", 1)[1].strip().strip('"')
            elif line.startswith("VITE_SUPABASE_ANON_KEY=") and not key:
                key = line.split("=", 1)[1].strip().strip('"')
    if not url or not key:
        raise SystemExit("Faltan SUPABASE_URL / SUPABASE_KEY")
    os.environ["SUPABASE_URL"] = url
    os.environ["SUPABASE_KEY"] = key


_prime_env()

import main  # noqa: E402,F401

from excel_formula_eval import FormulaBook  # noqa: E402
import informe_mes_items as imi  # noqa: E402
import informes as inf  # noqa: E402
from informes import (  # noqa: E402
    _cc_mes_002_acta_completo_ctx,
    _cc_mes_002_acta_completo_excel_bytes,
    _cc_mes_002_pdf_completo_bytes,
    _cc_mes_002_item_excel_bytes,
    _cc_mes_integral_excel_bytes,
    _contexto_acta_mes_conciliacion,
    _row,
    _sb,
    ccd_items_conciliacion_acta,
    matriz_params_contrato,
    pdf_cc_mes_002_acta,
)
from openpyxl import load_workbook  # noqa: E402

CONTRATO_NUMERO = "ICCU-CTO-1614-2025"
CONTRATO_ID = 3
ACTA_ID = 620
ACTA = {
    "id": ACTA_ID,
    "contrato_id": CONTRATO_ID,
    "consecutivo": 1,
    "numero_rpo": "1",
    "fecha_inicio": None,
    "fecha_fin": None,
    "tipo_grupo": "RPO",
}
CONTRATO = {
    "id": CONTRATO_ID,
    "numero": CONTRATO_NUMERO,
    "objeto": "",
    "contratista": "",
    "nit": "",
    "interventoria": "",
    "logo_contratista": None,
    "aiu": None,
    "iva": None,
    "anticipo": None,
    "amortizacion_pct": None,
}
PMT = (
    "1. PERSONAL OPERATIVO",
    "2. SEÑALIZACIÓN Y CANALIZACIÓN",
    "3. COMUNICACIONES",
    "4. OPERACIÓN PMT",
)


def _paginar(hacer):
    rows = []
    offset = 0
    while True:
        batch = hacer(offset, offset + 999).execute().data or []
        rows.extend(batch)
        if len(batch) < 1000:
            break
        offset += 1000
    return rows


def _listado_vivo():
    try:
        return _paginar(
            lambda ini, fin: _sb.table("listado_precios")
            .select(imi._SELECT_LISTADO)
            .eq("contrato_id", CONTRATO_ID)
            .order("id")
            .range(ini, fin)
        )
    except Exception as exc:
        print(f"listado_precios no legible: {exc}", flush=True)
        return []


def _presupuesto():
    return _paginar(
        lambda ini, fin: _sb.table("presupuesto")
        .select("id, capitulo, item, descripcion, und, cant_total, tipo_ejecucion, dado_de_baja")
        .eq("contrato_id", CONTRATO_ID)
        .order("id")
        .range(ini, fin)
    )


def _registros():
    return _paginar(
        lambda ini, fin: _sb.table("so_registros")
        .select(
            "id, numero_registro, capitulo, item_numero, item_descripcion, unidad, "
            "cantidad_total, vlr_unitario, costo_directo, nivel1_estado, nivel2_estado, nivel3_estado"
        )
        .eq("contrato_id", CONTRATO_ID)
        .eq("acta_rpo_id", ACTA_ID)
        .order("numero_registro")
        .range(ini, fin)
    )


def _sintetizar_listado(presupuesto, registros):
    """Presupuesto de obra y claves del acta, marcadas Aprobado.

    La clave anon no lee listado_precios. El estado Aprobado es el supuesto
    de esta síntesis; un ítem se pasa a Pendiente más abajo para probar el corte.
    Los capítulos de PMT se excluyen igual, aunque queden Aprobado.
    """
    filas = []
    n = 0
    for raw in presupuesto:
        if raw.get("dado_de_baja") is True:
            continue
        if str(raw.get("tipo_ejecucion") or "").strip() != "Presupuesto de Obra":
            continue
        n += 1
        filas.append(
            {
                "id": int(raw.get("id") or n),
                "capitulo": raw.get("capitulo"),
                "item_numero": raw.get("item"),
                "descripcion": raw.get("descripcion") or "",
                "unidad": raw.get("und") or "",
                "precio_unitario": 0,
                "estado_precio": "Aprobado",
                "especificacion_tecnica": "",
            }
        )
    conocidas = {imi.clave_item_exacta(f["capitulo"], f["item_numero"]) for f in filas}
    extra = 0
    for raw in registros:
        clave = imi.clave_item_exacta(raw.get("capitulo"), raw.get("item_numero"))
        if not clave[0] or not clave[1] or clave in conocidas:
            continue
        conocidas.add(clave)
        extra += 1
        n += 1
        filas.append(
            {
                "id": 1_000_000 + n,
                "capitulo": clave[0],
                "item_numero": clave[1],
                "descripcion": raw.get("item_descripcion") or "",
                "unidad": raw.get("unidad") or "",
                "precio_unitario": raw.get("vlr_unitario") or 0,
                "estado_precio": "Aprobado",
                "especificacion_tecnica": "",
            }
        )
    return filas, extra


def _row_visible(table, select, **eq):
    if table == "actas" and int(eq.get("id") or 0) == ACTA_ID:
        return dict(ACTA)
    if table == "contratos" and int(eq.get("id") or 0) == CONTRATO_ID:
        return dict(CONTRATO)
    return _row(table, select, **eq)


def _claves_ejecucion(items):
    return {
        (it["capitulo"], it["item_numero"])
        for it in items
        if it.get("tiene_registros_acta")
    }


def _hojas_memoria(wb):
    return [name for name in wb.sheetnames if name != "CC-MES-001"]


def _formulas_presente(ws):
    out = []
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("=ROUND(") and "!H" in value:
                out.append((cell.row, cell.column, value))
    return out


def main() -> int:
    t0 = time.time()
    registros = _registros()
    presupuesto = _presupuesto()
    vivo = _listado_vivo()
    if vivo:
        filas = vivo
        fuente = "listado_precios"
        extras = 0
    else:
        filas, extras = _sintetizar_listado(presupuesto, registros)
        fuente = "presupuesto_de_obra_mas_claves_del_acta"
    pmt_regs = [r for r in registros if imi.es_capitulo_pmt(r.get("capitulo"))]
    print(
        f"registros={len(registros)} pmt_en_acta={len(pmt_regs)} "
        f"presupuesto={len(presupuesto)} listado={len(filas)} fuente={fuente} extras_acta={extras}",
        flush=True,
    )
    try:
        _campo, niveles = matriz_params_contrato(_sb, CONTRATO_ID)
        nivel = max(niveles) if niveles else None
    except Exception as exc:
        print(f"matriz_params: {exc}; nivel=None", flush=True)
        nivel = None
    print(f"nivel_aprobacion={nivel}", flush=True)

    user = {"nombre": "QA", "apellidos": "Memorias", "cargo_nombre": "Verificación", "id": 1}
    errores = []

    def _listado(_sb, _cid):
        return filas

    with (
        patch.object(inf, "_row", side_effect=_row_visible),
        patch.object(inf, "_perm_informes_ccd"),
        patch.object(inf, "_perm_informes_ccd_lectura"),
        patch.object(inf, "_get_firma_cfg_para_documento", return_value={}),
        patch.object(inf, "_firma_data_uri_para_slot_contexto", return_value=None),
        patch.object(inf, "_memoria_prepare_media_uri", return_value=""),
        patch.object(inf, "_memoria_media_img_html", return_value=""),
        patch.object(imi, "fetch_listado_contrato", side_effect=_listado),
    ):
        ctx = _contexto_acta_mes_conciliacion(CONTRATO_ID, ACTA_ID, user, nivel_aprobacion=nivel)
        items = list(ctx.get("items") or [])
        claves = _claves_ejecucion(items)
        print(f"informe={len(items)} con_ejecucion={len(claves)} en {time.time() - t0:.1f}s", flush=True)
        lista = ccd_items_conciliacion_acta(CONTRATO_ID, ACTA_ID, nivel, user)
        nums_preview = {(it.get("item_numero") or "").strip() for it in lista["items"]}
        nums_informe = {it["item_numero"] for it in items if it.get("tiene_registros_acta")}
        if nums_preview != nums_informe:
            errores.append(
                f"vista previa {sorted(nums_preview)} != informe {sorted(nums_informe)}"
            )

        item_pdf = sorted(claves, key=lambda c: (c[0], c[1]))[0]
        resp = pdf_cc_mes_002_acta(CONTRATO_ID, ACTA_ID, item_pdf[1], nivel, user)
        cuerpo = bytes(resp.body or b"")
        if not cuerpo.startswith(b"%PDF"):
            errores.append("el PDF individual no empieza por %PDF")
        print(f"pdf item={item_pdf[1]!r} bytes={len(cuerpo)}", flush=True)
        ctx_pdf = _cc_mes_002_acta_completo_ctx(
            CONTRATO_ID, ACTA_ID, user, nivel_aprobacion=nivel
        )
        if ctx_pdf["n_items"] != len(nums_informe):
            errores.append(
                f"PDF completo ítems={ctx_pdf['n_items']} informe={len(nums_informe)}"
            )
        if set(ctx_pdf["numeros"]) != nums_informe:
            errores.append("PDF completo no trae los mismos códigos que el informe")
        unido = ctx_pdf["html_joined"]
        for nombre in PMT:
            if nombre in unido:
                errores.append(f"PDF completo menciona el capítulo PMT {nombre}")
        pdf_todos = _cc_mes_002_pdf_completo_bytes(ctx_pdf)
        if not pdf_todos.startswith(b"%PDF"):
            errores.append("el PDF de todas las memorias no empieza por %PDF")
        print(
            f"pdf completo items={ctx_pdf['n_items']} regs={ctx_pdf['n_regs']} bytes={len(pdf_todos)}",
            flush=True,
        )

        pmt_item = next((r for r in pmt_regs if str(r.get("item_numero") or "").strip()), None)
        if pmt_item is None:
            errores.append("el acta no trae registros de capítulos PMT para probar la exclusión")
        else:
            try:
                pdf_cc_mes_002_acta(
                    CONTRATO_ID, ACTA_ID, str(pmt_item.get("item_numero")), nivel, user
                )
            except Exception as exc:
                codigo = getattr(exc, "status_code", None)
                if codigo != 404:
                    # El mismo código puede existir también fuera de PMT.
                    otros = [
                        c for c in claves if c[1] == str(pmt_item.get("item_numero") or "").strip()
                    ]
                    if otros and codigo is None:
                        errores.append(f"PDF PMT no respondió 404 ni un PDF del ítem del informe: {exc}")
                    elif not otros:
                        errores.append(f"PDF de ítem PMT {pmt_item.get('item_numero')!r}: {exc}")
                else:
                    print(
                        f"pdf pmt item={pmt_item.get('item_numero')!r} cap={pmt_item.get('capitulo')!r} -> 404",
                        flush=True,
                    )
            else:
                otros = [c for c in claves if c[1] == str(pmt_item.get("item_numero") or "").strip()]
                if not otros:
                    errores.append(
                        f"se generó PDF para el ítem PMT {pmt_item.get('item_numero')!r}"
                    )
                else:
                    print(
                        f"pdf código {pmt_item.get('item_numero')!r} se genera por el capítulo del informe {otros}",
                        flush=True,
                    )

        print("excel completo…", flush=True)
        raw_todos = _cc_mes_002_acta_completo_excel_bytes(
            CONTRATO_ID, ACTA_ID, user, nivel_aprobacion=nivel
        )
        wb_todos = load_workbook(BytesIO(raw_todos))
        nums_excel = {str(ws["B5"].value or "").strip() for ws in wb_todos.worksheets}
        if nums_excel != nums_informe:
            errores.append(f"excel completo {sorted(nums_excel)} != informe {sorted(nums_informe)}")
        texto_todos = " ".join(
            str(c.value)
            for ws in wb_todos.worksheets
            for row in ws.iter_rows(max_row=6, max_col=10)
            for c in row
            if c.value is not None
        )
        for nombre in PMT:
            if nombre in texto_todos:
                errores.append(f"excel completo menciona el capítulo PMT {nombre}")

        print("libro integral…", flush=True)
        t1 = time.time()
        raw_int = _cc_mes_integral_excel_bytes(
            CONTRATO_ID, ACTA_ID, user, nivel_aprobacion=nivel
        )
        print(f"integral bytes={len(raw_int)} en {time.time() - t1:.1f}s", flush=True)
        wb = load_workbook(BytesIO(raw_int))
        if wb.sheetnames[0] != "CC-MES-001":
            errores.append(f"primera hoja {wb.sheetnames[0]!r}")
        hojas = _hojas_memoria(wb)
        if len(hojas) != len(claves):
            errores.append(f"pestañas={len(hojas)} claves_con_ejecucion={len(claves)}")
        vistas = set()
        for name in hojas:
            ws = wb[name]
            cap = str(ws["G4"].value or "").strip()
            num = str(ws["B5"].value or "").strip()
            if imi.es_capitulo_pmt(cap):
                errores.append(f"pestaña PMT {name} cap={cap}")
            clave = (cap, num)
            if clave not in claves:
                errores.append(f"pestaña sin ítem del informe: {clave}")
            vistas.add(clave)
        if vistas != claves:
            errores.append(f"faltan pestañas: {sorted(claves - vistas)[:8]}")
        formulas = _formulas_presente(wb["CC-MES-001"])
        if len(formulas) != len(claves):
            errores.append(f"fórmulas presente={len(formulas)} claves={len(claves)}")
        book = FormulaBook(wb)
        for row, col, formula in formulas:
            nombre = formula.split("'", 2)[1]
            if nombre not in hojas:
                errores.append(f"fórmula apunta a una hoja que no existe: {formula}")
            try:
                valor = book.eval_cell("CC-MES-001", row, col)
            except Exception as exc:
                errores.append(f"fórmula no calcula {formula}: {exc}")
                continue
            if isinstance(valor, str):
                errores.append(f"fórmula devolvió texto {formula} -> {valor}")
        print(
            f"preview={len(nums_preview)} excel_hojas={len(wb_todos.sheetnames)} "
            f"integral_memorias={len(hojas)} formulas={len(formulas)}",
            flush=True,
        )

        # El mismo listado, con un ítem ejecutado pasado a Pendiente.
        elegido = sorted(claves)[0]
        mutadas = []
        for fila in filas:
            copia = dict(fila)
            if imi.clave_item_exacta(copia.get("capitulo"), copia.get("item_numero")) == elegido:
                copia["estado_precio"] = "Pendiente"
            mutadas.append(copia)

        def _listado_mutado(_sb, _cid):
            return mutadas

        with patch.object(imi, "fetch_listado_contrato", side_effect=_listado_mutado):
            ctx2 = _contexto_acta_mes_conciliacion(
                CONTRATO_ID, ACTA_ID, user, nivel_aprobacion=nivel
            )
            lista2 = ccd_items_conciliacion_acta(CONTRATO_ID, ACTA_ID, nivel, user)
            try:
                _cc_mes_002_item_excel_bytes(
                    CONTRATO_ID, ACTA_ID, elegido[1], user, nivel_aprobacion=nivel
                )
                excel_mutado = "generado"
            except Exception as exc:
                excel_mutado = f"{getattr(exc, 'status_code', type(exc).__name__)}"
        claves2 = _claves_ejecucion(ctx2.get("items") or [])
        nums2 = {(it.get("item_numero") or "").strip() for it in lista2["items"]}
        print(
            f"tras Pendiente {elegido}: informe_sin_el={elegido not in claves2} "
            f"preview_sin_codigo={elegido[1] not in nums2} excel={excel_mutado}",
            flush=True,
        )
        if elegido in claves2:
            errores.append(f"el informe conservó {elegido} después de marcarlo Pendiente")
        if elegido[1] in nums2:
            otros = [c for c in claves2 if c[1] == elegido[1]]
            if not otros:
                errores.append(f"la vista previa conservó {elegido[1]} sin otro capítulo aprobado")
        if excel_mutado == "generado":
            otros = [c for c in claves2 if c[1] == elegido[1]]
            if not otros:
                errores.append(f"el Excel individual se generó para {elegido} en Pendiente")

    print(f"total {time.time() - t0:.1f}s errores={len(errores)}", flush=True)
    for err in errores:
        print(f"ERROR {err}", flush=True)
    return 1 if errores else 0


if __name__ == "__main__":
    raise SystemExit(main())
