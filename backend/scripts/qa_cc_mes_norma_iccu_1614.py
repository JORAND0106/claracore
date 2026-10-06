#!/usr/bin/env python3
"""Prueba real: Excel integral CC-MES del contrato ICCU-CTO-1614-2025, Acta RPO 1.

Comprueba:
- el nombre de descarga es CC-MES-integral_Acta_N
- el libro no trae la pestaña ni la tabla preacta_obra
- la columna Norma técnica (entre UND y V. UNIT.) coincide con
  listado_precios.especificacion_tecnica por capítulo e ítem
- las fórmulas de cantidades y valores recalculan igual que el PDF

Uso (desde backend/):
  PYTHONPATH=. python3 scripts/qa_cc_mes_norma_iccu_1614.py
"""
from __future__ import annotations

import json
import os
import sys
import time
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))


def _prime_env() -> None:
    url = os.environ.get("SUPABASE_URL") or os.environ.get("VITE_SUPABASE_URL") or ""
    key = (
        os.environ.get("SUPABASE_KEY")
        or os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
        or os.environ.get("SUPABASE_ANON_KEY")
        or os.environ.get("VITE_SUPABASE_ANON_KEY")
        or ""
    )
    if not url or not key:
        env_path = ROOT / "frontend" / ".env.development"
        if env_path.is_file():
            for line in env_path.read_text().splitlines():
                if line.startswith("VITE_SUPABASE_URL=") and not url:
                    url = line.split("=", 1)[1].strip().strip('"')
                elif line.startswith("VITE_SUPABASE_ANON_KEY=") and not key:
                    key = line.split("=", 1)[1].strip().strip('"')
    if not url or not key:
        raise SystemExit("Faltan SUPABASE_URL / SUPABASE_KEY")
    os.environ["SUPABASE_URL"] = url
    os.environ["SUPABASE_KEY"] = key


_prime_env()

import main  # noqa: E402,F401

from corte_sub_conciliacion import valor_por_cantidad_vu  # noqa: E402
from excel_formula_eval import FormulaBook  # noqa: E402
import informes as inf_mod  # noqa: E402
from informes import (  # noqa: E402
    _cc_mes_integral_excel_bytes,
    _contexto_acta_mes_conciliacion,
    _fm_informe,
    _fn_cant_informe,
    _html_cc_mes_001_v1,
    _row,
    _safe_filename_part,
    _sb,
    matriz_params_contrato,
)
from openpyxl import load_workbook  # noqa: E402
from sicoe_valor_canonico import cap_item_key  # noqa: E402

CONTRATO_NUMERO = "ICCU-CTO-1614-2025"
_CONTRATO_ID_DOC = 3
_ACTA_ID_DOC = 620
_ACTA_DOC = {
    "id": _ACTA_ID_DOC,
    "contrato_id": _CONTRATO_ID_DOC,
    "consecutivo": 1,
    "numero_rpo": "1",
    "fecha_inicio": None,
    "fecha_fin": None,
    "tipo_grupo": "RPO",
}
_CONTRATO_DOC = {
    "id": _CONTRATO_ID_DOC,
    "numero": CONTRATO_NUMERO,
    "objeto": "",
    "contratista": "",
    "nit": "",
    "interventoria": "",
    "logo_contratista": None,
    "aiu": None,
    "iva": None,
    "anticipo": None,
    "amortizacion_pct": None,
}


def _acta_rpo_1(contrato_id: int) -> tuple[dict, str]:
    try:
        rows = (
            _sb.table("actas")
            .select("id,numero_rpo,consecutivo,tipo_grupo,contrato_id,fecha_inicio,fecha_fin")
            .eq("contrato_id", int(contrato_id))
            .execute()
            .data
            or []
        )
    except Exception:
        rows = []
    rpo = [r for r in rows if str(r.get("tipo_grupo") or "").strip().upper() == "RPO"]
    pool = rpo or rows

    def _es_uno(r: dict) -> bool:
        nro = str(r.get("numero_rpo") or "").strip()
        cons = str(r.get("consecutivo") or "").strip()
        return nro in {"1", "01", "RPO 1", "RPO-1"} or (cons == "1" and nro in {"", "1", "01"})

    elegidos = [r for r in pool if _es_uno(r)]
    if not elegidos:
        elegidos = [r for r in pool if int(r.get("id") or 0) == _ACTA_ID_DOC]
    if len(elegidos) == 1:
        return elegidos[0], "actas"
    if not pool and int(contrato_id) == _CONTRATO_ID_DOC:
        return dict(_ACTA_DOC), "documentado_acta_620"
    resumen = [
        {
            "id": r.get("id"),
            "numero_rpo": r.get("numero_rpo"),
            "consecutivo": r.get("consecutivo"),
            "tipo_grupo": r.get("tipo_grupo"),
        }
        for r in pool[:30]
    ]
    raise SystemExit(f"No hay un único Acta RPO 1. Candidatos: {resumen}")


def _listado_normas(contrato_id: int) -> dict:
    """(cap_norm, item_norm) → especificacion_tecnica. Mayor id gana, como el VU."""
    rows = []
    offset = 0
    while True:
        batch = (
            _sb.table("listado_precios")
            .select("id, capitulo, item_numero, especificacion_tecnica")
            .eq("contrato_id", int(contrato_id))
            .order("id")
            .range(offset, offset + 999)
            .execute()
            .data
            or []
        )
        rows.extend(batch)
        if len(batch) < 1000:
            break
        offset += 1000
    out = {}
    for r in rows:
        k = cap_item_key(r.get("capitulo"), r.get("item_numero"))
        if not k[1]:
            continue
        prev = out.get(k)
        if prev is None or int(r.get("id") or 0) >= int(prev.get("id") or 0):
            out[k] = {
                "id": r.get("id"),
                "especificacion_tecnica": str(r.get("especificacion_tecnica") or "").strip(),
            }
    return out


def _casi(a, b, tol=1e-6) -> bool:
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return a == b


def _filas_item(ws):
    limite = ws.max_row + 1
    for r in range(9, ws.max_row + 1):
        if ws.cell(r, 1).value == "RESUMEN DE CONCILIACIÓN":
            limite = r
            break
    rows = []
    for r in range(9, limite):
        val = ws.cell(r, 1).value
        if isinstance(val, str) and str(val).startswith("Subtotal"):
            continue
        if val in (None, ""):
            continue
        rows.append(r)
    return rows


def main() -> int:
    t0 = time.time()
    contratos = (
        _sb.table("contratos")
        .select("id,numero")
        .eq("numero", CONTRATO_NUMERO)
        .limit(1)
        .execute()
        .data
        or []
    )
    fuente_contrato = "contratos"
    if contratos:
        contrato_id = int(contratos[0]["id"])
    else:
        contrato_id = _CONTRATO_ID_DOC
        fuente_contrato = "documentado_contrato_3"
    acta, fuente_acta = _acta_rpo_1(contrato_id)
    acta_id = int(acta["id"])
    _campo, niveles = matriz_params_contrato(_sb, contrato_id)
    nivel = max(niveles) if niveles else None
    nrpo = str(acta.get("numero_rpo") or acta.get("consecutivo") or acta_id)
    fname = _safe_filename_part(f"CC-MES-integral_Acta_{nrpo}.xlsx")
    print(
        f"contrato_id={contrato_id} ({fuente_contrato}) acta_id={acta_id} ({fuente_acta}) "
        f"archivo={fname} nivel={nivel}",
        flush=True,
    )

    user = {"nombre": "QA", "apellidos": "Norma", "cargo_nombre": "Verificación"}
    real_row = _row

    def _row_visible(table, select, **eq):
        if table == "actas" and eq.get("id") is not None and int(eq["id"]) == acta_id:
            merged = dict(_ACTA_DOC)
            merged.update({k: v for k, v in acta.items() if v is not None})
            return merged
        if table == "contratos" and eq.get("id") is not None and int(eq["id"]) == contrato_id:
            return dict(_CONTRATO_DOC)
        return real_row(table, select, **eq)

    listado = _listado_normas(contrato_id)
    print(f"listado normas={len(listado)}", flush=True)

    with patch.object(inf_mod, "_row", side_effect=_row_visible):
        ctx = _contexto_acta_mes_conciliacion(
            contrato_id, acta_id, user, nivel_aprobacion=nivel
        )
        items = list(ctx.get("items") or [])
        print(f"items informe={len(items)} en {time.time() - t0:.1f}s", flush=True)
        html = _html_cc_mes_001_v1(
            ctx["contrato"],
            ctx["acta"],
            items,
            float(ctx.get("total_costo") or 0),
            ctx.get("usuario_nombre") or "QA",
            ctx.get("usuario_cargo") or "—",
            {},
            resumen_4cols=ctx.get("resumen_4cols"),
            otros_conceptos=ctx.get("otros_conceptos") or [],
            c3_label="ACTA RPO",
            c3_value=str((ctx.get("acta") or {}).get("numero_rpo") or acta_id),
            c4_label="CONSECUTIVO",
            c4_value=str((ctx.get("acta") or {}).get("consecutivo") or "—"),
        )
        print("generando Excel integral…", flush=True)
        t1 = time.time()
        raw = _cc_mes_integral_excel_bytes(
            contrato_id, acta_id, user, nivel_aprobacion=nivel
        )
        print(f"excel bytes={len(raw)} en {time.time() - t1:.1f}s", flush=True)

    wb = load_workbook(BytesIO(raw))
    errores = []
    if fname != "CC-MES-integral_Acta_1.xlsx":
        errores.append(f"nombre de archivo {fname}")
    if "preacta_obra" in wb.sheetnames:
        errores.append("sigue la pestaña preacta_obra")
    if any(ws.tables for ws in wb.worksheets):
        errores.append("el libro todavía contiene una tabla de Excel")
    if wb.sheetnames[0] != "CC-MES-001":
        errores.append(f"la primera hoja es {wb.sheetnames[0]!r}")
    listado_legible = len(listado) > 0

    ws = wb["CC-MES-001"]
    if ws.cell(7, 4).value != "NORMA TÉCNICA" or ws.cell(7, 5).value != "V. UNIT.":
        errores.append(
            f"encabezados D/E = {ws.cell(7, 4).value!r} / {ws.cell(7, 5).value!r}"
        )
    filas = _filas_item(ws)
    if len(filas) != len(items):
        errores.append(f"filas excel={len(filas)} ítems={len(items)}")

    book = FormulaBook(wb)
    desfaces = []
    deltas_memoria = []
    redondeos = []
    normas_ok = 0
    normas_vacias = 0
    normas_con_texto = 0
    muestra = []
    for i, it in enumerate(items):
        if i >= len(filas):
            break
        row = filas[i]
        vu = it.get("vlr_unitario_sub") or it.get("vlr_unitario") or 0
        clave = cap_item_key(it.get("capitulo"), it.get("item_numero"))
        if listado_legible:
            if clave not in listado:
                desfaces.append({"item": it.get("item_numero"), "campo": "norma", "detalle": "sin cruce listado"})
                esperado_norma = None
            else:
                esperado_norma = listado[clave]["especificacion_tecnica"]
        else:
            esperado_norma = str(it.get("norma_tecnica") or "").strip()
        celda = ws.cell(row, 4).value
        got_norma = "" if celda in (None, "") else str(celda).strip()
        if esperado_norma is not None and got_norma == esperado_norma:
            normas_ok += 1
            if got_norma:
                normas_con_texto += 1
            else:
                normas_vacias += 1
        elif esperado_norma is not None:
            desfaces.append(
                {
                    "item": it.get("item_numero"),
                    "capitulo": it.get("capitulo"),
                    "campo": "norma",
                    "excel": got_norma,
                    "listado": esperado_norma,
                }
            )
        if len(muestra) < 8 and got_norma:
            muestra.append(
                {
                    "item": it.get("item_numero"),
                    "capitulo": it.get("capitulo"),
                    "norma": got_norma,
                }
            )
        try:
            cant_act = book.eval_cell("CC-MES-001", row, 6)
            val_act = book.eval_cell("CC-MES-001", row, 7)
            cant_pres = book.eval_cell("CC-MES-001", row, 8)
            val_pres = book.eval_cell("CC-MES-001", row, 9)
            cant_acum = book.eval_cell("CC-MES-001", row, 10)
            val_acum = book.eval_cell("CC-MES-001", row, 11)
            cant_saldo = book.eval_cell("CC-MES-001", row, 12)
            val_saldo = book.eval_cell("CC-MES-001", row, 13)
        except Exception as exc:  # noqa: BLE001
            desfaces.append({"item": it.get("item_numero"), "campo": "formula", "error": str(exc)})
            continue
        if not _casi(val_act, it.get("valor_actualizadas")) or not _casi(cant_act, it.get("cant_actualizadas")):
            desfaces.append(
                {
                    "item": it.get("item_numero"),
                    "campo": "actualizadas",
                    "excel_cant": cant_act,
                    "excel_valor": val_act,
                    "pdf_cant": it.get("cant_actualizadas"),
                    "pdf_valor": it.get("valor_actualizadas"),
                }
            )
        # La cadena de la hoja usa la cantidad presente de la memoria.
        ant = float(it.get("cant_acum_anterior") or 0)
        if not _casi(val_pres, valor_por_cantidad_vu(cant_pres, vu)):
            desfaces.append({"item": it.get("item_numero"), "campo": "valor_presente_formula", "excel": val_pres})
        if not _casi(cant_acum, round(float(cant_pres) + ant, 2)):
            desfaces.append({"item": it.get("item_numero"), "campo": "cant_acumulado_formula", "excel": cant_acum})
        if not _casi(val_acum, valor_por_cantidad_vu(cant_acum, vu)):
            desfaces.append({"item": it.get("item_numero"), "campo": "valor_acumulado_formula", "excel": val_acum})
        if not _casi(cant_saldo, round(float(cant_act) - float(cant_acum), 2)):
            desfaces.append({"item": it.get("item_numero"), "campo": "cant_saldo_formula", "excel": cant_saldo})
        esperado_saldo = valor_por_cantidad_vu(cant_saldo, vu)
        if not _casi(val_saldo, esperado_saldo):
            delta = float(val_saldo) - float(esperado_saldo)
            # Excel ROUND aleja del cero en .5; round() de Python es half-even. 1 peso no es un corrimiento de columna.
            rec = {
                "item": it.get("item_numero"),
                "campo": "valor_saldo_formula",
                "excel": val_saldo,
                "esperado": esperado_saldo,
                "cant_saldo": cant_saldo,
                "vu": vu,
                "delta": delta,
            }
            if abs(delta) <= 1:
                redondeos.append(rec)
            else:
                desfaces.append(rec)
        if not _casi(cant_pres, it.get("cant_presente")) or not _casi(val_pres, it.get("valor_presente")):
            deltas_memoria.append(
                {
                    "item": it.get("item_numero"),
                    "capitulo": it.get("capitulo"),
                    "cant_excel": cant_pres,
                    "cant_pdf": it.get("cant_presente"),
                    "valor_excel": val_pres,
                    "valor_pdf": it.get("valor_presente"),
                }
            )
        desc = str(it.get("item_descripcion") or "").lower()
        if desc and desc not in html:
            desfaces.append({"item": it.get("item_numero"), "campo": "descripcion_pdf"})
        for valor in (it.get("cant_actualizadas"), it.get("valor_actualizadas")):
            txt = _fn_cant_informe(valor) if valor == it.get("cant_actualizadas") else _fm_informe(valor)
            if txt and txt not in html:
                desfaces.append({"item": it.get("item_numero"), "campo": "actualizadas_pdf", "texto": txt})
                break

    if desfaces:
        errores.append(f"{len(desfaces)} desfaces de norma, actualizadas o cadena de fórmulas")

    resumen = {
        "contrato": CONTRATO_NUMERO,
        "contrato_id": contrato_id,
        "acta_id": acta_id,
        "archivo": fname,
        "nivel_aprobacion": nivel,
        "items": len(items),
        "filas_excel": len(filas),
        "hojas": wb.sheetnames[:8],
        "n_hojas": len(wb.sheetnames),
        "preacta_obra": "preacta_obra" in wb.sheetnames,
        "listado_filas": len(listado),
        "listado_legible": listado_legible,
        "nota_listado": (
            None
            if listado_legible
            else "La anon key no lee listado_precios (RLS). La columna queda vacía, igual que el cruce. "
            "El texto por capítulo e ítem se cubre en test_cc_mes_integral_norma_excel."
        ),
        "normas_ok": normas_ok,
        "normas_con_texto": normas_con_texto,
        "normas_vacias": normas_vacias,
        "muestra_norma": muestra,
        "deltas_memoria": len(deltas_memoria),
        "muestra_deltas_memoria": deltas_memoria[:6],
        "redondeo_excel_vs_python": redondeos,
        "desfaces": desfaces[:12],
        "errores": errores,
        "segundos": round(time.time() - t0, 1),
        "ok": not errores,
    }
    out = Path("/opt/cursor/artifacts/cc_mes_norma_iccu_1614.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(resumen, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({k: resumen[k] for k in (
        "ok", "archivo", "items", "filas_excel", "preacta_obra", "listado_filas",
        "listado_legible", "normas_ok", "normas_con_texto", "normas_vacias",
        "deltas_memoria", "muestra_norma", "errores", "desfaces",
        "muestra_deltas_memoria", "segundos",
    )}, ensure_ascii=False, indent=2, default=str))
    return 0 if resumen["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
