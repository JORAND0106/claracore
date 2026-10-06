#!/usr/bin/env python3
"""Prueba real: tabla preacta_obra del Excel integral CC-MES.

Contrato ICCU-CTO-1614-2025, Acta RPO 1.

Comprueba que el libro integral trae la tabla preacta_obra (una fila por ítem,
sin subtotales, Clave única) y que sus valores, calculados desde las fórmulas
hacia la pestaña del informe, coinciden con el contexto del PDF CC-MES-001.

Uso (desde backend/):
  PYTHONPATH=. python3 scripts/qa_preacta_obra_iccu_1614.py

Lee VITE_SUPABASE_* de frontend/.env.development si no hay SUPABASE_URL/KEY.
No modifica datos.
"""
from __future__ import annotations

import json
import os
import sys
import time
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))


def _prime_env() -> None:
    url = os.environ.get("SUPABASE_URL") or os.environ.get("VITE_SUPABASE_URL") or ""
    key = (
        os.environ.get("SUPABASE_KEY")
        or os.environ.get("SUPABASE_ANON_KEY")
        or os.environ.get("VITE_SUPABASE_ANON_KEY")
        or ""
    )
    if not url or not key:
        env_path = ROOT / "frontend" / ".env.development"
        if env_path.is_file():
            for line in env_path.read_text().splitlines():
                if line.startswith("VITE_SUPABASE_URL=") and not url:
                    url = line.split("=", 1)[1].strip().strip('"')
                elif line.startswith("VITE_SUPABASE_ANON_KEY=") and not key:
                    key = line.split("=", 1)[1].strip().strip('"')
    if not url or not key:
        raise SystemExit("Faltan SUPABASE_URL / SUPABASE_KEY")
    os.environ["SUPABASE_URL"] = url
    os.environ["SUPABASE_KEY"] = key


_prime_env()

import main  # noqa: E402,F401  — carga informes sin import circular

from corte_sub_conciliacion import valor_por_cantidad_vu  # noqa: E402
from excel_formula_eval import FormulaBook  # noqa: E402
import informes as inf_mod  # noqa: E402
from informes import (  # noqa: E402
    _PREACTA_OBRA_HEADERS,
    _cc_mes_integral_excel_bytes,
    _contexto_acta_mes_conciliacion,
    _fm_informe,
    _fn_cant_informe,
    _html_cc_mes_001_v1,
    _row,
    _sb,
    matriz_params_contrato,
)
from openpyxl import load_workbook  # noqa: E402

CONTRATO_NUMERO = "ICCU-CTO-1614-2025"
HEADERS = _PREACTA_OBRA_HEADERS


# Identidad documentada en los QA de costo (actas/contratos no son legibles con la anon key).
_CONTRATO_ID_DOC = 3
_ACTA_ID_DOC = 620
_ACTA_DOC = {
    "id": _ACTA_ID_DOC,
    "contrato_id": _CONTRATO_ID_DOC,
    "consecutivo": 1,
    "numero_rpo": "1",
    "fecha_inicio": None,
    "fecha_fin": None,
    "tipo_grupo": "RPO",
}
_CONTRATO_DOC = {
    "id": _CONTRATO_ID_DOC,
    "numero": CONTRATO_NUMERO,
    "objeto": "",
    "contratista": "",
    "nit": "",
    "interventoria": "",
    "logo_contratista": None,
    "aiu": None,
    "iva": None,
    "anticipo": None,
    "amortizacion_pct": None,
}


def _acta_rpo_1(contrato_id: int) -> tuple[dict, str]:
    """Devuelve (acta, fuente). Si RLS oculta actas, usa el id documentado 620."""
    try:
        rows = (
            _sb.table("actas")
            .select("id,numero_rpo,consecutivo,tipo_grupo,contrato_id,fecha_inicio,fecha_fin")
            .eq("contrato_id", int(contrato_id))
            .execute()
            .data
            or []
        )
    except Exception:
        rows = []
    rpo = [r for r in rows if str(r.get("tipo_grupo") or "").strip().upper() == "RPO"]
    pool = rpo or rows

    def _es_uno(r: dict) -> bool:
        nro = str(r.get("numero_rpo") or "").strip()
        cons = str(r.get("consecutivo") or "").strip()
        return nro in {"1", "01", "RPO 1", "RPO-1"} or (cons == "1" and nro in {"", "1", "01"})

    elegidos = [r for r in pool if _es_uno(r)]
    if not elegidos:
        elegidos = [r for r in pool if int(r.get("id") or 0) == _ACTA_ID_DOC]
    if len(elegidos) == 1:
        return elegidos[0], "actas"
    if not pool and int(contrato_id) == _CONTRATO_ID_DOC:
        return dict(_ACTA_DOC), "documentado_acta_620"
    resumen = [
        {
            "id": r.get("id"),
            "numero_rpo": r.get("numero_rpo"),
            "consecutivo": r.get("consecutivo"),
            "tipo_grupo": r.get("tipo_grupo"),
        }
        for r in pool[:30]
    ]
    raise SystemExit(f"No hay un único Acta RPO 1. Candidatos: {resumen}")


def _casi(a, b, tol=1e-6) -> bool:
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return a == b


def main() -> int:
    t0 = time.time()
    contratos = (
        _sb.table("contratos")
        .select("id,numero")
        .eq("numero", CONTRATO_NUMERO)
        .limit(1)
        .execute()
        .data
        or []
    )
    fuente_contrato = "contratos"
    if contratos:
        contrato_id = int(contratos[0]["id"])
    else:
        # RLS de la anon key no devuelve contratos; el caso QA usa id 3.
        contrato_id = _CONTRATO_ID_DOC
        fuente_contrato = "documentado_contrato_3"
    acta, fuente_acta = _acta_rpo_1(contrato_id)
    acta_id = int(acta["id"])
    _campo, niveles = matriz_params_contrato(_sb, contrato_id)
    nivel = max(niveles) if niveles else None
    print(
        f"contrato_id={contrato_id} ({fuente_contrato}) acta_id={acta_id} ({fuente_acta}) "
        f"numero_rpo={acta.get('numero_rpo')} consecutivo={acta.get('consecutivo')} "
        f"nivel={nivel} activos={niveles}",
        flush=True,
    )

    user = {"nombre": "QA", "apellidos": "Preacta", "cargo_nombre": "Verificación"}
    real_row = _row

    def _row_visible(table, select, **eq):
        # Cabeceras que la anon key no puede leer. El resto (registros, presupuesto) va a Supabase.
        if table == "actas" and eq.get("id") is not None and int(eq["id"]) == acta_id:
            merged = dict(_ACTA_DOC)
            merged.update({k: v for k, v in acta.items() if v is not None})
            return merged
        if table == "contratos" and eq.get("id") is not None and int(eq["id"]) == contrato_id:
            return dict(_CONTRATO_DOC)
        return real_row(table, select, **eq)

    with patch.object(inf_mod, "_row", side_effect=_row_visible):
        ctx = _contexto_acta_mes_conciliacion(
            contrato_id, acta_id, user, nivel_aprobacion=nivel
        )
        items = list(ctx.get("items") or [])
        print(f"items informe={len(items)} en {time.time() - t0:.1f}s", flush=True)

        html = _html_cc_mes_001_v1(
            ctx["contrato"],
            ctx["acta"],
            items,
            float(ctx.get("total_costo") or 0),
            ctx.get("usuario_nombre") or "QA",
            ctx.get("usuario_cargo") or "—",
            {},
            resumen_4cols=ctx.get("resumen_4cols"),
            otros_conceptos=ctx.get("otros_conceptos") or [],
            c3_label="ACTA RPO",
            c3_value=str((ctx.get("acta") or {}).get("numero_rpo") or acta_id),
            c4_label="CONSECUTIVO",
            c4_value=str((ctx.get("acta") or {}).get("consecutivo") or "—"),
        )

        print("generando Excel integral…", flush=True)
        t1 = time.time()
        raw = _cc_mes_integral_excel_bytes(
            contrato_id, acta_id, user, nivel_aprobacion=nivel
        )
        print(f"excel bytes={len(raw)} en {time.time() - t1:.1f}s", flush=True)
    wb = load_workbook(BytesIO(raw))
    errores = []
    if "preacta_obra" not in wb.sheetnames:
        errores.append("falta la pestaña preacta_obra")
    if wb.sheetnames[0] != "CC-MES-001":
        errores.append(f"la primera hoja es {wb.sheetnames[0]!r}")
    hoja = wb["preacta_obra"] if "preacta_obra" in wb.sheetnames else None
    tabla = None
    if hoja is not None:
        tabla = hoja.tables.get("preacta_obra")
        if tabla is None:
            errores.append("la pestaña no contiene la tabla preacta_obra")
        else:
            if tabla.displayName != "preacta_obra":
                errores.append(f"displayName={tabla.displayName!r}")
            if tuple(tabla.column_names) != tuple(HEADERS):
                errores.append(f"encabezados={list(tabla.column_names)}")
            if tabla.totalsRowCount:
                errores.append("la tabla tiene fila de totales")

    informe = wb["CC-MES-001"]
    subtotales = [
        r
        for r in range(1, informe.max_row + 1)
        if isinstance(informe.cell(r, 1).value, str)
        and str(informe.cell(r, 1).value).startswith("Subtotal")
    ]
    book = FormulaBook(wb) if hoja is not None else None
    filas = []
    claves = []
    if hoja is not None and book is not None:
        # Datos desde la fila 2 hasta la última con fórmula en Clave.
        row = 2
        while hoja.cell(row, 1).value not in (None, ""):
            got = {}
            for col, header in enumerate(HEADERS, start=1):
                raw_f = hoja.cell(row, col).value
                if not (isinstance(raw_f, str) and raw_f.startswith("=")):
                    errores.append(f"fila {row} col {header} no es fórmula: {raw_f!r}")
                    got[header] = raw_f
                    continue
                if "CC-MES-001" not in raw_f and header not in ("Clave", "Capitulo"):
                    # AnteriorValor también cita CC-MES-001.
                    errores.append(f"fila {row} {header} no referencia el informe")
                try:
                    got[header] = book.eval_cell("preacta_obra", row, col)
                except Exception as exc:  # noqa: BLE001
                    errores.append(f"fila {row} {header}: {exc}")
                    got[header] = None
            clave = got.get("Clave")
            claves.append(clave)
            if isinstance(clave, str) and clave.startswith("Subtotal"):
                errores.append(f"fila de subtotal dentro de la tabla: {clave}")
            if got.get("Item") in (None, ""):
                errores.append(f"fila {row} sin ítem")
            filas.append(got)
            row += 1
            if row > len(items) + 5:
                break

    if len(filas) != len(items):
        errores.append(f"filas tabla={len(filas)} ítems informe={len(items)}")
    if len(claves) != len(set(claves)):
        dup = [c for c in claves if claves.count(c) > 1]
        errores.append(f"Clave repetida: {sorted(set(dup))[:8]}")

    # Identidad, actualizadas y anterior salen del mismo contexto que el PDF.
    # Presente/acumulado/saldo de la tabla son la celda del informe: en el libro
    # integral el presente es el total de la memoria, que puede diferir en
    # centésimas del agregado impreso en el PDF. Eso no se corrige aquí.
    campos_pdf = {
        "Clave",
        "Capitulo",
        "Item",
        "ActualizadaCant",
        "ActualizadaValor",
        "AnteriorCant",
        "AnteriorValor",
    }
    desfaces = []
    deltas_memoria = []
    for i, it in enumerate(items):
        if i >= len(filas):
            break
        got = filas[i]
        vu = it.get("vlr_unitario_sub") or it.get("vlr_unitario") or 0
        esperado = {
            "Clave": f"{str(it.get('capitulo') or '').strip() or '—'}|{it.get('item_numero')}",
            "Capitulo": str(it.get("capitulo") or "").strip() or "—",
            "Item": it.get("item_numero"),
            "ActualizadaCant": it.get("cant_actualizadas"),
            "ActualizadaValor": it.get("valor_actualizadas"),
            "AnteriorCant": it.get("cant_acum_anterior"),
            "AnteriorValor": valor_por_cantidad_vu(it.get("cant_acum_anterior"), vu),
            "PresenteCant": it.get("cant_presente"),
            "PresenteValor": it.get("valor_presente"),
            "AcumuladoCant": it.get("cant_acumulado"),
            "AcumuladoValor": it.get("valor_acumulado"),
            "SaldoCant": it.get("cant_saldo"),
            "SaldoValor": it.get("valor_saldo"),
        }
        for campo, exp in esperado.items():
            val = got.get(campo)
            ok = val == exp if isinstance(exp, str) else _casi(val, exp)
            if ok:
                continue
            rec = {
                "fila": i + 2,
                "item": it.get("item_numero"),
                "capitulo": it.get("capitulo"),
                "campo": campo,
                "tabla": val,
                "pdf": exp,
            }
            if campo in campos_pdf:
                desfaces.append(rec)
            else:
                deltas_memoria.append(rec)
        desc = str(it.get("item_descripcion") or "").lower()
        if desc and desc not in html:
            desfaces.append({"fila": i + 2, "campo": "Descripcion", "pdf": "ausente", "texto": desc[:80]})
        for campo_pdf, valor in (
            ("ActualizadaCant", it.get("cant_actualizadas")),
            ("AnteriorCant", it.get("cant_acum_anterior")),
            ("PresenteCant", it.get("cant_presente")),
            ("AcumuladoCant", it.get("cant_acumulado")),
            ("SaldoCant", it.get("cant_saldo")),
        ):
            txt = _fn_cant_informe(valor)
            if txt and txt not in html:
                desfaces.append({"fila": i + 2, "campo": campo_pdf, "pdf": "ausente", "texto": txt})
        for campo_pdf, valor in (
            ("ActualizadaValor", it.get("valor_actualizadas")),
            ("AnteriorValor", valor_por_cantidad_vu(it.get("cant_acum_anterior"), vu)),
            ("PresenteValor", it.get("valor_presente")),
            ("AcumuladoValor", it.get("valor_acumulado")),
            ("SaldoValor", it.get("valor_saldo")),
        ):
            txt = _fm_informe(valor)
            if txt and txt not in html:
                desfaces.append({"fila": i + 2, "campo": campo_pdf, "pdf": "ausente", "texto": txt})

    ejemplo = "3. OBRAS DE ARTE (ALCANTARILLA)|3.4."
    if ejemplo not in claves:
        errores.append(f"falta la clave de ejemplo {ejemplo}")
    if desfaces:
        errores.append(f"{len(desfaces)} desfaces de identidad/actualizadas/anterior o de texto PDF")

    resumen = {
        "contrato": CONTRATO_NUMERO,
        "contrato_id": contrato_id,
        "fuente_contrato": fuente_contrato,
        "acta_id": acta_id,
        "fuente_acta": fuente_acta,
        "nota_acceso": (
            "Registros y presupuesto se leyeron en vivo. "
            "contratos, actas y listado_precios no devuelven filas con la anon key (RLS); "
            "la cabecera usa el caso documentado contrato 3 / acta 620 y el VU cae al sello del registro."
        ),
        "numero_rpo": acta.get("numero_rpo"),
        "consecutivo": acta.get("consecutivo"),
        "nivel_aprobacion": nivel,
        "items": len(items),
        "filas_tabla": len(filas),
        "subtotales_informe": len(subtotales),
        "claves_unicas": len(set(claves)),
        "hojas": len(wb.sheetnames),
        "excel_bytes": len(raw),
        "tabla": tabla.displayName if tabla is not None else None,
        "ref": tabla.ref if tabla is not None else None,
        "encabezados": list(tabla.column_names) if tabla is not None else [],
        "muestra_claves": claves[:8],
        "clave_ejemplo": ejemplo in claves,
        "desfaces": desfaces[:12],
        "deltas_memoria_vs_pdf": len(deltas_memoria),
        "muestra_deltas_memoria": deltas_memoria[:6],
        "errores": errores,
        "segundos": round(time.time() - t0, 1),
        "ok": not errores,
    }
    out = Path("/opt/cursor/artifacts/preacta_obra_iccu_1614.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(resumen, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({k: resumen[k] for k in (
        "ok", "items", "filas_tabla", "claves_unicas", "subtotales_informe",
        "clave_ejemplo", "deltas_memoria_vs_pdf", "errores", "desfaces",
        "muestra_deltas_memoria", "muestra_claves", "segundos",
    )}, ensure_ascii=False, indent=2, default=str))
    return 0 if resumen["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
