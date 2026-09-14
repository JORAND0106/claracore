"""
Exportación consolidada de Bitácora por rango de fechas.
Formatos: PDF (merge de hojas diarias), DOCX, Markdown.
"""
from __future__ import annotations

import io
import logging
from datetime import date
from typing import Any, List, Tuple

from bitacora_service import (
    _label_tramo,
    clima_label,
    list_entradas,
)
from bitacora_pdf import (
    _fotos_diario,
    _plain_from_html,
    _prepare_img_asset,
    generar_pdf_bitacora_dia,
)

_log = logging.getLogger("claracore.bitacora.export_rango")

MAX_DIAS_RANGO = 120
_FORMATOS = frozenset({"pdf", "docx", "md", "markdown"})


def _parse_rango(fecha_desde: str, fecha_hasta: str) -> Tuple[date, date]:
    d0 = date.fromisoformat(str(fecha_desde)[:10])
    d1 = date.fromisoformat(str(fecha_hasta)[:10])
    if d1 < d0:
        raise ValueError("La fecha hasta debe ser posterior o igual a la fecha desde.")
    if (d1 - d0).days > MAX_DIAS_RANGO:
        raise ValueError(
            f"El rango no puede superar {MAX_DIAS_RANGO} días. "
            "Divida la exportación en periodos más cortos."
        )
    return d0, d1


def list_diarios_rango(
    sb,
    contrato_id: int,
    fecha_desde: str,
    fecha_hasta: str,
) -> List[dict]:
    """Diarios del rango (ascendente por fecha, luego tramo/id). Omite días vacíos."""
    _parse_rango(fecha_desde, fecha_hasta)
    rows = list_entradas(
        sb,
        contrato_id,
        fecha_desde=str(fecha_desde)[:10],
        fecha_hasta=str(fecha_hasta)[:10],
        tipo="diario",
    )
    rows = [r for r in (rows or []) if isinstance(r, dict) and r.get("id") is not None]
    rows.sort(key=lambda r: (
        str(r.get("fecha") or "")[:10],
        str(r.get("tramo") or ""),
        int(r.get("id") or 0),
    ))
    return rows


def generar_pdf_bitacora_rango(
    sb,
    contrato_id: int,
    fecha_desde: str,
    fecha_hasta: str,
) -> bytes:
    """PDF único: concatena el PDF individual de cada diario del rango."""
    from pypdf import PdfReader, PdfWriter

    diarios = list_diarios_rango(sb, contrato_id, fecha_desde, fecha_hasta)
    if not diarios:
        raise ValueError("No hay Reportes Diarios en el rango seleccionado.")

    writer = PdfWriter()
    for d in diarios:
        try:
            chunk = generar_pdf_bitacora_dia(
                sb,
                contrato_id,
                str(d.get("fecha"))[:10],
                tramo=d.get("tramo"),
                entrada_id=int(d["id"]),
            )
        except Exception as exc:
            _log.warning("pdf rango omitiendo diario %s: %s", d.get("id"), exc)
            continue
        if not chunk:
            continue
        try:
            reader = PdfReader(io.BytesIO(chunk))
            for page in reader.pages:
                writer.add_page(page)
        except Exception as exc:
            _log.warning("pdf rango merge diario %s: %s", d.get("id"), exc)

    if len(writer.pages) == 0:
        raise ValueError("No se pudo generar páginas PDF para el rango.")

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _md_escape(text: Any) -> str:
    return str(text or "").replace("|", "\\|").replace("\n", " ").strip()


def _md_table(headers: List[str], rows: List[List[str]]) -> str:
    if not rows:
        return ""
    head = "| " + " | ".join(headers) + " |"
    sep = "| " + " | ".join("---" for _ in headers) + " |"
    body = "\n".join("| " + " | ".join(_md_escape(c) for c in r) + " |" for r in rows)
    return f"{head}\n{sep}\n{body}\n"


def _horario_asist(a: dict) -> str:
    ini = str(a.get("hora_ingreso") or "")[:5]
    fin = str(a.get("hora_salida") or "")[:5]
    if ini and fin:
        return f"{ini} – {fin}"
    return ini or fin or "—"


def _bloque_md_diario(d: dict) -> str:
    fecha = str(d.get("fecha") or "")[:10]
    tramo = _label_tramo(d.get("tramo"))
    parts = [
        f"## {fecha} · {tramo}",
        "",
        f"- **Estado:** {d.get('estado') or '—'}",
        f"- **Hora inicio:** {str(d.get('hora_inicio_labores') or '')[:5] or '—'}",
        f"- **Clima:** {d.get('clima_descripcion') or clima_label(d.get('clima_codigo')) or '—'} "
        f"({d.get('clima_temp_c') if d.get('clima_temp_c') is not None else '—'} °C)",
        f"- **Elaborado por:** {d.get('created_by_nombre') or '—'}",
        "",
    ]

    personal = d.get("personal") if isinstance(d.get("personal"), list) else []
    pers_rows = [
        [str(p.get("cargo") or ""), str(p.get("cantidad") or "")]
        for p in personal
        if isinstance(p, dict) and (p.get("cantidad") or 0)
    ]
    if pers_rows:
        parts.append("### Personal (resumen por cargo)")
        parts.append(_md_table(["Cargo", "Cant."], pers_rows))

    asist = d.get("asistencia_colaboradores") if isinstance(d.get("asistencia_colaboradores"), list) else []
    as_rows = [
        [
            str(a.get("nombre") or ""),
            str(a.get("cargo") or ""),
            str(a.get("subcontratista_nombre") or ""),
            _horario_asist(a),
        ]
        for a in asist if isinstance(a, dict) and a.get("nombre")
    ]
    if as_rows:
        parts.append("### Asistencia")
        parts.append(_md_table(["Nombre", "Cargo", "Empresa", "Horario"], as_rows))

    usos = d.get("equipos_uso") if isinstance(d.get("equipos_uso"), list) else []
    uso_rows = [
        [
            str(u.get("equipo_nombre") or ""),
            str(u.get("operador") or "—"),
            str(u.get("cantidad") if u.get("cantidad") is not None else "1"),
            f"{str(u.get('hora_inicio') or '')[:5]}–{str(u.get('hora_fin') or '')[:5]}".strip("–"),
        ]
        for u in usos if isinstance(u, dict) and u.get("equipo_nombre")
    ]
    if uso_rows:
        parts.append("### Maquinaria")
        parts.append(_md_table(["Equipo", "Operador", "Cant.", "Horario"], uso_rows))

    mats = d.get("materiales") if isinstance(d.get("materiales"), list) else []
    mat_rows = [
        [
            str(m.get("movimiento") or ""),
            str(m.get("tipo_material") or ""),
            str(m.get("proveedor") or ""),
            str(m.get("cantidad") or ""),
            str(m.get("numeros_vale") or ""),
        ]
        for m in mats
        if isinstance(m, dict) and (
            m.get("tipo_material") or m.get("proveedor") or (m.get("cantidad") or 0)
        )
    ]
    if mat_rows:
        parts.append("### Materiales")
        parts.append(_md_table(["Mov.", "Tipo", "Proveedor", "Cant.", "Vale"], mat_rows))

    obs = _plain_from_html(d.get("cuerpo_html") or "").strip()
    if obs:
        parts.append("### Observaciones")
        parts.append(obs)
        parts.append("")

    fotos = _fotos_diario(d)
    if fotos:
        parts.append("### Registro fotográfico")
        for i, im in enumerate(fotos, 1):
            cap = str((im or {}).get("pie") or (im or {}).get("caption") or f"Foto {i}")
            path = str((im or {}).get("blob_path") or (im or {}).get("path") or "")
            parts.append(f"- {cap}" + (f" (`{path}`)" if path else ""))
        parts.append("")

    eventos = d.get("eventos") if isinstance(d.get("eventos"), list) else []
    for ev in eventos:
        if not isinstance(ev, dict):
            continue
        et = str(ev.get("evento_tipo") or "evento")
        parts.append(f"### Evento: {et}")
        cuerpo = _plain_from_html(ev.get("cuerpo_html") or "").strip()
        if cuerpo:
            parts.append(cuerpo)
            parts.append("")
        det = ev.get("evento_detalle") if isinstance(ev.get("evento_detalle"), dict) else {}
        acts = det.get("actividades") if isinstance(det.get("actividades"), list) else []
        act_rows = []
        for a in acts:
            if not isinstance(a, dict):
                continue
            act_rows.append([
                str(a.get("actividad") or a.get("descripcion") or ""),
                str(a.get("cantidad") or ""),
                str(a.get("unidad") or ""),
                str(a.get("ubicacion_tramo") or a.get("tramo") or ""),
            ])
        if act_rows:
            parts.append(_md_table(["Actividad", "Cant.", "Und.", "Ubicación"], act_rows))

    parts.append("---")
    parts.append("")
    return "\n".join(parts)


def generar_markdown_bitacora_rango(
    sb,
    contrato_id: int,
    fecha_desde: str,
    fecha_hasta: str,
) -> str:
    diarios = list_diarios_rango(sb, contrato_id, fecha_desde, fecha_hasta)
    if not diarios:
        raise ValueError("No hay Reportes Diarios en el rango seleccionado.")
    d0, d1 = _parse_rango(fecha_desde, fecha_hasta)
    lines = [
        f"# Bitácora consolidada",
        "",
        f"- **Contrato ID:** {contrato_id}",
        f"- **Rango:** {d0.isoformat()} → {d1.isoformat()}",
        f"- **Reportes:** {len(diarios)}",
        "",
        "---",
        "",
    ]
    for d in diarios:
        lines.append(_bloque_md_diario(d))
    return "\n".join(lines).rstrip() + "\n"


def _docx_add_table(doc, headers: List[str], rows: List[List[str]]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    if not rows:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val or "")
            for p in table.rows[ri + 1].cells[ci].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph("")


def _docx_try_add_image(doc, im: dict, contrato_id: int) -> None:
    from docx.shared import Cm

    try:
        asset = _prepare_img_asset(im, contrato_id)
        if not asset:
            return
        uri, _w, _h = asset
        # data URI → bytes
        if not str(uri).startswith("data:"):
            return
        b64 = str(uri).split(",", 1)[-1]
        import base64
        raw = base64.b64decode(b64)
        stream = io.BytesIO(raw)
        doc.add_picture(stream, width=Cm(7.5))
        cap = str((im or {}).get("pie") or (im or {}).get("caption") or "").strip()
        if cap:
            p = doc.add_paragraph(cap)
            if p.runs:
                p.runs[0].italic = True
    except Exception as exc:
        _log.debug("docx foto skip: %s", exc)


def generar_docx_bitacora_rango(
    sb,
    contrato_id: int,
    fecha_desde: str,
    fecha_hasta: str,
) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise ValueError(
            "Exportación Word no disponible: falta la dependencia python-docx."
        ) from exc

    diarios = list_diarios_rango(sb, contrato_id, fecha_desde, fecha_hasta)
    if not diarios:
        raise ValueError("No hay Reportes Diarios en el rango seleccionado.")
    d0, d1 = _parse_rango(fecha_desde, fecha_hasta)

    doc = Document()
    doc.add_heading("Bitácora consolidada", level=0)
    meta = doc.add_paragraph(
        f"Contrato {contrato_id} · {d0.isoformat()} → {d1.isoformat()} · {len(diarios)} reporte(s)"
    )
    for run in meta.runs:
        run.font.size = Pt(10)

    for d in diarios:
        fecha = str(d.get("fecha") or "")[:10]
        tramo = _label_tramo(d.get("tramo"))
        doc.add_heading(f"{fecha} · {tramo}", level=1)
        doc.add_paragraph(
            f"Estado: {d.get('estado') or '—'} · "
            f"Inicio: {str(d.get('hora_inicio_labores') or '')[:5] or '—'} · "
            f"Clima: {d.get('clima_descripcion') or '—'} "
            f"({d.get('clima_temp_c') if d.get('clima_temp_c') is not None else '—'} °C) · "
            f"Elaborado: {d.get('created_by_nombre') or '—'}"
        )

        personal = d.get("personal") if isinstance(d.get("personal"), list) else []
        pers_rows = [
            [str(p.get("cargo") or ""), str(p.get("cantidad") or "")]
            for p in personal if isinstance(p, dict) and (p.get("cantidad") or 0)
        ]
        if pers_rows:
            doc.add_heading("Personal", level=2)
            _docx_add_table(doc, ["Cargo", "Cant."], pers_rows)

        asist = d.get("asistencia_colaboradores") if isinstance(d.get("asistencia_colaboradores"), list) else []
        as_rows = [
            [
                str(a.get("nombre") or ""),
                str(a.get("cargo") or ""),
                str(a.get("subcontratista_nombre") or ""),
                _horario_asist(a),
            ]
            for a in asist if isinstance(a, dict) and a.get("nombre")
        ]
        if as_rows:
            doc.add_heading("Asistencia", level=2)
            _docx_add_table(doc, ["Nombre", "Cargo", "Empresa", "Horario"], as_rows)

        usos = d.get("equipos_uso") if isinstance(d.get("equipos_uso"), list) else []
        uso_rows = [
            [
                str(u.get("equipo_nombre") or ""),
                str(u.get("operador") or "—"),
                str(u.get("cantidad") if u.get("cantidad") is not None else "1"),
            ]
            for u in usos if isinstance(u, dict) and u.get("equipo_nombre")
        ]
        if uso_rows:
            doc.add_heading("Maquinaria", level=2)
            _docx_add_table(doc, ["Equipo", "Operador", "Cant."], uso_rows)

        mats = d.get("materiales") if isinstance(d.get("materiales"), list) else []
        mat_rows = [
            [
                str(m.get("movimiento") or ""),
                str(m.get("tipo_material") or ""),
                str(m.get("proveedor") or ""),
                str(m.get("cantidad") or ""),
            ]
            for m in mats
            if isinstance(m, dict) and (m.get("tipo_material") or m.get("cantidad"))
        ]
        if mat_rows:
            doc.add_heading("Materiales", level=2)
            _docx_add_table(doc, ["Mov.", "Tipo", "Proveedor", "Cant."], mat_rows)

        obs = _plain_from_html(d.get("cuerpo_html") or "").strip()
        if obs:
            doc.add_heading("Observaciones", level=2)
            doc.add_paragraph(obs)

        fotos = _fotos_diario(d)
        if fotos:
            doc.add_heading("Registro fotográfico", level=2)
            for im in fotos[:4]:
                _docx_try_add_image(doc, im, int(contrato_id))

        eventos = d.get("eventos") if isinstance(d.get("eventos"), list) else []
        for ev in eventos:
            if not isinstance(ev, dict):
                continue
            doc.add_heading(f"Evento: {ev.get('evento_tipo') or 'evento'}", level=2)
            cuerpo = _plain_from_html(ev.get("cuerpo_html") or "").strip()
            if cuerpo:
                doc.add_paragraph(cuerpo)
            det = ev.get("evento_detalle") if isinstance(ev.get("evento_detalle"), dict) else {}
            acts = det.get("actividades") if isinstance(det.get("actividades"), list) else []
            act_rows = [
                [
                    str(a.get("actividad") or a.get("descripcion") or ""),
                    str(a.get("cantidad") or ""),
                    str(a.get("unidad") or ""),
                ]
                for a in acts if isinstance(a, dict)
            ]
            if act_rows:
                _docx_add_table(doc, ["Actividad", "Cant.", "Und."], act_rows)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def exportar_bitacora_rango(
    sb,
    contrato_id: int,
    fecha_desde: str,
    fecha_hasta: str,
    formato: str = "pdf",
) -> Tuple[bytes, str, str]:
    """
    Returns (content_bytes, media_type, filename).
    """
    fmt = str(formato or "pdf").strip().lower()
    if fmt == "markdown":
        fmt = "md"
    if fmt not in _FORMATOS:
        raise ValueError("Formato inválido. Use pdf, docx o md.")

    d0, d1 = _parse_rango(fecha_desde, fecha_hasta)
    base = f"bitacora_{contrato_id}_{d0.isoformat()}_{d1.isoformat()}"

    if fmt == "pdf":
        data = generar_pdf_bitacora_rango(sb, contrato_id, fecha_desde, fecha_hasta)
        return data, "application/pdf", f"{base}.pdf"
    if fmt == "docx":
        data = generar_docx_bitacora_rango(sb, contrato_id, fecha_desde, fecha_hasta)
        return (
            data,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{base}.docx",
        )
    md = generar_markdown_bitacora_rango(sb, contrato_id, fecha_desde, fecha_hasta)
    return md.encode("utf-8"), "text/markdown; charset=utf-8", f"{base}.md"
