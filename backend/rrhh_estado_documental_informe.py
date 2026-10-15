"""
Informe «Estado documental de colaboradores» (Excel + PDF).

Código de calidad (ajustable):
  CC-RRHH-DOC-001 · Versión 1.0

Encabezado: misma distribución que Precios Subcontratistas
  (logo | título del formato | código + versión + fecha).
"""
from __future__ import annotations

import io
import logging
import re
import unicodedata
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence, Tuple
from zoneinfo import ZoneInfo

_log = logging.getLogger("claracore.rrhh.estado_documental")

TZ_BOGOTA = ZoneInfo("America/Bogota")

# ── Identificación de calidad (sistema CCD). Cambiar aquí si Calidad redefine. ─
ESTADO_DOCUMENTAL_FORMATO = {
    "codigo": "CC-RRHH-DOC-001",
    "version": "1.0",
    "nombre": "Estado documental de colaboradores",
    "nombre_mayus": "ESTADO DOCUMENTAL DE COLABORADORES",
}

FILTROS = frozenset({"todos", "aprobado", "sin_aprobar"})

# Orden de columnas documentales (excluye bancario: no está en los 4 grupos del TAB Documentación).
_DOC_GROUPS_ORDER = ("contrato_laboral", "soporte", "ingreso", "afiliacion")

# Íconos de estado (Excel/PDF). Equivalentes semánticos si el motor no renderiza emoji.
_CELDA_OK = "✅"       # cargado
_CELDA_FALTA = "❌"    # obligatorio faltante
_CELDA_OPCIONAL = "🔘"  # opcional no cargado (no resta avance)

LEYENDA_ICONOS = (
    f"{_CELDA_OK} Cargado  ·  "
    f"{_CELDA_FALTA} Obligatorio faltante  ·  "
    f"{_CELDA_OPCIONAL} Opcional no cargado"
)


def _slug_tipo(label: str) -> str:
    s = unicodedata.normalize("NFD", str(label or "").strip().lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_") or "documento"
    return f"ext_{s[:72]}"


def _ahora_bogota() -> datetime:
    return datetime.now(TZ_BOGOTA)


def _fecha_calidad(d: Optional[datetime] = None) -> str:
    dt = d or _ahora_bogota()
    return dt.strftime("%d/%m/%Y")


def _nombre_informe(t: dict) -> str:
    """Primer nombre + primer apellido (columna de identificación del informe)."""
    from rrhh_docs_obligatoriedad import nombre_corto_colaborador

    return nombre_corto_colaborador(t.get("nombres"), t.get("apellidos"))


def _empresa_informe(t: dict) -> str:
    from rrhh_docs_obligatoriedad import iniciales_empresa

    return iniciales_empresa(
        t.get("empresa_nombre") or t.get("empresa") or t.get("empresa_key") or "",
    )


def _documento_identidad(t: dict) -> str:
    td = str(t.get("tipo_documento") or "CC").strip().upper()
    num = str(t.get("numero_documento") or "").strip()
    return f"{td} {num}".strip() or "—"


def _es_aprobado(t: dict) -> bool:
    return str(t.get("doc_validacion_estado") or "").strip().lower() == "aprobado"


def _filtrar_trabajadores(rows: List[dict], filtro: str) -> List[dict]:
    f = (filtro or "todos").strip().lower()
    if f not in FILTROS:
        f = "todos"
    if f == "aprobado":
        return [r for r in rows if _es_aprobado(r)]
    if f == "sin_aprobar":
        return [r for r in rows if not _es_aprobado(r)]
    return list(rows)


def _custom_tipos(sb, contrato_id: int) -> List[Tuple[str, str, str]]:
    """[(categoria, tipo_slug, label), ...] desde catálogo."""
    out: List[Tuple[str, str, str]] = []
    try:
        from rrhh_service import list_catalogo
    except Exception:
        return out
    for cat, key in (("soporte", "doc_soporte"), ("ingreso", "doc_ingreso")):
        try:
            rows = list_catalogo(sb, int(contrato_id), key) or []
        except Exception:
            rows = []
        for r in rows:
            label = str(r.get("valor") or "").strip()
            if label:
                out.append((cat, _slug_tipo(label), label))
    return out


def _fetch_docs_map(
    sb, contrato_id: int, ids: List[int],
) -> Tuple[Dict[int, set], Dict[Tuple[str, str], str]]:
    """
    Returns (by_trabajador, labels).
    labels: (categoria, tipo) → etiqueta legible (p. ej. tipos «Otro» / ext_*).
    """
    by: Dict[int, set] = {i: set() for i in ids}
    labels: Dict[Tuple[str, str], str] = {}
    if not ids:
        return by, labels
    try:
        rows = (
            sb.table("rrhh_trabajador_documentos")
            .select("trabajador_id, categoria, tipo, tipo_otro_texto, vigente")
            .eq("contrato_id", int(contrato_id))
            .in_("trabajador_id", ids)
            .eq("vigente", True)
            .is_("eliminado_en", "null")
            .execute()
            .data
            or []
        )
    except Exception as exc:
        _log.warning("estado documental docs: %s", exc)
        return by, labels
    for r in rows:
        try:
            tid = int(r["trabajador_id"])
        except (TypeError, ValueError, KeyError):
            continue
        cat = str(r.get("categoria") or "").strip().lower()
        tipo = str(r.get("tipo") or "").strip().lower()
        otro_txt = str(r.get("tipo_otro_texto") or "").strip()
        if not cat or not tipo:
            continue
        if tipo == "otro":
            if not otro_txt:
                continue
            tipo = _slug_tipo(otro_txt)
        if tid in by:
            by[tid].add((cat, tipo))
        if otro_txt and (cat, tipo) not in labels:
            labels[(cat, tipo)] = otro_txt
    return by, labels


def _fetch_contratos_set(sb, contrato_id: int, ids: List[int]) -> set:
    if not ids:
        return set()
    try:
        rows = (
            sb.table("rrhh_contratos_generados")
            .select("trabajador_id")
            .eq("contrato_id", int(contrato_id))
            .in_("trabajador_id", ids)
            .eq("vigente", True)
            .is_("eliminado_en", "null")
            .execute()
            .data
            or []
        )
    except Exception as exc:
        _log.warning("estado documental contratos: %s", exc)
        return set()
    out = set()
    for r in rows:
        try:
            out.add(int(r["trabajador_id"]))
        except (TypeError, ValueError, KeyError):
            continue
    return out


# Tipos base alineados con rrhh_docs_service (sin el slot UI «otro»).
_TIPOS_SOPORTE = (
    ("cedula", "Cédula / Documento de identidad"),
    ("hoja_vida", "Hoja de vida"),
    ("certificados", "Certificados"),
)
_TIPOS_INGRESO = (
    ("induccion", "Constancia de inducción"),
    ("reglamento", "Entrega de reglamentos"),
    ("examen_medico", "Examen médico de ingreso"),
)
_TIPOS_AFILIACION = (
    ("cert_eps", "Certificación EPS"),
    ("cert_pension", "Certificación Pensión / AFP"),
    ("cert_arl", "Certificación ARL"),
    ("cert_cesantias", "Certificación Cesantías"),
    ("cert_caja", "Certificación Caja de Compensación"),
)


def _columnas_checklist(custom: List[Tuple[str, str, str]]) -> List[dict]:
    """Columnas documentales en orden de grupos (sin bancario)."""
    cols: List[dict] = [
        {
            "key": "contrato_laboral",
            "categoria": "contrato_laboral",
            "tipo": "contrato_laboral",
            "label": "Contrato laboral",
            "grupo": "Contrato",
        },
    ]
    seen = {("contrato_laboral", "contrato_laboral")}

    def add(cat: str, tipo: str, label: str, grupo: str) -> None:
        if tipo == "otro":
            return
        key = (cat, tipo)
        if key in seen:
            return
        seen.add(key)
        cols.append({
            "key": f"{cat}:{tipo}",
            "categoria": cat,
            "tipo": tipo,
            "label": label,
            "grupo": grupo,
        })

    for tipo, label in _TIPOS_SOPORTE:
        add("soporte", tipo, label, "Documentos de soporte")
    for cat, tipo, label in custom:
        if cat == "soporte":
            add(cat, tipo, label, "Documentos de soporte")
    for tipo, label in _TIPOS_INGRESO:
        add("ingreso", tipo, label, "Documentos de ingreso")
    for cat, tipo, label in custom:
        if cat == "ingreso":
            add(cat, tipo, label, "Documentos de ingreso")
    for tipo, label in _TIPOS_AFILIACION:
        add("afiliacion", tipo, label, "Certificaciones de afiliación")

    return cols


def _celda_doc(
    trab: dict,
    col: dict,
    *,
    vigentes: set,
    tiene_contrato: bool,
    aprobado: bool,
) -> str:
    """
    ✅ cargado · ❌ obligatorio faltante · 🔘 opcional no cargado.
    EPS/ARL/Pensión/Cesantías nunca son 🔘.
    """
    from rrhh_docs_obligatoriedad import es_doc_opcional

    cat = col["categoria"]
    tipo = col["tipo"]
    if cat == "contrato_laboral":
        if aprobado or tiene_contrato:
            return _CELDA_OK
        return _CELDA_FALTA

    if aprobado or (cat, tipo) in vigentes:
        return _CELDA_OK
    if es_doc_opcional(cat, tipo):
        return _CELDA_OPCIONAL
    return _CELDA_FALTA


def _meta_contrato(sb, contrato_id: int) -> dict:
    try:
        rows = (
            sb.table("contratos")
            .select("id, numero, contratista, nit, logo_contratista, objeto")
            .eq("id", int(contrato_id))
            .limit(1)
            .execute()
            .data
            or []
        )
        return rows[0] if rows else {}
    except Exception as exc:
        _log.warning("estado documental meta contrato: %s", exc)
        return {}


def build_estado_documental_matrix(
    sb,
    contrato_id: int,
    *,
    filtro: str = "todos",
) -> dict:
    """
    Construye la matriz del informe.
    Returns dict with formato, filtro, columns, rows, meta.
    """
    from rrhh_service import list_trabajadores

    filtro_n = (filtro or "todos").strip().lower()
    if filtro_n not in FILTROS:
        raise ValueError("Filtro inválido. Use: todos | aprobado | sin_aprobar")

    todos = list_trabajadores(sb, int(contrato_id)) or []
    trabajadores = _filtrar_trabajadores(todos, filtro_n)
    ids: List[int] = []
    for t in trabajadores:
        try:
            ids.append(int(t["id"]))
        except (TypeError, ValueError, KeyError):
            continue

    custom = _custom_tipos(sb, int(contrato_id))
    doc_cols = _columnas_checklist(custom)
    docs_map, doc_labels = _fetch_docs_map(sb, int(contrato_id), ids)
    contratos = _fetch_contratos_set(sb, int(contrato_id), ids)

    # Tipos «Otro» vigentes no listados en catálogo → agregar columnas al final del grupo.
    extra_keys = set()
    for _tid, vig in docs_map.items():
        for cat, tipo in vig:
            if cat in ("soporte", "ingreso") and tipo.startswith("ext_"):
                if not any(c["categoria"] == cat and c["tipo"] == tipo for c in doc_cols):
                    extra_keys.add((cat, tipo))
    for cat, tipo in sorted(extra_keys):
        grupo = "Documentos de soporte" if cat == "soporte" else "Documentos de ingreso"
        # Insertar antes de afiliación / al final del grupo
        insert_at = len(doc_cols)
        for i, c in enumerate(doc_cols):
            if cat == "soporte" and c["categoria"] == "ingreso":
                insert_at = i
                break
            if cat == "ingreso" and c["categoria"] == "afiliacion":
                insert_at = i
                break
        label = doc_labels.get((cat, tipo)) or (
            tipo.replace("ext_", "").replace("_", " ").strip().title() or tipo
        )
        doc_cols.insert(insert_at, {
            "key": f"{cat}:{tipo}",
            "categoria": cat,
            "tipo": tipo,
            "label": label,
            "grupo": grupo,
        })

    identity_cols = [
        {"key": "nombre", "label": "Nombre"},
        {"key": "empresa", "label": "Empresa"},
        {"key": "cargo", "label": "Cargo"},
        {"key": "documento", "label": "Documento de identidad"},
    ]

    rows_out: List[dict] = []
    for t in trabajadores:
        try:
            tid = int(t["id"])
        except (TypeError, ValueError, KeyError):
            continue
        aprobado = _es_aprobado(t)
        vigentes = docs_map.get(tid) or set()
        tiene_c = tid in contratos
        celdas = {}
        for col in doc_cols:
            celdas[col["key"]] = _celda_doc(
                t, col, vigentes=vigentes, tiene_contrato=tiene_c, aprobado=aprobado,
            )
        rows_out.append({
            "id": tid,
            "nombre": _nombre_informe(t),
            "empresa": _empresa_informe(t),
            "cargo": str(t.get("cargo_aspira") or "").strip() or "—",
            "documento": _documento_identidad(t),
            "doc_validacion_estado": str(t.get("doc_validacion_estado") or "pendiente").lower(),
            "celdas": celdas,
        })

    meta = _meta_contrato(sb, int(contrato_id))
    now = _ahora_bogota()
    return {
        "formato": dict(ESTADO_DOCUMENTAL_FORMATO),
        "filtro": filtro_n,
        "generado_en": now.isoformat(),
        "fecha_calidad": _fecha_calidad(now),
        "contrato": {
            "id": int(contrato_id),
            "numero": meta.get("numero") or "",
            "contratista": meta.get("contratista") or "",
            "nit": meta.get("nit") or "",
            "objeto": meta.get("objeto") or "",
            "logo_contratista": meta.get("logo_contratista") or "",
        },
        "identity_columns": identity_cols,
        "doc_columns": doc_cols,
        "rows": rows_out,
        "total": len(rows_out),
        "leyenda": LEYENDA_ICONOS,
    }


def _filtro_label(filtro: str) -> str:
    return {
        "todos": "Todos",
        "aprobado": "Personal aprobado",
        "sin_aprobar": "Personal sin aprobar",
    }.get(filtro, filtro)


# ── Excel ────────────────────────────────────────────────────────────────────

def build_estado_documental_xlsx(matrix: dict) -> Tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    fmt = matrix.get("formato") or ESTADO_DOCUMENTAL_FORMATO
    identity = matrix.get("identity_columns") or []
    doc_cols = matrix.get("doc_columns") or []
    rows = matrix.get("rows") or []
    contrato = matrix.get("contrato") or {}
    n_cols = max(5, len(identity) + len(doc_cols))

    wb = Workbook()
    ws = wb.active
    ws.title = "Estado documental"

    fill_title = PatternFill("solid", fgColor="0F2942")
    fill_head = PatternFill("solid", fgColor="DDEFF8")
    fill_ok = PatternFill("solid", fgColor="DCFCE7")
    fill_falta = PatternFill("solid", fgColor="FEE2E2")
    fill_na = PatternFill("solid", fgColor="F1F5F9")
    font_title = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    font_calidad = Font(name="Calibri", bold=True, size=9, color="FFFFFF")
    font_head = Font(name="Calibri", bold=True, size=9, color="0F2942")
    font_body = Font(name="Calibri", size=9, color="0F2942")
    thin = Border(
        left=Side(style="thin", color="94A3B8"),
        right=Side(style="thin", color="94A3B8"),
        top=Side(style="thin", color="94A3B8"),
        bottom=Side(style="thin", color="94A3B8"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Encabezado 3 secciones: A | B:D | E…n  (mismo layout que Precios Subcontratistas)
    title_end = min(4, n_cols)
    right_start = min(5, n_cols)
    for c in range(1, n_cols + 1):
        cell = ws.cell(1, c, "")
        cell.fill = fill_title
        cell.border = thin
        cell.alignment = center
    if title_end >= 2:
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=title_end)
    if right_start <= n_cols and right_start > title_end:
        ws.merge_cells(start_row=1, start_column=right_start, end_row=1, end_column=n_cols)

    ws.cell(1, 1).value = "ClaraCore"
    ws.cell(1, 1).font = font_calidad
    titulo_cell = ws.cell(1, 2 if n_cols >= 2 else 1)
    titulo_cell.value = fmt.get("nombre_mayus") or fmt.get("nombre") or "ESTADO DOCUMENTAL"
    titulo_cell.font = font_title
    if right_start <= n_cols:
        cal = ws.cell(1, right_start)
        cal.value = (
            f"{fmt.get('codigo') or 'CC-RRHH-DOC-001'}\n"
            f"Versión {fmt.get('version') or '1.0'}\n"
            f"Fecha {matrix.get('fecha_calidad') or _fecha_calidad()}"
        )
        cal.font = font_calidad
        cal.alignment = center
    ws.row_dimensions[1].height = 48

    # Meta contrato
    ws.cell(2, 1, "Contrato").font = font_head
    ws.cell(2, 2, str(contrato.get("numero") or "—")).font = font_body
    if n_cols >= 4:
        ws.cell(2, 3, "Objeto").font = font_head
        ws.merge_cells(start_row=2, start_column=4, end_row=2, end_column=n_cols)
        ws.cell(2, 4, str(contrato.get("objeto") or "—")).font = font_body
    ws.cell(3, 1, "Contratista").font = font_head
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=max(2, min(4, n_cols)))
    ws.cell(3, 2, str(contrato.get("contratista") or "—")).font = font_body
    if n_cols >= 5:
        ws.cell(3, 5, "NIT").font = font_head
        ws.merge_cells(start_row=3, start_column=6, end_row=3, end_column=n_cols) if n_cols >= 6 else None
        ws.cell(3, min(6, n_cols), str(contrato.get("nit") or "—")).font = font_body
    ws.cell(4, 1, "Filtro").font = font_head
    ws.cell(4, 2, _filtro_label(matrix.get("filtro") or "todos")).font = font_body
    ws.cell(4, 3, "Colaboradores").font = font_head
    ws.cell(4, 4, int(matrix.get("total") or 0)).font = font_body
    ws.cell(5, 1, "Leyenda").font = font_head
    leyenda = matrix.get("leyenda") or LEYENDA_ICONOS
    if n_cols >= 2:
        ws.merge_cells(start_row=5, start_column=2, end_row=5, end_column=n_cols)
    ws.cell(5, 2, leyenda).font = font_body

    n_id = max(1, len(identity))
    # Header tabla
    header_row = 7
    all_headers = [c["label"] for c in identity] + [c["label"] for c in doc_cols]
    for i, label in enumerate(all_headers, start=1):
        cell = ws.cell(header_row, i, label)
        cell.fill = fill_head
        cell.font = font_head
        cell.border = thin
        cell.alignment = center
    ws.row_dimensions[header_row].height = 36

    for ri, row in enumerate(rows):
        r = header_row + 1 + ri
        values = [
            row.get("nombre"),
            row.get("empresa"),
            row.get("cargo"),
            row.get("documento"),
        ]
        for col in doc_cols:
            values.append((row.get("celdas") or {}).get(col["key"], _CELDA_FALTA))
        for ci, val in enumerate(values, start=1):
            cell = ws.cell(r, ci, val)
            cell.font = font_body
            cell.border = thin
            cell.alignment = left if ci <= n_id else center
            if ci > n_id:
                s = str(val)
                if s == _CELDA_OK:
                    cell.fill = fill_ok
                elif s == _CELDA_FALTA:
                    cell.fill = fill_falta
                elif s == _CELDA_OPCIONAL:
                    cell.fill = fill_na

    # Nombre más ancho; documentos ~50% del ancho previo (12 → 6).
    widths = [36, 10, 16, 16] + [6] * len(doc_cols)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A8"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    codigo = (fmt.get("codigo") or "CC-RRHH-DOC-001").replace("/", "-")
    fname = f"{codigo}_estado_documental_{matrix.get('filtro') or 'todos'}.xlsx"
    return data, fname


# ── PDF ──────────────────────────────────────────────────────────────────────

def build_estado_documental_pdf(matrix: dict) -> Tuple[bytes, str]:
    from topografia_utils import to_pdf_bytes

    fmt = matrix.get("formato") or ESTADO_DOCUMENTAL_FORMATO
    identity = matrix.get("identity_columns") or []
    doc_cols = matrix.get("doc_columns") or []
    rows = matrix.get("rows") or []
    contrato = matrix.get("contrato") or {}

    def esc(s: Any) -> str:
        return (
            str(s if s is not None else "")
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )

    # Equivalentes tipográficos para PDF (Helvetica no trae emoji).
    pdf_icon = {
        _CELDA_OK: ("ok", "&#10004;"),       # ✓
        _CELDA_FALTA: ("falta", "&#10008;"),  # ✘
        _CELDA_OPCIONAL: ("na", "&#9679;"),   # ●
    }
    leyenda_pdf = (
        "&#10004; Cargado &nbsp;&middot;&nbsp; "
        "&#10008; Obligatorio faltante &nbsp;&middot;&nbsp; "
        "&#9679; Opcional no cargado"
    )

    # Nombre ~2× más ancho; docs ~50% (reparto del espacio liberado).
    name_w = 18
    emp_w = 7
    cargo_w = 10
    docid_w = 11
    doc_w = max(3.2, (100 - name_w - emp_w - cargo_w - docid_w) / max(1, len(doc_cols)))

    th_id = (
        f"<th class='name' style='width:{name_w}%'>{esc(identity[0]['label']) if identity else 'Nombre'}</th>"
        f"<th style='width:{emp_w}%'>{esc(identity[1]['label']) if len(identity) > 1 else 'Empresa'}</th>"
        f"<th style='width:{cargo_w}%'>{esc(identity[2]['label']) if len(identity) > 2 else 'Cargo'}</th>"
        f"<th style='width:{docid_w}%'>{esc(identity[3]['label']) if len(identity) > 3 else 'Documento'}</th>"
    )
    th_doc = "".join(
        f"<th class='doc' style='width:{doc_w:.1f}%'>{esc(c['label'])}</th>" for c in doc_cols
    )

    body_rows = []
    for row in rows:
        tds = [
            f"<td class='name'>{esc(row.get('nombre'))}</td>",
            f"<td class='emp'>{esc(row.get('empresa'))}</td>",
            f"<td>{esc(row.get('cargo'))}</td>",
            f"<td>{esc(row.get('documento'))}</td>",
        ]
        for col in doc_cols:
            v = (row.get("celdas") or {}).get(col["key"], _CELDA_FALTA)
            cls, glyph = pdf_icon.get(v, ("falta", "&#10008;"))
            tds.append(f"<td class='{cls}'>{glyph}</td>")
        body_rows.append("<tr>" + "".join(tds) + "</tr>")

    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"/>
<style>
  @page {{ size: letter landscape; margin: 10mm; }}
  body {{ font-family: Helvetica, Arial, sans-serif; font-size: 8pt; color: #0F2942; }}
  .hdr {{ width: 100%; border-collapse: collapse; margin-bottom: 8px; }}
  .hdr td {{ border: 1px solid #334155; background: #0F2942; color: #fff; padding: 8px; vertical-align: middle; }}
  .hdr .logo {{ width: 18%; text-align: center; font-size: 9pt; font-weight: bold; }}
  .hdr .title {{ width: 52%; text-align: center; font-size: 13pt; font-weight: bold; }}
  .hdr .calidad {{ width: 30%; text-align: center; font-size: 8pt; white-space: pre-line; }}
  .meta {{ width: 100%; border-collapse: collapse; margin-bottom: 6px; font-size: 8pt; }}
  .meta td {{ border: 1px solid #94A3B8; padding: 3px 6px; }}
  .meta .lbl {{ font-weight: bold; background: #F1F5F9; width: 14%; }}
  .leyenda {{ margin: 0 0 8px 0; font-size: 7.5pt; color: #334155; }}
  table.grid {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
  table.grid th, table.grid td {{ border: 1px solid #94A3B8; padding: 2px 3px; word-wrap: break-word; }}
  table.grid th {{ background: #DDEFF8; font-size: 6.5pt; }}
  table.grid th.doc {{ font-size: 5.5pt; }}
  table.grid td.name {{ font-size: 7.5pt; }}
  table.grid td.emp {{ font-size: 7pt; text-align: center; font-weight: bold; }}
  .ok {{ background: #DCFCE7; text-align: center; font-weight: bold; color: #166534; font-size: 9pt; }}
  .falta {{ background: #FEE2E2; text-align: center; font-weight: bold; color: #991B1B; font-size: 9pt; }}
  .na {{ background: #F1F5F9; text-align: center; color: #64748B; font-size: 9pt; }}
</style></head><body>
<table class="hdr"><tr>
  <td class="logo">ClaraCore</td>
  <td class="title">{esc(fmt.get('nombre_mayus') or fmt.get('nombre'))}</td>
  <td class="calidad">{esc(fmt.get('codigo'))}
Versión {esc(fmt.get('version'))}
Fecha {esc(matrix.get('fecha_calidad'))}</td>
</tr></table>
<table class="meta">
<tr>
  <td class="lbl">Contrato</td><td>{esc(contrato.get('numero'))}</td>
  <td class="lbl">Contratista</td><td>{esc(contrato.get('contratista'))}</td>
  <td class="lbl">NIT</td><td>{esc(contrato.get('nit'))}</td>
</tr>
<tr>
  <td class="lbl">Filtro</td><td>{esc(_filtro_label(matrix.get('filtro') or 'todos'))}</td>
  <td class="lbl">Colaboradores</td><td colspan="3">{esc(matrix.get('total') or 0)}</td>
</tr>
</table>
<div class="leyenda">{leyenda_pdf}</div>
<table class="grid">
<thead><tr>{th_id}{th_doc}</tr></thead>
<tbody>
{''.join(body_rows) if body_rows else '<tr><td colspan="99">Sin colaboradores para el filtro seleccionado.</td></tr>'}
</tbody>
</table>
</body></html>"""

    pdf = to_pdf_bytes(html, landscape=True)
    codigo = (fmt.get("codigo") or "CC-RRHH-DOC-001").replace("/", "-")
    fname = f"{codigo}_estado_documental_{matrix.get('filtro') or 'todos'}.pdf"
    return pdf, fname
