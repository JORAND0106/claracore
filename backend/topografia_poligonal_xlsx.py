"""Exportación Excel de poligonal con fórmulas vivas (cierre + Bowditch + esquema).

Hoja «Cartera»:
  - Encabezado tipo PDF (contrato, equipo, etc.)
  - Fila de arranque + filas de estaciones (ángulos/distancias = entrada)
  - Columnas calculadas con fórmulas Excel (ángulo derivado, azimut corregido,
    proyecciones, correcciones Bowditch, coordenadas ajustadas)
  - Bloques de cierre angular y lineal con semáforo CUMPLE/NO CUMPLE

Hoja «Esquema»:
  - Tabla Este/Norte ajustados (fórmulas que referencian Cartera)
  - Gráfico de dispersión que se actualiza al editar datos de campo
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── estilos ──────────────────────────────────────────────────────────────────
_SIDE = Side(style="thin", color="94A3B8")
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
_FILL_HDR = PatternFill("solid", fgColor="1E40AF")
_FILL_HDR2 = PatternFill("solid", fgColor="475569")
_FILL_INPUT = PatternFill("solid", fgColor="FEF9C3")  # amarillo: editable
_FILL_CALC = PatternFill("solid", fgColor="F1F5F9")
_FILL_OK = PatternFill("solid", fgColor="DCFCE7")
_FILL_BAD = PatternFill("solid", fgColor="FEE2E2")
_FILL_ARRANQUE = PatternFill("solid", fgColor="DCFCE7")
_FILL_TITLE = PatternFill("solid", fgColor="EFF6FF")
_FONT_HDR = Font(bold=True, color="FFFFFF", size=9)
_FONT_TITLE = Font(bold=True, size=14, color="1E40AF")
_FONT_SUB = Font(bold=True, size=10, color="0F172A")
_FONT_MUTED = Font(size=8, color="64748B")
_FONT_CELL = Font(size=9, color="0F172A")
_AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
_AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
_AL_R = Alignment(horizontal="right", vertical="center")


def _cell(ws, row: int, col: int, value=None, *, fill=None, font=None, align=None, border=True, num_fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    if fill:
        c.fill = fill
    if font:
        c.font = font
    if align:
        c.alignment = align
    if border:
        c.border = _BORDER
    if num_fmt:
        c.number_format = num_fmt
    return c


def _legs_estacion(estaciones: list) -> List[dict]:
    """Tramos de la poligonal: estaciones con distancia > 0, en orden."""
    out = []
    for e in sorted(estaciones or [], key=lambda x: int(x.get("orden") or 0)):
        if (e.get("tipo_punto") or "auxiliar") != "estacion":
            continue
        if float(e.get("distancia") or 0) <= 1e-9:
            continue
        out.append(e)
    return out


def _es_azimut_directo(legs: list) -> bool:
    return any(
        (e.get("metodo_azimut") == "coordenadas") or bool(e.get("angulo_derivado_para_cierre"))
        for e in (legs or [])
    )


def _az_campo_lectura(e: dict, *, azimut_directo: bool) -> Optional[float]:
    """Lectura de campo para la columna de entrada D.

    Azimut directo: ``angulo_medido`` es la lectura de azimut del equipo.
    Ángulos (ceros atrás): ``angulo_medido`` es el ángulo horizontal observado.
    """
    if e.get("angulo_medido") is not None:
        try:
            return float(e["angulo_medido"]) % 360.0
        except (TypeError, ValueError):
            pass
    if azimut_directo and e.get("azimut") is not None:
        try:
            return float(e["azimut"]) % 360.0
        except (TypeError, ValueError):
            pass
    return None


def _base_azimut_arranque(legs: list, armadas: Optional[list]) -> float:
    """Azimut base de la 1ª armada (reciproco / amarre), editable en B14."""
    if armadas:
        arms = sorted(armadas, key=lambda a: int(a.get("orden") or 0))
        if arms and arms[0].get("base_azimut") is not None:
            try:
                return float(arms[0]["base_azimut"]) % 360.0
            except (TypeError, ValueError):
                pass
    if not legs:
        return 0.0
    a0 = legs[0]
    az = None
    if a0.get("angulo_medido") is not None:
        try:
            az = float(a0["angulo_medido"]) % 360.0
        except (TypeError, ValueError):
            az = None
    if az is None and a0.get("azimut") is not None:
        try:
            az = float(a0["azimut"]) % 360.0
        except (TypeError, ValueError):
            az = None
    if a0.get("angulo_derivado") is not None and az is not None:
        return (az - float(a0["angulo_derivado"])) % 360.0
    if a0.get("base_azimut") is not None:
        try:
            return float(a0["base_azimut"]) % 360.0
        except (TypeError, ValueError):
            pass
    return float(az or 0.0)


def _hi_por_armada(armadas: Optional[list], e: dict) -> float:
    if not armadas:
        return float(e.get("altura_instrumento") or 0)
    by_id = {a.get("id"): a for a in armadas if a.get("id")}
    arm = by_id.get(e.get("armada_id"))
    if arm is None:
        for a in armadas:
            if a.get("orden") is not None and a.get("orden") == e.get("armada_orden"):
                arm = a
                break
    if arm and arm.get("altura_instrumento") is not None:
        return float(arm["altura_instrumento"])
    return float(e.get("altura_instrumento") or 0)


def build_poligonal_xlsx_bytes(
    *,
    contrato: dict,
    pol: dict,
    estaciones: list,
    punto_inicial: Optional[dict],
    punto_final: Optional[dict] = None,
    cierre: Optional[dict] = None,
    armadas: Optional[list] = None,
) -> bytes:
    """Genera .xlsx con fórmulas vivas. No depende de red."""
    del punto_final  # reservado; cierre abierto puede ampliarse después
    cierre = cierre or {}
    legs = _legs_estacion(estaciones)
    n = len(legs)
    pi = punto_inicial or {}
    sentido = (cierre.get("sentido") or pol.get("sentido") or "antihorario").lower()
    antihorario = sentido != "horario"
    tol_rel = int(cierre.get("tolerancia_relativa") or pol.get("tolerancia_relativa") or 30000)
    prec_ang = float(cierre.get("precision_angular_equipo_seg") or pol.get("precision_angular_seg") or 10.0)
    azimut_directo = _es_azimut_directo(legs)
    base_az = _base_azimut_arranque(legs, armadas)
    # En modo ángulos, B14 guarda el azimut del 1º tramo (se conserva en compensación).
    az1_seed = 0.0
    if legs:
        if azimut_directo:
            az1_seed = _az_campo_lectura(legs[0], azimut_directo=True) or 0.0
        elif legs[0].get("azimut") is not None:
            az1_seed = float(legs[0]["azimut"]) % 360.0
        else:
            az1_seed = (base_az + (_az_campo_lectura(legs[0], azimut_directo=False) or 0.0)) % 360.0

    wb = Workbook()
    ws = wb.active
    ws.title = "Cartera"

    # ── Encabezado (estilo PDF) ──────────────────────────────────────────────
    titulo = f"Poligonal trigonométrica — {pol.get('nombre') or ''}"
    ws.merge_cells("A1:T1")
    _cell(ws, 1, 1, titulo, fill=_FILL_TITLE, font=_FONT_TITLE, align=_AL_L, border=False)

    equipo = " / ".join(
        x
        for x in [
            pol.get("equipo_marca"),
            pol.get("equipo_referencia"),
            pol.get("equipo_serial") and f"S/N {pol.get('equipo_serial')}",
        ]
        if x
    ) or pol.get("equipo") or "—"

    meta = [
        ("Contrato", str(contrato.get("numero") or "—")),
        ("Objeto", str(contrato.get("objeto") or "—")[:120]),
        ("Contratista", str(contrato.get("contratista") or "—")),
        ("Interventoría", str(contrato.get("interventoria") or "—")),
        ("Equipo", str(equipo)),
        ("Operador / profesional", str(pol.get("operador") or "—")),
        ("Fecha campo", str(pol.get("fecha_campo") or "—")),
        ("Tipo", str(pol.get("tipo") or "cerrada")),
    ]
    for i, (lab, val) in enumerate(meta):
        r = 2 + i
        _cell(ws, r, 1, lab, font=_FONT_SUB, fill=_FILL_CALC, align=_AL_L)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        _cell(ws, r, 2, val, font=_FONT_CELL, align=_AL_L)

    # Parámetros de cálculo
    # B11 sentido_anti, B12 tol_rel, B13 prec_ang, B14 base/Az1, B15 n, B16 modo
    _cell(ws, 10, 1, "PARÁMETROS DE CÁLCULO", fill=_FILL_HDR2, font=_FONT_HDR, align=_AL_L)
    ws.merge_cells("A10:D10")
    param_b14_label = (
        "Azimut base arranque (°)" if azimut_directo else "Azimut 1º tramo (°)"
    )
    params = [
        (11, "Sentido antihorario (1=sí, 0=horario)", 1 if antihorario else 0),
        (12, "Tolerancia relativa plan (1:N)", tol_rel),
        (13, "Precisión angular equipo (\")", prec_ang),
        (14, param_b14_label, round(base_az if azimut_directo else az1_seed, 6)),
        (15, "Nº tramos (estaciones)", n),
        (16, "Modo azimut directo (1=sí, 0=ángulos)", 1 if azimut_directo else 0),
    ]
    for r, lab, val in params:
        _cell(ws, r, 1, lab, font=_FONT_MUTED, fill=_FILL_CALC, align=_AL_L)
        fill = _FILL_INPUT if r in (11, 12, 13, 14, 16) else _FILL_CALC
        num_fmt = "0" if r in (11, 12, 15, 16) else "0.######"
        _cell(ws, r, 2, val, fill=fill, font=_FONT_CELL, align=_AL_R, num_fmt=num_fmt)

    _cell(
        ws,
        17,
        1,
        "Celdas amarillas = entrada editable. El resto son fórmulas: al cambiar ángulo/distancia "
        "se recalculan cierre angular, lineal, Bowditch y el esquema (hoja Esquema).",
        font=_FONT_MUTED,
        border=False,
        align=_AL_L,
    )
    ws.merge_cells("A17:T17")

    # ── Tabla principal ─────────────────────────────────────────────────────
    HDR_ROW = 18
    START_ROW = 19
    FIRST_LEG = 20
    LAST_LEG = FIRST_LEG + n - 1 if n else FIRST_LEG - 1

    col_d_hdr = "Az campo °" if azimut_directo else "Ang.obs °"
    headers = [
        "#",  # A
        "Armada",  # B
        "Punto",  # C
        col_d_hdr,  # D INPUT
        "Dist m",  # E INPUT
        "Ang.vert °",  # F INPUT
        "HI m",  # G INPUT
        "HT m",  # H INPUT
        "Ang.deriv °",  # I
        "Ang.corr °",  # J
        "Az corr °",  # K
        "ΔN m",  # L
        "ΔE m",  # M
        "ΔZ m",  # N
        "Corr.N m",  # O
        "Corr.E m",  # P
        "Corr.Z m",  # Q
        "Norte aj.",  # R
        "Este aj.",  # S
        "Cota aj.",  # T
    ]
    for c, h in enumerate(headers, 1):
        _cell(ws, HDR_ROW, c, h, fill=_FILL_HDR, font=_FONT_HDR, align=_AL_C)

    # Arranque
    _cell(ws, START_ROW, 1, "—", fill=_FILL_ARRANQUE, align=_AL_C)
    _cell(ws, START_ROW, 2, "—", fill=_FILL_ARRANQUE, align=_AL_C)
    _cell(
        ws,
        START_ROW,
        3,
        str(pi.get("nombre") or "Arranque"),
        fill=_FILL_ARRANQUE,
        font=Font(bold=True, size=9),
        align=_AL_L,
    )
    for c in range(4, 18):
        _cell(ws, START_ROW, c, "—", fill=_FILL_ARRANQUE, align=_AL_C, font=_FONT_MUTED)
    _cell(
        ws,
        START_ROW,
        18,
        float(pi["norte"]) if pi.get("norte") is not None else 0.0,
        fill=_FILL_INPUT,
        align=_AL_R,
        num_fmt="0.0000",
    )
    _cell(
        ws,
        START_ROW,
        19,
        float(pi["este"]) if pi.get("este") is not None else 0.0,
        fill=_FILL_INPUT,
        align=_AL_R,
        num_fmt="0.0000",
    )
    _cell(
        ws,
        START_ROW,
        20,
        float(pi["cota"]) if pi.get("cota") is not None else None,
        fill=_FILL_INPUT,
        align=_AL_R,
        num_fmt="0.0000",
    )

    # Referencias fijas del bloque de cierres (columna W):
    # W19 Diff angular (°), W21 ErrN=ΣΔN, W22 ErrE=ΣΔE, W23 Peri, W24 ErrZ
    for i, e in enumerate(legs):
        r = FIRST_LEG + i
        lectura = _az_campo_lectura(e, azimut_directo=azimut_directo)
        dist = float(e.get("distancia") or 0)
        ang_v = e.get("angulo_vertical")
        hi = _hi_por_armada(armadas, e)
        ht = e.get("altura_objetivo")
        if ht is None:
            ht = 0

        _cell(ws, r, 1, e.get("orden") or (i + 1), fill=_FILL_CALC, align=_AL_C)
        _cell(ws, r, 2, e.get("armada_orden") or "—", fill=_FILL_CALC, align=_AL_C)
        _cell(
            ws,
            r,
            3,
            str(e.get("nombre_punto") or f"P{i + 1}"),
            fill=_FILL_CALC,
            font=Font(bold=True, size=9),
            align=_AL_L,
        )
        _cell(
            ws,
            r,
            4,
            round(lectura, 8) if lectura is not None else 0.0,
            fill=_FILL_INPUT,
            align=_AL_R,
            num_fmt="0.000000",
        )
        _cell(ws, r, 5, dist, fill=_FILL_INPUT, align=_AL_R, num_fmt="0.000")
        _cell(
            ws,
            r,
            6,
            float(ang_v) if ang_v is not None else 90.0,
            fill=_FILL_INPUT,
            align=_AL_R,
            num_fmt="0.000000",
        )
        _cell(ws, r, 7, float(hi), fill=_FILL_INPUT, align=_AL_R, num_fmt="0.000")
        _cell(ws, r, 8, float(ht), fill=_FILL_INPUT, align=_AL_R, num_fmt="0.000")

        # I: Ang.deriv
        if azimut_directo:
            if i == 0:
                _cell(ws, r, 9, f"=MOD(D{r}-$B$14,360)", fill=_FILL_CALC, align=_AL_R, num_fmt="0.000000")
            else:
                prev = r - 1
                _cell(
                    ws,
                    r,
                    9,
                    f"=MOD(D{r}-D{prev}-180,360)",
                    fill=_FILL_CALC,
                    align=_AL_R,
                    num_fmt="0.000000",
                )
        else:
            # Ángulo observado de campo (entrada D)
            _cell(ws, r, 9, f"=D{r}", fill=_FILL_CALC, align=_AL_R, num_fmt="0.000000")

        # J: Ang.corr = Ang − Diff/n   (corr = −Diff/n)
        _cell(ws, r, 10, f"=I{r}-$W$19/$B$15", fill=_FILL_CALC, align=_AL_R, num_fmt="0.000000")

        # K: Az corr — conserva 1º; luego MOD(prev+180+Ang.corr,360)
        if azimut_directo:
            if i == 0:
                _cell(ws, r, 11, f"=D{r}", fill=_FILL_CALC, align=_AL_R, num_fmt="0.000000")
            else:
                prev = r - 1
                _cell(
                    ws,
                    r,
                    11,
                    f"=MOD(K{prev}+180+J{r},360)",
                    fill=_FILL_CALC,
                    align=_AL_R,
                    num_fmt="0.000000",
                )
        else:
            if i == 0:
                _cell(ws, r, 11, "=$B$14", fill=_FILL_CALC, align=_AL_R, num_fmt="0.000000")
            else:
                prev = r - 1
                _cell(
                    ws,
                    r,
                    11,
                    f"=MOD(K{prev}+180+J{r},360)",
                    fill=_FILL_CALC,
                    align=_AL_R,
                    num_fmt="0.000000",
                )

        _cell(ws, r, 12, f"=E{r}*COS(RADIANS(K{r}))", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")
        _cell(ws, r, 13, f"=E{r}*SIN(RADIANS(K{r}))", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")
        _cell(
            ws,
            r,
            14,
            f"=IF(ABS(SIN(RADIANS(F{r})))<1E-9,0,G{r}+E{r}*COS(RADIANS(F{r}))/SIN(RADIANS(F{r}))-H{r})",
            fill=_FILL_CALC,
            align=_AL_R,
            num_fmt="0.0000",
        )
        _cell(ws, r, 15, f"=IF($W$23<=0,0,-$W$21*E{r}/$W$23)", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")
        _cell(ws, r, 16, f"=IF($W$23<=0,0,-$W$22*E{r}/$W$23)", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")
        _cell(ws, r, 17, f"=IF($W$23<=0,0,-$W$24*E{r}/$W$23)", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")

        prev_r = START_ROW if i == 0 else (r - 1)
        _cell(ws, r, 18, f"=R{prev_r}+L{r}+O{r}", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")
        _cell(ws, r, 19, f"=S{prev_r}+M{r}+P{r}", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")
        _cell(
            ws,
            r,
            20,
            f'=IF(T{prev_r}="","",T{prev_r}+N{r}+Q{r})',
            fill=_FILL_CALC,
            align=_AL_R,
            num_fmt="0.0000",
        )

    if n == 0:
        FIRST_LEG = START_ROW
        LAST_LEG = START_ROW

    # ── Bloque auxiliar W (referencias internas) ─────────────────────────────
    BLK = 22  # V labels, W values
    _cell(ws, 18, BLK, "REF. FÓRMULAS", fill=_FILL_HDR2, font=_FONT_HDR, align=_AL_L)
    ws.merge_cells(start_row=18, start_column=BLK, end_row=18, end_column=BLK + 1)

    if n > 0:
        rng_I = f"I{FIRST_LEG}:I{LAST_LEG}"
        rng_L = f"L{FIRST_LEG}:L{LAST_LEG}"
        rng_M = f"M{FIRST_LEG}:M{LAST_LEG}"
        rng_N = f"N{FIRST_LEG}:N{LAST_LEG}"
        rng_E = f"E{FIRST_LEG}:E{LAST_LEG}"
    else:
        rng_I = rng_L = rng_M = rng_N = rng_E = "I19"

    # W19 Diff — usado por Ang.corr
    _cell(ws, 19, BLK, "Diff angular (°)", fill=_FILL_CALC, font=_FONT_MUTED, align=_AL_L)
    _cell(
        ws,
        19,
        BLK + 1,
        f"=IF($B$15<=0,0,SUM({rng_I})-IF($B$11=1,($B$15-2)*180,($B$15+2)*180))",
        fill=_FILL_CALC,
        font=_FONT_CELL,
        align=_AL_R,
        num_fmt="0.000000",
    )
    _cell(ws, 20, BLK, "(reservado)", fill=_FILL_CALC, font=_FONT_MUTED, align=_AL_L)
    _cell(ws, 20, BLK + 1, None, fill=_FILL_CALC)

    _cell(ws, 21, BLK, "ErrN = ΣΔN (m)", fill=_FILL_CALC, font=_FONT_MUTED, align=_AL_L)
    _cell(
        ws,
        21,
        BLK + 1,
        f"=IF($B$15<=0,0,SUM({rng_L}))",
        fill=_FILL_CALC,
        font=_FONT_CELL,
        align=_AL_R,
        num_fmt="0.0000",
    )
    _cell(ws, 22, BLK, "ErrE = ΣΔE (m)", fill=_FILL_CALC, font=_FONT_MUTED, align=_AL_L)
    _cell(
        ws,
        22,
        BLK + 1,
        f"=IF($B$15<=0,0,SUM({rng_M}))",
        fill=_FILL_CALC,
        font=_FONT_CELL,
        align=_AL_R,
        num_fmt="0.0000",
    )
    _cell(ws, 23, BLK, "Perímetro (m)", fill=_FILL_CALC, font=_FONT_MUTED, align=_AL_L)
    _cell(
        ws,
        23,
        BLK + 1,
        f"=IF($B$15<=0,0,SUM({rng_E}))",
        fill=_FILL_CALC,
        font=_FONT_CELL,
        align=_AL_R,
        num_fmt="0.000",
    )
    _cell(ws, 24, BLK, "ErrZ = ΣΔZ (m)", fill=_FILL_CALC, font=_FONT_MUTED, align=_AL_L)
    _cell(
        ws,
        24,
        BLK + 1,
        f"=IF($B$15<=0,0,SUM({rng_N}))",
        fill=_FILL_CALC,
        font=_FONT_CELL,
        align=_AL_R,
        num_fmt="0.0000",
    )

    # ── Cierre angular (columna Y) ───────────────────────────────────────────
    YA = 25
    _cell(ws, 18, YA, "CIERRE ANGULAR", fill=_FILL_HDR2, font=_FONT_HDR, align=_AL_L)
    ws.merge_cells(start_row=18, start_column=YA, end_row=18, end_column=YA + 1)
    # Valores en columna Z (YA+1)
    ang_rows = [
        (19, "Sentido", '=IF($B$11=1,"Antihorario (int.)","Horario (ext.)")', None),
        (20, "Áng. / Vért.", '=$B$15&" / "&$B$15', None),
        (21, "Σ Observada (°)", f"=IF($B$15<=0,0,SUM({rng_I}))", "0.000000"),
        (22, "Σ Teórica (°)", '=IF($B$11=1,($B$15-2)*180,($B$15+2)*180)', "0.000000"),
        (23, "Diferencia (\")", "=$W$19*3600", "0.00"),
        (24, "Tolerancia (\")", "=$B$13*SQRT(MAX($B$15,1))", "0.0"),
        (25, "Estado", '=IF(ABS(Z23)<=Z24,"CUMPLE","REVISAR")', None),
    ]
    for r, lab, formul, fmt in ang_rows:
        _cell(ws, r, YA, lab, fill=_FILL_CALC, font=_FONT_MUTED, align=_AL_L)
        cell = _cell(
            ws,
            r,
            YA + 1,
            formul,
            fill=_FILL_CALC,
            font=Font(bold=True, size=9) if r == 25 else _FONT_CELL,
            align=_AL_R,
            num_fmt=fmt,
        )
        del cell
    ws.conditional_formatting.add(
        f"{get_column_letter(YA + 1)}25",
        FormulaRule(formula=[f'{get_column_letter(YA + 1)}25="CUMPLE"'], fill=_FILL_OK),
    )
    ws.conditional_formatting.add(
        f"{get_column_letter(YA + 1)}25",
        FormulaRule(formula=[f'{get_column_letter(YA + 1)}25="REVISAR"'], fill=_FILL_BAD),
    )

    # ── Cierre lineal (columna AA) — ΔN/ΔE como en plataforma (−Σ proyecciones)
    LA = 27
    _cell(ws, 18, LA, "CIERRE LINEAL (campo)", fill=_FILL_HDR2, font=_FONT_HDR, align=_AL_L)
    ws.merge_cells(start_row=18, start_column=LA, end_row=18, end_column=LA + 1)
    # Valores en columna AB (LA+1)
    lin_rows = [
        (19, "Perímetro (m)", "=$W$23", "0.000"),
        (20, "ΔN (m)", "=-W21", "0.0000"),
        (21, "ΔE (m)", "=-W22", "0.0000"),
        (22, "ΔZ (m)", "=-W24", "0.0000"),
        (23, "Error lineal (m)", "=SQRT(W21^2+W22^2)", "0.0000"),
        (24, "Cierre 1:N", "=IF(AB23<=1E-9,1E9,ROUND(W23/AB23,0))", "0"),
        (25, "Tolerancia plan", '="1:"&$B$12', None),
        (26, "Estado", '=IF(AB24>=$B$12,"CUMPLE","NO CUMPLE")', None),
    ]
    for r, lab, formul, fmt in lin_rows:
        _cell(ws, r, LA, lab, fill=_FILL_CALC, font=_FONT_MUTED, align=_AL_L)
        _cell(
            ws,
            r,
            LA + 1,
            formul,
            fill=_FILL_CALC,
            font=Font(bold=True, size=9) if r == 26 else _FONT_CELL,
            align=_AL_R,
            num_fmt=fmt,
        )
    ws.conditional_formatting.add(
        f"{get_column_letter(LA + 1)}26",
        FormulaRule(formula=[f'{get_column_letter(LA + 1)}26="CUMPLE"'], fill=_FILL_OK),
    )
    ws.conditional_formatting.add(
        f"{get_column_letter(LA + 1)}26",
        FormulaRule(formula=[f'{get_column_letter(LA + 1)}26="NO CUMPLE"'], fill=_FILL_BAD),
    )

    note_row = max(LAST_LEG + 2, 28)
    _cell(
        ws,
        note_row,
        1,
        "Compensación Bowditch: Corr.N/E/Z = −ΣΔ × Dist / Perímetro; "
        "Norte/Este/Cota aj. acumulan Δ + corrección desde el arranque. "
        "Ángulos en grados decimales. Diff angular = Σobs − Σteórica; "
        "Ang.corr = Ang.deriv − Diff/n (reparte −Diff/n).",
        font=_FONT_MUTED,
        border=False,
        align=_AL_L,
    )
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=20)

    widths = {
        "A": 5,
        "B": 7,
        "C": 12,
        "D": 11,
        "E": 9,
        "F": 10,
        "G": 7,
        "H": 7,
        "I": 11,
        "J": 11,
        "K": 10,
        "L": 9,
        "M": 9,
        "N": 9,
        "O": 9,
        "P": 9,
        "Q": 9,
        "R": 11,
        "S": 11,
        "T": 10,
        "V": 18,
        "W": 12,
        "Y": 16,
        "Z": 14,
        "AA": 18,
        "AB": 12,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A19"
    ws.print_title_rows = "1:18"

    # ── Hoja Esquema ─────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Esquema")
    _cell(
        ws2,
        1,
        1,
        "Esquema de la poligonal (coordenadas ajustadas)",
        fill=_FILL_TITLE,
        font=_FONT_TITLE,
        border=False,
    )
    ws2.merge_cells("A1:D1")
    _cell(
        ws2,
        2,
        1,
        "Este/Norte referencian la hoja Cartera. Si edita azimut/ángulo o distancia allí, "
        "el gráfico se actualiza al recalcular en Excel.",
        font=_FONT_MUTED,
        border=False,
    )
    ws2.merge_cells("A2:F2")

    for c, h in enumerate(["#", "Punto", "Este aj. (m)", "Norte aj. (m)", "Cota aj. (m)"], 1):
        _cell(ws2, 4, c, h, fill=_FILL_HDR, font=_FONT_HDR, align=_AL_C)

    _cell(ws2, 5, 1, 0, fill=_FILL_ARRANQUE, align=_AL_C)
    _cell(ws2, 5, 2, f"=Cartera!C{START_ROW}", fill=_FILL_ARRANQUE, align=_AL_L)
    _cell(ws2, 5, 3, f"=Cartera!S{START_ROW}", fill=_FILL_ARRANQUE, align=_AL_R, num_fmt="0.0000")
    _cell(ws2, 5, 4, f"=Cartera!R{START_ROW}", fill=_FILL_ARRANQUE, align=_AL_R, num_fmt="0.0000")
    _cell(ws2, 5, 5, f"=Cartera!T{START_ROW}", fill=_FILL_ARRANQUE, align=_AL_R, num_fmt="0.0000")

    for i in range(n):
        r = 6 + i
        src = FIRST_LEG + i
        _cell(ws2, r, 1, i + 1, fill=_FILL_CALC, align=_AL_C)
        _cell(ws2, r, 2, f"=Cartera!C{src}", fill=_FILL_CALC, align=_AL_L)
        _cell(ws2, r, 3, f"=Cartera!S{src}", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")
        _cell(ws2, r, 4, f"=Cartera!R{src}", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")
        _cell(ws2, r, 5, f"=Cartera!T{src}", fill=_FILL_CALC, align=_AL_R, num_fmt="0.0000")

    close_r = 6 + n
    if n > 0:
        _cell(ws2, close_r, 1, n + 1, fill=_FILL_ARRANQUE, align=_AL_C)
        _cell(ws2, close_r, 2, f'=Cartera!C{START_ROW}&" (cierre)"', fill=_FILL_ARRANQUE, align=_AL_L)
        _cell(ws2, close_r, 3, f"=Cartera!S{START_ROW}", fill=_FILL_ARRANQUE, align=_AL_R, num_fmt="0.0000")
        _cell(ws2, close_r, 4, f"=Cartera!R{START_ROW}", fill=_FILL_ARRANQUE, align=_AL_R, num_fmt="0.0000")
        _cell(ws2, close_r, 5, f"=Cartera!T{START_ROW}", fill=_FILL_ARRANQUE, align=_AL_R, num_fmt="0.0000")

    for col, w in [("A", 5), ("B", 18), ("C", 14), ("D", 14), ("E", 12)]:
        ws2.column_dimensions[col].width = w

    if n > 0:
        chart = ScatterChart()
        chart.title = "Esquema — Este vs Norte (ajustado)"
        chart.style = 10
        chart.x_axis.title = "Este (m)"
        chart.y_axis.title = "Norte (m)"
        chart.height = 14
        chart.width = 18
        end_r = close_r
        xvalues = Reference(ws2, min_col=3, min_row=5, max_row=end_r)
        yvalues = Reference(ws2, min_col=4, min_row=5, max_row=end_r)
        series = Series(yvalues, xvalues, title="Poligonal")
        series.marker.symbol = "circle"
        series.marker.size = 7
        series.graphicalProperties.line.width = 15000
        chart.series.append(series)
        chart.legend = None
        ws2.add_chart(chart, "G4")

    _cell(ws2, close_r + 2, 1, "Leyenda Cartera:", font=_FONT_SUB, border=False)
    _cell(
        ws2,
        close_r + 3,
        1,
        "Amarillo = editable (Az/Ang campo, Dist, Ang.vert, HI, HT, arranque N/E/Z, parámetros)",
        font=_FONT_MUTED,
        border=False,
    )
    ws2.merge_cells(start_row=close_r + 3, start_column=1, end_row=close_r + 3, end_column=5)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
