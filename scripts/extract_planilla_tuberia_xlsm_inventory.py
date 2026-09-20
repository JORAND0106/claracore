#!/usr/bin/env python3
"""
Inventario literal de Planilla_Tuberia_original.xlsm.

Lee el archivo con openpyxl (data_only=False) y emite JSON estructurado:
hojas, merges, anchos/altos, celdas (valor/fórmula literal + formato),
validaciones, formato condicional y gráficos embebidos.

Uso:
  python3 scripts/extract_planilla_tuberia_xlsm_inventory.py \\
    [--src docs/topografia/planillas_tuberia/Planilla_Tuberia_original.xlsm] \\
    [--out docs/topografia/planillas_tuberia/inventario_planilla_tuberia_xlsm.json]
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string


def _rgb(color) -> Optional[str]:
    if color is None:
        return None
    try:
        rgb = getattr(color, "rgb", None)
        if rgb and isinstance(rgb, str) and rgb not in ("00000000", "None"):
            return rgb
        theme = getattr(color, "theme", None)
        tint = getattr(color, "tint", None)
        indexed = getattr(color, "indexed", None)
        out: dict[str, Any] = {}
        if theme is not None:
            out["theme"] = theme
        if tint not in (None, 0, 0.0):
            out["tint"] = tint
        if indexed is not None:
            out["indexed"] = indexed
        return out or None
    except Exception:
        return None


def _font(f) -> Optional[dict]:
    if f is None:
        return None
    d = {
        "name": f.name,
        "size": f.size,
        "bold": bool(f.bold),
        "italic": bool(f.italic),
        "underline": f.underline,
        "strike": bool(f.strike) if f.strike is not None else False,
        "color": _rgb(f.color),
    }
    if not any(v not in (None, False, 0) for k, v in d.items() if k != "name"):
        # still keep name/size if present
        if d["name"] is None and d["size"] is None:
            return None
    return d


def _fill(fl) -> Optional[dict]:
    if fl is None:
        return None
    pattern = getattr(fl, "patternType", None) or getattr(fl, "fill_type", None)
    if not pattern or pattern == "none":
        return None
    return {
        "patternType": pattern,
        "fgColor": _rgb(getattr(fl, "fgColor", None)),
        "bgColor": _rgb(getattr(fl, "bgColor", None)),
    }


def _alignment(a) -> Optional[dict]:
    if a is None:
        return None
    d = {
        "horizontal": a.horizontal,
        "vertical": a.vertical,
        "wrapText": bool(a.wrapText) if a.wrapText is not None else False,
        "textRotation": a.textRotation or 0,
        "indent": a.indent or 0,
        "shrinkToFit": bool(a.shrinkToFit) if a.shrinkToFit is not None else False,
    }
    if not any(d.values()):
        return None
    return d


def _border_side(s) -> Optional[dict]:
    if s is None or not s.style:
        return None
    return {"style": s.style, "color": _rgb(s.color)}


def _border(b) -> Optional[dict]:
    if b is None:
        return None
    d = {
        "left": _border_side(b.left),
        "right": _border_side(b.right),
        "top": _border_side(b.top),
        "bottom": _border_side(b.bottom),
        "diagonal": _border_side(b.diagonal),
    }
    if not any(d.values()):
        return None
    return d


def _jsonable(v: Any) -> Any:
    """Convierte tipos openpyxl / Excel a tipos JSON nativos."""
    if v is None or isinstance(v, (bool, str)):
        return v
    if isinstance(v, (int, float)):
        return v
    # openpyxl usa Integer/Float wrappers a veces
    try:
        from openpyxl.cell.cell import Cell  # noqa: F401
    except Exception:
        pass
    tname = type(v).__name__
    if tname in ("Integer", "Long", "Float", "Double"):
        try:
            return int(v) if tname in ("Integer", "Long") else float(v)
        except Exception:
            return str(v)
    if isinstance(v, dict):
        return {str(k): _jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [_jsonable(x) for x in v]
    try:
        return v.isoformat()
    except Exception:
        pass
    try:
        return int(v)
    except Exception:
        pass
    try:
        return float(v)
    except Exception:
        pass
    return str(v)


def _cell_value(v) -> Any:
    return _jsonable(v)


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _cell_relevant(cell) -> bool:
    """Incluir celdas con valor/fórmula o formato no trivial."""
    if cell.value is not None:
        return True
    if _fill(cell.fill):
        return True
    if _border(cell.border):
        return True
    nf = cell.number_format
    if nf and nf not in ("General", "0", "@"):
        return True
    return False


def extract_validations(ws) -> list[dict]:
    out = []
    dv = getattr(ws, "data_validations", None)
    if not dv:
        return out
    for v in getattr(dv, "dataValidation", []) or []:
        out.append({
            "type": v.type,
            "operator": v.operator,
            "formula1": v.formula1,
            "formula2": v.formula2,
            "allow_blank": v.allow_blank,
            "showDropDown": v.showDropDown,
            "showErrorMessage": v.showErrorMessage,
            "errorTitle": v.errorTitle,
            "error": v.error,
            "promptTitle": v.promptTitle,
            "prompt": v.prompt,
            "sqref": str(v.sqref) if v.sqref else None,
        })
    return out


def extract_conditional_formatting(ws) -> list[dict]:
    out = []
    cf = getattr(ws, "conditional_formatting", None)
    if not cf:
        return out
    try:
        items = list(cf._cf_rules.items())  # type: ignore[attr-defined]
    except Exception:
        try:
            items = [(str(r), cf[r]) for r in cf]
        except Exception:
            return out
    for rng, rules in items:
        rules_ser = []
        for rule in rules:
            rules_ser.append({
                "type": getattr(rule, "type", None),
                "operator": getattr(rule, "operator", None),
                "formula": list(getattr(rule, "formula", []) or []),
                "dxf": _serialize_dxf(getattr(rule, "dxf", None)),
                "priority": getattr(rule, "priority", None),
                "stopIfTrue": getattr(rule, "stopIfTrue", None),
                "text": getattr(rule, "text", None),
            })
        out.append({"range": str(rng), "rules": rules_ser})
    return out


def _serialize_dxf(dxf) -> Optional[dict]:
    if dxf is None:
        return None
    return {
        "font": _font(getattr(dxf, "font", None)),
        "fill": _fill(getattr(dxf, "fill", None)),
        "border": _border(getattr(dxf, "border", None)),
        "alignment": _alignment(getattr(dxf, "alignment", None)),
        "number_format": getattr(getattr(dxf, "numFmt", None), "formatCode", None)
        if getattr(dxf, "numFmt", None) else getattr(dxf, "number_format", None),
    }


def extract_charts(ws) -> list[dict]:
    out = []
    charts = list(getattr(ws, "_charts", []) or [])
    for i, ch in enumerate(charts):
        info: dict[str, Any] = {
            "index": i,
            "type": type(ch).__name__,
            "title": _chart_title(ch),
            "style": getattr(ch, "style", None),
            "anchor": _chart_anchor(ch),
            "series": [],
        }
        for si, ser in enumerate(getattr(ch, "series", []) or []):
            info["series"].append({
                "index": si,
                "title": _series_title(ser),
                "val": _data_ref(getattr(ser, "val", None)),
                "cat": _data_ref(getattr(ser, "cat", None)),
                "tx": _data_ref(getattr(ser, "tx", None)),
            })
        # axes
        for ax_name in ("x_axis", "y_axis", "z_axis"):
            ax = getattr(ch, ax_name, None)
            if ax is not None:
                info[ax_name] = {
                    "title": _chart_title(ax) if hasattr(ax, "title") else None,
                    "scaling": {
                        "min": getattr(getattr(ax, "scaling", None), "min", None),
                        "max": getattr(getattr(ax, "scaling", None), "max", None),
                    } if getattr(ax, "scaling", None) else None,
                }
        out.append(info)
    return out


def _chart_title(obj) -> Any:
    t = getattr(obj, "title", None)
    if t is None:
        return None
    try:
        # openpyxl ChartLabel / Title
        if hasattr(t, "tx") and t.tx is not None:
            rich = getattr(t.tx, "rich", None)
            if rich and getattr(rich, "p", None):
                texts = []
                for p in rich.p:
                    for r in getattr(p, "r", []) or []:
                        if getattr(r, "t", None):
                            texts.append(r.t)
                if texts:
                    return "".join(texts)
            st = getattr(t.tx, "strRef", None)
            if st is not None:
                return {"strRef": getattr(st, "f", None)}
        if isinstance(t, str):
            return t
        return str(t)
    except Exception:
        return str(t)


def _series_title(ser) -> Any:
    tx = getattr(ser, "tx", None)
    if tx is None:
        return None
    try:
        v = getattr(tx, "v", None)
        if v is not None:
            return v
        st = getattr(tx, "strRef", None)
        if st is not None:
            return {"strRef": getattr(st, "f", None)}
        return str(tx)
    except Exception:
        return str(tx)


def _data_ref(ref) -> Any:
    if ref is None:
        return None
    try:
        num = getattr(ref, "numRef", None)
        if num is not None:
            return {"numRef": getattr(num, "f", None)}
        st = getattr(ref, "strRef", None)
        if st is not None:
            return {"strRef": getattr(st, "f", None)}
        multi = getattr(ref, "multiLvlStrRef", None)
        if multi is not None:
            return {"multiLvlStrRef": getattr(multi, "f", None)}
        return str(ref)
    except Exception:
        return str(ref)


def _chart_anchor(ch) -> Any:
    anchor = getattr(ch, "anchor", None)
    if anchor is None:
        return None
    try:
        return str(anchor)
    except Exception:
        return repr(anchor)


def extract_images(ws) -> list[dict]:
    out = []
    for i, img in enumerate(getattr(ws, "_images", []) or []):
        out.append({
            "index": i,
            "path": getattr(img, "path", None),
            "anchor": str(getattr(img, "anchor", None)),
            "width": getattr(img, "width", None),
            "height": getattr(img, "height", None),
        })
    return out


def extract_sheet(ws) -> dict:
    dims = ws.dimensions
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    col_widths = {}
    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        dim = ws.column_dimensions.get(letter)
        width = dim.width if dim and dim.width is not None else None
        hidden = bool(dim.hidden) if dim else False
        col_widths[letter] = {"width": width, "hidden": hidden}

    row_heights = {}
    for row_idx in range(1, max_row + 1):
        dim = ws.row_dimensions.get(row_idx)
        height = dim.height if dim and dim.height is not None else None
        hidden = bool(dim.hidden) if dim else False
        if height is not None or hidden:
            row_heights[str(row_idx)] = {"height": height, "hidden": hidden}

    merges = [str(m) for m in ws.merged_cells.ranges]

    cells = []
    formulas = []
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            if not _cell_relevant(cell):
                continue
            val = cell.value
            entry = {
                "addr": cell.coordinate,
                "row": cell.row,
                "col": cell.column,
                "value": _cell_value(val),
                "is_formula": _is_formula(val),
                "number_format": cell.number_format,
                "font": _font(cell.font),
                "fill": _fill(cell.fill),
                "alignment": _alignment(cell.alignment),
                "border": _border(cell.border),
            }
            cells.append(entry)
            if entry["is_formula"]:
                formulas.append({"addr": cell.coordinate, "formula": val})

    # print areas / freeze
    freeze = ws.freeze_panes
    print_area = ws.print_area
    print_title_rows = ws.print_title_rows
    print_title_cols = ws.print_title_cols

    return {
        "title": ws.title,
        "sheet_state": ws.sheet_state,  # visible | hidden | veryHidden
        "dimensions": dims,
        "max_row": max_row,
        "max_column": max_col,
        "merged_cells": merges,
        "column_widths": col_widths,
        "row_heights": row_heights,
        "freeze_panes": freeze,
        "print_area": print_area,
        "print_title_rows": print_title_rows,
        "print_title_cols": print_title_cols,
        "cells": cells,
        "formulas": formulas,
        "data_validations": extract_validations(ws),
        "conditional_formatting": extract_conditional_formatting(ws),
        "charts": extract_charts(ws),
        "images": extract_images(ws),
    }


def extract_defined_names(wb) -> list[dict]:
    out = []
    try:
        dns = wb.defined_names
        for name in dns.definedName:
            out.append({
                "name": name.name,
                "attr_text": name.attr_text,
                "hidden": bool(name.hidden),
            })
    except Exception as exc:
        out.append({"error": str(exc)})
    return out


def build_inventory(src: Path) -> dict:
    wb = load_workbook(src, data_only=False, keep_vba=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        sheets.append(extract_sheet(ws))

    inv = {
        "source_file": str(src.name),
        "source_path": str(src),
        "openpyxl_version": __import__("openpyxl").__version__,
        "sheetnames": wb.sheetnames,
        "defined_names": extract_defined_names(wb),
        "sheets": sheets,
        "notes": [
            "Valores de celda son literales (fórmulas con data_only=False).",
            "Colores theme/tint no se resuelven a RGB absoluto (dependen del tema del libro).",
            "Gráficos: tipo, series y referencias numRef/strRef extraídos de openpyxl._charts.",
        ],
    }
    return inv


def build_human_summary(inv: dict) -> str:
    lines = []
    lines.append(f"# Inventario Planilla Tubería — {inv['source_file']}")
    lines.append("")
    lines.append(f"Hojas ({len(inv['sheetnames'])}): {', '.join(inv['sheetnames'])}")
    lines.append("")
    for sh in inv["sheets"]:
        lines.append(f"## Hoja: {sh['title']}  (state={sh['sheet_state']}, {sh['max_row']}×{sh['max_column']})")
        lines.append(f"- Merges ({len(sh['merged_cells'])}): {', '.join(sh['merged_cells'][:40])}"
                     + (" ..." if len(sh['merged_cells']) > 40 else ""))
        lines.append(f"- Fórmulas: {len(sh['formulas'])}")
        lines.append(f"- Validaciones: {len(sh['data_validations'])}")
        lines.append(f"- Formato condicional: {len(sh['conditional_formatting'])}")
        lines.append(f"- Gráficos: {len(sh['charts'])}")
        lines.append(f"- Imágenes: {len(sh['images'])}")
        lines.append("")
        # Labels: non-formula text cells
        labels = [
            c for c in sh["cells"]
            if isinstance(c["value"], str) and not c["is_formula"] and c["value"].strip()
        ]
        lines.append("### Rótulos / textos (literal)")
        for c in labels:
            lines.append(f"- {c['addr']}: {c['value']!r}  nf={c['number_format']!r}")
        lines.append("")
        lines.append("### Fórmulas (literal)")
        for f in sh["formulas"]:
            lines.append(f"- {f['addr']}: {f['formula']}")
        lines.append("")
        if sh["data_validations"]:
            lines.append("### Validaciones")
            for v in sh["data_validations"]:
                lines.append(f"- {v}")
            lines.append("")
        if sh["conditional_formatting"]:
            lines.append("### Formato condicional")
            for cf in sh["conditional_formatting"]:
                lines.append(f"- {cf}")
            lines.append("")
        if sh["charts"]:
            lines.append("### Gráficos")
            for ch in sh["charts"]:
                lines.append(f"- {ch}")
            lines.append("")
        # Column widths summary
        lines.append("### Anchos de columna (no default)")
        for letter, meta in sh["column_widths"].items():
            if meta["width"] is not None or meta["hidden"]:
                lines.append(f"- {letter}: width={meta['width']} hidden={meta['hidden']}")
        lines.append("")
        lines.append("### Altos de fila (explicitos)")
        for r, meta in sh["row_heights"].items():
            lines.append(f"- R{r}: height={meta['height']} hidden={meta['hidden']}")
        lines.append("")
    return "\n".join(lines)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--src",
        type=Path,
        default=Path("docs/topografia/planillas_tuberia/Planilla_Tuberia_original.xlsm"),
    )
    ap.add_argument(
        "--out",
        type=Path,
        default=Path("docs/topografia/planillas_tuberia/inventario_planilla_tuberia_xlsm.json"),
    )
    ap.add_argument(
        "--summary",
        type=Path,
        default=Path("docs/topografia/planillas_tuberia/inventario_planilla_tuberia_xlsm.md"),
    )
    args = ap.parse_args()
    if not args.src.exists():
        print(f"No existe: {args.src}", file=sys.stderr)
        return 1
    inv = build_inventory(args.src)
    inv = _jsonable(inv)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(inv, ensure_ascii=False, indent=2), encoding="utf-8")
    args.summary.write_text(build_human_summary(inv), encoding="utf-8")
    print(f"JSON → {args.out} ({args.out.stat().st_size} bytes)")
    print(f"MD   → {args.summary} ({args.summary.stat().st_size} bytes)")
    print(f"Hojas: {inv['sheetnames']}")
    for sh in inv["sheets"]:
        print(
            f"  - {sh['title']}: cells={len(sh['cells'])} formulas={len(sh['formulas'])} "
            f"merges={len(sh['merged_cells'])} charts={len(sh['charts'])} "
            f"validations={len(sh['data_validations'])} cf={len(sh['conditional_formatting'])}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
